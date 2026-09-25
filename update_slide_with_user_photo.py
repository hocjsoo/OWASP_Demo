import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor

pptx_path = "/media/hocjsoo/New Volume/OWASP_Demo/Slide_OWASP_SQL_Injection.pptx"
prs = Presentation(pptx_path)

# Slide 6 is index 5
slide6 = prs.slides[5]

# Let's inspect shapes on slide 6
print(f"Slide 6 has {len(slide6.shapes)} shapes")
