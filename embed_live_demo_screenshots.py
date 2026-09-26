import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor

pptx_path = "/media/hocjsoo/New Volume/OWASP_Demo/Slide_OWASP_Top3_Nhom_WNC_G01.pptx"
prs = Presentation(pptx_path)

# Slide indices (0-based):
# Slide 6 is index 5 (SQLi Live Demo)
# Slide 12 is index 11 (XSS Live Demo)
# Slide 17 is index 16 (CSRF Live Demo)

p_sqli = "/media/hocjsoo/New Volume/OWASP_Demo/screenshots/only_cards.png"
p_xss = "/media/hocjsoo/New Volume/OWASP_Demo/screenshots/crop_xss_ui.png"
p_csrf = "/media/hocjsoo/New Volume/OWASP_Demo/screenshots/crop_csrf_ui.png"

# Update Slide 6 (SQLi)
s6 = prs.slides[5]
# Let's check shapes in s6: shape 1 is b6_v, shape 2 is b6_s, shape 3 is callout6
# We can remove b6_v and b6_s text or adjust and place the actual screenshot!
print("Slide 6 shapes:", len(s6.shapes))

