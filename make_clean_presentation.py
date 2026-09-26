import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)

# Modern Tech Palette
DARK_BG = RGBColor(15, 23, 42)        # Slate 900
DARK_CARD = RGBColor(30, 41, 59)      # Slate 800
WHITE = RGBColor(255, 255, 255)
LIGHT_BG = RGBColor(248, 250, 252)    # Slate 50
TEXT_MAIN = RGBColor(30, 41, 59)      # Slate 800
TEXT_MUTED = RGBColor(100, 116, 139)  # Slate 500
BORDER_COLOR = RGBColor(226, 232, 240)

ACCENT_RED = RGBColor(225, 29, 72)     # Rose 600 (Danger)
RED_BG = RGBColor(255, 241, 242)       # Rose 50
ACCENT_GREEN = RGBColor(16, 185, 129)  # Emerald 500 (Secure)
GREEN_BG = RGBColor(236, 253, 245)     # Emerald 50
ACCENT_BLUE = RGBColor(37, 99, 235)    # Blue 600 (Primary)
BLUE_BG = RGBColor(239, 246, 255)      # Blue 50
ACCENT_AMBER = RGBColor(217, 119, 6)   # Amber 600

blank_layout = prs.slide_layouts[6]

def add_header(slide, title_text, category="OWASP TOP 10 - A03:2021 (INJECTION)"):
    # Header Background
    header = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(1.15))
    header.fill.solid()
    header.fill.fore_color.rgb = DARK_BG
    header.line.fill.background()

    # Category Tag
    tx = slide.shapes.add_textbox(Inches(0.8), Inches(0.12), Inches(11.7), Inches(0.3))
    tf = tx.text_frame
    tf.word_wrap = True
    p0 = tf.paragraphs[0]
    p0.text = category.upper()
    p0.font.size = Pt(11)
    p0.font.bold = True
    p0.font.color.rgb = ACCENT_RED

    # Title
    p1 = tf.add_paragraph()
    p1.text = title_text
    p1.font.size = Pt(19)
    p1.font.bold = True
    p1.font.color.rgb = WHITE

    # Footer
    footer = slide.shapes.add_textbox(Inches(0.8), Inches(7.1), Inches(11.7), Inches(0.3))
    ft = footer.text_frame
    p_ft = ft.paragraphs[0]
    p_ft.text = "Nhóm WNC.G01 (HOU) • Môn Lập Trình Web Nâng Cao • ThS. Lê Hữu Dũng"
    p_ft.font.size = Pt(9.5)
    p_ft.font.color.rgb = TEXT_MUTED

# ==============================================================================
# SLIDE 1: BÌA BÁO CÁO
# ==============================================================================
s1 = prs.slides.add_slide(blank_layout)
bg1 = s1.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(7.5))
bg1.fill.solid()
bg1.fill.fore_color.rgb = DARK_BG
bg1.line.fill.background()

tb1 = s1.shapes.add_textbox(Inches(1.0), Inches(1.1), Inches(11.333), Inches(5.2))
tf1 = tb1.text_frame
tf1.word_wrap = True

p = tf1.paragraphs[0]
p.text = "TRƯỜNG ĐẠI HỌC MỞ HÀ NỘI — KHOA CÔNG NGHỆ THÔNG TIN"
p.font.size = Pt(14)
p.font.bold = True
p.font.color.rgb = RGBColor(148, 163, 184)

p = tf1.add_paragraph()
p.text = "BÁO CÁO THỰC NGHIỆM AN TOÀN WEB & BẢO MẬT HỆ THỐNG"
p.font.size = Pt(17)
p.font.bold = True
p.font.color.rgb = ACCENT_RED
p.space_before = Pt(14)

p = tf1.add_paragraph()
p.text = "LỖ HỔNG SQL INJECTION (OWASP A03:2021)\nKịch Bản Khai Thác Thực Tế & Cơ Chế Phòng Thủ Trong ASP.NET Core"
p.font.size = Pt(26)
p.font.bold = True
p.font.color.rgb = WHITE
p.space_before = Pt(8)

p = tf1.add_paragraph()
p.text = "Giảng viên hướng dẫn: ThS. Lê Hữu Dũng | Học phần: Lập trình Web nâng cao"
p.font.size = Pt(14)
p.font.color.rgb = RGBColor(203, 213, 225)
p.space_before = Pt(18)

p = tf1.add_paragraph()
p.text = "Nhóm sinh viên thực hiện (WNC.G01):\n1. Nguyễn Danh Học  — MSV: 23A1001D0158 (Báo cáo & Phụ trách demo chính)\n2. Nguyễn Thanh Bình — MSV: 23A1001D0041\n3. Nguyễn Minh Cường — MSV: 23A1001D0058"
p.font.size = Pt(13)
p.font.color.rgb = RGBColor(148, 163, 184)
p.space_before = Pt(14)


# ==============================================================================
# SLIDE 2: BẢN CHẤT LỖ HỔNG
# ==============================================================================
s2 = prs.slides.add_slide(blank_layout)
add_header(s2, "1. Bản Chất Kỹ Thuật Của Lỗ Hổng SQL Injection")

# Left Column: Theory
tb2_1 = s2.shapes.add_textbox(Inches(0.8), Inches(1.35), Inches(5.6), Inches(5.4))
tf2_1 = tb2_1.text_frame
tf2_1.word_wrap = True

p = tf2_1.paragraphs[0]
p.text = "Nguyên Nhân Cốt Lõi"
p.font.size = Pt(16)
p.font.bold = True
p.font.color.rgb = ACCENT_RED

p = tf2_1.add_paragraph()
p.text = "• Không phân tách giữa DỮ LIỆU (Data) và CÂU LỆNH (Code).\n• Ghép chuỗi trực tiếp từ Request vào câu lệnh SQL khiến hệ CSDL không phân biệt được đâu là dữ liệu của user và đâu là cấu trúc lệnh của lập trình viên.\n• Kẻ tấn công lợi dụng các ký tự điều khiển (', --, /*, ;) để 'thoát' khỏi ngữ cảnh dữ liệu và ép SQL Server thực thi logic của hacker."
p.font.size = Pt(13)
p.font.color.rgb = TEXT_MAIN
p.space_before = Pt(8)

p = tf2_1.add_paragraph()
p.text = "Phân Loại Theo Chuẩn OWASP WSTG"
p.font.size = Pt(16)
p.font.bold = True
p.font.color.rgb = ACCENT_BLUE
p.space_before = Pt(16)

p = tf2_1.add_paragraph()
p.text = "1. In-band SQLi (Khai thác trực diện): Error-based, Union-based.\n2. Inferential SQLi (Blind SQLi / Mù): Boolean-based, Time-based.\n3. Out-of-band SQLi: Kích hoạt DNS/HTTP request từ CSDL."
p.font.size = Pt(12.5)
p.font.color.rgb = TEXT_MAIN
p.space_before = Pt(6)

# Right Column: Code Vulnerable
box2 = s2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.8), Inches(1.35), Inches(5.7), Inches(5.4))
box2.fill.solid()
box2.fill.fore_color.rgb = LIGHT_BG
box2.line.color.rgb = BORDER_COLOR

tb2_2 = s2.shapes.add_textbox(Inches(7.0), Inches(1.5), Inches(5.3), Inches(5.0))
tf2_2 = tb2_2.text_frame
tf2_2.word_wrap = True

p = tf2_2.paragraphs[0]
p.text = "Đoạn Code Gây Lỗi Cụ Thể (C# / ASP.NET Core)"
p.font.size = Pt(14)
p.font.bold = True
p.font.color.rgb = ACCENT_RED

p = tf2_2.add_paragraph()
p.text = "// Nối chuỗi trực tiếp từ Request:\nstring query = $\"SELECT * FROM Accounts \" +\n               $\"WHERE Username = '{user}' \" +\n               $\"AND Password = '{pass}'\";\n\nvar users = _context.Accounts\n                    .FromSqlRaw(query)\n                    .ToList();"
p.font.name = "Courier New"
p.font.size = Pt(11)
p.font.color.rgb = RGBColor(185, 28, 28)
p.space_before = Pt(8)

p = tf2_2.add_paragraph()
p.text = "⚠️ Điểm mù nguy hiểm:"
p.font.size = Pt(13)
p.font.bold = True
p.font.color.rgb = TEXT_MAIN
p.space_before = Pt(14)

p = tf2_2.add_paragraph()
p.text = "Lập trình viên mặc định tin rằng '{user}' chỉ là chuỗi thông thường. Khi hacker truyền vào ký tự nháy đơn ('), toàn bộ cấu trúc cú pháp của câu lệnh bị bẻ gãy."
p.font.size = Pt(12)
p.font.color.rgb = TEXT_MUTED
p.space_before = Pt(4)


# ==============================================================================
# SLIDE 3: ĐỘ NGUY HIỂM
# ==============================================================================
s3 = prs.slides.add_slide(blank_layout)
add_header(s3, "2. Độ Nguy Hiểm Thực Tế Của SQL Injection (Tác Động Nghiêm Trọng)")

col_w = Inches(3.6)
gap = Inches(0.4)

c1 = s3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.4), col_w, Inches(5.3))
c1.fill.solid()
c1.fill.fore_color.rgb = RED_BG
c1.line.color.rgb = ACCENT_RED
t1 = s3.shapes.add_textbox(Inches(0.95), Inches(1.6), col_w - Inches(0.3), Inches(4.9))
tf = t1.text_frame
tf.word_wrap = True
p = tf.paragraphs[0]
p.text = "🚨 CHIẾM QUYỀN HỆ THỐNG\n(Authentication Bypass)"
p.font.size = Pt(14)
p.font.bold = True
p.font.color.rgb = ACCENT_RED
p = tf.add_paragraph()
p.text = "• Không cần biết mật khẩu của bất kỳ ai.\n• Chỉ với 1 chuỗi payload ngắn, kẻ tấn công đăng nhập thẳng vào tài khoản Quản trị viên (Admin).\n• Chiếm đoạt phiên làm việc, thay đổi phân quyền và khóa tài khoản người dùng khác."
p.font.size = Pt(12)
p.font.color.rgb = TEXT_MAIN
p.space_before = Pt(10)

c2 = s3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8) + col_w + gap, Inches(1.4), col_w, Inches(5.3))
c2.fill.solid()
c2.fill.fore_color.rgb = RGBColor(255, 251, 235)
c2.line.color.rgb = ACCENT_AMBER
t2 = s3.shapes.add_textbox(Inches(0.95) + col_w + gap, Inches(1.6), col_w - Inches(0.3), Inches(4.9))
tf = t2.text_frame
tf.word_wrap = True
p = tf.paragraphs[0]
p.text = "📂 ĐÁNH CẮP TOÀN BỘ CSDL\n(Data Exfiltration)"
p.font.size = Pt(14)
p.font.bold = True
p.font.color.rgb = ACCENT_AMBER
p = tf.add_paragraph()
p.text = "• Trích xuất thông tin khách hàng, hồ sơ chuyên gia, số dư tài khoản ngân hàng.\n• Rò rỉ mật khẩu và ghi chú bảo mật cá nhân (Secret Notes).\n• Vi phạm nghiêm trọng luật an ninh mạng, làm sụp đổ hoàn toàn uy tín nền tảng."
p.font.size = Pt(12)
p.font.color.rgb = TEXT_MAIN
p.space_before = Pt(10)

c3 = s3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8) + (col_w + gap)*2, Inches(1.4), col_w, Inches(5.3))
c3.fill.solid()
c3.fill.fore_color.rgb = BLUE_BG
c3.line.color.rgb = DARK_BG
t3 = s3.shapes.add_textbox(Inches(0.95) + (col_w + gap)*2, Inches(1.6), col_w - Inches(0.3), Inches(4.9))
tf = t3.text_frame
tf.word_wrap = True
p = tf.paragraphs[0]
p.text = "💣 PHÁ HỦY / THỰC THI LỆNH\n(System Takeover)"
p.font.size = Pt(14)
p.font.bold = True
p.font.color.rgb = DARK_BG
p = tf.add_paragraph()
p.text = "• Xóa sổ dữ liệu với lệnh DROP TABLE hoặc TRUNCATE.\n• Sửa đổi dữ liệu, giả mạo giao dịch tài chính.\n• Có thể kích hoạt quyền thực thi lệnh hệ điều hành (xp_cmdshell) nếu CSDL cấu hình lỏng lẻo."
p.font.size = Pt(12)
p.font.color.rgb = TEXT_MAIN
p.space_before = Pt(10)


# ==============================================================================
# SLIDE 4: PHÂN TÍCH CƠ CHẾ BẺ GÃY CÚ PHÁP
# ==============================================================================
s4 = prs.slides.add_slide(blank_layout)
add_header(s4, "3. Kịch Bản Tấn Công: Phân Tích Cơ Chế Bẻ Gãy Cú Pháp SQL")

tb4 = s4.shapes.add_textbox(Inches(0.8), Inches(1.35), Inches(11.7), Inches(5.4))
tf4 = tb4.text_frame
tf4.word_wrap = True

p = tf4.paragraphs[0]
p.text = "Kịch bản: Vượt qua xác thực đăng nhập (Authentication Bypass)"
p.font.size = Pt(16)
p.font.bold = True
p.font.color.rgb = ACCENT_RED

p = tf4.add_paragraph()
p.text = "Payload đưa vào Username:   ' OR '1'='1' --\nMật khẩu đưa vào Password: (Nhập bất kỳ hoặc để trống)"
p.font.name = "Courier New"
p.font.size = Pt(13)
p.font.bold = True
p.font.color.rgb = ACCENT_BLUE
p.space_before = Pt(8)

p = tf4.add_paragraph()
p.text = "Cấu trúc câu truy vấn được SQL Server phân tích cú pháp:"
p.font.size = Pt(15)
p.font.bold = True
p.font.color.rgb = TEXT_MAIN
p.space_before = Pt(14)

p = tf4.add_paragraph()
p.text = "SELECT * FROM Accounts WHERE Username = '' OR '1'='1' --' AND Password = 'xyz'"
p.font.name = "Courier New"
p.font.size = Pt(13.5)
p.font.bold = True
p.font.color.rgb = ACCENT_RED
p.space_before = Pt(6)

p = tf4.add_paragraph()
p.text = "3 Thành Phần Cốt Lõi Của Chuỗi Khai Thác:"
p.font.size = Pt(15)
p.font.bold = True
p.font.color.rgb = TEXT_MAIN
p.space_before = Pt(14)

p = tf4.add_paragraph()
p.text = "1. Dấu nháy đơn (') : Đóng sớm chuỗi ký tự hợp lệ của trường Username, đưa ngữ cảnh về câu lệnh SQL.\n2. Mệnh đề OR '1'='1' : Mệnh đề chân lý logic. Vì '1'='1' luôn đúng, mệnh đề WHERE trả về TRUE cho toàn bộ bản ghi trong bảng.\n3. Ký tự chú thích (--) : Biến toàn bộ đoạn kiểm tra mật khẩu phía sau thành chú thích vô hiệu lực."
p.font.size = Pt(12.5)
p.font.color.rgb = TEXT_MAIN
p.space_before = Pt(6)


# ==============================================================================
# SLIDE 5: ĐỐI CHỨNG THỰC NGHIỆM (CLEAN CARDS - NO SCREENSHOT CUTOFF)
# ==============================================================================
s5 = prs.slides.add_slide(blank_layout)
add_header(s5, "4. Đối Chứng Thực Nghiệm: Nhánh Dính Lỗi vs Nhánh Đã Phòng Thủ")

card_w = Inches(5.65)

# Left Card: Vulnerable
box_v = s5.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.35), card_w, Inches(4.3))
box_v.fill.solid()
box_v.fill.fore_color.rgb = RED_BG
box_v.line.color.rgb = ACCENT_RED

tb_v = s5.shapes.add_textbox(Inches(0.95), Inches(1.45), card_w - Inches(0.3), Inches(4.0))
tf_v = tb_v.text_frame
tf_v.word_wrap = True

p = tf_v.paragraphs[0]
p.text = "❌ 1. NHÁNH BỊ LỖI (Nối Chuỗi SQL Thuần)"
p.font.size = Pt(14)
p.font.bold = True
p.font.color.rgb = ACCENT_RED

p = tf_v.add_paragraph()
p.text = "• Payload thử nghiệm: ' OR '1'='1' --\n• Câu lệnh thực thi:\n  SELECT * FROM Accounts WHERE Username = '' OR '1'='1' --'\n• Cơ chế thông dịch: Điều kiện WHERE luôn TRUE.\n• Hậu quả thực tế:\n  - Đăng nhập thành công mà KHÔNG CẦN mật khẩu.\n  - Rò rỉ toàn bộ 4 tài khoản trong CSDL (Admin, Specialist, Customer).\n  - Lộ mật khẩu gốc và số dư tài khoản ngân hàng."
p.font.size = Pt(12)
p.font.color.rgb = TEXT_MAIN
p.space_before = Pt(8)

# Right Card: Secure
box_s = s5.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.85), Inches(1.35), card_w, Inches(4.3))
box_s.fill.solid()
box_s.fill.fore_color.rgb = GREEN_BG
box_s.line.color.rgb = ACCENT_GREEN

tb_s = s5.shapes.add_textbox(Inches(7.0), Inches(1.45), card_w - Inches(0.3), Inches(4.0))
tf_s = tb_s.text_frame
tf_s.word_wrap = True

p = tf_s.paragraphs[0]
p.text = "✅ 2. NHÁNH AN TOÀN (EF Core Parameterized)"
p.font.size = Pt(14)
p.font.bold = True
p.font.color.rgb = ACCENT_GREEN

p = tf_s.add_paragraph()
p.text = "• Cùng Payload thử nghiệm: ' OR '1'='1' --\n• Câu lệnh thực thi:\n  SELECT * FROM Accounts WHERE Username = @p0 AND Pass = @p1\n• Cơ chế thông dịch:\n  - CSDL biên dịch cây truy vấn (Execution Plan) TRƯỚC.\n  - Toàn bộ chuỗi payload được coi là String Literal thuần túy.\n• Kết quả phòng thủ:\n  - Hệ thống chặn đứng, trả về 0 bản ghi.\n  - Báo lỗi 'Sai tài khoản/mật khẩu', bảo vệ an toàn 100%."
p.font.size = Pt(12)
p.font.color.rgb = TEXT_MAIN
p.space_before = Pt(8)

# Bottom Callout: Live Demo Indicator
callout = s5.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(5.8), Inches(11.7), Inches(1.0))
callout.fill.solid()
callout.fill.fore_color.rgb = DARK_CARD
callout.line.color.rgb = ACCENT_BLUE

tb_c = s5.shapes.add_textbox(Inches(1.0), Inches(5.85), Inches(11.3), Inches(0.9))
tf_c = tb_c.text_frame
tf_c.word_wrap = True
p = tf_c.paragraphs[0]
p.text = "💻 [LIVE DEMO TRÊN MÁY]: Chuyển sang Trình duyệt Web (http://localhost:5076) & VS Code"
p.font.size = Pt(13)
p.font.bold = True
p.font.color.rgb = RGBColor(56, 189, 248)
p = tf_c.add_paragraph()
p.text = "Trình diễn trực tiếp 3 bước: Bấm nạp payload → Kiểm tra kết quả trả về → Đối chiếu Log SQL Server."
p.font.size = Pt(11.5)
p.font.color.rgb = WHITE
p.space_before = Pt(2)


# ==============================================================================
# SLIDE 6: ĐỐI CHỨNG TẦNG SÂU (API CLI & SQL SERVER 2025)
# ==============================================================================
s6 = prs.slides.add_slide(blank_layout)
add_header(s6, "5. Đối Chứng Tầng Sâu: Tầng API (cURL) & Microsoft SQL Server 2025")

# Box 1: API Proof
b_api = s6.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.35), card_w, Inches(4.3))
b_api.fill.solid()
b_api.fill.fore_color.rgb = LIGHT_BG
b_api.line.color.rgb = BORDER_COLOR

tb_api = s6.shapes.add_textbox(Inches(1.0), Inches(1.5), card_w - Inches(0.4), Inches(4.0))
tf_api = tb_api.text_frame
tf_api.word_wrap = True
p = tf_api.paragraphs[0]
p.text = "Minh Chứng Tầng API / Dòng Lệnh (cURL)"
p.font.size = Pt(14)
p.font.bold = True
p.font.color.rgb = ACCENT_BLUE

p = tf_api.add_paragraph()
p.text = "$ curl -X POST http://localhost:5076/SqlInjection/LoginVulnerable \\\n       -F \"username=' OR '1'='1' --\" -F \"password=any\"\n\nPhản hồi JSON từ Server:\n{\n  \"success\": true,\n  \"isExploited\": true,\n  \"count\": 4,\n  \"accounts\": [ ... 4 tài khoản bị rò rỉ ... ]\n}"
p.font.name = "Courier New"
p.font.size = Pt(10)
p.font.color.rgb = TEXT_MAIN
p.space_before = Pt(6)

p = tf_api.add_paragraph()
p.text = "📌 Ý nghĩa: Chứng minh lỗ hổng nằm hoàn toàn ở logic C# Backend, kẻ tấn công có thể dùng script tự động quét và khai thác mà không cần mở trình duyệt."
p.font.size = Pt(11.5)
p.font.color.rgb = TEXT_MUTED
p.space_before = Pt(8)

# Box 2: Database Proof
b_db = s6.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.85), Inches(1.35), card_w, Inches(4.3))
b_db.fill.solid()
b_db.fill.fore_color.rgb = LIGHT_BG
b_db.line.color.rgb = BORDER_COLOR

tb_db = s6.shapes.add_textbox(Inches(7.0), Inches(1.5), card_w - Inches(0.4), Inches(4.0))
tf_db = tb_db.text_frame
tf_db.word_wrap = True
p = tf_db.paragraphs[0]
p.text = "Minh Chứng Tầng CSDL (SQL Server 2025)"
p.font.size = Pt(14)
p.font.bold = True
p.font.color.rgb = ACCENT_GREEN

p = tf_db.add_paragraph()
p.text = "-- Database: OwaspDemoDB | Bảng: dbo.Accounts\n-- 1. Câu lệnh dính lỗi:\nSELECT * FROM Accounts WHERE Username='' OR '1'='1' --'\n>> Kết quả: 4 rows returned (Bypass thành công)\n\n-- 2. Câu lệnh tham số hóa chuẩn:\nEXEC sp_executesql N'SELECT * FROM Accounts WHERE...',\n     N'@p0 nvarchar(50)', @p0 = N''' OR ''1''=''1'' --'\n>> Kết quả: 0 rows returned (An toàn 100%)"
p.font.name = "Courier New"
p.font.size = Pt(10)
p.font.color.rgb = TEXT_MAIN
p.space_before = Pt(6)

p = tf_db.add_paragraph()
p.text = "📌 Ý nghĩa: Đối chiếu trực tiếp trên máy chủ SQL Server 2025 chứng minh cơ chế phân tách rạch ròi giữa cú pháp lệnh và tham số giá trị."
p.font.size = Pt(11.5)
p.font.color.rgb = TEXT_MUTED
p.space_before = Pt(8)

# Bottom Callout 2
callout2 = s6.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(5.8), Inches(11.7), Inches(1.0))
callout2.fill.solid()
callout2.fill.fore_color.rgb = DARK_CARD
callout2.line.color.rgb = ACCENT_GREEN

tb_c2 = s6.shapes.add_textbox(Inches(1.0), Inches(5.85), Inches(11.3), Inches(0.9))
tf_c2 = tb_c2.text_frame
tf_c2.word_wrap = True
p = tf_c2.paragraphs[0]
p.text = "🗄️ [LIVE DEMO TRÊN MÁY]: Chuyển sang VS Code mở bảng dbo.Accounts & file verify_sql_server.sql"
p.font.size = Pt(13)
p.font.bold = True
p.font.color.rgb = RGBColor(74, 222, 128)
p = tf_c2.add_paragraph()
p.text = "Thao tác trực tiếp: Mở extension MSSQL trên VS Code, đối chiếu bảng dữ liệu gốc và chạy câu query đối chứng."
p.font.size = Pt(11.5)
p.font.color.rgb = WHITE
p.space_before = Pt(2)


# ==============================================================================
# SLIDE 7: PHÒNG THỦ TRONG ASP.NET CORE
# ==============================================================================
s7 = prs.slides.add_slide(blank_layout)
add_header(s7, "6. Biện Pháp Phòng Thủ Chuẩn Trong ASP.NET Core MVC")

box7 = s7.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.35), Inches(5.7), Inches(5.4))
box7.fill.solid()
box7.fill.fore_color.rgb = LIGHT_BG
box7.line.color.rgb = ACCENT_GREEN

tb7_1 = s7.shapes.add_textbox(Inches(1.0), Inches(1.5), Inches(5.3), Inches(5.0))
tf7_1 = tb7_1.text_frame
tf7_1.word_wrap = True

p = tf7_1.paragraphs[0]
p.text = "Cách Viết Code Chuẩn Hóa Với EF Core"
p.font.size = Pt(14)
p.font.bold = True
p.font.color.rgb = ACCENT_GREEN

p = tf7_1.add_paragraph()
p.text = "// Dùng LINQ Parameterized (KHUYÊN DÙNG):\nvar account = _context.Accounts\n    .FirstOrDefault(a => a.Username == user \n                      && a.Password == pass);\n\n// Hoặc dùng FromSqlInterpolated nếu viết SQL thuần:\nvar account = _context.Accounts\n    .FromSqlInterpolated($\"SELECT * FROM Accounts WHERE Username={user} AND Password={pass}\")\n    .FirstOrDefault();"
p.font.name = "Courier New"
p.font.size = Pt(10)
p.font.color.rgb = RGBColor(22, 101, 52)
p.space_before = Pt(8)

p = tf7_1.add_paragraph()
p.text = "Tại sao cách này an toàn tuyệt đối?"
p.font.size = Pt(12)
p.font.bold = True
p.font.color.rgb = TEXT_MAIN
p.space_before = Pt(12)

p = tf7_1.add_paragraph()
p.text = "SQL Server biên dịch cấu trúc lệnh TRƯỚC khi gán dữ liệu. Biến @p0 nhận toàn bộ payload như một chuỗi chữ bình thường, vô hiệu hóa hoàn toàn ý đồ chèn lệnh."
p.font.size = Pt(11.5)
p.font.color.rgb = TEXT_MUTED
p.space_before = Pt(4)

tb7_2 = s7.shapes.add_textbox(Inches(6.9), Inches(1.35), Inches(5.6), Inches(5.4))
tf7_2 = tb7_2.text_frame
tf7_2.word_wrap = True

p = tf7_2.paragraphs[0]
p.text = "Bộ Quy Tắc Phòng Thủ Toàn Diện (Defense in Depth)"
p.font.size = Pt(15)
p.font.bold = True
p.font.color.rgb = ACCENT_BLUE

p = tf7_2.add_paragraph()
p.text = "1. Luôn sử dụng Parameterized Query / ORM:\nTuyệt đối không ghép chuỗi SQL thủ công dưới bất kỳ hình thức nào."
p.font.size = Pt(12)
p.font.color.rgb = TEXT_MAIN
p.space_before = Pt(10)

p = tf7_2.add_paragraph()
p.text = "2. Nguyên tắc đặc quyền tối thiểu (Least Privilege):\nTài khoản kết nối CSDL của ứng dụng chỉ có quyền SELECT/INSERT/UPDATE trên các bảng cần thiết, không bao giờ dùng tài khoản 'sa' hay quyền DDL (DROP, ALTER)."
p.font.size = Pt(12)
p.font.color.rgb = TEXT_MAIN
p.space_before = Pt(8)

p = tf7_2.add_paragraph()
p.text = "3. Xác thực dữ liệu đầu vào (Input Validation):\nSử dụng Data Annotations ([RegularExpression], [StringLength]) kiểm tra chặt chẽ khuôn dạng dữ liệu."
p.font.size = Pt(12)
p.font.color.rgb = TEXT_MAIN
p.space_before = Pt(8)

p = tf7_2.add_paragraph()
p.text = "4. Mã hóa mật khẩu một chiều (Password Hashing):\nSử dụng ASP.NET Core Identity (PBKDF2/BCrypt) để nếu CSDL có bị lộ, mật khẩu vẫn được bảo vệ an toàn."
p.font.size = Pt(12)
p.font.color.rgb = TEXT_MAIN
p.space_before = Pt(8)


# ==============================================================================
# SLIDE 8: KẾT LUẬN & LIÊN HỆ BTL
# ==============================================================================
s8 = prs.slides.add_slide(blank_layout)
add_header(s8, "7. Tổng Kết & Bài Học Rút Ra Cho Dự Án Website Tư Vấn Trực Tuyến")

tb8 = s8.shapes.add_textbox(Inches(1.0), Inches(1.5), Inches(11.333), Inches(5.0))
tf8 = tb8.text_frame
tf8.word_wrap = True

p = tf8.paragraphs[0]
p.text = "BÀI HỌC VÀ CAM KẾT CHUẨN AN TOÀN CHO ĐỀ TÀI BTL (WNC.G01):"
p.font.size = Pt(17)
p.font.bold = True
p.font.color.rgb = ACCENT_RED

p = tf8.add_paragraph()
p.text = "1. Nhận thức an toàn: SQL Injection là lỗ hổng nghiêm trọng ở mức độ CRITICAL (9.8/10 theo CVSS v3.1). Lập trình viên phải luôn coi mọi dữ liệu đầu vào từ người dùng là 'không an toàn'."
p.font.size = Pt(13.5)
p.font.color.rgb = TEXT_MAIN
p.space_before = Pt(12)

p = tf8.add_paragraph()
p.text = "2. Phòng thủ triệt để: Trong đồ án Website cung cấp dịch vụ tư vấn trực tuyến, nhóm cam kết áp dụng 100% Entity Framework Core Code-First với Parameterized LINQ cho toàn bộ các thao tác đặt lịch, quản lý chuyên gia và đánh giá chất lượng."
p.font.size = Pt(13.5)
p.font.color.rgb = TEXT_MAIN
p.space_before = Pt(10)

p = tf8.add_paragraph()
p.text = "3. Bảo mật đa tầng (Security View - Buổi 8): Tích hợp chặt chẽ ASP.NET Core Identity để băm mật khẩu, AntiForgeryToken chống CSRF trên toàn bộ form và kiểm soát phân quyền 3 Role (Customer, Specialist, Admin)."
p.font.size = Pt(13.5)
p.font.color.rgb = TEXT_MAIN
p.space_before = Pt(10)

p = tf8.add_paragraph()
p.text = "XIN TRÂN TRỌNG CẢM ƠN THẦY VÀ CÁC BẠN ĐÃ THEO DÕI!\nNhóm WNC.G01 sẵn sàng lắng nghe câu hỏi và nhận xét từ giảng viên."
p.font.size = Pt(15.5)
p.font.bold = True
p.font.color.rgb = ACCENT_BLUE
p.space_before = Pt(22)

out_file = "/media/hocjsoo/New Volume/OWASP_Demo/Slide_OWASP_SQL_Injection.pptx"
prs.save(out_file)
print(f"Clean presentation successfully saved to: {out_file}")
