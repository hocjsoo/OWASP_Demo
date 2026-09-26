import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)

# ==============================================================================
# DESIGN SYSTEM: ACADEMIC LIGHT / CYBERSECURITY
# ==============================================================================
BG_WHITE = RGBColor(255, 255, 255)
WHITE = RGBColor(255, 255, 255)
CARD_BG = RGBColor(248, 250, 252)          # Slate 50
CARD_BORDER = RGBColor(226, 232, 240)      # Slate 200

TEXT_HEAD = RGBColor(15, 23, 42)           # Slate 900
TEXT_BODY = RGBColor(51, 65, 85)           # Slate 700
TEXT_MUTED = RGBColor(100, 116, 139)       # Slate 500

ACCENT_BLUE = RGBColor(29, 78, 216)        # Blue 700 (Primary)
BLUE_LIGHT = RGBColor(239, 246, 255)       # Blue 50
BLUE_BORDER = RGBColor(191, 219, 254)

ACCENT_RED = RGBColor(190, 18, 60)         # Rose 700 (Vulnerable / Attack)
RED_LIGHT = RGBColor(255, 241, 242)        # Rose 50
RED_BORDER = RGBColor(254, 205, 211)

ACCENT_GREEN = RGBColor(4, 120, 87)        # Emerald 700 (Secure / Fixed)
GREEN_LIGHT = RGBColor(236, 253, 245)      # Emerald 50
GREEN_BORDER = RGBColor(167, 243, 208)

ACCENT_AMBER = RGBColor(180, 83, 9)        # Amber 700 (Warning)
AMBER_LIGHT = RGBColor(255, 251, 235)      # Amber 50
AMBER_BORDER = RGBColor(254, 240, 138)

blank_layout = prs.slide_layouts[6]

def add_header(slide, title_text, category_tag="OWASP TOP 10 • THỰC NGHIỆM AN TOÀN WEB"):
    # Top Accent Line
    top_line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(0.08))
    top_line.fill.solid()
    top_line.fill.fore_color.rgb = ACCENT_BLUE
    top_line.line.fill.background()

    # Section Tag Pill
    badge = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(0.28), Inches(5.2), Inches(0.36))
    badge.fill.solid()
    badge.fill.fore_color.rgb = BLUE_LIGHT
    badge.line.color.rgb = BLUE_BORDER

    tf_b = badge.text_frame
    tf_b.word_wrap = False
    p_b = tf_b.paragraphs[0]
    p_b.text = category_tag.upper()
    p_b.font.name = "Arial"
    p_b.font.size = Pt(10)
    p_b.font.bold = True
    p_b.font.color.rgb = ACCENT_BLUE

    # Title
    tx = slide.shapes.add_textbox(Inches(0.8), Inches(0.68), Inches(11.7), Inches(0.6))
    tf = tx.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = title_text
    p.font.name = "Arial"
    p.font.size = Pt(20)
    p.font.bold = True
    p.font.color.rgb = TEXT_HEAD

    # Footer
    footer = slide.shapes.add_textbox(Inches(0.8), Inches(7.15), Inches(11.7), Inches(0.28))
    ft = footer.text_frame
    ft.margin_left = ft.margin_right = ft.margin_top = ft.margin_bottom = 0
    p_ft = ft.paragraphs[0]
    p_ft.text = "Nhóm WNC.G01 • Trường Đại học Mở Hà Nội (HOU) • Giảng viên: ThS. Lê Hữu Dũng"
    p_ft.font.name = "Arial"
    p_ft.font.size = Pt(9.5)
    p_ft.font.color.rgb = TEXT_MUTED

def add_p(tf, text, size=12, bold=False, color=TEXT_BODY, space_before=0, font_name="Arial"):
    p = tf.add_paragraph()
    p.text = text
    p.font.name = font_name
    p.font.size = Pt(size)
    p.font.bold = bold
    p.font.color.rgb = color
    if space_before:
        p.space_before = Pt(space_before)
    return p

card_w = Inches(5.65)
col_w = Inches(3.64)
gap = Inches(0.2)

# ==============================================================================
# SLIDE 1: COVER (MODERN CYBERSECURITY REDESIGN)
# ==============================================================================
s1 = prs.slides.add_slide(blank_layout)
bar1 = s1.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(0.12))
bar1.fill.solid()
bar1.fill.fore_color.rgb = ACCENT_BLUE
bar1.line.fill.background()

# Left Column: Title & Mission
tb1 = s1.shapes.add_textbox(Inches(1.0), Inches(0.8), Inches(6.8), Inches(3.6))
tf1 = tb1.text_frame
tf1.word_wrap = True
p = tf1.paragraphs[0]
p.text = "OWASP TOP 10 • THỰC NGHIỆM AN TOÀN WEB"
p.font.name = "Arial"
p.font.size = Pt(13)
p.font.bold = True
p.font.color.rgb = ACCENT_BLUE

add_p(tf1, "Nghiên Cứu Lỗ Hổng Web & Cơ Chế Phòng Thủ Đa Tầng", size=26, bold=True, color=TEXT_HEAD, space_before=8)
add_p(tf1, "SQL Injection • Stored XSS • CSRF Attack", size=16, bold=True, color=ACCENT_RED, space_before=6)
add_p(tf1, "Học phần: Lập trình Web nâng cao  •  GVHD: ThS. Lê Hữu Dũng\nKhoa Công nghệ Thông tin — Trường Đại học Mở Hà Nội (HOU)", size=12, color=TEXT_BODY, space_before=10)

# Right Column: Mental Model Architecture Card
c_overview = s1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(8.1), Inches(0.8), Inches(4.2), Inches(3.6))
c_overview.fill.solid()
c_overview.fill.fore_color.rgb = CARD_BG
c_overview.line.color.rgb = CARD_BORDER
tf_o = c_overview.text_frame
tf_o.word_wrap = True
p = tf_o.paragraphs[0]
p.text = "MÔ HÌNH 3 TẦNG & 3 ĐIỂM NGUY CƠ"
p.font.name = "Arial"
p.font.size = Pt(12)
p.font.bold = True
p.font.color.rgb = ACCENT_BLUE
add_p(tf_o, "1. Browser / Client DOM", size=11, bold=True, color=TEXT_HEAD, space_before=8)
add_p(tf_o, "   ➔ ⚠️ Nguy cơ: Stored XSS (Đánh cắp Cookie)", size=10.5, color=ACCENT_RED, space_before=2)
add_p(tf_o, "2. Web Server (ASP.NET Core)", size=11, bold=True, color=TEXT_HEAD, space_before=8)
add_p(tf_o, "   ➔ ⚠️ Nguy cơ: CSRF (Mượn quyền giao dịch)", size=10.5, color=ACCENT_RED, space_before=2)
add_p(tf_o, "3. Database Server (SQL Server 2025)", size=11, bold=True, color=TEXT_HEAD, space_before=8)
add_p(tf_o, "   ➔ ⚠️ Nguy cơ: SQL Injection (Bẻ gãy cú pháp)", size=10.5, color=ACCENT_RED, space_before=2)

# Bottom: 3 Member Cards
m1 = s1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.0), Inches(4.7), col_w, Inches(2.0))
m1.fill.solid()
m1.fill.fore_color.rgb = CARD_BG
m1.line.color.rgb = BLUE_BORDER
tf = m1.text_frame
p = tf.paragraphs[0]
p.text = "CHỦ ĐỀ 1: SQL INJECTION"
p.font.name = "Arial"
p.font.size = Pt(11)
p.font.bold = True
p.font.color.rgb = ACCENT_BLUE
add_p(tf, "Nguyễn Danh Học\nMSV: 23A1001D0158\nTrưởng nhóm • Phụ trách CSDL & API", size=12, color=TEXT_HEAD, space_before=4)

m2 = s1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.0) + col_w + gap, Inches(4.7), col_w, Inches(2.0))
m2.fill.solid()
m2.fill.fore_color.rgb = CARD_BG
m2.line.color.rgb = RED_BORDER
tf = m2.text_frame
p = tf.paragraphs[0]
p.text = "CHỦ ĐỀ 2: STORED XSS"
p.font.name = "Arial"
p.font.size = Pt(11)
p.font.bold = True
p.font.color.rgb = ACCENT_RED
add_p(tf, "Nguyễn Thanh Bình\nMSV: 23A1001D0041\nThành viên • Phụ trách Cookie & DOM", size=12, color=TEXT_HEAD, space_before=4)

m3 = s1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.0) + (col_w + gap)*2, Inches(4.7), col_w, Inches(2.0))
m3.fill.solid()
m3.fill.fore_color.rgb = CARD_BG
m3.line.color.rgb = GREEN_BORDER
tf = m3.text_frame
p = tf.paragraphs[0]
p.text = "CHỦ ĐỀ 3: CSRF ATTACK"
p.font.name = "Arial"
p.font.size = Pt(11)
p.font.bold = True
p.font.color.rgb = ACCENT_GREEN
add_p(tf, "Nguyễn Minh Cường\nMSV: 23A1001D0058\nThành viên • Phụ trách Giao dịch & Token", size=12, color=TEXT_HEAD, space_before=4)


# ==============================================================================
# SLIDE 2: MÔI TRƯỜNG & PHẠM VI (2 CỘT THOÁNG ĐẠT)
# ==============================================================================
s2 = prs.slides.add_slide(blank_layout)
add_header(s2, "Môi Trường Công Nghệ & Phạm Vi Thực Nghiệm")

c2_1 = s2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.4), card_w, Inches(5.4))
c2_1.fill.solid()
c2_1.fill.fore_color.rgb = CARD_BG
c2_1.line.color.rgb = CARD_BORDER
tf = c2_1.text_frame
p = tf.paragraphs[0]
p.text = "TECH STACK (MÔI TRƯỜNG KỸ THUẬT)"
p.font.name = "Arial"
p.font.size = Pt(14)
p.font.bold = True
p.font.color.rgb = ACCENT_BLUE

add_p(tf, "• Operating System:\n  Ubuntu 24.04 LTS / Linux Kernel 6.8", size=12, space_before=8)
add_p(tf, "• Backend Platform:\n  ASP.NET Core MVC / .NET 10.0", size=12, space_before=6)
add_p(tf, "• Database Engine:\n  Microsoft SQL Server 2025 Developer (Port 1433)", size=12, space_before=6)
add_p(tf, "• Frontend Engine:\n  Razor View, HTML5, CSS3, JavaScript", size=12, space_before=6)
add_p(tf, "• Development & Testing Tools:\n  VS Code, Chrome DevTools, cURL CLI, MSSQL Extension", size=12, space_before=6)

c2_2 = s2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.85), Inches(1.4), card_w, Inches(5.4))
c2_2.fill.solid()
c2_2.fill.fore_color.rgb = GREEN_LIGHT
c2_2.line.color.rgb = GREEN_BORDER
tf = c2_2.text_frame
p = tf.paragraphs[0]
p.text = "LAB SCOPE (PHẠM VI & NGUYÊN TẮC ĐẠO ĐỨC)"
p.font.name = "Arial"
p.font.size = Pt(14)
p.font.bold = True
p.font.color.rgb = ACCENT_GREEN

add_p(tf, "✓ Môi trường cô lập (Localhost):\nThực nghiệm hoàn toàn trên ứng dụng do nhóm tự xây dựng (http://localhost:5076), tuyệt đối không quét hay tấn công hệ thống bên ngoài.", size=12, space_before=8)
add_p(tf, "✓ Thiết kế đối chứng song song:\nMỗi lỗ hổng đều xây dựng 2 nhánh: Vulnerable (minh họa lỗi) vs Secure (chứng minh phòng thủ).", size=12, space_before=10)
add_p(tf, "✓ Dữ liệu mô phỏng an toàn:\nToàn bộ tài khoản, số dư và mật khẩu là dữ liệu giả định phục vụ học tập.", size=12, space_before=10)


# ==============================================================================
# PHẦN 1: SQL INJECTION — NGUYỄN DANH HỌC (SLIDES 3 - 7)
# ==============================================================================
# Slide 3: 1.1 Where & Why (Sơ đồ luồng + Code lỗi ngắn)
s3 = prs.slides.add_slide(blank_layout)
add_header(s3, "1.1. SQL Injection — Lỗi Xảy Ra Ở Đâu & Nguyên Nhân Cốt Lõi", "CHỦ ĐỀ 1 • NGUYỄN DANH HỌC (23A1001D0158)")

c3_1 = s3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.4), card_w, Inches(5.4))
c3_1.fill.solid()
c3_1.fill.fore_color.rgb = CARD_BG
c3_1.line.color.rgb = CARD_BORDER
tf = c3_1.text_frame
p = tf.paragraphs[0]
p.text = "Sơ Đồ Luồng Xử Lý & Root Cause"
p.font.name = "Arial"
p.font.size = Pt(15)
p.font.bold = True
p.font.color.rgb = ACCENT_RED

add_p(tf, "Browser (Client) ➔ Gửi Input (' OR '1'='1' --)", size=12, bold=True, color=ACCENT_BLUE, space_before=10)
add_p(tf, "            │\n            ▼", size=12, bold=True, color=TEXT_MUTED, space_before=2)
add_p(tf, "ASP.NET Core Controller ➔ Ghép chuỗi SQL trực tiếp", size=12, bold=True, color=ACCENT_RED, space_before=2)
add_p(tf, "            │\n            ▼", size=12, bold=True, color=TEXT_MUTED, space_before=2)
add_p(tf, "SQL Server 2025 ➔ Thực thi câu lệnh bị bẻ gãy cú pháp", size=12, bold=True, color=ACCENT_RED, space_before=2)

add_p(tf, "\n📌 ROOT CAUSE (Nguyên nhân cốt lõi):\nDữ liệu người dùng được ghép trực tiếp vào câu lệnh SQL. Hệ CSDL không phân định được ranh giới giữa Dữ liệu (Data) và Lệnh thực thi (Code).", size=12, color=TEXT_BODY, space_before=10)

c3_2 = s3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.85), Inches(1.4), card_w, Inches(5.4))
c3_2.fill.solid()
c3_2.fill.fore_color.rgb = RED_LIGHT
c3_2.line.color.rgb = RED_BORDER
tf = c3_2.text_frame
p = tf.paragraphs[0]
p.text = "🔴 VULNERABLE CODE (Ghép Chuỗi SQL)"
p.font.name = "Arial"
p.font.size = Pt(14)
p.font.bold = True
p.font.color.rgb = ACCENT_RED

add_p(tf, "// Nối chuỗi trực tiếp từ tham số đầu vào:\nstring rawSql = $\"SELECT * FROM Accounts \" +\n                $\"WHERE Username = '{username}' \" +\n                $\"AND Password = '{password}'\";\n\nvar accounts = _context.Accounts\n                       .FromSqlRaw(rawSql)\n                       .ToList();", size=11, color=RGBColor(159, 18, 57), space_before=10, font_name="Courier New")

add_p(tf, "Cơ chế phá vỡ:\nDấu nháy đơn (') đóng chuỗi sớm, biến mệnh đề kiểm tra mật khẩu thành chú thích (--), mở toang quyền truy cập CSDL.", size=12, bold=True, color=TEXT_HEAD, space_before=14)

# Slide 4: 1.2 3 Kịch bản khai thác (3 Card ngang)
s4 = prs.slides.add_slide(blank_layout)
add_header(s4, "1.2. SQL Injection — 3 Kịch Bản Khai Thác Thực Nghiệm", "CHỦ ĐỀ 1 • NGUYỄN DANH HỌC (23A1001D0158)")

# 3 Horizontal Cards
card_h = Inches(1.6)
gap_h = Inches(0.2)

# Card 1
c4_1 = s4.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.4), Inches(11.733), card_h)
c4_1.fill.solid()
c4_1.fill.fore_color.rgb = RED_LIGHT
c4_1.line.color.rgb = RED_BORDER
tf = c4_1.text_frame
p = tf.paragraphs[0]
p.text = "01 — BYPASS ĐĂNG NHẬP TỔNG QUÁT: ' OR '1'='1' --"
p.font.name = "Arial"
p.font.size = Pt(13)
p.font.bold = True
p.font.color.rgb = ACCENT_RED
add_p(tf, "• Cơ chế: Dấu nháy đơn (') đóng chuỗi; mệnh đề OR '1'='1' luôn đúng cho mọi dòng; ký tự '--' vô hiệu hóa bước kiểm tra mật khẩu.\n• Kết quả thực tế: Vượt qua xác thực thành công, đăng nhập ngay vào tài khoản đầu tiên trong bảng (Admin).", size=11.5, space_before=4)

# Card 2
c4_2 = s4.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.4) + card_h + gap_h, Inches(11.733), card_h)
c4_2.fill.solid()
c4_2.fill.fore_color.rgb = AMBER_LIGHT
c4_2.line.color.rgb = AMBER_BORDER
tf = c4_2.text_frame
p = tf.paragraphs[0]
p.text = "02 — TẤN CÔNG ĐÍCH DANH QUẢN TRỊ VIÊN: admin' --"
p.font.name = "Arial"
p.font.size = Pt(13)
p.font.bold = True
p.font.color.rgb = ACCENT_AMBER
add_p(tf, "• Cơ chế: Cố định Username = 'admin'; ký tự '--' cắt bỏ hoàn toàn mệnh đề 'AND Password = ...' phía sau.\n• Kết quả thực tế: Chiếm đoạt chính xác tài khoản Admin cấp cao nhất mà không cần nhập mật khẩu.", size=11.5, space_before=4)

# Card 3
c4_3 = s4.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.4) + (card_h + gap_h)*2, Inches(11.733), card_h)
c4_3.fill.solid()
c4_3.fill.fore_color.rgb = CARD_BG
c4_3.line.color.rgb = CARD_BORDER
tf = c4_3.text_frame
p = tf.paragraphs[0]
p.text = "03 — ĐĂNG NHẬP CHUẨN HỢP LỆ: admin / AdminPassword@2026"
p.font.name = "Arial"
p.font.size = Pt(13)
p.font.bold = True
p.font.color.rgb = ACCENT_BLUE
add_p(tf, "• Cơ chế: Dữ liệu người dùng bình thường, không chứa ký tự điều khiển hay can thiệp cú pháp.\n• Kết quả đối chứng: Đăng nhập thành công trên cả 2 nhánh (Lỗi & An toàn), chứng minh hệ thống hoạt động đúng nghiệp vụ.", size=11.5, space_before=4)

# Slide 5: 1.3 Minh chứng Web UI (Ảnh lớn)
s5 = prs.slides.add_slide(blank_layout)
add_header(s5, "1.3. SQL Injection — Minh Chứng Thực Nghiệm Trực Quan Trên Web UI", "CHỦ ĐỀ 1 • NGUYỄN DANH HỌC (23A1001D0158)")
p_sqli = "/media/hocjsoo/New Volume/OWASP_Demo/screenshots/final_sqli_proof.png"
if os.path.exists(p_sqli):
    s5.shapes.add_picture(p_sqli, Inches(2.26), Inches(1.35), width=Inches(8.8))

cap5 = s5.shapes.add_textbox(Inches(0.8), Inches(6.55), Inches(11.733), Inches(0.4))
tf5 = cap5.text_frame
p = tf5.paragraphs[0]
p.text = "🔴 Vulnerable: Trích xuất 4 tài khoản CSDL khi inject ' OR '1'='1' --  vs  🟢 Secure: Chặn đứng an toàn với EF Core (0 kết quả)"
p.font.name = "Arial"
p.font.size = Pt(11)
p.font.bold = True
p.font.color.rgb = ACCENT_BLUE
p.alignment = PP_ALIGN.CENTER

# Slide 6: 1.4 Minh chứng DevTools Network (Ảnh lớn)
s6 = prs.slides.add_slide(blank_layout)
add_header(s6, "1.4. SQL Injection — Minh Chứng Gói Tin Mạng (Chrome DevTools Network)", "CHỦ ĐỀ 1 • NGUYỄN DANH HỌC (23A1001D0158)")
p_dt_sqli = "/media/hocjsoo/New Volume/OWASP_Demo/screenshots/01_sqli_devtools_proof.png"
if os.path.exists(p_dt_sqli):
    s6.shapes.add_picture(p_dt_sqli, Inches(2.06), Inches(1.35), width=Inches(9.2))

cap6 = s6.shapes.add_textbox(Inches(0.8), Inches(6.55), Inches(11.733), Inches(0.4))
tf6 = cap6.text_frame
p = tf6.paragraphs[0]
p.text = "Soi gói tin mạng (Network Fetch/XHR): Payload gửi lên mang chuỗi ' OR '1'='1' -- và JSON trích xuất 4 tài khoản rò rỉ"
p.font.name = "Arial"
p.font.size = Pt(11)
p.font.bold = True
p.font.color.rgb = ACCENT_BLUE
p.alignment = PP_ALIGN.CENTER

# Slide 7: 1.5 Phòng thủ (Tách biệt Parameterized Query & Password Hashing)
s7 = prs.slides.add_slide(blank_layout)
add_header(s7, "1.5. SQL Injection — Cơ Chế Phòng Thủ Chuẩn Hóa", "CHỦ ĐỀ 1 • NGUYỄN DANH HỌC (23A1001D0158)")
c7_1 = s7.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.4), card_w, Inches(5.4))
c7_1.fill.solid()
c7_1.fill.fore_color.rgb = GREEN_LIGHT
c7_1.line.color.rgb = GREEN_BORDER
tf = c7_1.text_frame
p = tf.paragraphs[0]
p.text = "🟢 SECURE CODE: PARAMETERIZED QUERY"
p.font.name = "Arial"
p.font.size = Pt(14)
p.font.bold = True
p.font.color.rgb = ACCENT_GREEN

add_p(tf, "// DÙNG LINQ PARAMETERIZED TRONG EF CORE:\nvar account = _context.Accounts\n    .FirstOrDefault(a => a.Username == user \n                      && a.Password == pass);", size=11, color=RGBColor(6, 95, 70), space_before=8, font_name="Courier New")

add_p(tf, "Nguyên tắc cốt lõi: Data ≠ SQL Syntax", size=12.5, bold=True, color=TEXT_HEAD, space_before=14)
add_p(tf, "• Máy chủ SQL Server biên dịch cấu trúc cây truy vấn (Execution Plan) TRƯỚC khi nhận tham số.\n• Biến @p0, @p1 nhận toàn bộ payload như một giá trị văn bản (literal), vô hiệu hóa hoàn toàn ý đồ chèn lệnh.", size=12, space_before=6)

c7_2 = s7.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.85), Inches(1.4), card_w, Inches(5.4))
c7_2.fill.solid()
c7_2.fill.fore_color.rgb = CARD_BG
c7_2.line.color.rgb = CARD_BORDER
tf = c7_2.text_frame
p = tf.paragraphs[0]
p.text = "BẢO VỆ MẬT KHẨU (TÁCH BIỆT LOGIC)"
p.font.name = "Arial"
p.font.size = Pt(14)
p.font.bold = True
p.font.color.rgb = ACCENT_BLUE

add_p(tf, "Phân định rõ ràng trong báo cáo bảo mật:", size=12, bold=True, color=TEXT_HEAD, space_before=8)
add_p(tf, "1. Phòng thủ SQL Injection: Dùng Parameterized Query / EF Core LINQ.", size=12, bold=True, color=ACCENT_GREEN, space_before=6)
add_p(tf, "2. Bảo vệ mật khẩu (Defense-in-Depth): Dùng Password Hashing (ASP.NET Core Identity PBKDF2) kèm Salt ngẫu nhiên.", size=12, bold=True, color=ACCENT_BLUE, space_before=8)
add_p(tf, "3. Đặc quyền tối thiểu (Least Privilege): Ứng dụng chỉ được cấp quyền SELECT/INSERT/UPDATE trên bảng cần thiết, cấm quyền sa/DDL.", size=12, space_before=8)


# ==============================================================================
# PHẦN 2: STORED XSS — NGUYỄN THANH BÌNH (SLIDES 8 - 12)
# ==============================================================================
# Slide 8: 2.1 Where & Why (Sơ đồ luồng + Code lỗi @Html.Raw)
s8 = prs.slides.add_slide(blank_layout)
add_header(s8, "2.1. Stored XSS — Lỗi Phát Tán Như Thế Nào & Nguyên Nhân Gốc", "CHỦ ĐỀ 2 • NGUYỄN THANH BÌNH (23A1001D0041)")
c8_1 = s8.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.4), card_w, Inches(5.4))
c8_1.fill.solid()
c8_1.fill.fore_color.rgb = CARD_BG
c8_1.line.color.rgb = CARD_BORDER
tf = c8_1.text_frame
p = tf.paragraphs[0]
p.text = "Sơ Đồ Luồng Phát Tán Của Stored XSS"
p.font.name = "Arial"
p.font.size = Pt(15)
p.font.bold = True
p.font.color.rgb = ACCENT_RED

add_p(tf, "Attacker ➔ Nhập mã JavaScript vào Form đánh giá", size=12, bold=True, color=ACCENT_RED, space_before=8)
add_p(tf, "            │\n            ▼", size=12, bold=True, color=TEXT_MUTED, space_before=2)
add_p(tf, "Web Server ➔ Lưu nguyên văn mã độc vào CSDL (dbo.Comments)", size=12, bold=True, color=TEXT_HEAD, space_before=2)
add_p(tf, "            │\n            ▼", size=12, bold=True, color=TEXT_MUTED, space_before=2)
add_p(tf, "Victim Browser ➔ Tải trang, tự động biên dịch & thực thi script", size=12, bold=True, color=ACCENT_RED, space_before=2)

add_p(tf, "\n📌 ROOT CAUSE (Nguyên nhân cốt lõi):\nDữ liệu người dùng không tin cậy được render trực tiếp ra trình duyệt mà không qua cơ chế mã hóa đầu ra (Output Encoding).", size=12, color=TEXT_BODY, space_before=8)

c8_2 = s8.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.85), Inches(1.4), card_w, Inches(5.4))
c8_2.fill.solid()
c8_2.fill.fore_color.rgb = RED_LIGHT
c8_2.line.color.rgb = RED_BORDER
tf = c8_2.text_frame
p = tf.paragraphs[0]
p.text = "🔴 VULNERABLE: SỰ LẠM DỤNG @Html.Raw()"
p.font.name = "Arial"
p.font.size = Pt(14)
p.font.bold = True
p.font.color.rgb = ACCENT_RED

add_p(tf, "// CONTROLLER LƯU NGUYÊN VĂN NỘI DUNG VÀO CSDL:\nvar comment = new Comment {\n    Author = author,\n    Content = content // Chứa thẻ <script> độc hại\n};\n_context.Comments.Add(comment);\n\n// RAZOR VIEW HIỂN THỊ DỮ LIỆU THÔ BẰNG HTML.RAW:\n<div class=\"comment-body\">\n    @Html.Raw(comment.Content)\n</div>", size=10, color=RGBColor(159, 18, 57), space_before=6, font_name="Courier New")

add_p(tf, "Điểm mù kỹ thuật:\n@Html.Raw() vô hiệu hóa tính năng tự động mã hóa HTML của Razor View, mở toang cửa cho script độc hại chạy trên máy người dùng.", size=12, bold=True, color=TEXT_HEAD, space_before=10)

# Slide 9: 2.2 3 Kịch bản XSS (3 Card ngang)
s9 = prs.slides.add_slide(blank_layout)
add_header(s9, "2.2. Stored XSS — 3 Kịch Bản Tiêm Mã Độc Thực Nghiệm", "CHỦ ĐỀ 2 • NGUYỄN THANH BÌNH (23A1001D0041)")

# Card 1
c9_1 = s9.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.4), Inches(11.733), card_h)
c9_1.fill.solid()
c9_1.fill.fore_color.rgb = RED_LIGHT
c9_1.line.color.rgb = RED_BORDER
tf = c9_1.text_frame
p = tf.paragraphs[0]
p.text = "01 — ĐÁNH CẮP COOKIE / SESSION: <script>alert(document.cookie);</script>"
p.font.name = "Arial"
p.font.size = Pt(13)
p.font.bold = True
p.font.color.rgb = ACCENT_RED
add_p(tf, "• Cơ chế: Trình duyệt thực thi thẻ <script> và đọc toàn bộ giá trị lưu trong thuộc tính document.cookie.\n• Hậu quả: Hacker chiếm đoạt Session Token (Cookie Stealing) để đăng nhập mạo danh người dùng hợp lệ.", size=11.5, space_before=4)

# Card 2
c9_2 = s9.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.4) + card_h + gap_h, Inches(11.733), card_h)
c9_2.fill.solid()
c9_2.fill.fore_color.rgb = AMBER_LIGHT
c9_2.line.color.rgb = AMBER_BORDER
tf = c9_2.text_frame
p = tf.paragraphs[0]
p.text = "02 — BYPASS BỘ LỌC QUA EVENT HANDLER: <img src=x onerror=alert('XSS!')>"
p.font.name = "Arial"
p.font.size = Pt(13)
p.font.bold = True
p.font.color.rgb = ACCENT_AMBER
add_p(tf, "• Cơ chế: Khi đường dẫn ảnh bị lỗi (src=x), trình duyệt lập tức kích hoạt sự kiện onerror để chạy mã JavaScript.\n• Mục tiêu: Vượt qua các bộ lọc từ khóa đơn giản chỉ chặn chuỗi '<script>'.", size=11.5, space_before=4)

# Card 3
c9_3 = s9.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.4) + (card_h + gap_h)*2, Inches(11.733), card_h)
c9_3.fill.solid()
c9_3.fill.fore_color.rgb = CARD_BG
c9_3.line.color.rgb = CARD_BORDER
tf = c9_3.text_frame
p = tf.paragraphs[0]
p.text = "03 — ĐÁNH GIÁ CHUẨN HỢP LỆ: Dịch vụ tư vấn rất chuyên nghiệp, 5 sao!"
p.font.name = "Arial"
p.font.size = Pt(13)
p.font.bold = True
p.font.color.rgb = ACCENT_BLUE
add_p(tf, "• Cơ chế: Văn bản thuần túy không chứa thẻ HTML hay kịch bản đặc biệt.\n• Kết quả đối chứng: Hiển thị an toàn trên cả 2 nhánh (Lỗi & Đã phòng thủ).", size=11.5, space_before=4)

# Slide 10: 2.3 Minh chứng Web UI XSS (Ảnh lớn)
s10 = prs.slides.add_slide(blank_layout)
add_header(s10, "2.3. Stored XSS — Minh Chứng Thực Nghiệm Trực Quan Trên Web UI", "CHỦ ĐỀ 2 • NGUYỄN THANH BÌNH (23A1001D0041)")
p_xss = "/media/hocjsoo/New Volume/OWASP_Demo/screenshots/final_xss_proof.png"
if os.path.exists(p_xss):
    s10.shapes.add_picture(p_xss, Inches(2.86), Inches(1.35), width=Inches(7.6))

cap10 = s10.shapes.add_textbox(Inches(0.8), Inches(6.55), Inches(11.733), Inches(0.4))
tf10 = cap10.text_frame
p = tf10.paragraphs[0]
p.text = "🔴 Vulnerable: Script lưu CSDL tự động bật popup Alert  vs  🟢 Secure: Razor tự động mã hóa HTML thành chữ an toàn"
p.font.name = "Arial"
p.font.size = Pt(11)
p.font.bold = True
p.font.color.rgb = ACCENT_RED
p.alignment = PP_ALIGN.CENTER

# Slide 11: 2.4 Minh chứng DevTools Cookie & Console (Ảnh lớn + Chú thích Lab chuẩn)
s11 = prs.slides.add_slide(blank_layout)
add_header(s11, "2.4. Stored XSS — Minh Chứng Kiểm Thử Qua Chrome DevTools (Cookie & Console)", "CHỦ ĐỀ 2 • NGUYỄN THANH BÌNH (23A1001D0041)")
p_dt_xss = "/media/hocjsoo/New Volume/OWASP_Demo/screenshots/02_xss_devtools_proof.png"
if os.path.exists(p_dt_xss):
    s11.shapes.add_picture(p_dt_xss, Inches(2.06), Inches(1.35), width=Inches(9.2))

cap11 = s11.shapes.add_textbox(Inches(0.8), Inches(6.55), Inches(11.733), Inches(0.4))
tf11 = cap11.text_frame
p = tf11.paragraphs[0]
p.text = "Mô phỏng Lab: Nhánh vulnerable cấu hình HttpOnly = false để minh họa JavaScript đọc trộm Session Cookie"
p.font.name = "Arial"
p.font.size = Pt(11)
p.font.bold = True
p.font.color.rgb = ACCENT_RED
p.alignment = PP_ALIGN.CENTER

# Slide 12: 2.5 Phòng thủ XSS
s12 = prs.slides.add_slide(blank_layout)
add_header(s12, "2.5. Stored XSS — Cơ Chế Phòng Thủ Chuẩn Trong ASP.NET Core", "CHỦ ĐỀ 2 • NGUYỄN THANH BÌNH (23A1001D0041)")
c12_1 = s12.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.4), card_w, Inches(5.4))
c12_1.fill.solid()
c12_1.fill.fore_color.rgb = GREEN_LIGHT
c12_1.line.color.rgb = GREEN_BORDER
tf = c12_1.text_frame
p = tf.paragraphs[0]
p.text = "🟢 SECURE: TỰ ĐỘNG HTML ENCODING"
p.font.name = "Arial"
p.font.size = Pt(14)
p.font.bold = True
p.font.color.rgb = ACCENT_GREEN

add_p(tf, "// CÚ PHÁP RAZOR VIEW MẶC ĐỊNH (KHÔNG DÙNG HTML.RAW):\n<p>@comment.Content</p>\n\n// HOẶC MÃ HÓA TẠI CONTROLLER VỚI HTMLENCODER:\nstring encoded = HtmlEncoder.Default.Encode(content);\ncomment.Content = encoded;", size=10.5, color=RGBColor(6, 95, 70), space_before=8, font_name="Courier New")

add_p(tf, "Cơ chế biến đổi ký tự nhạy cảm:\n• '<' biến đổi thành &lt;\n• '>' biến đổi thành &gt;\nTrình duyệt chỉ hiển thị dưới dạng văn bản thuần, không thực thi script.", size=12, space_before=12)

c12_2 = s12.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.85), Inches(1.4), card_w, Inches(5.4))
c12_2.fill.solid()
c12_2.fill.fore_color.rgb = CARD_BG
c12_2.line.color.rgb = CARD_BORDER
tf = c12_2.text_frame
p = tf.paragraphs[0]
p.text = "CẤU HÌNH COOKIE AN TOÀN (PRODUCTION)"
p.font.name = "Arial"
p.font.size = Pt(14)
p.font.bold = True
p.font.color.rgb = ACCENT_BLUE

add_p(tf, "// BẢO VỆ COOKIE PHIÊN TRONG SẢN PHẨM THỰC TẾ:\nResponse.Cookies.Append(\"AuthSessionToken\", token, new CookieOptions\n{\n    HttpOnly = true, // CẤM JAVASCRIPT ĐỌC TRỘM\n    Secure = true,   // BẮT BUỘC TRUYỀN QUA HTTPS\n    SameSite = SameSiteMode.Strict\n});", size=10, color=TEXT_HEAD, space_before=8, font_name="Courier New")

add_p(tf, "Chính sách CSP (Content Security Policy):\nCấu hình HTTP Header 'script-src self' ngăn chặn trình duyệt chạy inline script lạ, tạo thêm lớp bảo vệ kiên cố.", size=12, space_before=12)


# ==============================================================================
# PHẦN 3: CSRF ATTACK — NGUYỄN MINH CƯỜNG (SLIDES 13 - 17)
# ==============================================================================
# Slide 13: 3.1 Where & Why (Sơ đồ luồng + Phân loại OWASP chuẩn)
s13 = prs.slides.add_slide(blank_layout)
add_header(s13, "3.1. CSRF Attack — Cơ Chế Mượn Quyền & Phân Loại Chuẩn OWASP", "CHỦ ĐỀ 3 • NGUYỄN MINH CƯỜNG (23A1001D0058)")
c13_1 = s13.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.4), card_w, Inches(5.4))
c13_1.fill.solid()
c13_1.fill.fore_color.rgb = CARD_BG
c13_1.line.color.rgb = CARD_BORDER
tf = c13_1.text_frame
p = tf.paragraphs[0]
p.text = "Sơ Đồ Luồng Tấn Công Mượn Quyền (CSRF)"
p.font.name = "Arial"
p.font.size = Pt(15)
p.font.bold = True
p.font.color.rgb = ACCENT_RED

add_p(tf, "Victim logged in ➔ Đang có Cookie phiên làm việc hợp lệ", size=12, bold=True, color=TEXT_HEAD, space_before=8)
add_p(tf, "            │\n            ▼", size=12, bold=True, color=TEXT_MUTED, space_before=2)
add_p(tf, "Mở Attacker Site ➔ Trang web độc hại kích hoạt POST ngầm", size=12, bold=True, color=ACCENT_RED, space_before=2)
add_p(tf, "            │\n            ▼", size=12, bold=True, color=TEXT_MUTED, space_before=2)
add_p(tf, "Target Server ➔ Tự động nhận Cookie, thực hiện trừ 20 triệu", size=12, bold=True, color=ACCENT_RED, space_before=2)

add_p(tf, "\n📌 PHÂN LOẠI CHUẨN OWASP:\nCSRF không phải mục độc lập trong OWASP Top 10 2021; nó liên quan trực tiếp đến A01:2021 – Broken Access Control và cơ chế bảo toàn yêu cầu phiên.", size=12, color=TEXT_BODY, space_before=8)

c13_2 = s13.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.85), Inches(1.4), card_w, Inches(5.4))
c13_2.fill.solid()
c13_2.fill.fore_color.rgb = RED_LIGHT
c13_2.line.color.rgb = RED_BORDER
tf = c13_2.text_frame
p = tf.paragraphs[0]
p.text = "🔴 VULNERABLE: ACTION THIẾU TOKEN"
p.font.name = "Arial"
p.font.size = Pt(14)
p.font.bold = True
p.font.color.rgb = ACCENT_RED

add_p(tf, "// ACTION NHẬN YÊU CẦU KHÔNG CÓ BẢO VỆ TOKEN:\n[HttpPost]\npublic IActionResult TransferVulnerable(decimal amount)\n{\n    var victim = _context.Wallets.Find(1);\n    var hacker = _context.Wallets.Find(2);\n    \n    victim.Balance -= amount;\n    hacker.Balance += amount;\n    _context.SaveChanges();\n    return RedirectToAction(\"Index\");\n}", size=10, color=RGBColor(159, 18, 57), space_before=8, font_name="Courier New")

add_p(tf, "Điểm mù kiểm soát:\nAction thiếu [ValidateAntiForgeryToken], máy chủ chỉ kiểm tra Cookie hợp lệ rồi trừ tiền ngay mà không xác minh yêu cầu xuất phát từ đâu!", size=12, bold=True, color=TEXT_HEAD, space_before=10)

# Slide 14: 3.2 Kịch bản Form ẩn Attacker Site
s14 = prs.slides.add_slide(blank_layout)
add_header(s14, "3.2. CSRF Attack — Kịch Bản Khai Thác Qua Biểu Mẫu Ẩn", "CHỦ ĐỀ 3 • NGUYỄN MINH CƯỜNG (23A1001D0058)")
c14_main = s14.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.4), Inches(11.733), Inches(5.3))
c14_main.fill.solid()
c14_main.fill.fore_color.rgb = RED_LIGHT
c14_main.line.color.rgb = RED_BORDER
tf = c14_main.text_frame
p = tf.paragraphs[0]
p.text = "Kịch Bản: Trang Web Giả Mạo Điều Khiển Yêu Cầu Chuyển Tiền"
p.font.name = "Arial"
p.font.size = Pt(16)
p.font.bold = True
p.font.color.rgb = ACCENT_RED
add_p(tf, "1. Nạn nhân có phiên làm việc đang hoạt động trên hệ thống tư vấn trực tuyến.\n2. Nạn nhân được dẫn dụ truy cập một trang web độc hại ở một tab khác (AttackerSite.cshtml).\n3. Trang web độc hại chứa một biểu mẫu ẩn trỏ trực tiếp tới địa chỉ nhận lệnh của ứng dụng mục tiêu:", size=13, space_before=8)
add_p(tf, "<!-- BIỂU MẪU ẨN TRÊN TRANG WEB CỦA KẺ TẤN CÔNG -->\n<form id=\"exploit\" action=\"http://localhost:5076/Csrf/TransferVulnerable\" method=\"POST\">\n    <input type=\"hidden\" name=\"amount\" value=\"20000000\" />\n</form>\n<script>document.getElementById('exploit').submit();</script>", size=11, bold=True, color=RGBColor(159, 18, 57), space_before=6, font_name="Courier New")
add_p(tf, "Hậu quả thực tế: Giao dịch chuyển tiền được thực thi hoàn toàn tự động mà không có sự đồng thuận chủ động của người dùng, dẫn đến thay đổi số dư ví tài khoản.", size=13, bold=True, color=ACCENT_RED, space_before=8)

# Slide 15: 3.3 Minh chứng Web UI CSRF (Ảnh thật)
s15 = prs.slides.add_slide(blank_layout)
add_header(s15, "3.3. CSRF Attack — Minh Chứng Thực Nghiệm Trực Quan Trên Web UI", "CHỦ ĐỀ 3 • NGUYỄN MINH CƯỜNG (23A1001D0058)")
p_csrf = "/media/hocjsoo/New Volume/OWASP_Demo/screenshots/final_csrf_proof.png"
if os.path.exists(p_csrf):
    s15.shapes.add_picture(p_csrf, Inches(2.76), Inches(1.35), width=Inches(7.8))

cap15 = s15.shapes.add_textbox(Inches(0.8), Inches(6.55), Inches(11.733), Inches(0.4))
tf15 = cap15.text_frame
p = tf15.paragraphs[0]
p.text = "Minh chứng đối chứng: Mô phỏng giao dịch trái phép 20.000.000 VNĐ trong lab (Trái) vs Chặn đứng yêu cầu thiếu Token (Phải)"
p.font.name = "Arial"
p.font.size = Pt(11)
p.font.bold = True
p.font.color.rgb = ACCENT_GREEN
p.alignment = PP_ALIGN.CENTER

# Slide 16: 3.4 Minh chứng DevTools Headers (Ảnh thật)
s16 = prs.slides.add_slide(blank_layout)
add_header(s16, "3.4. CSRF Attack — Minh Chứng Gói Tin Mạng (Chrome DevTools Headers)", "CHỦ ĐỀ 3 • NGUYỄN MINH CƯỜNG (23A1001D0058)")
p_dt_csrf = "/media/hocjsoo/New Volume/OWASP_Demo/screenshots/03_csrf_devtools_proof.png"
if os.path.exists(p_dt_csrf):
    s16.shapes.add_picture(p_dt_csrf, Inches(2.06), Inches(1.35), width=Inches(9.2))

cap16 = s16.shapes.add_textbox(Inches(0.8), Inches(6.55), Inches(11.733), Inches(0.4))
tf16 = cap16.text_frame
p = tf16.paragraphs[0]
p.text = "Soi gói tin HTTP POST gửi ngầm từ Attacker Site: Đính kèm Cookie tự động nhưng thiếu hoàn toàn Anti-Forgery Token"
p.font.name = "Arial"
p.font.size = Pt(11)
p.font.bold = True
p.font.color.rgb = ACCENT_GREEN
p.alignment = PP_ALIGN.CENTER

# Slide 17: 3.5 Phòng thủ CSRF Token & Fetch API
s17 = prs.slides.add_slide(blank_layout)
add_header(s17, "3.5. CSRF Attack — Cơ Chế Phòng Thủ Bằng Anti-Forgery Token", "CHỦ ĐỀ 3 • NGUYỄN MINH CƯỜNG (23A1001D0058)")
c17_1 = s17.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.4), card_w, Inches(5.4))
c17_1.fill.solid()
c17_1.fill.fore_color.rgb = GREEN_LIGHT
c17_1.line.color.rgb = GREEN_BORDER
tf = c17_1.text_frame
p = tf.paragraphs[0]
p.text = "🟢 SECURE: SYNCHRONIZER TOKEN PATTERN"
p.font.name = "Arial"
p.font.size = Pt(14)
p.font.bold = True
p.font.color.rgb = ACCENT_GREEN

add_p(tf, "// 1. TRONG RAZOR VIEW: NHÚNG TOKEN VÀO FORM\n<form asp-action=\"TransferSecure\" method=\"post\">\n    @Html.AntiForgeryToken()\n    <input type=\"number\" name=\"amount\" />\n    <button type=\"submit\">Chuyển tiền</button>\n</form>\n\n// 2. TRONG CONTROLLER: BẮT BUỘC KIỂM TRA TOKEN\n[HttpPost]\n[ValidateAntiForgeryToken]\npublic IActionResult TransferSecure(decimal amount)\n{\n    // Chỉ thực thi khi Token từ Form khớp với Token trong Cookie\n}", size=9.5, color=RGBColor(6, 95, 70), space_before=6, font_name="Courier New")

add_p(tf, "Cơ chế bảo vệ:\nTrang web của hacker ở tên miền khác bị chặn đọc token bí mật bởi chính sách Same-Origin Policy (SOP). Mọi request thiếu token đều bị chặn đứng.", size=11.5, space_before=6)

c17_2 = s17.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.85), Inches(1.4), card_w, Inches(5.4))
c17_2.fill.solid()
c17_2.fill.fore_color.rgb = CARD_BG
c17_2.line.color.rgb = CARD_BORDER
tf = c17_2.text_frame
p = tf.paragraphs[0]
p.text = "TÍCH HỢP TOKEN TRONG AJAX FETCH API"
p.font.name = "Arial"
p.font.size = Pt(14)
p.font.bold = True
p.font.color.rgb = ACCENT_BLUE

add_p(tf, "// TRÍCH XUẤT VÀ GỬI TOKEN QUA HTTP HEADER:\nconst token = document.querySelector('input[name=\"__RequestVerificationToken\"]').value;\n\nfetch('/Csrf/TransferAjax', {\n    method: 'POST',\n    headers: {\n        'Content-Type': 'application/json',\n        'RequestVerificationToken': token // HEADER BẢO VỆ\n    },\n    body: JSON.stringify({ amount: 5000000 })\n})\n.then(response => response.json());", size=9.5, color=TEXT_HEAD, space_before=6, font_name="Courier New")

add_p(tf, "Ứng dụng trong thực tế:\nĐối với các tác vụ gọi AJAX không tải lại trang, token được đọc từ thẻ ẩn và truyền qua HTTP Header để đảm bảo vừa tiện dụng vừa an toàn.", size=11.5, space_before=6)


# ==============================================================================
# PHẦN 4: BÀI HỌC RÚT RA (3 CARDS LỚN THEO ĐÚNG GỢI Ý)
# ==============================================================================
s18 = prs.slides.add_slide(blank_layout)
add_header(s18, "Bài Học Rút Ra: Bảng Đối Chiếu 3 Lỗ Hổng & Cơ Chế Phòng Thủ")

card3_w = Inches(3.64)
gap3 = Inches(0.39)

# Card 1: SQLi
k1 = s18.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.4), card3_w, Inches(5.3))
k1.fill.solid()
k1.fill.fore_color.rgb = BLUE_LIGHT
k1.line.color.rgb = BLUE_BORDER
tf = k1.text_frame
p = tf.paragraphs[0]
p.text = "SQL INJECTION"
p.font.name = "Arial"
p.font.size = Pt(15)
p.font.bold = True
p.font.color.rgb = ACCENT_BLUE

add_p(tf, "Nguyên nhân cốt lõi:", size=12, bold=True, color=TEXT_HEAD, space_before=10)
add_p(tf, "Ghép chuỗi SQL trực tiếp từ tham số người dùng.", size=12, color=ACCENT_RED, space_before=2)

add_p(tf, "      │\n      ▼", size=14, bold=True, color=TEXT_MUTED, space_before=6)

add_p(tf, "Cơ chế phòng thủ:", size=12, bold=True, color=TEXT_HEAD, space_before=6)
add_p(tf, "Parameterized Query\n(Entity Framework Core LINQ)", size=12, bold=True, color=ACCENT_GREEN, space_before=2)

add_p(tf, "\nPhụ trách: Nguyễn Danh Học\nMSV: 23A1001D0158", size=11, color=TEXT_MUTED, space_before=14)

# Card 2: XSS
k2 = s18.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8) + card3_w + gap3, Inches(1.4), card3_w, Inches(5.3))
k2.fill.solid()
k2.fill.fore_color.rgb = RED_LIGHT
k2.line.color.rgb = RED_BORDER
tf = k2.text_frame
p = tf.paragraphs[0]
p.text = "STORED XSS"
p.font.name = "Arial"
p.font.size = Pt(15)
p.font.bold = True
p.font.color.rgb = ACCENT_RED

add_p(tf, "Nguyên nhân cốt lõi:", size=12, bold=True, color=TEXT_HEAD, space_before=10)
add_p(tf, "Lưu script thô vào CSDL & render bằng @Html.Raw().", size=12, color=ACCENT_RED, space_before=2)

add_p(tf, "      │\n      ▼", size=14, bold=True, color=TEXT_MUTED, space_before=6)

add_p(tf, "Cơ chế phòng thủ:", size=12, bold=True, color=TEXT_HEAD, space_before=6)
add_p(tf, "Context-Aware HTML Encoding\n(@comment.Content + HttpOnly)", size=12, bold=True, color=ACCENT_GREEN, space_before=2)

add_p(tf, "\nPhụ trách: Nguyễn Thanh Bình\nMSV: 23A1001D0041", size=11, color=TEXT_MUTED, space_before=14)

# Card 3: CSRF
k3 = s18.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8) + (card3_w + gap3)*2, Inches(1.4), card3_w, Inches(5.3))
k3.fill.solid()
k3.fill.fore_color.rgb = GREEN_LIGHT
k3.line.color.rgb = GREEN_BORDER
tf = k3.text_frame
p = tf.paragraphs[0]
p.text = "CSRF ATTACK"
p.font.name = "Arial"
p.font.size = Pt(15)
p.font.bold = True
p.font.color.rgb = ACCENT_GREEN

add_p(tf, "Nguyên nhân cốt lõi:", size=12, bold=True, color=TEXT_HEAD, space_before=10)
add_p(tf, "Chỉ kiểm tra Cookie mà thiếu xác thực nguồn gốc Request.", size=12, color=ACCENT_RED, space_before=2)

add_p(tf, "      │\n      ▼", size=14, bold=True, color=TEXT_MUTED, space_before=6)

add_p(tf, "Cơ chế phòng thủ:", size=12, bold=True, color=TEXT_HEAD, space_before=6)
add_p(tf, "Synchronizer Token Pattern\n(@Html.AntiForgeryToken)", size=12, bold=True, color=ACCENT_GREEN, space_before=2)

add_p(tf, "\nPhụ trách: Nguyễn Minh Cường\nMSV: 23A1001D0058", size=11, color=TEXT_MUTED, space_before=14)


# ==============================================================================
# SLIDE 19: KẾT LUẬN & Q&A (SÚC TÍCH, KHÔNG SÁO RỖNG)
# ==============================================================================
s19 = prs.slides.add_slide(blank_layout)
add_header(s19, "Kết Luận Thực Nghiệm & Q&A")
main_c19 = s19.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.0), Inches(1.5), Inches(11.333), Inches(5.0))
main_c19.fill.solid()
main_c19.fill.fore_color.rgb = CARD_BG
main_c19.line.color.rgb = CARD_BORDER
tf = main_c19.text_frame
p = tf.paragraphs[0]
p.text = "TỔNG KẾT: TỪ KHAI THÁC THỰC NGHIỆM ĐẾN PHÒNG THỦ ĐA TẦNG"
p.font.name = "Arial"
p.font.size = Pt(16)
p.font.bold = True
p.font.color.rgb = ACCENT_BLUE

add_p(tf, "01 — Tái lập thành công lỗ hổng:", size=13, bold=True, color=ACCENT_RED, space_before=14)
add_p(tf, "Nhóm đã biểu diễn thành công kịch bản tấn công trên 3 bề mặt: Giao diện Web, Gói tin mạng (Chrome DevTools) và CSDL SQL Server 2025.", size=12.5, color=TEXT_BODY, space_before=2)

add_p(tf, "02 — Kiểm chứng giải pháp phòng thủ triệt để:", size=13, bold=True, color=ACCENT_GREEN, space_before=12)
add_p(tf, "Áp dụng Parameterized Query, HTML Encoding và Anti-Forgery Token chặn đứng 100% các vector tấn công.", size=12.5, color=TEXT_BODY, space_before=2)

add_p(tf, "03 — Triết lý an ninh cốt lõi (Security Principle):", size=13, bold=True, color=ACCENT_BLUE, space_before=12)
add_p(tf, "• Never trust user-controlled input (Không bao giờ tin tưởng dữ liệu đầu vào).\n• Separate data from executable instructions (Luôn phân tách dữ liệu khỏi cú pháp lệnh).", size=12.5, bold=True, color=TEXT_HEAD, space_before=4)

add_p(tf, "\nXIN TRÂN TRỌNG CẢM ƠN THẦY VÀ CÁC BẠN! — Q&A", size=16, bold=True, color=ACCENT_RED, space_before=20)

out_file = "/media/hocjsoo/New Volume/OWASP_Demo/Slide_OWASP_Top3_Nhom_WNC_G01.pptx"
prs.save(out_file)
print(f"Master presentation v3 successfully saved to: {out_file}")
