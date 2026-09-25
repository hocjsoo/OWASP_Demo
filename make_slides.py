import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)

# Color Palette
DARK_BG = RGBColor(26, 32, 44)       # #1A202C
WHITE = RGBColor(255, 255, 255)
LIGHT_GRAY = RGBColor(241, 245, 249)
TEXT_DARK = RGBColor(30, 41, 59)
TEXT_MUTED = RGBColor(100, 116, 139)
ACCENT_RED = RGBColor(220, 38, 38)     # Danger / Hack
ACCENT_GREEN = RGBColor(22, 163, 74)   # Defense / Secure
ACCENT_BLUE = RGBColor(37, 99, 235)    # Brand / Primary
CARD_BG = RGBColor(248, 250, 252)

blank_layout = prs.slide_layouts[6]

def add_header(slide, title_text, category_text="OWASP TOP 10 - A03:2021 (INJECTION)"):
    # Header bar
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(1.2))
    shape.fill.solid()
    shape.fill.fore_color.rgb = DARK_BG
    shape.line.fill.background()

    # Category tag
    tx_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.15), Inches(11.7), Inches(0.35))
    tf = tx_box.text_frame
    tf.word_wrap = True
    p0 = tf.paragraphs[0]
    p0.text = category_text.upper()
    p0.font.size = Pt(11)
    p0.font.bold = True
    p0.font.color.rgb = ACCENT_RED

    # Title
    p1 = tf.add_paragraph()
    p1.text = title_text
    p1.font.size = Pt(20)
    p1.font.bold = True
    p1.font.color.rgb = WHITE

    # Footer
    footer = slide.shapes.add_textbox(Inches(0.8), Inches(7.05), Inches(11.7), Inches(0.35))
    ft = footer.text_frame
    p_ft = ft.paragraphs[0]
    p_ft.text = "Nhóm WNC.G01 (HOU) • Môn Lập Trình Web Nâng Cao • ThS. Lê Hữu Dũng"
    p_ft.font.size = Pt(10)
    p_ft.font.color.rgb = TEXT_MUTED

# ==========================================
# SLIDE 1: TRANG TIÊU ĐỀ
# ==========================================
s1 = prs.slides.add_slide(blank_layout)
bg1 = s1.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(7.5))
bg1.fill.solid()
bg1.fill.fore_color.rgb = DARK_BG
bg1.line.fill.background()

tb1 = s1.shapes.add_textbox(Inches(1.0), Inches(1.2), Inches(11.333), Inches(5.0))
tf1 = tb1.text_frame
tf1.word_wrap = True

p = tf1.paragraphs[0]
p.text = "TRƯỜNG ĐẠI HỌC MỞ HÀ NỘI — KHOA CÔNG NGHỆ THÔNG TIN"
p.font.size = Pt(14)
p.font.bold = True
p.font.color.rgb = RGBColor(148, 163, 184)

p = tf1.add_paragraph()
p.text = "BÁO CÁO THỰC NGHIỆM LỖ HỔNG AN TOÀN WEB"
p.font.size = Pt(18)
p.font.bold = True
p.font.color.rgb = ACCENT_RED
p.space_before = Pt(14)

p = tf1.add_paragraph()
p.text = "SQL INJECTION (OWASP A03:2021)\nKịch Bản Tấn Công & Cơ Chế Phòng Thủ Trực Quan Trên ASP.NET Core"
p.font.size = Pt(28)
p.font.bold = True
p.font.color.rgb = WHITE
p.space_before = Pt(8)

p = tf1.add_paragraph()
p.text = "Giảng viên hướng dẫn: ThS. Lê Hữu Dũng\nMôn học: Lập trình Web nâng cao"
p.font.size = Pt(14)
p.font.color.rgb = RGBColor(203, 213, 225)
p.space_before = Pt(20)

p = tf1.add_paragraph()
p.text = "Nhóm thực hiện: WNC.G01\n1. Nguyễn Danh Học (MSV: 23A1001D0158) — Báo cáo & Thực nghiệm chính\n2. Nguyễn Thanh Bình (MSV: 23A1001D0041)\n3. Nguyễn Minh Cường (MSV: 23A1001D0058)"
p.font.size = Pt(13)
p.font.color.rgb = RGBColor(148, 163, 184)
p.space_before = Pt(16)


# ==========================================
# SLIDE 2: BẢN CHẤT LỖ HỔNG
# ==========================================
s2 = prs.slides.add_slide(blank_layout)
add_header(s2, "1. Bản Chất Kỹ Thuật Của Lỗ Hổng SQL Injection")

tb2_1 = s2.shapes.add_textbox(Inches(0.8), Inches(1.5), Inches(5.6), Inches(5.2))
tf2_1 = tb2_1.text_frame
tf2_1.word_wrap = True

p = tf2_1.paragraphs[0]
p.text = "Nguyên Nhân Cốt Lõi"
p.font.size = Pt(16)
p.font.bold = True
p.font.color.rgb = ACCENT_RED

p = tf2_1.add_paragraph()
p.text = "• Không phân tách giữa DỮ LIỆU (Data) và CÂU LỆNH (Code).\n• Lập trình viên ghép chuỗi đầu vào của người dùng trực tiếp vào câu lệnh SQL.\n• Kẻ tấn công lợi dụng các ký tự điều khiển (', --, /*, ;) để 'thoát' khỏi ngữ cảnh dữ liệu và ép SQL Server thực thi logic do hacker chỉ định."
p.font.size = Pt(13)
p.font.color.rgb = TEXT_DARK
p.space_before = Pt(8)

p = tf2_1.add_paragraph()
p.text = "Phân Loại Theo Tiêu Chuẩn OWASP WSTG"
p.font.size = Pt(16)
p.font.bold = True
p.font.color.rgb = ACCENT_BLUE
p.space_before = Pt(16)

p = tf2_1.add_paragraph()
p.text = "1. In-band SQLi (Khai thác trực diện):\n   - Error-based: Ép CSDL sinh lỗi để đọc dữ liệu.\n   - Union-based: Dùng UNION trích xuất các bảng khác.\n2. Inferential SQLi (Blind SQLi / Mù):\n   - Boolean-based: Đoán dữ liệu qua phản hồi Đúng/Sai.\n   - Time-based: Dùng WAITFOR DELAY đo thời gian.\n3. Out-of-band SQLi: Kích hoạt DNS/HTTP request từ CSDL."
p.font.size = Pt(12)
p.font.color.rgb = TEXT_DARK
p.space_before = Pt(6)

# Right Box: Code comparison
box2 = s2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.8), Inches(1.5), Inches(5.7), Inches(5.2))
box2.fill.solid()
box2.fill.fore_color.rgb = CARD_BG
box2.line.color.rgb = RGBColor(226, 232, 240)

tb2_2 = s2.shapes.add_textbox(Inches(7.0), Inches(1.6), Inches(5.3), Inches(5.0))
tf2_2 = tb2_2.text_frame
tf2_2.word_wrap = True

p = tf2_2.paragraphs[0]
p.text = "Đoạn Code Gây Lỗi Cụ Thể (C# / ASP.NET Core)"
p.font.size = Pt(14)
p.font.bold = True
p.font.color.rgb = ACCENT_RED

p = tf2_2.add_paragraph()
p.text = "// Nối chuỗi trực tiếp từ Request:\nstring query = $\"SELECT * FROM Accounts \" +\n               $\"WHERE Username = '{user}' \" +\n               $\"AND Password = '{pass}'\";\n\nvar users = _context.Accounts\n                    .FromSqlRaw(query)\n                    .ToList();"
p.font.name = "Courier New"
p.font.size = Pt(11)
p.font.color.rgb = RGBColor(185, 28, 28)
p.space_before = Pt(8)

p = tf2_2.add_paragraph()
p.text = "⚠️ Điểm mù nguy hiểm:"
p.font.size = Pt(13)
p.font.bold = True
p.font.color.rgb = TEXT_DARK
p.space_before = Pt(14)

p = tf2_2.add_paragraph()
p.text = "Lập trình viên tin tưởng rằng '{user}' chỉ chứa chữ cái thường. Khi hacker truyền vào ký tự nháy đơn ('), cấu trúc cú pháp của cả câu lệnh bị bẻ gãy hoàn toàn."
p.font.size = Pt(12)
p.font.color.rgb = TEXT_MUTED
p.space_before = Pt(4)


# ==========================================
# SLIDE 3: ĐỘ NGUY HIỂM THỰC TẾ
# ==========================================
s3 = prs.slides.add_slide(blank_layout)
add_header(s3, "2. Độ Nguy Hiểm Thực Tế Của SQL Injection (Tác Động Nghiêm Trọng)")

col_w = Inches(3.6)
gap = Inches(0.4)

# Danger Card 1
c1 = s3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.6), col_w, Inches(5.0))
c1.fill.solid()
c1.fill.fore_color.rgb = RGBColor(254, 242, 242)
c1.line.color.rgb = ACCENT_RED

t1 = s3.shapes.add_textbox(Inches(0.95), Inches(1.8), col_w - Inches(0.3), Inches(4.6))
tf = t1.text_frame
tf.word_wrap = True
p = tf.paragraphs[0]
p.text = "🚨 CHIẾM QUYỀN HỆ THỐNG\n(Authentication Bypass)"
p.font.size = Pt(14)
p.font.bold = True
p.font.color.rgb = ACCENT_RED
p = tf.add_paragraph()
p.text = "• Không cần biết mật khẩu của bất kỳ ai.\n• Chỉ với 1 chuỗi payload ngắn gọn, kẻ tấn công đăng nhập thẳng vào tài khoản Quản trị viên (Admin).\n• Chiếm đoạt phiên làm việc, thay đổi phân quyền và khóa tài khoản người dùng khác."
p.font.size = Pt(12)
p.font.color.rgb = TEXT_DARK
p.space_before = Pt(10)

# Danger Card 2
c2 = s3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8) + col_w + gap, Inches(1.6), col_w, Inches(5.0))
c2.fill.solid()
c2.fill.fore_color.rgb = RGBColor(255, 251, 235)
c2.line.color.rgb = RGBColor(217, 119, 6)

t2 = s3.shapes.add_textbox(Inches(0.95) + col_w + gap, Inches(1.8), col_w - Inches(0.3), Inches(4.6))
tf = t2.text_frame
tf.word_wrap = True
p = tf.paragraphs[0]
p.text = "📂 ĐÁNH CẮP TOÀN BỘ CSDL\n(Data Exfiltration)"
p.font.size = Pt(14)
p.font.bold = True
p.font.color.rgb = RGBColor(180, 83, 9)
p = tf.add_paragraph()
p.text = "• Trích xuất thông tin khách hàng, hồ sơ chuyên gia, số dư tài khoản ngân hàng.\n• Rò rỉ mật khẩu và ghi chú bảo mật cá nhân (Secret Notes).\n• Vi phạm nghiêm trọng luật bảo vệ dữ liệu cá nhân, làm sụp đổ uy tín của nền tảng tư vấn trực tuyến."
p.font.size = Pt(12)
p.font.color.rgb = TEXT_DARK
p.space_before = Pt(10)

# Danger Card 3
c3 = s3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8) + (col_w + gap)*2, Inches(1.6), col_w, Inches(5.0))
c3.fill.solid()
c3.fill.fore_color.rgb = RGBColor(241, 245, 249)
c3.line.color.rgb = DARK_BG

t3 = s3.shapes.add_textbox(Inches(0.95) + (col_w + gap)*2, Inches(1.8), col_w - Inches(0.3), Inches(4.6))
tf = t3.text_frame
tf.word_wrap = True
p = tf.paragraphs[0]
p.text = "💣 PHÁ HỦY / THỰC THI LỆNH\n(System Takeover)"
p.font.size = Pt(14)
p.font.bold = True
p.font.color.rgb = DARK_BG
p = tf.add_paragraph()
p.text = "• Xóa sổ toàn bộ cơ sở dữ liệu với lệnh DROP TABLE hoặc TRUNCATE.\n• Thay đổi số dư tiền trong tài khoản trái phép.\n• Trong các CSDL cấu hình lỏng lẻo, hacker có thể bật xp_cmdshell để chiếm quyền điều khiển toàn bộ máy chủ hệ điều hành."
p.font.size = Pt(12)
p.font.color.rgb = TEXT_DARK
p.space_before = Pt(10)


# ==========================================
# SLIDE 4: PHÂN TÍCH KỊCH BẢN TẤN CÔNG
# ==========================================
s4 = prs.slides.add_slide(blank_layout)
add_header(s4, "3. Kịch Bản Tấn Công: Phân Tích Cơ Chế Bẻ Gãy Cú Pháp")

tb4 = s4.shapes.add_textbox(Inches(0.8), Inches(1.5), Inches(11.7), Inches(5.3))
tf4 = tb4.text_frame
tf4.word_wrap = True

p = tf4.paragraphs[0]
p.text = "Kịch bản: Bypass xác thực đăng nhập Quản Trị Viên"
p.font.size = Pt(16)
p.font.bold = True
p.font.color.rgb = ACCENT_RED

p = tf4.add_paragraph()
p.text = "Input đưa vào ô Username:   ' OR '1'='1' --\nInput đưa vào ô Password:   (Nhập chuỗi bất kỳ hoặc bỏ trống)"
p.font.name = "Courier New"
p.font.size = Pt(14)
p.font.bold = True
p.font.color.rgb = ACCENT_BLUE
p.space_before = Pt(8)

p = tf4.add_paragraph()
p.text = "Cấu trúc câu truy vấn được SQL Server thông dịch:"
p.font.size = Pt(15)
p.font.bold = True
p.font.color.rgb = TEXT_DARK
p.space_before = Pt(16)

p = tf4.add_paragraph()
p.text = "SELECT * FROM Accounts WHERE Username = '' OR '1'='1' --' AND Password = 'xyz'"
p.font.name = "Courier New"
p.font.size = Pt(14)
p.font.bold = True
p.font.color.rgb = ACCENT_RED
p.space_before = Pt(6)

p = tf4.add_paragraph()
p.text = "3 Thành Phần Của Payload Khai Thác:"
p.font.size = Pt(15)
p.font.bold = True
p.font.color.rgb = TEXT_DARK
p.space_before = Pt(16)

p = tf4.add_paragraph()
p.text = "1. Dấu nháy đơn (') : Đóng sớm chuỗi ký tự hợp lệ của trường Username, đưa ngữ cảnh về câu lệnh SQL.\n2. Mệnh đề OR '1'='1' : Mệnh đề luận lý chân lý. Vì '1'='1' luôn đúng, toàn bộ mệnh đề WHERE trả về TRUE cho TẤT CẢ các dòng trong bảng.\n3. Ký tự chú thích (--) : Trong SQL Server/SQLite, '--' biến toàn bộ đoạn mã phía sau thành ghi chú vô hiệu, triệt tiêu hoàn toàn bước kiểm tra mật khẩu (AND Password = '...')."
p.font.size = Pt(13)
p.font.color.rgb = TEXT_DARK
p.space_before = Pt(6)


# ==========================================
# SLIDE 5: THỰC NGHIỆM DEMO TRỰC TIẾP
# ==========================================
s5 = prs.slides.add_slide(blank_layout)
add_header(s5, "4. Minh Chứng Thực Nghiệm Demo Đối Chứng Trực Tiếp")

# Left: Exploit screenshot
img_path1 = "/media/hocjsoo/New Volume/OWASP_Demo/screenshots/02_Tan_Cong_Bypass_Thanh_Cong.png"
if os.path.exists(img_path1):
    s5.shapes.add_picture(img_path1, Inches(0.8), Inches(1.5), Inches(5.6), Inches(4.5))

lbl1 = s5.shapes.add_textbox(Inches(0.8), Inches(6.1), Inches(5.6), Inches(0.8))
tf = lbl1.text_frame
tf.word_wrap = True
p = tf.paragraphs[0]
p.text = "❌ BỊ LỖI: Nhập payload → Vượt mặt đăng nhập thành công, lộ toàn bộ CSDL và mật khẩu."
p.font.size = Pt(11)
p.font.bold = True
p.font.color.rgb = ACCENT_RED

# Right: Defense screenshot
img_path2 = "/media/hocjsoo/New Volume/OWASP_Demo/screenshots/03_So_Sanh_Phong_Thu_Thanh_Cong.png"
if os.path.exists(img_path2):
    s5.shapes.add_picture(img_path2, Inches(6.8), Inches(1.5), Inches(5.7), Inches(4.5))

lbl2 = s5.shapes.add_textbox(Inches(6.8), Inches(6.1), Inches(5.7), Inches(0.8))
tf = lbl2.text_frame
tf.word_wrap = True
p = tf.paragraphs[0]
p.text = "✅ ĐÃ VÁ: Cùng payload sang nhánh EF Core → Bị chặn đứng, bảo vệ dữ liệu an toàn 100%."
p.font.size = Pt(11)
p.font.bold = True
p.font.color.rgb = ACCENT_GREEN


# ==========================================
# SLIDE 6: BIỆN PHÁP PHÒNG THỦ
# ==========================================
s6 = prs.slides.add_slide(blank_layout)
add_header(s6, "5. Biện Pháp Phòng Thủ Chuẩn Trong ASP.NET Core MVC")

# Card Left: Code phòng thủ
box6 = s6.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.5), Inches(5.7), Inches(5.2))
box6.fill.solid()
box6.fill.fore_color.rgb = CARD_BG
box6.line.color.rgb = ACCENT_GREEN

tb6_1 = s6.shapes.add_textbox(Inches(1.0), Inches(1.6), Inches(5.3), Inches(5.0))
tf6_1 = tb6_1.text_frame
tf6_1.word_wrap = True

p = tf6_1.paragraphs[0]
p.text = "Cách Viết Code Chuẩn Hóa Với EF Core"
p.font.size = Pt(14)
p.font.bold = True
p.font.color.rgb = ACCENT_GREEN

p = tf6_1.add_paragraph()
p.text = "// Dùng LINQ Parameterized (KHUYÊN DÙNG):\nvar account = _context.Accounts\n    .FirstOrDefault(a => a.Username == user \n                      && a.Password == pass);\n\n// Hoặc dùng FromSqlInterpolated nếu viết SQL thuần:\nvar account = _context.Accounts\n    .FromSqlInterpolated($\"SELECT * FROM Accounts WHERE Username={user} AND Password={pass}\")\n    .FirstOrDefault();"
p.font.name = "Courier New"
p.font.size = Pt(10.5)
p.font.color.rgb = RGBColor(22, 101, 52)
p.space_before = Pt(8)

p = tf6_1.add_paragraph()
p.text = "Tại sao cách này an toàn tuyệt đối?"
p.font.size = Pt(12)
p.font.bold = True
p.font.color.rgb = TEXT_DARK
p.space_before = Pt(12)

p = tf6_1.add_paragraph()
p.text = "SQL Server biên dịch cấu trúc lệnh TRƯỚC khi gán dữ liệu. Biến @p0 nhận toàn bộ payload như một chuỗi chữ bình thường, vô hiệu hóa hoàn toàn ý đồ chèn lệnh."
p.font.size = Pt(11)
p.font.color.rgb = TEXT_MUTED
p.space_before = Pt(4)

# Card Right: Bộ nguyên tắc phòng thủ
tb6_2 = s6.shapes.add_textbox(Inches(6.9), Inches(1.5), Inches(5.6), Inches(5.2))
tf6_2 = tb6_2.text_frame
tf6_2.word_wrap = True

p = tf6_2.paragraphs[0]
p.text = "Bộ Quy Tắc Phòng Thủ Toàn Diện (Defense in Depth)"
p.font.size = Pt(15)
p.font.bold = True
p.font.color.rgb = ACCENT_BLUE

p = tf6_2.add_paragraph()
p.text = "1. Luôn sử dụng Parameterized Query / ORM:\nTuyệt đối không ghép chuỗi SQL thủ công dưới bất kỳ hình thức nào."
p.font.size = Pt(12)
p.font.color.rgb = TEXT_DARK
p.space_before = Pt(10)

p = tf6_2.add_paragraph()
p.text = "2. Nguyên tắc đặc quyền tối thiểu (Least Privilege):\nTài khoản kết nối CSDL của ứng dụng chỉ có quyền SELECT/INSERT/UPDATE trên các bảng cần thiết, không bao giờ dùng tài khoản 'sa' hay quyền DDL (DROP, ALTER)."
p.font.size = Pt(12)
p.font.color.rgb = TEXT_DARK
p.space_before = Pt(8)

p = tf6_2.add_paragraph()
p.text = "3. Xác thực dữ liệu đầu vào (Input Validation):\nSử dụng Data Annotations ([RegularExpression], [StringLength]) kiểm tra chặt chẽ khuôn dạng dữ liệu."
p.font.size = Pt(12)
p.font.color.rgb = TEXT_DARK
p.space_before = Pt(8)

p = tf6_2.add_paragraph()
p.text = "4. Mã hóa mật khẩu một chiều (Password Hashing):\nSử dụng ASP.NET Core Identity (PBKDF2/BCrypt) để nếu CSDL có bị lộ, mật khẩu vẫn được bảo vệ an toàn."
p.font.size = Pt(12)
p.font.color.rgb = TEXT_DARK
p.space_before = Pt(8)

output_file = "/media/hocjsoo/New Volume/OWASP_Demo/Slide_OWASP_SQL_Injection.pptx"
prs.save(output_file)
print(f"Slide saved successfully to {output_file}")
