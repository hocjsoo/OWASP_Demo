import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)

# Bảng màu Gamma / Canva Light Theme
BG_WHITE = RGBColor(255, 255, 255)
CARD_BG = RGBColor(248, 250, 252)          # Slate 50
CARD_BORDER = RGBColor(226, 232, 240)      # Slate 200

TEXT_HEAD = RGBColor(15, 23, 42)           # Slate 900
TEXT_BODY = RGBColor(51, 65, 85)           # Slate 700
TEXT_MUTED = RGBColor(100, 116, 139)       # Slate 500

ACCENT_BLUE = RGBColor(29, 78, 216)        # Blue 700 (HOU)
BLUE_LIGHT = RGBColor(239, 246, 255)       # Blue 50
BLUE_BORDER = RGBColor(191, 219, 254)

ACCENT_RED = RGBColor(190, 18, 60)         # Rose 700 (Danger)
RED_LIGHT = RGBColor(255, 241, 242)        # Rose 50
RED_BORDER = RGBColor(254, 205, 211)

ACCENT_GREEN = RGBColor(4, 120, 87)        # Emerald 700 (Secure)
GREEN_LIGHT = RGBColor(236, 253, 245)      # Emerald 50
GREEN_BORDER = RGBColor(167, 243, 208)

ACCENT_AMBER = RGBColor(180, 83, 9)        # Amber 700
AMBER_LIGHT = RGBColor(255, 251, 235)      # Amber 50
AMBER_BORDER = RGBColor(254, 240, 138)

blank_layout = prs.slide_layouts[6]

def add_header(slide, title_text, category_tag="OWASP TOP 10 • BÁO CÁO THỰC NGHIỆM"):
    top_line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(0.08))
    top_line.fill.solid()
    top_line.fill.fore_color.rgb = ACCENT_BLUE
    top_line.line.fill.background()

    badge = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(0.35), Inches(4.5), Inches(0.38))
    badge.fill.solid()
    badge.fill.fore_color.rgb = BLUE_LIGHT
    badge.line.color.rgb = BLUE_BORDER

    tf_b = badge.text_frame
    tf_b.word_wrap = False
    p_b = tf_b.paragraphs[0]
    p_b.text = category_tag.upper()
    p_b.font.size = Pt(10)
    p_b.font.bold = True
    p_b.font.color.rgb = ACCENT_BLUE

    tx = slide.shapes.add_textbox(Inches(0.8), Inches(0.75), Inches(11.7), Inches(0.6))
    tf = tx.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = title_text
    p.font.size = Pt(21)
    p.font.bold = True
    p.font.color.rgb = TEXT_HEAD

    footer = slide.shapes.add_textbox(Inches(0.8), Inches(7.1), Inches(11.7), Inches(0.3))
    ft = footer.text_frame
    p_ft = ft.paragraphs[0]
    p_ft.text = "Nhóm WNC.G01 (Học - Bình - Cường) • Khoa CNTT - Đại học Mở Hà Nội • GVHD: ThS. Lê Hữu Dũng"
    p_ft.font.size = Pt(9.5)
    p_ft.font.color.rgb = TEXT_MUTED

# ==============================================================================
# SLIDE 1: BÌA BÁO CÁO NHÓM
# ==============================================================================
s1 = prs.slides.add_slide(blank_layout)

bar1 = s1.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(0.12))
bar1.fill.solid()
bar1.fill.fore_color.rgb = ACCENT_BLUE
bar1.line.fill.background()

u_badge = s1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.0), Inches(0.75), Inches(5.8), Inches(0.42))
u_badge.fill.solid()
u_badge.fill.fore_color.rgb = BLUE_LIGHT
u_badge.line.color.rgb = BLUE_BORDER
tf = u_badge.text_frame
p = tf.paragraphs[0]
p.text = "TRƯỜNG ĐẠI HỌC MỞ HÀ NỘI — KHOA CÔNG NGHỆ THÔNG TIN"
p.font.size = Pt(10.5)
p.font.bold = True
p.font.color.rgb = ACCENT_BLUE

tb1 = s1.shapes.add_textbox(Inches(1.0), Inches(1.3), Inches(11.333), Inches(3.2))
tf1 = tb1.text_frame
tf1.word_wrap = True

p = tf1.paragraphs[0]
p.text = "BÁO CÁO THỰC NGHIỆM AN TOÀN WEB (OWASP TOP 10)"
p.font.size = Pt(16)
p.font.bold = True
p.font.color.rgb = ACCENT_RED

p = tf1.add_paragraph()
p.text = "Bộ 3 Lỗ Hổng Kinh Điển & Cơ Chế Phòng Thủ\n(SQL Injection • Stored XSS • CSRF Attack)"
p.font.size = Pt(30)
p.font.bold = True
p.font.color.rgb = TEXT_HEAD
p.space_before = Pt(8)

p = tf1.add_paragraph()
p.text = "Môn học: Lập trình Web nâng cao  •  Giảng viên hướng dẫn: ThS. Lê Hữu Dũng"
p.font.size = Pt(14)
p.font.color.rgb = TEXT_BODY
p.space_before = Pt(10)

# Card 3 Thành viên
col_w = Inches(3.64)
gap = Inches(0.2)

m1 = s1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.0), Inches(4.7), col_w, Inches(2.0))
m1.fill.solid()
m1.fill.fore_color.rgb = CARD_BG
m1.line.color.rgb = BLUE_BORDER
tf = m1.text_frame
tf.word_wrap = True
p = tf.paragraphs[0]
p.text = "CHỦ ĐỀ 1: SQL INJECTION"
p.font.size = Pt(11)
p.font.bold = True
p.font.color.rgb = ACCENT_BLUE
p = tf.add_paragraph()
p.text = "Nguyễn Danh Học\nMSV: 23A1001D0158\nTrưởng nhóm • Demo CSDL & API"
p.font.size = Pt(12)
p.font.color.rgb = TEXT_HEAD
p.space_before = Pt(4)

m2 = s1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.0) + col_w + gap, Inches(4.7), col_w, Inches(2.0))
m2.fill.solid()
m2.fill.fore_color.rgb = CARD_BG
m2.line.color.rgb = RED_BORDER
tf = m2.text_frame
tf.word_wrap = True
p = tf.paragraphs[0]
p.text = "CHỦ ĐỀ 2: STORED XSS"
p.font.size = Pt(11)
p.font.bold = True
p.font.color.rgb = ACCENT_RED
p = tf.add_paragraph()
p.text = "Nguyễn Thanh Bình\nMSV: 23A1001D0041\nThành viên • Demo Cookie & DOM"
p.font.size = Pt(12)
p.font.color.rgb = TEXT_HEAD
p.space_before = Pt(4)

m3 = s1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.0) + (col_w + gap)*2, Inches(4.7), col_w, Inches(2.0))
m3.fill.solid()
m3.fill.fore_color.rgb = CARD_BG
m3.line.color.rgb = GREEN_BORDER
tf = m3.text_frame
tf.word_wrap = True
p = tf.paragraphs[0]
p.text = "CHỦ ĐỀ 3: CSRF ATTACK"
p.font.size = Pt(11)
p.font.bold = True
p.font.color.rgb = ACCENT_GREEN
p = tf.add_paragraph()
p.text = "Nguyễn Minh Cường\nMSV: 23A1001D0058\nThành viên • Demo Form ẩn giả mạo"
p.font.size = Pt(12)
p.font.color.rgb = TEXT_HEAD
p.space_before = Pt(4)


# ==============================================================================
# PHẦN 1: SQL INJECTION (NGUYỄN DANH HỌC)
# ==============================================================================
s2 = prs.slides.add_slide(blank_layout)
add_header(s2, "Phần 1: SQL Injection (OWASP A03:2021) — Nguyễn Danh Học", "CHỦ ĐỀ 1 • NGUYỄN DANH HỌC (23A1001D0158)")

# 2 cards
card_w = Inches(5.65)
b1_v = s2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.45), card_w, Inches(4.2))
b1_v.fill.solid()
b1_v.fill.fore_color.rgb = RED_LIGHT
b1_v.line.color.rgb = RED_BORDER

tf = b1_v.text_frame
tf.word_wrap = True
p = tf.paragraphs[0]
p.text = "❌ BẢN CHẤT LỖ HỔNG & KHAI THÁC"
p.font.size = Pt(14)
p.font.bold = True
p.font.color.rgb = ACCENT_RED
p = tf.add_paragraph()
p.text = "• Nguyên nhân: Ghép chuỗi đầu vào trực tiếp vào câu lệnh SQL.\n  query = $\"SELECT * FROM Accounts WHERE User='{u}' AND Pass='{p}'\"\n• Payload tấn công: ' OR '1'='1' --\n• Cơ chế bẻ gãy:\n  - Dấu nháy đơn (') đóng chuỗi sớm.\n  - Mệnh đề OR '1'='1' biến điều kiện WHERE thành luôn TRUE.\n  - Ký tự '--' triệt tiêu toàn bộ bước kiểm tra mật khẩu phía sau.\n• Hậu quả: Bypass đăng nhập Admin không cần mật khẩu, rò rỉ CSDL."
p.font.size = Pt(12)
p.font.color.rgb = TEXT_BODY
p.space_before = Pt(8)

b1_s = s2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.85), Inches(1.45), card_w, Inches(4.2))
b1_s.fill.solid()
b1_s.fill.fore_color.rgb = GREEN_LIGHT
b1_s.line.color.rgb = GREEN_BORDER

tf = b1_s.text_frame
tf.word_wrap = True
p = tf.paragraphs[0]
p.text = "✅ CƠ CHẾ PHÒNG THỦ CHUẨN EF CORE"
p.font.size = Pt(14)
p.font.bold = True
p.font.color.rgb = ACCENT_GREEN
p = tf.add_paragraph()
p.text = "• Giải pháp: Sử dụng Parameterized Query / LINQ trong EF Core.\n  _context.Accounts.Where(a => a.Username == u && a.Password == p)\n• Cơ chế an toàn tuyệt đối:\n  - CSDL SQL Server biên dịch cấu trúc lệnh (Execution Plan) TRƯỚC.\n  - Giá trị payload được đưa vào qua tham số @p0, @p1 riêng biệt.\n  - Toàn bộ chuỗi payload bị cô lập thành String Literal thuần túy.\n• Kết quả: Trả về 0 dòng, chặn đứng tấn công 100%."
p.font.size = Pt(12)
p.font.color.rgb = TEXT_BODY
p.space_before = Pt(8)

callout1 = s2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(5.8), Inches(11.733), Inches(0.95))
callout1.fill.solid()
callout1.fill.fore_color.rgb = BLUE_LIGHT
callout1.line.color.rgb = BLUE_BORDER
tf = callout1.text_frame
p = tf.paragraphs[0]
p.text = "💻 [HỌC DEMO TRỰC TIẾP]: Chuyển sang Web UI Tab 1 & VS Code SQL Server 2025"
p.font.size = Pt(12.5)
p.font.bold = True
p.font.color.rgb = ACCENT_BLUE
p = tf.add_paragraph()
p.text = "Thao tác: Bấm nạp payload ' OR '1'='1' -- trên Web -> Show bảng dữ liệu rò rỉ -> Chuyển VS Code đối chiếu SQL."
p.font.size = Pt(11)
p.font.color.rgb = TEXT_BODY
p.space_before = Pt(2)


# ==============================================================================
# PHẦN 2: STORED XSS (NGUYỄN THANH BÌNH)
# ==============================================================================
s3 = prs.slides.add_slide(blank_layout)
add_header(s3, "Phần 2: Stored Cross-Site Scripting (XSS) — Nguyễn Thanh Bình", "CHỦ ĐỀ 2 • NGUYỄN THANH BÌNH (23A1001D0041)")

b2_v = s3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.45), card_w, Inches(4.2))
b2_v.fill.solid()
b2_v.fill.fore_color.rgb = RED_LIGHT
b2_v.line.color.rgb = RED_BORDER
tf = b2_v.text_frame
tf.word_wrap = True
p = tf.paragraphs[0]
p.text = "❌ BẢN CHẤT LỖ HỔNG & KHAI THÁC"
p.font.size = Pt(14)
p.font.bold = True
p.font.color.rgb = ACCENT_RED
p = tf.add_paragraph()
p.text = "• Ngữ cảnh: Chức năng nhận xét / đánh giá chất lượng chuyên gia.\n• Cơ chế lỗi: Máy chủ lưu nguyên văn chuỗi HTML chứa script của hacker vào CSDL, sau đó View dùng @Html.Raw() để render ra màn hình.\n• Payload tấn công:\n  <script>alert('Lộ Cookie: ' + document.cookie);</script>\n  <img src=x onerror=alert('XSS!')>\n• Hậu quả: Bất kỳ ai vào xem đánh giá đều bị trình duyệt tự động kích hoạt mã script, bị đánh cắp Cookie, Session Token và thông tin cá nhân."
p.font.size = Pt(12)
p.font.color.rgb = TEXT_BODY
p.space_before = Pt(8)

b2_s = s3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.85), Inches(1.45), card_w, Inches(4.2))
b2_s.fill.solid()
b2_s.fill.fore_color.rgb = GREEN_LIGHT
b2_s.line.color.rgb = GREEN_BORDER
tf = b2_s.text_frame
tf.word_wrap = True
p = tf.paragraphs[0]
p.text = "✅ CƠ CHẾ PHÒNG THỦ TRONG ASP.NET CORE"
p.font.size = Pt(14)
p.font.bold = True
p.font.color.rgb = ACCENT_GREEN
p = tf.add_paragraph()
p.text = "• Cơ chế phòng thủ 1 (Mặc định của Razor View):\n  Sử dụng @comment.Content (tự động HTML Encode).\n  Các ký tự nguy hiểm <, >, \", ' tự động biến thành &lt;, &gt;, &quot;.\n• Cơ chế phòng thủ 2 (Mã hóa ở tầng Controller):\n  Sử dụng HtmlEncoder.Default.Encode(content) trước khi lưu.\n• Cơ chế phòng thủ 3 (Bảo vệ Cookie):\n  Bật cờ HttpOnly = true cho toàn bộ Cookie nhạy cảm, ngăn cấm hoàn toàn JavaScript truy cập document.cookie."
p.font.size = Pt(12)
p.font.color.rgb = TEXT_BODY
p.space_before = Pt(8)

callout2 = s3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(5.8), Inches(11.733), Inches(0.95))
callout2.fill.solid()
callout2.fill.fore_color.rgb = RED_LIGHT
callout2.line.color.rgb = RED_BORDER
tf = callout2.text_frame
p = tf.paragraphs[0]
p.text = "💬 [BÌNH DEMO TRỰC TIẾP]: Chuyển sang Web UI Tab 2 (Stored XSS)"
p.font.size = Pt(12.5)
p.font.bold = True
p.font.color.rgb = ACCENT_RED
p = tf.add_paragraph()
p.text = "Thao tác: Bấm nạp payload đánh cắp Cookie -> Đăng bình luận bên Lỗi (bật popup alert lộ Session) -> Đăng bên An toàn (chữ hiển thị an toàn)."
p.font.size = Pt(11)
p.font.color.rgb = TEXT_BODY
p.space_before = Pt(2)


# ==============================================================================
# PHẦN 3: CSRF ATTACK (NGUYỄN MINH CƯỜNG)
# ==============================================================================
s4 = prs.slides.add_slide(blank_layout)
add_header(s4, "Phần 3: Cross-Site Request Forgery (CSRF) — Nguyễn Minh Cường", "CHỦ ĐỀ 3 • NGUYỄN MINH CƯỜNG (23A1001D0058)")

b3_v = s4.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.45), card_w, Inches(4.2))
b3_v.fill.solid()
b3_v.fill.fore_color.rgb = RED_LIGHT
b3_v.line.color.rgb = RED_BORDER
tf = b3_v.text_frame
tf.word_wrap = True
p = tf.paragraphs[0]
p.text = "❌ BẢN CHẤT LỖ HỔNG & KHAI THÁC"
p.font.size = Pt(14)
p.font.bold = True
p.font.color.rgb = ACCENT_RED
p = tf.add_paragraph()
p.text = "• Ngữ cảnh: Chức năng chuyển tiền ví hoặc đổi mật khẩu tài khoản.\n• Bản chất: Lừa nạn nhân (đang đăng nhập) gửi một HTTP POST ngầm mà nạn nhân không hề hay biết.\n• Kịch bản tấn công:\n  1. Nạn nhân đăng nhập vào web tư vấn trực tuyến (phiên hợp lệ).\n  2. Nạn nhân mở một trang web bẫy trúng thưởng do hacker dựng.\n  3. Trang web bẫy tự động kích hoạt form POST chuyển 20.000.000đ từ ví nạn nhân sang ví hacker.\n• Hậu quả: Mất tiền, mất tài khoản trong khi máy chủ vẫn tưởng là do nạn nhân tự gửi."
p.font.size = Pt(12)
p.font.color.rgb = TEXT_BODY
p.space_before = Pt(8)

b3_s = s4.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.85), Inches(1.45), card_w, Inches(4.2))
b3_s.fill.solid()
b3_s.fill.fore_color.rgb = GREEN_LIGHT
b3_s.line.color.rgb = GREEN_BORDER
tf = b3_s.text_frame
tf.word_wrap = True
p = tf.paragraphs[0]
p.text = "✅ CƠ CHẾ PHÒNG THỦ ANTI-FORGERY TOKEN"
p.font.size = Pt(14)
p.font.bold = True
p.font.color.rgb = ACCENT_GREEN
p = tf.add_paragraph()
p.text = "• Giải pháp bắt buộc trong ASP.NET Core:\n  1. Trong View: Nhúng @Html.AntiForgeryToken() sinh token bí mật.\n  2. Trong Controller: Đặt thuộc tính [ValidateAntiForgeryToken].\n• Cơ chế hoạt động của Token:\n  - Server phát hành một cặp token: 1 lưu trong Cookie, 1 nhúng vào Form ẩn.\n  - Trang web của hacker ở tên miền khác KHÔNG THỂ đọc trộm token trong form này (do chính sách Same-Origin Policy).\n  - Request giả mạo thiếu token bị máy chủ từ chối ngay (Lỗi 400 Bad Request)."
p.font.size = Pt(12)
p.font.color.rgb = TEXT_BODY
p.space_before = Pt(8)

callout3 = s4.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(5.8), Inches(11.733), Inches(0.95))
callout3.fill.solid()
callout3.fill.fore_color.rgb = GREEN_LIGHT
callout3.line.color.rgb = GREEN_BORDER
tf = callout3.text_frame
p = tf.paragraphs[0]
p.text = "💳 [CƯỜNG DEMO TRỰC TIẾP]: Chuyển sang Web UI Tab 3 (CSRF)"
p.font.size = Pt(12.5)
p.font.bold = True
p.font.color.rgb = ACCENT_GREEN
p = tf.add_paragraph()
p.text = "Thao tác: Mở trang web trúng thưởng của hacker -> Kích hoạt form ẩn -> Số dư ví nạn nhân bị trừ 20 triệu -> Thử lại với form có Token (chặn đứng)."
p.font.size = Pt(11)
p.font.color.rgb = TEXT_BODY
p.space_before = Pt(2)


# ==============================================================================
# SLIDE 5: TỔNG KẾT & CAM KẾT ĐỒ ÁN BTL
# ==============================================================================
s5 = prs.slides.add_slide(blank_layout)
add_header(s5, "Tổng Kết: Áp Dụng Chuẩn An Toàn Vào Đồ Án Website Tư Vấn Trực Tuyến")

main_c5 = s5.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.45), Inches(11.733), Inches(5.3))
main_c5.fill.solid()
main_c5.fill.fore_color.rgb = CARD_BG
main_c5.line.color.rgb = CARD_BORDER

tf = main_c5.text_frame
tf.word_wrap = True

p = tf.paragraphs[0]
p.text = "BẢNG TỔNG HỢP GIẢI PHÁP PHÒNG THỦ TRONG ĐỒ ÁN BTL (WNC.G01):"
p.font.size = Pt(16)
p.font.bold = True
p.font.color.rgb = ACCENT_BLUE

p = tf.add_paragraph()
p.text = "1. Phòng thủ SQL Injection (Nguyễn Danh Học):\n   Áp dụng 100% Entity Framework Core Code-First với Parameterized LINQ cho toàn bộ các thao tác đặt lịch, quản lý chuyên gia và đánh giá chất lượng."
p.font.size = Pt(13)
p.font.color.rgb = TEXT_BODY
p.space_before = Pt(10)

p = tf.add_paragraph()
p.text = "2. Phòng thủ Stored XSS (Nguyễn Thanh Bình):\n   Toàn bộ phản hồi tư vấn và đánh giá của khách hàng được Razor View tự động HTML Encode, kết hợp cấu hình Cookie an toàn (HttpOnly = true, SameSite = Strict)."
p.font.size = Pt(13)
p.font.color.rgb = TEXT_BODY
p.space_before = Pt(10)

p = tf.add_paragraph()
p.text = "3. Phòng thủ CSRF (Nguyễn Minh Cường):\n   Bắt buộc nhúng AntiForgeryToken và kiểm tra [ValidateAntiForgeryToken] trên 100% các form POST giao dịch (đặt lịch tư vấn, nạp tiền ví, đổi mật khẩu)."
p.font.size = Pt(13)
p.font.color.rgb = TEXT_BODY
p.space_before = Pt(10)

p = tf.add_paragraph()
p.text = "XIN TRÂN TRỌNG CẢM ƠN THẦY VÀ CÁC BẠN ĐÃ THEO DÕI!\nNhóm WNC.G01 sẵn sàng nhận câu hỏi phản biện từ giảng viên."
p.font.size = Pt(15)
p.font.bold = True
p.font.color.rgb = ACCENT_RED
p.space_before = Pt(18)

out_file = "/media/hocjsoo/New Volume/OWASP_Demo/Slide_OWASP_Top3_Nhom_WNC_G01.pptx"
prs.save(out_file)
print(f"Master 3-topics slide saved to: {out_file}")
