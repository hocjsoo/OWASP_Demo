import os
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

doc = docx.Document()

# Thiết lập khổ giấy A4
for section in doc.sections:
    section.page_width = Inches(8.27)
    section.page_height = Inches(11.69)
    section.top_margin = Inches(0.79)
    section.bottom_margin = Inches(0.79)
    section.left_margin = Inches(0.98)
    section.right_margin = Inches(0.79)

COLOR_PRIMARY = RGBColor(29, 78, 216)    # Blue 700 (HOU)
COLOR_DANGER = RGBColor(190, 18, 60)     # Rose 700
COLOR_SUCCESS = RGBColor(4, 120, 87)     # Emerald 700
COLOR_TEXT = RGBColor(15, 23, 42)        # Slate 900
COLOR_MUTED = RGBColor(100, 116, 139)    # Slate 500

def set_cell_shading(cell, color_hex):
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{color_hex}"/>')
    cell._tc.get_or_add_tcPr().append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'''
        <w:tcMar {nsdecls("w")}>
            <w:top w:w="{top}" w:type="dxa"/>
            <w:bottom w:w="{bottom}" w:type="dxa"/>
            <w:left w:w="{left}" w:type="dxa"/>
            <w:right w:w="{right}" w:type="dxa"/>
        </w:tcMar>
    ''')
    tcPr.append(tcMar)

def add_card(doc, title, text, border_color="1D4ED8", bg_color="EFF6FF"):
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = tbl.cell(0, 0)
    set_cell_shading(cell, bg_color)
    set_cell_margins(cell, top=120, bottom=120, left=180, right=180)
    
    tcPr = cell._tc.get_or_add_tcPr()
    tcBorders = parse_xml(f'''
        <w:tcBorders {nsdecls("w")}>
            <w:top w:val="none"/>
            <w:left w:val="single" w:sz="24" w:space="0" w:color="{border_color}"/>
            <w:bottom w:val="none"/>
            <w:right w:val="none"/>
        </w:tcBorders>
    ''')
    tcPr.append(tcBorders)
    
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.15
    run_t = p.add_run(f"{title}\n")
    run_t.font.name = "Arial"
    run_t.font.size = Pt(10)
    run_t.font.bold = True
    if border_color == "1D4ED8":
        run_t.font.color.rgb = COLOR_PRIMARY
    elif border_color == "BE123C":
        run_t.font.color.rgb = COLOR_DANGER
    elif border_color == "047857":
        run_t.font.color.rgb = COLOR_SUCCESS
    else:
        run_t.font.color.rgb = COLOR_TEXT
        
    run_b = p.add_run(text)
    run_b.font.name = "Arial"
    run_b.font.size = Pt(9.5)
    run_b.font.color.rgb = COLOR_TEXT

# ==============================================================================
# HEADER TRANG BÌA
# ==============================================================================
p_top = doc.add_paragraph()
p_top.alignment = WD_ALIGN_PARAGRAPH.CENTER
p_top.paragraph_format.space_after = Pt(2)
r_top1 = p_top.add_run("TRƯỜNG ĐẠI HỌC MỞ HÀ NỘI — KHOA CÔNG NGHỆ THÔNG TIN\n")
r_top1.font.name = "Arial"
r_top1.font.size = Pt(11)
r_top1.font.bold = True
r_top1.font.color.rgb = COLOR_PRIMARY

r_top2 = p_top.add_run("BỘ MÔN: LẬP TRÌNH WEB NÂNG CAO (ThS. LÊ HỮU DŨNG)")
r_top2.font.name = "Arial"
r_top2.font.size = Pt(10)
r_top2.font.bold = True
r_top2.font.color.rgb = COLOR_MUTED

p_main = doc.add_paragraph()
p_main.alignment = WD_ALIGN_PARAGRAPH.CENTER
p_main.paragraph_format.space_before = Pt(12)
p_main.paragraph_format.space_after = Pt(4)
r_m = p_main.add_run("KỊCH BẢN THUYẾT TRÌNH & HƯỚNG DẪN BẢO VỆ CHUYÊN ĐỀ OWASP\n(BẢN NÓI MIỆNG TỰ NHIÊN — CẤU TRÚC 3 BƯỚC CHUẨN 10/10)")
r_m.font.name = "Arial"
r_m.font.size = Pt(15)
r_m.font.bold = True
r_m.font.color.rgb = COLOR_DANGER

p_sub = doc.add_paragraph()
p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
p_sub.paragraph_format.space_after = Pt(12)
r_sub = p_sub.add_run("Đề tài BTL liên hệ: Website Cung Cấp Dịch Vụ Tư Vấn Trực Tuyến • Nhóm WNC.G01\n(Học: SQL Injection • Bình: Stored XSS • Cường: CSRF Attack)")
r_sub.font.name = "Arial"
r_sub.font.size = Pt(10)
r_sub.font.italic = True
r_sub.font.color.rgb = COLOR_TEXT

# ==============================================================================
# HÀM TẠO MỤC SLIDE 3 PHẦN
# ==============================================================================
def add_presentation_slide(doc, slide_no, title, speaker, speech_text, visual_action, qa_tip):
    h = doc.add_heading(level=2)
    h.paragraph_format.space_before = Pt(12)
    h.paragraph_format.space_after = Pt(4)
    r = h.add_run(f"SLIDE {slide_no:02d}: {title} (Phụ trách: {speaker})")
    r.font.name = "Arial"
    r.font.size = Pt(11)
    r.font.bold = True
    r.font.color.rgb = COLOR_PRIMARY

    # 1. Nói gì
    add_card(doc, "1. Lời nói tự nhiên (2–4 câu cốt lõi, không đọc slide):", speech_text, border_color="1D4ED8", bg_color="EFF6FF")
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    
    # 2. Chỉ vào đâu
    add_card(doc, "2. Thao tác trỏ màn hình / Demo khi chiếu:", visual_action, border_color="047857", bg_color="ECFDF5")
    doc.add_paragraph().paragraph_format.space_after = Pt(2)

    # 3. Phản xạ Q&A
    add_card(doc, "3. Phản xạ khi thầy hỏi xoáy:", qa_tip, border_color="B45309", bg_color="FFFBEB")
    doc.add_paragraph().paragraph_format.space_after = Pt(6)

# ==============================================================================
# DỮ LIỆU 19 SLIDE ĐƯỢC TINH CHẾ GỌN GÀNG, TỰ NHIÊN
# ==============================================================================
slides_data = [
    (1, "TRANG BÌA — TỔNG QUAN BÁO CÁO THỰC NGHIỆM", "Nguyễn Danh Học",
     "\"Kính chào ThS. Lê Hữu Dũng và các bạn! Nhóm WNC.G01 xin trình bày báo cáo thực nghiệm 3 lỗ hổng bảo mật: SQL Injection, Stored XSS và CSRF. Nhóm chúng em đặt cả 3 lỗ hổng này vào bài toán đồ án BTL 'Website tư vấn trực tuyến' tương ứng 3 bề mặt: Em - Nguyễn Danh Học phụ trách CSDL; bạn Nguyễn Thanh Bình phụ trách Client DOM & Cookie; và bạn Nguyễn Minh Cường phụ trách Giao dịch & Token.\"",
     "Chỉ tay vào khung bên phải slide 'Mô hình: 3 bề mặt nguy cơ ➔ 3 cơ chế phòng thủ' để người nghe nắm tổng thể.",
     "Nếu thầy hỏi: 'Tại sao nhóm chọn 3 đề tài này?' ➔ Trả lời: 'Thưa thầy, đây là 3 lỗ hổng kinh điển đại diện cho 3 tầng trọng yếu: CSDL, Trình duyệt và Máy chủ ứng dụng trong đồ án BTL của nhóm ạ.'"),

    (2, "MÔI TRƯỜNG CÔNG NGHỆ & PHẠM VI THỰC NGHIỆM", "Nguyễn Danh Học",
     "\"Về môi trường lab, nhóm xây dựng ứng dụng mẫu chạy trên máy chủ nội bộ localhost:5076 bằng ASP.NET Core .NET 10, kết nối trực tiếp Microsoft SQL Server 2025 Developer. Toàn bộ thực nghiệm diễn ra trong môi trường cô lập, sử dụng dữ liệu giả lập có kiểm soát và tuyệt đối không quét hay can thiệp hệ thống bên ngoài.\"",
     "Chỉ lướt 2 cột: Cột trái Tech Stack (.NET 10 / SQL Server 2025), Cột phải Lab Scope (Localhost / Dữ liệu giả định).",
     "Nếu thầy hỏi: 'Dùng công cụ gì để test?' ➔ Trả lời: 'Nhóm dùng trực tiếp Chrome DevTools để soi gói tin mạng, cURL CLI để test API Backend, và VS Code MSSQL Extension để kiểm tra CSDL ạ.'"),

    (3, "1.1. SQL INJECTION — LỖI XẢY RA Ở ĐÂU & ROOT CAUSE", "Nguyễn Danh Học",
     "\"Với SQL Injection, trong bài toán BTL lỗi xuất hiện ở chức năng Đăng nhập cho cả 3 vai trò. Nguyên nhân cốt lõi do hệ thống ghép trực tiếp dữ liệu người dùng vào câu SQL bằng phép cộng chuỗi. Khi đó, dữ liệu đầu vào làm thay đổi cú pháp câu lệnh, khiến SQL Server không còn phân biệt được ranh giới giữa Dữ liệu và Mã lệnh.\"",
     "Chỉ vào sơ đồ luồng (Browser ➔ Controller ➔ SQL Server) và chỉ vào dòng code C# `FromSqlRaw` dính lỗi ghép chuỗi.",
     "Nếu thầy hỏi: 'Ký tự nháy đơn đóng vai trò gì?' ➔ Trả lời: 'Ký tự nháy đơn đóng chuỗi giá trị Username sớm, còn dấu hai gạch ngang (--) biến vế kiểm tra mật khẩu phía sau thành chú thích, giúp bypass xác thực ạ.'"),

    (4, "1.2. SQL INJECTION — 3 KỊCH BẢN KHAI THÁC THỰC NGHIỆM", "Nguyễn Danh Học",
     "\"Nhóm xây dựng 3 kịch bản kiểm thử: Thứ nhất là Bypass tổng quát với chuỗi nháy đơn OR 1=1 gạch gạch, lợi dụng mệnh đề luôn đúng để đăng nhập vào tài khoản đầu tiên là Admin. Thứ hai là Tấn công có chủ đích với chuỗi admin' --, nhắm thẳng quyền Quản trị viên tối cao. Thứ ba là kịch bản đăng nhập chuẩn để đối chứng hệ thống hoạt động đúng nghiệp vụ.\"",
     "Chỉ lần lượt qua 3 thẻ card ngang 01, 02, 03 trên slide.",
     "Nếu thầy hỏi: 'Tại sao lại là tài khoản Admin?' ➔ Trả lời: 'Do câu lệnh SELECT không có ORDER BY, máy chủ sẽ trả về bản ghi đầu tiên được quét trúng trong bảng Accounts, đó chính là tài khoản Admin có Id=1 ạ.'"),

    (5, "1.3. SQL INJECTION — MINH CHỨNG WEB UI", "Nguyễn Danh Học",
     "\"Đây là kết quả thực tế trên giao diện: Ở nhánh dính lỗi bên trái, khi nạp chuỗi bypass, hệ thống xác thực thành công và trích xuất danh sách 4 tài khoản CSDL gồm cả Admin, Chuyên gia và Khách hàng kèm số dư ví. Ngược lại ở nhánh an toàn bên phải dùng EF Core LINQ, chuỗi payload được coi là văn bản thuần túy, truy vấn trả về 0 kết quả và từ chối đăng nhập.\"",
     "Chỉ vào bảng 4 tài khoản bị lộ ở cột đỏ (trái) và đối chiếu với nhãn xanh an toàn (phải).",
     "Nếu thầy hỏi: 'Mật khẩu hiển thị trên kia là gì?' ➔ Trả lời: 'Trong lab nhóm để plain-text để thấy rõ hậu quả rò rỉ; còn ở sản phẩm thực tế bắt buộc phải băm bằng ASP.NET Core Identity ạ.'"),

    (6, "1.4. SQL INJECTION — MINH CHỨNG CHROME DEVTOOLS NETWORK", "Nguyễn Danh Học",
     "\"Nhìn sâu vào gói tin mạng qua Chrome DevTools: Request POST gửi lên mang nguyên văn chuỗi payload khai thác. Phản hồi JSON trả về mã 200 OK chứa mảng 4 tài khoản bị rò rỉ. Điều này chứng minh lỗi nằm ở tầng C# Backend, kẻ tấn công hoàn toàn có thể dùng cURL hoặc Postman để trích xuất dữ liệu mà không cần thông qua giao diện Web.\"",
     "Chỉ vào tab Network: Dòng request `LoginVulnerable 200 OK`, tab Payload chứa chuỗi inject và tab Response chứa JSON.",
     "Nếu thầy hỏi: 'DevTools chứng minh được điều gì?' ➔ Trả lời: 'Chứng minh lỗ hổng xảy ra ở tầng nhận và xử lý request của Controller, độc lập hoàn toàn với việc kiểm tra form ở Client-side ạ.'"),

    (7, "1.5. SQL INJECTION — CƠ CHẾ PHÒNG THỦ CHUẨN TRONG BTL", "Nguyễn Danh Học",
     "\"Để phòng thủ, nhóm dùng LINQ Parameterized Query trong EF Core: FirstOrDefault theo user và pass. SQL Server sẽ biên dịch Execution Plan trước khi gán tham số @p0, đảm bảo Data tách biệt hoàn toàn với cú pháp lệnh. Nhóm xin nhấn mạnh: Parameterized Query giúp giải quyết lỗi SQLi, còn Password Hashing bằng Identity là lớp phòng thủ bổ sung để bảo vệ mật khẩu nếu CSDL bị rò rỉ.\"",
     "Chỉ vào code LINQ màu xanh bên trái và 3 nguyên tắc bảo mật (Parameterized, Least Privilege, Hashing) bên phải.",
     "Nếu thầy hỏi: 'Dùng FromSqlRaw có an toàn không?' ➔ Trả lời: 'Nếu dùng FromSqlRaw thì bắt buộc phải truyền kèm SqlParameter, hoặc dùng FromSqlInterpolated để EF Core tự tham số hóa chuỗi ạ.'"),

    (8, "2.1. STORED XSS — LỖI PHÁT TÁN NHƯ THẾ NÀO & ROOT CAUSE", "Nguyễn Thanh Bình",
     "\"Em là Nguyễn Thanh Bình, phụ trách Stored XSS. Trong BTL, lỗi này xuất hiện ở chức năng đánh giá, nhận xét chuyên gia. Kẻ tấn công gửi mã độc JavaScript, máy chủ lưu nguyên văn vào CSDL. Khi người dùng khác vào xem nhận xét, trình duyệt tự động tải và thực thi đoạn script đó. Nguyên nhân cốt lõi do tầng View lạm dụng cú pháp @Html.Raw(), vô tình vô hiệu hóa cơ chế mã hóa HTML tự động của Razor.\"",
     "Chỉ sơ đồ luồng 4 bước (Attacker ➔ Web Server ➔ CSDL ➔ Nạn nhân) và dòng code lỗi `@Html.Raw()`.",
     "Nếu thầy hỏi: 'Tại sao lập trình viên lại dùng Html.Raw?' ➔ Trả lời: 'Lập trình viên muốn hỗ trợ người dùng gõ chữ in đậm, in nghiêng hoặc xuống dòng, nhưng lại không lọc thẻ nguy hiểm trước khi hiển thị ạ.'"),

    (9, "2.2. STORED XSS — 3 KỊCH BẢN TIÊM MÃ ĐỘC THỰC NGHIỆM", "Nguyễn Thanh Bình",
     "\"Nhóm kiểm thử 3 kịch bản: Thứ nhất là Đánh cắp Cookie bằng đoạn script alert document.cookie để trích xuất token phiên làm việc. Thứ hai là Vượt qua bộ lọc từ khóa bằng thẻ ảnh img với thuộc tính onerror. Và thứ ba là đánh giá hợp lệ để đối chứng hệ thống hiển thị bình thường khi người dùng không chèn mã độc.\"",
     "Chỉ qua 3 thẻ card 01, 02, 03 trên slide.",
     "Nếu thầy hỏi: 'Filter Bypass qua thẻ ảnh hoạt động ra sao?' ➔ Trả lời: 'Khi đường dẫn ảnh bị sai (src=x), trình duyệt kích hoạt sự kiện onerror để chạy mã JavaScript bên trong, qua mặt các bộ lọc chỉ chặn từ khóa <script> ạ.'"),

    (10, "2.3. STORED XSS — MINH CHỨNG WEB UI", "Nguyễn Thanh Bình",
     "\"Trên giao diện thực nghiệm: Ở nhánh dính lỗi bên trái, khi tải trang, script lưu trong CSDL lập tức kích hoạt hộp thoại Alert đọc trộm Cookie phiên. Trong khi ở nhánh an toàn bên phải, Razor View tự động chuyển đổi các ký tự nhạy cảm thành thực thể an toàn, đoạn mã chỉ hiển thị như văn bản thuần túy và không thể thực thi.\"",
     "Chỉ vào popup Alert ở ảnh trái và dòng chữ an toàn không bị kích hoạt ở ảnh phải.",
     "Nếu thầy hỏi: 'Tại sao bên phải không bị chạy script?' ➔ Trả lời: 'Vì Razor View mặc định tự động encode ký tự < thành &lt; và > thành &gt;, trình duyệt hiểu đó là văn bản hiển thị chứ không phải thẻ lệnh ạ.'"),

    (11, "2.4. STORED XSS — MINH CHỨNG CHROME DEVTOOLS COOKIE & CONSOLE", "Nguyễn Thanh Bình",
     "\"Soi chi tiết qua Chrome DevTools: Trong tab Application mục Cookies, nhóm cố ý cấu hình cờ HttpOnly bằng false trong lab để minh họa nguy cơ. Khi sang tab Console gõ document.cookie, toàn bộ token phiên làm việc bị lộ ra. Đây chính là cách kẻ tấn công có thể chiếm đoạt phiên đăng nhập nếu Cookie không được bảo vệ đúng cách.\"",
     "Chỉ vào dòng HttpOnly = false trong bảng Cookies và lệnh `document.cookie` in ra token nhạy cảm trong Console.",
     "Nếu thầy hỏi: 'Tại sao Cookie lại đọc được bằng document.cookie?' ➔ Trả lời: 'Dạ thưa thầy, đây là cấu hình mô phỏng trong phòng lab. Khi triển khai thực tế của BTL, Cookie xác thực bắt buộc phải bật HttpOnly=true để cấm hoàn toàn JavaScript truy cập ạ.'"),

    (12, "2.5. STORED XSS — CƠ CHẾ PHÒNG THỦ CHUẨN TRONG BTL", "Nguyễn Thanh Bình",
     "\"Giải pháp phòng thủ của nhóm gồm 2 lớp: Ở tầng View, luôn dùng cú pháp mặc định @comment.Content hoặc hàm HtmlEncoder để mã hóa ngữ cảnh. Ở tầng quản lý Cookie trong sản phẩm thực tế, bắt buộc cấu hình HttpOnly bằng true, Secure bằng true và SameSite bằng Strict. Đồng thời nhóm thiết lập thêm Header CSP ngăn chặn chạy inline script lạ.\"",
     "Chỉ vào cú pháp Razor chuẩn bên trái và đoạn cấu hình CookieOptions bên phải.",
     "Nếu thầy hỏi: 'Chỉ dùng HttpOnly đã đủ chống XSS chưa?' ➔ Trả lời: 'HttpOnly chỉ ngăn chặn đánh cắp cookie, còn kẻ tấn công vẫn có thể deface giao diện hoặc keylogging. Vì vậy bắt buộc phải kết hợp cả Output Encoding ở View và CSP Header ạ.'"),

    (13, "3.1. CSRF ATTACK — CƠ CHẾ MƯỢN QUYỀN & PHÂN LOẠI OWASP", "Nguyễn Minh Cường",
     "\"Em là Nguyễn Minh Cường, phụ trách CSRF. Nhóm xin lưu ý: CSRF không phải là mục độc lập trong OWASP Top 10 2021, mà nhóm liên hệ nó với A01 – Broken Access Control để phân tích rủi ro mượn quyền người dùng. Nạn nhân đã đăng nhập hợp lệ nhưng bị lừa truy cập một trang web độc hại ở tab bên cạnh. Trang web độc hại gửi request POST ngầm, và trình duyệt tự động đính kèm Cookie của nạn nhân mà máy chủ không hề kiểm tra nguồn gốc phát sinh request.\"",
     "Chỉ vào sơ đồ luồng 4 bước và dòng code Action thiếu token kiểm soát.",
     "Nếu thầy hỏi: 'Tại sao trình duyệt lại tự động gửi Cookie sang?' ➔ Trả lời: 'Theo cơ chế mặc định của trình duyệt, mọi request gửi đến tên miền mục tiêu đều tự động đính kèm cookie gắn liền với tên miền đó, trừ khi được kiểm soát bởi thuộc tính SameSite ạ.'"),

    (14, "3.2. CSRF ATTACK — KỊCH BẢN KHAI THÁC QUA BIỂU MẪU ẨN", "Nguyễn Minh Cường",
     "\"Kịch bản thực tế: Nạn nhân đang mở web tư vấn và bị lừa bấm vào trang nhận thưởng ở tab khác. Trang bẫy này chứa một biểu mẫu ẩn trỏ thẳng đến địa chỉ chuyển tiền của hệ thống với số tiền 20.000.000 VNĐ. Khi form được submit ngầm, vì Action thiếu cơ chế kiểm tra token, máy chủ ngộ nhận đây là lệnh của nạn nhân và thực hiện trừ tiền ngay lập tức.\"",
     "Chỉ vào khối code HTML form ẩn và lệnh submit form tự động.",
     "Nếu thầy hỏi: 'Hacker có cần biết mật khẩu của nạn nhân không?' ➔ Trả lời: 'Hoàn toàn không cần ạ. Kẻ tấn công chỉ lợi dụng phiên đăng nhập và Cookie đang còn hiệu lực trên trình duyệt của nạn nhân để gửi request mượn quyền.'"),

    (15, "3.3. CSRF ATTACK — MINH CHỨNG WEB UI", "Nguyễn Minh Cường",
     "\"Minh chứng trực quan trên giao diện: Trước khi kích hoạt, ví nạn nhân có 50 triệu và ví hacker có 0 đồng. Sau khi kích hoạt form ẩn từ trang bẫy, số dư ví nạn nhân bị trừ còn 30 triệu và ví hacker tăng lên 20 triệu. Ngược lại ở nhánh an toàn, form chính chủ có nhúng token bảo vệ, request giả mạo không có token hợp lệ sẽ bị máy chủ từ chối ngay lập tức.\"",
     "Chỉ vào biến động số dư 50M ➔ 30M trên ảnh và thông báo từ chối ở nhánh có token.",
     "Nếu thầy hỏi: 'Máy chủ từ chối thì trả về lỗi gì?' ➔ Trả lời: 'Dạ trả về lỗi HTTP 400 Bad Request kèm thông báo thiếu hoặc sai mã Antiforgery Token ạ.'"),

    (16, "3.4. CSRF ATTACK — MINH CHỨNG CHROME DEVTOOLS HEADERS", "Nguyễn Minh Cường",
     "\"Kiểm tra gói tin mạng qua DevTools: Header Origin xuất phát từ trang bẫy của hacker, nhưng Cookie phiên của nạn nhân vẫn được đính kèm tự động. Điểm yếu mấu chốt được khoanh đỏ là gói tin hoàn toàn không có trường RequestVerificationToken. Vì Action thiếu kiểm tra, hệ thống đã thực hiện giao dịch trái phép mà người dùng không hề hay biết.\"",
     "Chỉ vào phần Headers: Origin attacker, Cookie nạn nhân và khung cảnh báo đỏ 'Anti-Forgery Token: ABSENT'.",
     "Nếu thầy hỏi: 'Header Origin do ai sinh ra?' ➔ Trả lời: 'Do trình duyệt tự động gán vào các request POST cross-origin để chỉ rõ nguồn gốc trang web phát sinh yêu cầu, lập trình viên phía client không thể giả mạo header này ạ.'"),

    (17, "3.5. CSRF ATTACK — CƠ CHẾ PHÒNG THỦ ANTI-FORGERY TOKEN", "Nguyễn Minh Cường",
     "\"Giải pháp phòng thủ chuẩn là Synchronizer Token Pattern: Nhúng @Html.AntiForgeryToken() vào Form và gắn [ValidateAntiForgeryToken] trên Controller. Nhờ chính sách Same-Origin Policy, trang web ngoài không thể đọc trộm mã token bí mật này. Đặc biệt bám sát Buổi 9, với các thao tác AJAX Fetch API đặt lịch tư vấn, nhóm đọc token từ thẻ ẩn và truyền qua HTTP Header RequestVerificationToken để bảo vệ giao dịch không reload trang.\"",
     "Chỉ rõ 2 phần: Bên trái là cặp Token trong Razor Form & Controller; Bên phải là code Fetch API đính kèm qua Header.",
     "Nếu thầy hỏi: 'Tại sao hacker không đọc trộm được Token trong form?' ➔ Trả lời: 'Nhờ chính sách Same-Origin Policy (SOP), trình duyệt ngăn cấm JavaScript ở tên miền attacker-site.com đọc nội dung DOM của trang web tư vấn localhost:5076 ạ.'"),

    (18, "TỔNG HỢP: BÀI HỌC RÚT RA TỪ 3 LỖ HỔNG", "Nguyễn Danh Học",
     "\"Tổng kết lại toàn bộ bài thực nghiệm qua 3 thẻ card cốt lõi: 1. SQL Injection do ghép chuỗi SQL trực tiếp, phòng thủ bằng Parameterized Query với EF Core LINQ. 2. Stored XSS do render dữ liệu thô qua @Html.Raw(), phòng thủ bằng Context-Aware HTML Encoding kết hợp HttpOnly Cookie. 3. CSRF do thiếu xác thực nguồn gốc request, phòng thủ bằng Synchronizer Token Pattern cho cả Form và AJAX Fetch API.\"",
     "Chỉ tay lần lượt qua 3 Card lớn: SQLi (Xanh) ➔ Stored XSS (Đỏ) ➔ CSRF (Xanh lá).",
     "Nếu thầy hỏi: 'Trong 3 lỗi này lỗi nào nguy hiểm nhất?' ➔ Trả lời: 'Cả 3 đều nguy hiểm tùy ngữ cảnh, nhưng SQL Injection có điểm CVSS cao nhất (9.8 Critical) vì nó phá vỡ tầng CSDL và có thể chiếm toàn quyền hệ thống máy chủ ạ.'"),

    (19, "KẾT LUẬN & CAM KẾT CHUẨN BẢO MẬT BTL", "Nguyễn Danh Học",
     "\"Qua thực nghiệm, nhóm WNC.G01 đã tái lập thành công lỗ hổng trên cả 3 bề mặt Web UI, DevTools và Database, đồng thời kiểm chứng hiệu quả phòng thủ triệt để. Nhóm đúc kết 2 nguyên tắc an ninh cốt lõi: Không bao giờ tin tưởng dữ liệu người dùng, và luôn phân tách dữ liệu khỏi cú pháp lệnh. Toàn bộ giải pháp này đã được tích hợp thành Security View trong đồ án BTL Website tư vấn trực tuyến. Nhóm xin trân trọng cảm ơn thầy và sẵn sàng bước vào phần Q&A ạ!\"",
     "Cả 3 thành viên đứng nghiêm túc, hướng về phía giảng viên và sẵn sàng chuyển sang phần trả lời câu hỏi.",
     "Sẵn sàng cho phần hỏi đáp của thầy.")
]

for s in slides_data:
    s_no, s_title, s_spk, s_txt, s_act, s_qa = s
    add_presentation_slide(doc, s_no, s_title, s_spk, s_txt, s_act, s_qa)

# ==============================================================================
# PHỤ LỤC: 10 CÂU HỎI VẤN ĐÁP PHẢN BIỆN CHUẨN XÁC
# ==============================================================================
doc.add_page_break()

h_qa = doc.add_heading(level=1)
h_qa.paragraph_format.space_before = Pt(14)
h_qa.paragraph_format.space_after = Pt(8)
r_qa = h_qa.add_run("PHỤ LỤC: 10 CÂU HỎI VẤN ĐÁP PHẢN BIỆN DỰ KIẾN CỦA GIẢNG VIÊN & ĐÁP ÁN MẪU")
r_qa.font.name = "Arial"
r_qa.font.size = Pt(13)
r_qa.font.bold = True
r_qa.font.color.rgb = COLOR_PRIMARY

qa_items = [
    ("Câu 1 (Học): Tại sao Entity Framework Core LINQ lại chống được SQL Injection?",
     "Thưa thầy, khi sử dụng LINQ (như .Where() hay .FirstOrDefault()), EF Core không ghép chuỗi mà tự động dịch thành câu lệnh SQL có tham số hóa (sp_executesql @p0). SQL Server sẽ phân tích cú pháp và lập Execution Plan trước, sau đó mới truyền giá trị người dùng vào tham số @p0 dưới dạng literal (văn bản thuần). Do đó mọi ký tự điều khiển đều không thể can thiệp vào logic truy vấn."),

    ("Câu 2 (Học): Nếu bắt buộc phải dùng câu lệnh SQL thuần (Raw SQL) trong EF Core thì làm sao để an toàn?",
     "Thưa thầy, chúng ta tuyệt đối không dùng phép nối chuỗi ($) trong FromSqlRaw. Thay vào đó, EF Core cung cấp hàm FromSqlInterpolated() hoặc FromSqlRaw() kèm mảng tham số SqlParameter. Với FromSqlInterpolated, EF Core tự động nhận diện các biến truyền vào trong chuỗi và bọc chúng thành tham số @p0, giúp truy vấn sử dụng cơ chế tham số hóa và tránh việc ghép dữ liệu trực tiếp vào SQL."),

    ("Câu 3 (Học): Password Hashing có ngăn chặn được SQL Injection không?",
     "Thưa thầy, hoàn toàn không ạ. Password Hashing là cơ chế bảo vệ mật khẩu (Defense-in-Depth) phòng trường hợp CSDL bị đánh cắp hoặc trích xuất ra ngoài thì kẻ tấn công không đọc được mật khẩu gốc. Còn lỗi SQL Injection xảy ra ở mệnh đề logic WHERE, kẻ tấn công bypass trực tiếp mà không cần khớp mật khẩu, nên để chặn SQLi bắt buộc phải dùng Parameterized Query."),

    ("Câu 4 (Bình): Phân biệt Stored XSS với Reflected XSS và DOM-based XSS?",
     "Thưa thầy: Reflected XSS là mã độc nằm trong URL và chỉ chạy khi nạn nhân bấm vào link đó; DOM-based XSS là lỗi hoàn toàn ở mã JavaScript phía client tự đọc và ghi dữ liệu không an toàn vào DOM; còn Stored XSS là mã độc được lưu trữ trực tiếp trong CSDL của máy chủ, bất kỳ ai truy cập trang hiển thị dữ liệu đó đều tự động bị tấn công mà không cần bấm link lạ."),

    ("Câu 5 (Bình): Nếu viết hàm kiểm tra, cứ thấy chữ '<script>' là xóa đi thì có chặn được XSS không?",
     "Thưa thầy là không an toàn ạ. Đây là cách tiếp cận Blacklist (danh sách đen) rất dễ bị vượt qua. Kẻ tấn công có thể chèn script qua các thẻ HTML khác có hỗ trợ Event Handler như <img src=x onerror=alert(1)>, <svg onload=...>, hoặc dùng kỹ thuật lồng thẻ. Cách phòng thủ chuẩn nhất là Whitelist Sanitization hoặc Context-Aware HTML Encoding ở đầu ra."),

    ("Câu 6 (Bình): Cờ HttpOnly trong Cookie hoạt động thế nào và tại sao nó quan trọng?",
     "Thưa thầy, cờ HttpOnly là một chỉ thị yêu cầu trình duyệt cấm tuyệt đối mọi đoạn mã JavaScript (như document.cookie) truy cập vào cookie đó. Cookie chỉ được đính kèm tự động trong các HTTP Request gửi về máy chủ. Kể cả khi trang web xuất hiện lỗi XSS, kẻ tấn công cũng không thể đọc trộm được Session Cookie để chiếm đoạt phiên làm việc."),

    ("Câu 7 (Cường): CSRF khác gì so với XSS?",
     "Thưa thầy: XSS là kẻ tấn công tiêm mã kịch bản độc hại để chạy TRỰC TIẾP trên trình duyệt của nạn nhân tại trang web mục tiêu; còn CSRF là kẻ tấn công tạo một trang web KHÁC (Cross-Site) để lừa trình duyệt của nạn nhân gửi request mượn quyền (kèm Cookie hợp lệ) sang trang web mục tiêu mà không cần chạy mã trên trang web mục tiêu."),

    ("Câu 8 (Cường): Tại sao trang web của hacker lại không đọc được Anti-Forgery Token của nạn nhân?",
     "Thưa thầy, đó là nhờ Chính sách Cùng nguồn gốc (Same-Origin Policy - SOP) của trình duyệt. Trình duyệt ngăn cấm mã JavaScript ở tên miền A (như attacker-site.com) đọc nội dung DOM hoặc tài nguyên của tên miền B (như website tư vấn localhost:5076). Hacker chỉ có thể kích hoạt gửi request POST sang chứ không thể đọc trộm mã token ẩn trong form của nạn nhân."),

    ("Câu 9 (Cường): Thuộc tính SameSite của Cookie có thay thế hoàn toàn được Anti-Forgery Token không?",
     "Thưa thầy, SameSite=Strict hoặc Lax là một lớp phòng thủ rất tốt ở tầng trình duyệt để hạn chế gửi cookie cross-site. Tuy nhiên, nó không thể thay thế hoàn toàn Token vì vẫn có rủi ro từ các trình duyệt cũ không hỗ trợ SameSite, hoặc các request dạng top-level navigation ở chế độ Lax. Do đó, việc kết hợp cả hai lớp bảo vệ mang lại độ an toàn cao nhất."),

    ("Câu 10 (Cả nhóm): Nhóm đã đưa các giải pháp an ninh này vào Đồ án BTL như thế nào?",
     "Thưa thầy, toàn bộ các giải pháp này đã được tích hợp thành Security View trong đồ án: Toàn bộ truy vấn CSDL trong BTL dùng 100% LINQ qua EF Core Code-First; Form đăng nhập và đánh giá chuyên gia dùng Razor HTML Encoding mặc định và mã hóa Identity; còn toàn bộ Form chuyển tiền ví và gọi AJAX Fetch API đặt lịch hẹn đều bắt buộc xác thực qua Anti-Forgery Token.")
]

for q, a in qa_items:
    p_q = doc.add_paragraph()
    p_q.paragraph_format.space_before = Pt(6)
    p_q.paragraph_format.space_after = Pt(2)
    r_q = p_q.add_run(f"❓ {q}")
    r_q.font.name = "Arial"
    r_q.font.size = Pt(10)
    r_q.font.bold = True
    r_q.font.color.rgb = COLOR_PRIMARY
    
    p_a = doc.add_paragraph()
    p_a.paragraph_format.space_before = Pt(0)
    p_a.paragraph_format.space_after = Pt(4)
    r_a = p_a.add_run(f"➔ Trả lời: {a}")
    r_a.font.name = "Arial"
    r_a.font.size = Pt(9.5)
    r_a.font.color.rgb = COLOR_TEXT

out_path = "/media/hocjsoo/New Volume/OWASP_Demo/Kich_Ban_Thuyet_Trinh_OWASP_WNC_G01.docx"
doc.save(out_path)
print(f"Updated natural speech Word document at: {out_path}")
