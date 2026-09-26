import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)

# Bảng màu Academic Light Theme chuẩn giảng đường
BG_WHITE = RGBColor(255, 255, 255)
WHITE = RGBColor(255, 255, 255)
CARD_BG = RGBColor(248, 250, 252)          # Slate 50
CARD_BORDER = RGBColor(226, 232, 240)      # Slate 200

TEXT_HEAD = RGBColor(15, 23, 42)           # Slate 900
TEXT_BODY = RGBColor(51, 65, 85)           # Slate 700
TEXT_MUTED = RGBColor(100, 116, 139)       # Slate 500

ACCENT_BLUE = RGBColor(29, 78, 216)        # Blue 700 (HOU Chuẩn)
BLUE_LIGHT = RGBColor(239, 246, 255)       # Blue 50
BLUE_BORDER = RGBColor(191, 219, 254)

ACCENT_RED = RGBColor(190, 18, 60)         # Rose 700 (Danger)
RED_LIGHT = RGBColor(255, 241, 242)        # Rose 50
RED_BORDER = RGBColor(254, 205, 211)

ACCENT_GREEN = RGBColor(4, 120, 87)        # Emerald 700 (Secure)
GREEN_LIGHT = RGBColor(236, 253, 245)      # Emerald 50
GREEN_BORDER = RGBColor(167, 243, 208)

ACCENT_AMBER = RGBColor(180, 83, 9)        # Amber 700
AMBER_LIGHT = RGBColor(255, 251, 235)      # Amber 50
AMBER_BORDER = RGBColor(254, 240, 138)

blank_layout = prs.slide_layouts[6]

def add_header(slide, title_text, category_tag="OWASP TOP 10 • THỰC NGHIỆM AN TOÀN WEB"):
    top_line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(0.08))
    top_line.fill.solid()
    top_line.fill.fore_color.rgb = ACCENT_BLUE
    top_line.line.fill.background()

    badge = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(0.28), Inches(5.2), Inches(0.36))
    badge.fill.solid()
    badge.fill.fore_color.rgb = BLUE_LIGHT
    badge.line.color.rgb = BLUE_BORDER

    tf_b = badge.text_frame
    tf_b.word_wrap = False
    p_b = tf_b.paragraphs[0]
    p_b.text = category_tag.upper()
    p_b.font.name = "Arial"
    p_b.font.size = Pt(10)
    p_b.font.bold = True
    p_b.font.color.rgb = ACCENT_BLUE

    tx = slide.shapes.add_textbox(Inches(0.8), Inches(0.68), Inches(11.7), Inches(0.6))
    tf = tx.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = title_text
    p.font.name = "Arial"
    p.font.size = Pt(20)
    p.font.bold = True
    p.font.color.rgb = TEXT_HEAD

    footer = slide.shapes.add_textbox(Inches(0.8), Inches(7.15), Inches(11.7), Inches(0.28))
    ft = footer.text_frame
    ft.margin_left = ft.margin_right = ft.margin_top = ft.margin_bottom = 0
    p_ft = ft.paragraphs[0]
    p_ft.text = "Nhóm WNC.G01 • Trường Đại học Mở Hà Nội (HOU) • Giảng viên: ThS. Lê Hữu Dũng"
    p_ft.font.name = "Arial"
    p_ft.font.size = Pt(9.5)
    p_ft.font.color.rgb = TEXT_MUTED

def add_p(tf, text, size=12, bold=False, color=TEXT_BODY, space_before=0, font_name="Arial"):
    p = tf.add_paragraph()
    p.text = text
    p.font.name = font_name
    p.font.size = Pt(size)
    p.font.bold = bold
    p.font.color.rgb = color
    if space_before:
        p.space_before = Pt(space_before)
    return p

card_w = Inches(5.65)
col_w = Inches(3.64)
gap = Inches(0.2)

# ==============================================================================
# SLIDE 1: BÌA BÁO CÁO (GỌN GÀNG, CHUYÊN NGHIỆP THEO GỢI Ý)
# ==============================================================================
s1 = prs.slides.add_slide(blank_layout)
bar1 = s1.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(0.12))
bar1.fill.solid()
bar1.fill.fore_color.rgb = ACCENT_BLUE
bar1.line.fill.background()

u_badge = s1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.0), Inches(0.7), Inches(5.8), Inches(0.42))
u_badge.fill.solid()
u_badge.fill.fore_color.rgb = BLUE_LIGHT
u_badge.line.color.rgb = BLUE_BORDER
p = u_badge.text_frame.paragraphs[0]
p.text = "TRƯỜNG ĐẠI HỌC MỞ HÀ NỘI — KHOA CÔNG NGHỆ THÔNG TIN"
p.font.name = "Arial"
p.font.size = Pt(10.5)
p.font.bold = True
p.font.color.rgb = ACCENT_BLUE

tb1 = s1.shapes.add_textbox(Inches(1.0), Inches(1.25), Inches(11.333), Inches(3.2))
tf1 = tb1.text_frame
tf1.word_wrap = True
p = tf1.paragraphs[0]
p.text = "BÁO CÁO THỰC NGHIỆM AN TOÀN WEB (OWASP TOP 10)"
p.font.name = "Arial"
p.font.size = Pt(15)
p.font.bold = True
p.font.color.rgb = ACCENT_RED

add_p(tf1, "SQL Injection • Stored XSS • CSRF Attack\nNghiên Cứu Lỗ Hổng — Thực Nghiệm — Cơ Chế Phòng Thủ", size=27, bold=True, color=TEXT_HEAD, space_before=8)
add_p(tf1, "Học phần: Lập trình Web nâng cao  •  Giảng viên hướng dẫn: ThS. Lê Hữu Dũng", size=13.5, color=TEXT_BODY, space_before=8)

# 3 Thẻ thành viên
m1 = s1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.0), Inches(4.7), col_w, Inches(2.0))
m1.fill.solid()
m1.fill.fore_color.rgb = CARD_BG
m1.line.color.rgb = BLUE_BORDER
tf = m1.text_frame
p = tf.paragraphs[0]
p.text = "CHỦ ĐỀ 1: SQL INJECTION"
p.font.name = "Arial"
p.font.size = Pt(11)
p.font.bold = True
p.font.color.rgb = ACCENT_BLUE
add_p(tf, "Nguyễn Danh Học\nMSV: 23A1001D0158\nTrưởng nhóm • Phụ trách CSDL & API", size=12, color=TEXT_HEAD, space_before=4)

m2 = s1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.0) + col_w + gap, Inches(4.7), col_w, Inches(2.0))
m2.fill.solid()
m2.fill.fore_color.rgb = CARD_BG
m2.line.color.rgb = RED_BORDER
tf = m2.text_frame
p = tf.paragraphs[0]
p.text = "CHỦ ĐỀ 2: STORED XSS"
p.font.name = "Arial"
p.font.size = Pt(11)
p.font.bold = True
p.font.color.rgb = ACCENT_RED
add_p(tf, "Nguyễn Thanh Bình\nMSV: 23A1001D0041\nThành viên • Phụ trách Cookie & DOM", size=12, color=TEXT_HEAD, space_before=4)

m3 = s1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.0) + (col_w + gap)*2, Inches(4.7), col_w, Inches(2.0))
m3.fill.solid()
m3.fill.fore_color.rgb = CARD_BG
m3.line.color.rgb = GREEN_BORDER
tf = m3.text_frame
p = tf.paragraphs[0]
p.text = "CHỦ ĐỀ 3: CSRF ATTACK"
p.font.name = "Arial"
p.font.size = Pt(11)
p.font.bold = True
p.font.color.rgb = ACCENT_GREEN
add_p(tf, "Nguyễn Minh Cường\nMSV: 23A1001D0058\nThành viên • Phụ trách Giao dịch & Token", size=12, color=TEXT_HEAD, space_before=4)


# ==============================================================================
# SLIDE 2: MÔI TRƯỜNG & PHẠM VI THỰC NGHIỆM (THÊM MỚI THEO GỢI Ý)
# ==============================================================================
s2 = prs.slides.add_slide(blank_layout)
add_header(s2, "Môi Trường Công Nghệ & Phạm Vi Thực Nghiệm")

c2_1 = s2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.4), card_w, Inches(5.4))
c2_1.fill.solid()
c2_1.fill.fore_color.rgb = CARD_BG
c2_1.line.color.rgb = CARD_BORDER
tf = c2_1.text_frame
p = tf.paragraphs[0]
p.text = "Thông Số Cấu Hình Hệ Thống Thực Nghiệm"
p.font.name = "Arial"
p.font.size = Pt(15)
p.font.bold = True
p.font.color.rgb = ACCENT_BLUE

add_p(tf, "• Hệ điều hành: Linux (Ubuntu 24.04 LTS / Zorin OS)", size=12, space_before=8)
add_p(tf, "• Nền tảng Backend: ASP.NET Core MVC (.NET 10.0)", size=12, space_before=6)
add_p(tf, "• Cơ sở dữ liệu: Microsoft SQL Server 2025 Developer (Port 1433)", size=12, space_before=6)
add_p(tf, "• Tầng giao diện: Razor View Engine, HTML5, CSS3, JavaScript", size=12, space_before=6)
add_p(tf, "• Môi trường phát triển: Visual Studio Code, C# DevKit", size=12, space_before=6)
add_p(tf, "• Công cụ kiểm thử: Chrome DevTools, cURL CLI, MSSQL Extension", size=12, space_before=6)

c2_2 = s2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.85), Inches(1.4), card_w, Inches(5.4))
c2_2.fill.solid()
c2_2.fill.fore_color.rgb = GREEN_LIGHT
c2_2.line.color.rgb = GREEN_BORDER
tf = c2_2.text_frame
p = tf.paragraphs[0]
p.text = "Phạm Vi & Nguyên Tắc Đạo Đức An Ninh"
p.font.name = "Arial"
p.font.size = Pt(15)
p.font.bold = True
p.font.color.rgb = ACCENT_GREEN

add_p(tf, "✓ Môi trường cô lập (Localhost):", size=12.5, bold=True, color=TEXT_HEAD, space_before=8)
add_p(tf, "Toàn bộ thực nghiệm được tiến hành trên ứng dụng mẫu do nhóm tự xây dựng (chạy trên localhost:5076), tuyệt đối không thực hiện trên các hệ thống bên ngoài.", size=12, space_before=4)

add_p(tf, "✓ Thiết kế đối chứng song song (Vulnerable vs Secure):", size=12.5, bold=True, color=TEXT_HEAD, space_before=12)
add_p(tf, "Mỗi lỗ hổng đều xây dựng 2 nhánh xử lý: Nhánh dính lỗi (để chứng minh sự nguy hiểm) và Nhánh an toàn (để kiểm chứng hiệu quả phòng thủ).", size=12, space_before=4)

add_p(tf, "✓ Dữ liệu mô phỏng an toàn:", size=12.5, bold=True, color=TEXT_HEAD, space_before=12)
add_p(tf, "Toàn bộ tài khoản, mật khẩu và số dư ví trong CSDL là dữ liệu mẫu giả định phục vụ mục đích học tập và nghiên cứu.", size=12, space_before=4)


# ==============================================================================
# PHẦN 1: SQL INJECTION — NGUYỄN DANH HỌC (SLIDES 3 - 7)
# ==============================================================================
# Slide 3: Bản chất & Sơ đồ luồng (Rút gọn chữ theo gợi ý)
s3 = prs.slides.add_slide(blank_layout)
add_header(s3, "1.1. SQL Injection: Sơ Đồ Luồng & Nguyên Nhân Cốt Lõi", "CHỦ ĐỀ 1 • NGUYỄN DANH HỌC (23A1001D0158)")

c3_1 = s3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.4), card_w, Inches(5.4))
c3_1.fill.solid()
c3_1.fill.fore_color.rgb = CARD_BG
c3_1.line.color.rgb = CARD_BORDER
tf = c3_1.text_frame
p = tf.paragraphs[0]
p.text = "Sơ Đồ Luồng Xử Lý Dữ Liệu"
p.font.name = "Arial"
p.font.size = Pt(15)
p.font.bold = True
p.font.color.rgb = ACCENT_RED

add_p(tf, "Browser (Client) ➔ Gửi Input (' OR '1'='1' --)", size=12, bold=True, color=ACCENT_BLUE, space_before=8)
add_p(tf, "            │\n            ▼", size=12, bold=True, color=TEXT_MUTED, space_before=2)
add_p(tf, "ASP.NET Core Controller ➔ Ghép chuỗi SQL trực tiếp", size=12, bold=True, color=ACCENT_RED, space_before=2)
add_p(tf, "            │\n            ▼", size=12, bold=True, color=TEXT_MUTED, space_before=2)
add_p(tf, "SQL Server 2025 ➔ Thông dịch cấu trúc lệnh bị tiêm nhiễm", size=12, bold=True, color=ACCENT_RED, space_before=2)

add_p(tf, "\nRoot Cause (Nguyên nhân cốt lõi):\nDữ liệu người dùng được ghép trực tiếp vào câu truy vấn. Hệ CSDL không phân định được đâu là Dữ liệu (Data) và đâu là Cú pháp lệnh (Code).", size=12, color=TEXT_BODY, space_before=8)

c3_2 = s3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.85), Inches(1.4), card_w, Inches(5.4))
c3_2.fill.solid()
c3_2.fill.fore_color.rgb = RED_LIGHT
c3_2.line.color.rgb = RED_BORDER
tf = c3_2.text_frame
p = tf.paragraphs[0]
p.text = "Đoạn Mã Gây Lỗi (C# Controller)"
p.font.name = "Arial"
p.font.size = Pt(15)
p.font.bold = True
p.font.color.rgb = ACCENT_RED
add_p(tf, "// NỐI CHUỖI TRỰC TIẾP TỪ THAM SỐ:\nstring rawSql = $\"SELECT * FROM Accounts \" +\n                $\"WHERE Username = '{username}' \" +\n                $\"AND Password = '{password}'\";\n\nvar accounts = _context.Accounts\n                       .FromSqlRaw(rawSql)\n                       .ToList();", size=10.5, color=RGBColor(159, 18, 57), space_before=8, font_name="Courier New")
add_p(tf, "Cơ chế phá vỡ cú pháp:\nDấu nháy đơn (') của kẻ tấn công đóng sớm chuỗi giá trị Username. Phần còn lại được máy chủ thông dịch thành lệnh điều khiển thay vì dữ liệu.", size=12, bold=True, color=TEXT_HEAD, space_before=14)

# Slide 4: 3 Kịch bản thực nghiệm (Rút gọn còn 3 kịch bản theo gợi ý)
s4 = prs.slides.add_slide(blank_layout)
add_header(s4, "1.2. SQL Injection: 3 Kịch Bản Khai Thác Thực Nghiệm", "CHỦ ĐỀ 1 • NGUYỄN DANH HỌC (23A1001D0158)")
t4_s = s4.shapes.add_table(4, 4, Inches(0.8), Inches(1.4), Inches(11.733), Inches(5.3))
t4 = t4_s.table
headers = ["Kịch bản kiểm thử", "Payload đưa vào ô Username", "Cơ chế bẻ gãy cú pháp", "Kết quả thực tế"]
for i, h in enumerate(headers):
    cell = t4.cell(0, i)
    cell.fill.solid()
    cell.fill.fore_color.rgb = ACCENT_BLUE
    p = cell.text_frame.paragraphs[0]
    p.text = h
    p.font.name = "Arial"
    p.font.bold = True
    p.font.size = Pt(11)
    p.font.color.rgb = WHITE

rows_s4 = [
    ("Kịch bản 1: Bypass tổng quát", "' OR '1'='1' --", "Nháy đơn (') đóng chuỗi; OR '1'='1' luôn đúng; '--' hủy vế kiểm tra mật khẩu.", "Bypass thành công, lấy tài khoản đầu tiên trong bảng (Admin)."),
    ("Kịch bản 2: Nhắm đích danh Admin", "admin' --", "Cố định Username = 'admin'; '--' hủy toàn bộ vế AND Password = ...", "Chiếm đoạt chính xác quyền Admin mà không cần biết mật khẩu."),
    ("Kịch bản chuẩn: Đăng nhập đúng", "admin / AdminPassword@2026", "Dữ liệu hợp lệ, không chứa ký tự điều khiển hay can thiệp cú pháp.", "Đăng nhập thành công trên cả 2 nhánh (Lỗi & Đã phòng thủ).")
]
for r_i, rdata in enumerate(rows_s4, start=1):
    for c_i, val in enumerate(rdata):
        cell = t4.cell(r_i, c_i)
        cell.fill.solid()
        cell.fill.fore_color.rgb = CARD_BG if r_i % 2 == 1 else WHITE
        p = cell.text_frame.paragraphs[0]
        p.text = val
        p.font.name = "Arial"
        p.font.size = Pt(11)
        p.font.color.rgb = TEXT_HEAD if c_i == 0 else TEXT_BODY
        if c_i == 1:
            p.font.name = "Courier New"
            p.font.bold = True
            p.font.color.rgb = ACCENT_RED if r_i < 3 else ACCENT_GREEN

# Slide 5: Minh chứng Web UI (Ảnh thật)
s5 = prs.slides.add_slide(blank_layout)
add_header(s5, "1.3. SQL Injection: Minh Chứng Thực Nghiệm Trực Quan Trên Web UI", "CHỦ ĐỀ 1 • NGUYỄN DANH HỌC (23A1001D0158)")
p_sqli = "/media/hocjsoo/New Volume/OWASP_Demo/screenshots/final_sqli_proof.png"
if os.path.exists(p_sqli):
    s5.shapes.add_picture(p_sqli, Inches(2.26), Inches(1.35), width=Inches(8.8))

cap5 = s5.shapes.add_textbox(Inches(0.8), Inches(6.55), Inches(11.733), Inches(0.4))
tf5 = cap5.text_frame
p = tf5.paragraphs[0]
p.text = "Minh chứng đối chứng: Trích xuất 4 tài khoản CSDL khi inject ' OR '1'='1' -- (Trái) vs Chặn đứng an toàn với EF Core (Phải)"
p.font.name = "Arial"
p.font.size = Pt(11)
p.font.bold = True
p.font.color.rgb = ACCENT_BLUE
p.alignment = PP_ALIGN.CENTER

# Slide 6: Minh chứng DevTools Network (Ảnh thật)
s6 = prs.slides.add_slide(blank_layout)
add_header(s6, "1.4. SQL Injection: Minh Chứng Gói Tin Mạng (Chrome DevTools Network)", "CHỦ ĐỀ 1 • NGUYỄN DANH HỌC (23A1001D0158)")
p_dt_sqli = "/media/hocjsoo/New Volume/OWASP_Demo/screenshots/01_sqli_devtools_proof.png"
if os.path.exists(p_dt_sqli):
    s6.shapes.add_picture(p_dt_sqli, Inches(2.06), Inches(1.35), width=Inches(9.2))

cap6 = s6.shapes.add_textbox(Inches(0.8), Inches(6.55), Inches(11.733), Inches(0.4))
tf6 = cap6.text_frame
p = tf6.paragraphs[0]
p.text = "Soi gói tin mạng (Network Fetch/XHR): Payload gửi lên mang chuỗi ' OR '1'='1' -- và JSON trích xuất 4 tài khoản rò rỉ"
p.font.name = "Arial"
p.font.size = Pt(11)
p.font.bold = True
p.font.color.rgb = ACCENT_BLUE
p.alignment = PP_ALIGN.CENTER

# Slide 7: Cơ chế phòng thủ chuẩn (Tách biệt SQLi và Hashing theo gợi ý)
s7 = prs.slides.add_slide(blank_layout)
add_header(s7, "1.5. SQL Injection: Cơ Chế Phòng Thủ Chuẩn Hóa", "CHỦ ĐỀ 1 • NGUYỄN DANH HỌC (23A1001D0158)")
c7_1 = s7.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.4), card_w, Inches(5.4))
c7_1.fill.solid()
c7_1.fill.fore_color.rgb = GREEN_LIGHT
c7_1.line.color.rgb = GREEN_BORDER
tf = c7_1.text_frame
p = tf.paragraphs[0]
p.text = "Phòng Thủ SQL Injection: Parameterized Query"
p.font.name = "Arial"
p.font.size = Pt(15)
p.font.bold = True
p.font.color.rgb = ACCENT_GREEN
add_p(tf, "// DÙNG LINQ PARAMETERIZED TRONG EF CORE:\nvar account = _context.Accounts\n    .FirstOrDefault(a => a.Username == user \n                      && a.Password == pass);", size=11, color=RGBColor(6, 95, 70), space_before=8, font_name="Courier New")
add_p(tf, "Nguyên lý kỹ thuật:\n• Data ≠ SQL Syntax: Đầu vào người dùng được truyền qua tham số riêng biệt (@p0, @p1).\n• Máy chủ SQL Server biên dịch cấu trúc cây truy vấn (Execution Plan) TRƯỚC khi gán dữ liệu.\n• Toàn bộ chuỗi payload bị cô lập thành giá trị văn bản (literal), không thể can thiệp vào logic thực thi.", size=12, space_before=12)

c7_2 = s7.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.85), Inches(1.4), card_w, Inches(5.4))
c7_2.fill.solid()
c7_2.fill.fore_color.rgb = CARD_BG
c7_2.line.color.rgb = CARD_BORDER
tf = c7_2.text_frame
p = tf.paragraphs[0]
p.text = "Bảo Vệ Mật Khẩu: Tách Biệt Với SQL Injection"
p.font.name = "Arial"
p.font.size = Pt(15)
p.font.bold = True
p.font.color.rgb = ACCENT_BLUE
add_p(tf, "Lưu ý logic bảo mật quan trọng:\n• Parameterized Query giải quyết triệt để lỗi SQL Injection.\n• Tuy nhiên, để bảo vệ mật khẩu người dùng phòng trường hợp CSDL bị rò rỉ, cần lớp phòng thủ bổ sung (Defense-in-Depth):", size=12, space_before=8)
add_p(tf, "1. Password Hashing (ASP.NET Core Identity):\n   Băm mật khẩu 1 chiều bằng thuật toán an toàn (PBKDF2/BCrypt) kèm Salt, tuyệt đối không lưu mật khẩu dạng plain-text.\n\n2. Phân quyền tối thiểu (Least Privilege):\n   Tài khoản kết nối CSDL của ứng dụng chỉ có quyền SELECT/INSERT/UPDATE trên các bảng nghiệp vụ, không dùng quyền sa/DDL.", size=12, space_before=8)


# ==============================================================================
# CHỦ ĐỀ 2: STORED XSS — NGUYỄN THANH BÌNH (SLIDES 8 - 12)
# ==============================================================================
# Slide 8: Bản chất XSS (Không viết 'nghiêm trọng nhất' theo gợi ý)
s8 = prs.slides.add_slide(blank_layout)
add_header(s8, "2.1. Stored XSS: Bản Chất Kỹ Thuật & Sơ Đồ Luồng", "CHỦ ĐỀ 2 • NGUYỄN THANH BÌNH (23A1001D0041)")
c8_1 = s8.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.4), card_w, Inches(5.4))
c8_1.fill.solid()
c8_1.fill.fore_color.rgb = CARD_BG
c8_1.line.color.rgb = CARD_BORDER
tf = c8_1.text_frame
p = tf.paragraphs[0]
p.text = "Sơ Đồ Luồng Phát Tán Của Stored XSS"
p.font.name = "Arial"
p.font.size = Pt(15)
p.font.bold = True
p.font.color.rgb = ACCENT_RED

add_p(tf, "Kẻ tấn công ➔ Nhập mã độc JavaScript vào form đánh giá", size=12, bold=True, color=ACCENT_RED, space_before=8)
add_p(tf, "            │\n            ▼", size=12, bold=True, color=TEXT_MUTED, space_before=2)
add_p(tf, "Web Server ➔ Lưu nguyên văn mã độc vào Database (dbo.Comments)", size=12, bold=True, color=TEXT_HEAD, space_before=2)
add_p(tf, "            │\n            ▼", size=12, bold=True, color=TEXT_MUTED, space_before=2)
add_p(tf, "Nạn nhân truy cập trang ➔ Trình duyệt tự động thực thi script độc hại", size=12, bold=True, color=ACCENT_RED, space_before=2)

add_p(tf, "\nĐặc điểm kỹ thuật:\nStored XSS là dạng XSS lưu trữ: mã độc được lưu trên máy chủ CSDL và có thể được kích hoạt tự động mỗi khi người dùng khác truy cập trang nội dung bị nhiễm.", size=12, color=TEXT_BODY, space_before=8)

c8_2 = s8.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.85), Inches(1.4), card_w, Inches(5.4))
c8_2.fill.solid()
c8_2.fill.fore_color.rgb = RED_LIGHT
c8_2.line.color.rgb = RED_BORDER
tf = c8_2.text_frame
p = tf.paragraphs[0]
p.text = "Nguyên Nhân: Sự Lạm Dụng @@Html.Raw()"
p.font.name = "Arial"
p.font.size = Pt(15)
p.font.bold = True
p.font.color.rgb = ACCENT_RED
add_p(tf, "// 1. CONTROLLER LƯU NGUYÊN VĂN NỘI DUNG VÀO CSDL:\nvar comment = new Comment {\n    Author = author,\n    Content = content // Chứa thẻ <script> độc hại\n};\n_context.Comments.Add(comment);\n\n// 2. RAZOR VIEW HIỂN THỊ DỮ LIỆU THÔ BẰNG HTML.RAW:\n<div class=\"comment-body\">\n    @Html.Raw(comment.Content)\n</div>", size=10, color=RGBColor(159, 18, 57), space_before=6, font_name="Courier New")
add_p(tf, "Điểm mù kỹ thuật:\n@Html.Raw() vô hiệu hóa hoàn toàn cơ chế mã hóa HTML tự động của Razor Engine, ép trình duyệt phải biên dịch và thực thi chuỗi dữ liệu như mã lệnh.", size=12, bold=True, color=TEXT_HEAD, space_before=10)

# Slide 9: 3 Kịch bản XSS
s9 = prs.slides.add_slide(blank_layout)
add_header(s9, "2.2. Stored XSS: 3 Kịch Bản Tiêm Mã Độc Thực Nghiệm", "CHỦ ĐỀ 2 • NGUYỄN THANH BÌNH (23A1001D0041)")
t9_s = s9.shapes.add_table(4, 4, Inches(0.8), Inches(1.4), Inches(11.733), Inches(5.3))
t9 = t9_s.table
headers_xss = ["Kịch bản kiểm thử", "Mã độc Payload đưa vào", "Cơ chế kích hoạt trên Trình duyệt", "Mục tiêu khai thác"]
for i, h in enumerate(headers_xss):
    cell = t9.cell(0, i)
    cell.fill.solid()
    cell.fill.fore_color.rgb = ACCENT_RED
    p = cell.text_frame.paragraphs[0]
    p.text = h
    p.font.name = "Arial"
    p.font.bold = True
    p.font.size = Pt(11)
    p.font.color.rgb = WHITE

rows_s9 = [
    ("Kịch bản 1: Đánh cắp Cookie", "<script>alert(document.cookie);</script>", "Trình duyệt thực thi thẻ <script> và đọc toàn bộ giá trị lưu trong document.cookie.", "Chiếm đoạt Session Token (Cookie Stealing) để mạo danh phiên làm việc."),
    ("Kịch bản 2: Bypass qua thẻ <img>", "<img src=x onerror=alert('XSS!')>", "Khi đường dẫn ảnh bị lỗi (src=x), trình duyệt lập tức kích hoạt sự kiện onerror.", "Vượt qua các bộ lọc từ khóa đơn giản chỉ chặn chuỗi '<script>'."),
    ("Kịch bản chuẩn: Đánh giá hợp lệ", "Dịch vụ tư vấn rất chuyên nghiệp, 5 sao!", "Văn bản thuần túy không chứa thẻ HTML hay kịch bản đặc biệt.", "Hiển thị bình thường trên cả 2 nhánh (Lỗi & Đã phòng thủ).")
]
for r_i, rdata in enumerate(rows_s9, start=1):
    for c_i, val in enumerate(rdata):
        cell = t9.cell(r_i, c_i)
        cell.fill.solid()
        cell.fill.fore_color.rgb = CARD_BG if r_i % 2 == 1 else WHITE
        p = cell.text_frame.paragraphs[0]
        p.text = val
        p.font.name = "Arial"
        p.font.size = Pt(11)
        p.font.color.rgb = TEXT_HEAD if c_i == 0 else TEXT_BODY
        if c_i == 1:
            p.font.name = "Courier New"
            p.font.bold = True
            p.font.color.rgb = ACCENT_RED if r_i < 3 else ACCENT_GREEN

# Slide 10: Minh chứng Web UI XSS (Ảnh thật)
s10 = prs.slides.add_slide(blank_layout)
add_header(s10, "2.3. Stored XSS: Minh Chứng Thực Nghiệm Trực Quan Trên Web UI", "CHỦ ĐỀ 2 • NGUYỄN THANH BÌNH (23A1001D0041)")
p_xss = "/media/hocjsoo/New Volume/OWASP_Demo/screenshots/final_xss_proof.png"
if os.path.exists(p_xss):
    s10.shapes.add_picture(p_xss, Inches(2.86), Inches(1.35), width=Inches(7.6))

cap10 = s10.shapes.add_textbox(Inches(0.8), Inches(6.55), Inches(11.733), Inches(0.4))
tf10 = cap10.text_frame
p = tf10.paragraphs[0]
p.text = "Minh chứng đối chứng chức năng Đánh giá: Nhánh dính lỗi kích hoạt popup Alert (Trái) vs Nhánh an toàn mã hóa HTML (Phải)"
p.font.name = "Arial"
p.font.size = Pt(11)
p.font.bold = True
p.font.color.rgb = ACCENT_RED
p.alignment = PP_ALIGN.CENTER

# Slide 11: Minh chứng DevTools Application & Console (Ảnh thật + chú thích lab rõ ràng)
s11 = prs.slides.add_slide(blank_layout)
add_header(s11, "2.4. Stored XSS: Minh Chứng Kiểm Thử Qua Chrome DevTools (Cookie & Console)", "CHỦ ĐỀ 2 • NGUYỄN THANH BÌNH (23A1001D0041)")
p_dt_xss = "/media/hocjsoo/New Volume/OWASP_Demo/screenshots/02_xss_devtools_proof.png"
if os.path.exists(p_dt_xss):
    s11.shapes.add_picture(p_dt_xss, Inches(2.06), Inches(1.35), width=Inches(9.2))

cap11 = s11.shapes.add_textbox(Inches(0.8), Inches(6.55), Inches(11.733), Inches(0.4))
tf11 = cap11.text_frame
p = tf11.paragraphs[0]
p.text = "Mô phỏng Lab: Nhánh vulnerable cấu hình HttpOnly = false để chứng minh JavaScript đọc trộm được Session Cookie"
p.font.name = "Arial"
p.font.size = Pt(11)
p.font.bold = True
p.font.color.rgb = ACCENT_RED
p.alignment = PP_ALIGN.CENTER

# Slide 12: Phòng thủ chuẩn XSS
s12 = prs.slides.add_slide(blank_layout)
add_header(s12, "2.5. Stored XSS: Cơ Chế Phòng Thủ Chuẩn Trong ASP.NET Core", "CHỦ ĐỀ 2 • NGUYỄN THANH BÌNH (23A1001D0041)")
c12_1 = s12.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.4), card_w, Inches(5.4))
c12_1.fill.solid()
c12_1.fill.fore_color.rgb = GREEN_LIGHT
c12_1.line.color.rgb = GREEN_BORDER
tf = c12_1.text_frame
p = tf.paragraphs[0]
p.text = "Phòng Thủ Tầng View: Context-Aware Encoding"
p.font.name = "Arial"
p.font.size = Pt(15)
p.font.bold = True
p.font.color.rgb = ACCENT_GREEN
add_p(tf, "// CÚ PHÁP RAZOR VIEW MẶC ĐỊNH (KHÔNG DÙNG HTML.RAW):\n<p>@comment.Content</p>\n\n// HOẶC MÃ HÓA TẠI CONTROLLER VỚI HTMLENCODER:\nstring encoded = HtmlEncoder.Default.Encode(content);\ncomment.Content = encoded;", size=10.5, color=RGBColor(6, 95, 70), space_before=8, font_name="Courier New")
add_p(tf, "Cơ chế chuyển đổi thực thể HTML:\n• Ký tự '<' biến đổi thành &lt;\n• Ký tự '>' biến đổi thành &gt;\nTrình duyệt hiển thị dưới dạng văn bản thuần túy, loại bỏ hoàn toàn khả năng thực thi mã kịch bản.", size=12, space_before=12)

c12_2 = s12.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.85), Inches(1.4), card_w, Inches(5.4))
c12_2.fill.solid()
c12_2.fill.fore_color.rgb = CARD_BG
c12_2.line.color.rgb = CARD_BORDER
tf = c12_2.text_frame
p = tf.paragraphs[0]
p.text = "Phòng Thủ Tầng Cookie & Header (Khuyến Nghị)"
p.font.name = "Arial"
p.font.size = Pt(15)
p.font.bold = True
p.font.color.rgb = ACCENT_BLUE
add_p(tf, "// CẤU HÌNH COOKIE AN TOÀN TRONG SẢN PHẨM THẬT:\nResponse.Cookies.Append(\"AuthSessionToken\", token, new CookieOptions\n{\n    HttpOnly = true, // CẤM JAVASCRIPT ĐỌC TRỘM\n    Secure = true,   // BẮT BUỘC TRUYỀN QUA HTTPS\n    SameSite = SameSiteMode.Strict\n});", size=10, color=TEXT_HEAD, space_before=6, font_name="Courier New")
add_p(tf, "Chính sách Content Security Policy (CSP Header):\nThiết lập HTTP Header 'script-src self' ngăn cấm trình duyệt thực thi inline script, tạo thêm lớp bảo vệ vững chắc ngay cả khi có sơ hở ở tầng View.", size=12, space_before=10)


# ==============================================================================
# CHỦ ĐỀ 3: CSRF ATTACK — NGUYỄN MINH CƯỜNG (SLIDES 13 - 17)
# ==============================================================================
# Slide 13: Bản chất CSRF (Sửa phân loại OWASP chuẩn theo gợi ý)
s13 = prs.slides.add_slide(blank_layout)
add_header(s13, "3.1. CSRF Attack: Bản Chất Kỹ Thuật & Phân Loại OWASP", "CHỦ ĐỀ 3 • NGUYỄN MINH CƯỜNG (23A1001D0058)")
c13_1 = s13.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.4), card_w, Inches(5.4))
c13_1.fill.solid()
c13_1.fill.fore_color.rgb = CARD_BG
c13_1.line.color.rgb = CARD_BORDER
tf = c13_1.text_frame
p = tf.paragraphs[0]
p.text = "Bản Chất Kỹ Thuật Của Tấn Công CSRF"
p.font.name = "Arial"
p.font.size = Pt(15)
p.font.bold = True
p.font.color.rgb = ACCENT_RED
add_p(tf, "• Phân loại chuẩn OWASP: CSRF không phải một mục riêng lẻ trong OWASP Top 10 2021; nó liên quan trực tiếp đến A01:2021 – Broken Access Control và cơ chế xác thực phiên.", size=12.5, bold=True, color=ACCENT_RED, space_before=8)
add_p(tf, "• Cơ chế mượn quyền:\n  1. Nạn nhân có phiên làm việc hợp lệ (Cookie xác thực còn hiệu lực trên trình duyệt).\n  2. Trình duyệt tự động đính kèm Cookie xác thực theo mọi HTTP Request gửi tới ứng dụng mục tiêu.\n  3. Kẻ tấn công lừa nạn nhân truy cập một trang web nguồn gốc chéo (Cross-Origin Attacker Site).\n• Điểm mù: Máy chủ chỉ xác thực Cookie mà không kiểm tra nguồn gốc phát sinh Request.", size=12, space_before=8)

c13_2 = s13.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.85), Inches(1.4), card_w, Inches(5.4))
c13_2.fill.solid()
c13_2.fill.fore_color.rgb = RED_LIGHT
c13_2.line.color.rgb = RED_BORDER
tf = c13_2.text_frame
p = tf.paragraphs[0]
p.text = "Action Xử Lý Giao Dịch Thiếu Kiểm Soát"
p.font.name = "Arial"
p.font.size = Pt(15)
p.font.bold = True
p.font.color.rgb = ACCENT_RED
add_p(tf, "// ACTION NHẬN YÊU CẦU KHÔNG CÓ BẢO VỆ TOKEN:\n[HttpPost]\npublic IActionResult TransferVulnerable(decimal amount)\n{\n    var victim = _context.Wallets.Find(1);\n    var hacker = _context.Wallets.Find(2);\n    \n    victim.Balance -= amount;\n    hacker.Balance += amount;\n    _context.SaveChanges();\n    return RedirectToAction(\"Index\");\n}", size=10, color=RGBColor(159, 18, 57), space_before=8, font_name="Courier New")
add_p(tf, "Khi Action thiếu thuộc tính [ValidateAntiForgeryToken], yêu cầu POST bắt nguồn từ một website khác vẫn được thực thi do trình duyệt tự động kẹp theo Cookie của nạn nhân.", size=12, bold=True, color=TEXT_HEAD, space_before=10)

# Slide 14: Kịch bản lừa đảo bẫy trúng thưởng
s14 = prs.slides.add_slide(blank_layout)
add_header(s14, "3.2. CSRF Attack: Kịch Bản Khai Thác Qua Biểu Mẫu Ẩn", "CHỦ ĐỀ 3 • NGUYỄN MINH CƯỜNG (23A1001D0058)")
c14_main = s14.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.4), Inches(11.733), Inches(5.3))
c14_main.fill.solid()
c14_main.fill.fore_color.rgb = RED_LIGHT
c14_main.line.color.rgb = RED_BORDER
tf = c14_main.text_frame
p = tf.paragraphs[0]
p.text = "Kịch Bản: Trang Web Giả Mạo Điều Khiển Yêu Cầu Chuyển Tiền"
p.font.name = "Arial"
p.font.size = Pt(16)
p.font.bold = True
p.font.color.rgb = ACCENT_RED
add_p(tf, "1. Nạn nhân có phiên làm việc đang hoạt động trên hệ thống tư vấn trực tuyến.\n2. Nạn nhân được dẫn dụ truy cập một trang web độc hại ở một tab khác (AttackerSite.cshtml).\n3. Trang web độc hại chứa một biểu mẫu ẩn trỏ trực tiếp tới địa chỉ nhận lệnh của ứng dụng mục tiêu:", size=13, space_before=8)
add_p(tf, "<!-- BIỂU MẪU ẨN TRÊN TRANG WEB CỦA KẺ TẤN CÔNG -->\n<form id=\"exploit\" action=\"http://localhost:5076/Csrf/TransferVulnerable\" method=\"POST\">\n    <input type=\"hidden\" name=\"amount\" value=\"20000000\" />\n</form>\n<script>document.getElementById('exploit').submit();</script>", size=11, bold=True, color=RGBColor(159, 18, 57), space_before=6, font_name="Courier New")
add_p(tf, "Hậu quả thực tế: Giao dịch chuyển tiền được thực thi hoàn toàn tự động mà không có sự đồng thuận chủ động của người dùng, làm thay đổi số dư ví tài khoản.", size=13, bold=True, color=ACCENT_RED, space_before=8)

# Slide 15: Minh chứng Web UI CSRF (Ảnh thật + cách diễn đạt chuyên nghiệp theo gợi ý)
s15 = prs.slides.add_slide(blank_layout)
add_header(s15, "3.3. CSRF Attack: Minh Chứng Thực Nghiệm Trực Quan Trên Web UI", "CHỦ ĐỀ 3 • NGUYỄN MINH CƯỜNG (23A1001D0058)")
p_csrf = "/media/hocjsoo/New Volume/OWASP_Demo/screenshots/final_csrf_proof.png"
if os.path.exists(p_csrf):
    s15.shapes.add_picture(p_csrf, Inches(2.76), Inches(1.35), width=Inches(7.8))

cap15 = s15.shapes.add_textbox(Inches(0.8), Inches(6.55), Inches(11.733), Inches(0.4))
tf15 = cap15.text_frame
p = tf15.paragraphs[0]
p.text = "Minh chứng đối chứng: Mô phỏng giao dịch trái phép 20.000.000 VNĐ trong lab (Trái) vs Chặn đứng yêu cầu thiếu Token (Phải)"
p.font.name = "Arial"
p.font.size = Pt(11)
p.font.bold = True
p.font.color.rgb = ACCENT_GREEN
p.alignment = PP_ALIGN.CENTER

# Slide 16: Minh chứng DevTools Headers (Ảnh thật)
s16 = prs.slides.add_slide(blank_layout)
add_header(s16, "3.4. CSRF Attack: Minh Chứng Gói Tin Mạng (Chrome DevTools Headers)", "CHỦ ĐỀ 3 • NGUYỄN MINH CƯỜNG (23A1001D0058)")
p_dt_csrf = "/media/hocjsoo/New Volume/OWASP_Demo/screenshots/03_csrf_devtools_proof.png"
if os.path.exists(p_dt_csrf):
    s16.shapes.add_picture(p_dt_csrf, Inches(2.06), Inches(1.35), width=Inches(9.2))

cap16 = s16.shapes.add_textbox(Inches(0.8), Inches(6.55), Inches(11.733), Inches(0.4))
tf16 = cap16.text_frame
p = tf16.paragraphs[0]
p.text = "Soi gói tin HTTP POST gửi ngầm từ Attacker Site: Đính kèm Cookie tự động nhưng thiếu hoàn toàn Anti-Forgery Token"
p.font.name = "Arial"
p.font.size = Pt(11)
p.font.bold = True
p.font.color.rgb = ACCENT_GREEN
p.alignment = PP_ALIGN.CENTER

# Slide 17: Phòng thủ CSRF Token & Fetch API
s17 = prs.slides.add_slide(blank_layout)
add_header(s17, "3.5. CSRF Attack: Cơ Chế Phòng Thủ Bằng Anti-Forgery Token", "CHỦ ĐỀ 3 • NGUYỄN MINH CƯỜNG (23A1001D0058)")
c17_1 = s17.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.4), card_w, Inches(5.4))
c17_1.fill.solid()
c17_1.fill.fore_color.rgb = GREEN_LIGHT
c17_1.line.color.rgb = GREEN_BORDER
tf = c17_1.text_frame
p = tf.paragraphs[0]
p.text = "Mô Hình Synchronizer Token Pattern"
p.font.name = "Arial"
p.font.size = Pt(15)
p.font.bold = True
p.font.color.rgb = ACCENT_GREEN
add_p(tf, "// 1. TRONG RAZOR VIEW: NHÚNG TOKEN VÀO FORM\n<form asp-action=\"TransferSecure\" method=\"post\">\n    @Html.AntiForgeryToken()\n    <input type=\"number\" name=\"amount\" />\n    <button type=\"submit\">Chuyển tiền</button>\n</form>\n\n// 2. TRONG CONTROLLER: BẮT BUỘC KIỂM TRA TOKEN\n[HttpPost]\n[ValidateAntiForgeryToken]\npublic IActionResult TransferSecure(decimal amount)\n{\n    // Chỉ thực thi khi Token từ Form khớp với Token trong Cookie\n}", size=9.5, color=RGBColor(6, 95, 70), space_before=8, font_name="Courier New")
add_p(tf, "Cơ chế bảo vệ:\nMáy chủ phát hành một cặp token: một lưu trong Cookie và một nhúng vào biểu mẫu. Trang web độc hại bị chặn đọc token bởi chính sách Same-Origin Policy (SOP).", size=11.5, space_before=8)

c17_2 = s17.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.85), Inches(1.4), card_w, Inches(5.4))
c17_2.fill.solid()
c17_2.fill.fore_color.rgb = CARD_BG
c17_2.line.color.rgb = CARD_BORDER
tf = c17_2.text_frame
p = tf.paragraphs[0]
p.text = "Tích Hợp Token Trong Fetch API (Slide Buổi 9)"
p.font.name = "Arial"
p.font.size = Pt(15)
p.font.bold = True
p.font.color.rgb = ACCENT_BLUE
add_p(tf, "// TRÍCH XUẤT VÀ GỬI TOKEN QUA HTTP HEADER:\nconst token = document.querySelector('input[name=\"__RequestVerificationToken\"]').value;\n\nfetch('/Csrf/TransferAjax', {\n    method: 'POST',\n    headers: {\n        'Content-Type': 'application/json',\n        'RequestVerificationToken': token // HEADER BẢO VỆ\n    },\n    body: JSON.stringify({ amount: 5000000 })\n})\n.then(response => response.json());", size=9.5, color=TEXT_HEAD, space_before=8, font_name="Courier New")
add_p(tf, "Ứng dụng trong thực tế:\nĐối với các tác vụ gọi AJAX không tải lại trang, token được đọc từ thẻ ẩn và truyền qua HTTP Header để đảm bảo vừa tiện dụng vừa an toàn.", size=11.5, space_before=8)


# ==============================================================================
# TỔNG KẾT: BÀI HỌC RÚT RA & LỜI KẾT (SLIDES 18 - 19)
# ==============================================================================
# Slide 18: Bảng bài học rút ra (Thêm mới theo gợi ý)
s18 = prs.slides.add_slide(blank_layout)
add_header(s18, "Tổng Kết: Bài Học Kỹ Thuật & Bảng Đối Chiếu Phòng Thủ")
t18_s = s18.shapes.add_table(4, 4, Inches(0.8), Inches(1.5), Inches(11.733), Inches(5.2))
t18 = t18_s.table

headers_sum = ["Lỗ hổng an ninh", "Nguyên nhân cốt lõi", "Cơ chế phòng thủ chuẩn", "Thành viên phụ trách"]
for i, h in enumerate(headers_sum):
    cell = t18.cell(0, i)
    cell.fill.solid()
    cell.fill.fore_color.rgb = ACCENT_BLUE
    p = cell.text_frame.paragraphs[0]
    p.text = h
    p.font.name = "Arial"
    p.font.bold = True
    p.font.size = Pt(12)
    p.font.color.rgb = WHITE

rows_s18 = [
    ("SQL Injection\n(A03:2021)", "Ghép chuỗi đầu vào trực tiếp vào câu lệnh SQL.", "Parameterized Query (EF Core LINQ) + Nguyên tắc đặc quyền tối thiểu.", "Nguyễn Danh Học\n(23A1001D0158)"),
    ("Stored XSS\n(A03:2021)", "Lưu trữ script độc hại vào CSDL & render thô bằng @Html.Raw().", "Context-Aware HTML Encoding (@model.Property) + HttpOnly Cookie.", "Nguyễn Thanh Bình\n(23A1001D0041)"),
    ("CSRF Attack\n(A01:2021 / Session)", "Máy chủ chỉ kiểm tra Cookie mà không xác thực nguồn gốc Request.", "Synchronizer Token Pattern (@Html.AntiForgeryToken() / Header).", "Nguyễn Minh Cường\n(23A1001D0058)")
]

for r_i, rdata in enumerate(rows_s18, start=1):
    for c_i, val in enumerate(rdata):
        cell = t18.cell(r_i, c_i)
        cell.fill.solid()
        cell.fill.fore_color.rgb = CARD_BG if r_i % 2 == 1 else WHITE
        p = cell.text_frame.paragraphs[0]
        p.text = val
        p.font.name = "Arial"
        p.font.size = Pt(11)
        p.font.color.rgb = TEXT_HEAD if c_i in (0, 3) else TEXT_BODY
        if c_i == 0:
            p.font.bold = True
        elif c_i == 2:
            p.font.bold = True
            p.font.color.rgb = ACCENT_GREEN

# Slide 19: Lời kết
s19 = prs.slides.add_slide(blank_layout)
add_header(s19, "Kết Luận Báo Cáo Thực Nghiệm OWASP Top 10")
main_c19 = s19.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.0), Inches(1.5), Inches(11.333), Inches(5.0))
main_c19.fill.solid()
main_c19.fill.fore_color.rgb = CARD_BG
main_c19.line.color.rgb = CARD_BORDER
tf = main_c19.text_frame
p = tf.paragraphs[0]
p.text = "KẾT LUẬN & CAM KẾT CHUẨN BẢO MẬT (NHÓM WNC.G01)"
p.font.name = "Arial"
p.font.size = Pt(16)
p.font.bold = True
p.font.color.rgb = ACCENT_BLUE

add_p(tf, "1. Thực nghiệm trực quan: Nhóm đã chứng minh rõ cơ chế khai thác và mức độ rủi ro của từng lỗ hổng trên cả 3 tầng: Giao diện Web, Gói tin mạng (Chrome DevTools) và Cơ sở dữ liệu (SQL Server 2025).", size=13, space_before=14)
add_p(tf, "2. Phòng thủ đa tầng (Security View): Các cơ chế phòng thủ (EF Core Parameterized, Razor HTML Encoding, Anti-Forgery Token) được tích hợp trực tiếp vào kiến trúc nền tảng của Đồ án BTL Website Tư Vấn Trực Tuyến.", size=13, space_before=12)
add_p(tf, "3. Đóng góp đồng đều: Cả 3 thành viên đều trực tiếp xây dựng mã nguồn, kiểm thử đối chứng và phân tích chuyên môn độc lập.", size=13, space_before=12)
add_p(tf, "\nKính chúc thầy sức khỏe và công tác tốt. Nhóm WNC.G01 xin trân trọng cảm ơn!", size=15, bold=True, color=ACCENT_RED, space_before=24)

out_file = "/media/hocjsoo/New Volume/OWASP_Demo/Slide_OWASP_Top3_Nhom_WNC_G01.pptx"
prs.save(out_file)
print(f"Master presentation v2 successfully saved to: {out_file}")
