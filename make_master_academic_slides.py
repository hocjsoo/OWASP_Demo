import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)

# Bảng màu chuẩn học thuật (Academic Light Theme)
BG_WHITE = RGBColor(255, 255, 255)
WHITE = RGBColor(255, 255, 255)
CARD_BG = RGBColor(248, 250, 252)          # Slate 50
CARD_BORDER = RGBColor(226, 232, 240)      # Slate 200

TEXT_HEAD = RGBColor(15, 23, 42)           # Slate 900 (Đen kỹ thuật)
TEXT_BODY = RGBColor(51, 65, 85)           # Slate 700 (Xám đậm dễ đọc)
TEXT_MUTED = RGBColor(100, 116, 139)       # Slate 500

ACCENT_BLUE = RGBColor(29, 78, 216)        # Blue 700 (HOU Chuẩn)
BLUE_LIGHT = RGBColor(239, 246, 255)       # Blue 50
BLUE_BORDER = RGBColor(191, 219, 254)

ACCENT_RED = RGBColor(190, 18, 60)         # Rose 700 (Cảnh báo lỗi)
RED_LIGHT = RGBColor(255, 241, 242)        # Rose 50
RED_BORDER = RGBColor(254, 205, 211)

ACCENT_GREEN = RGBColor(4, 120, 87)        # Emerald 700 (An toàn)
GREEN_LIGHT = RGBColor(236, 253, 245)      # Emerald 50
GREEN_BORDER = RGBColor(167, 243, 208)

ACCENT_AMBER = RGBColor(180, 83, 9)        # Amber 700
AMBER_LIGHT = RGBColor(255, 251, 235)      # Amber 50
AMBER_BORDER = RGBColor(254, 240, 138)

blank_layout = prs.slide_layouts[6]

def add_header(slide, title_text, category_tag="OWASP TOP 10 • THỰC NGHIỆM AN TOÀN WEB"):
    top_line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(0.08))
    top_line.fill.solid()
    top_line.fill.fore_color.rgb = ACCENT_BLUE
    top_line.line.fill.background()

    badge = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(0.32), Inches(4.8), Inches(0.38))
    badge.fill.solid()
    badge.fill.fore_color.rgb = BLUE_LIGHT
    badge.line.color.rgb = BLUE_BORDER

    tf_b = badge.text_frame
    tf_b.word_wrap = False
    p_b = tf_b.paragraphs[0]
    p_b.text = category_tag.upper()
    p_b.font.size = Pt(10)
    p_b.font.bold = True
    p_b.font.color.rgb = ACCENT_BLUE

    tx = slide.shapes.add_textbox(Inches(0.8), Inches(0.72), Inches(11.7), Inches(0.6))
    tf = tx.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = title_text
    p.font.size = Pt(20)
    p.font.bold = True
    p.font.color.rgb = TEXT_HEAD

    footer = slide.shapes.add_textbox(Inches(0.8), Inches(7.1), Inches(11.7), Inches(0.3))
    ft = footer.text_frame
    p_ft = ft.paragraphs[0]
    p_ft.text = "Nhóm WNC.G01 • Trường Đại học Mở Hà Nội (HOU) • Giảng viên: ThS. Lê Hữu Dũng"
    p_ft.font.size = Pt(9.5)
    p_ft.font.color.rgb = TEXT_MUTED

card_w = Inches(5.65)
col_w = Inches(3.64)
gap = Inches(0.2)

# ==============================================================================
# SLIDE 1: BÌA
# ==============================================================================
s1 = prs.slides.add_slide(blank_layout)
bar1 = s1.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(0.12))
bar1.fill.solid()
bar1.fill.fore_color.rgb = ACCENT_BLUE
bar1.line.fill.background()

u_badge = s1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.0), Inches(0.7), Inches(5.8), Inches(0.42))
u_badge.fill.solid()
u_badge.fill.fore_color.rgb = BLUE_LIGHT
u_badge.line.color.rgb = BLUE_BORDER
p = u_badge.text_frame.paragraphs[0]
p.text = "TRƯỜNG ĐẠI HỌC MỞ HÀ NỘI — KHOA CÔNG NGHỆ THÔNG TIN"
p.font.size = Pt(10.5)
p.font.bold = True
p.font.color.rgb = ACCENT_BLUE

tb1 = s1.shapes.add_textbox(Inches(1.0), Inches(1.25), Inches(11.333), Inches(3.2))
tf1 = tb1.text_frame
tf1.word_wrap = True
p = tf1.paragraphs[0]
p.text = "BÁO CÁO THỰC NGHIỆM AN TOÀN WEB (OWASP TOP 10)"
p.font.size = Pt(15)
p.font.bold = True
p.font.color.rgb = ACCENT_RED

p = tf1.add_paragraph()
p.text = "Nghiên Cứu Thực Nghiệm Lỗ Hổng Web & Cơ Chế Phòng Thủ\n(SQL Injection • Stored XSS • CSRF Attack)"
p.font.size = Pt(28)
p.font.bold = True
p.font.color.rgb = TEXT_HEAD
p.space_before = Pt(8)

p = tf1.add_paragraph()
p.text = "Học phần: Lập trình Web nâng cao  •  Giảng viên hướng dẫn: ThS. Lê Hữu Dũng"
p.font.size = Pt(13.5)
p.font.color.rgb = TEXT_BODY
p.space_before = Pt(8)

m1 = s1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.0), Inches(4.7), col_w, Inches(2.0))
m1.fill.solid()
m1.fill.fore_color.rgb = CARD_BG
m1.line.color.rgb = BLUE_BORDER
tf = m1.text_frame
tf.word_wrap = True
p = tf.paragraphs[0]
p.text = "CHỦ ĐỀ 1: SQL INJECTION"
p.font.size = Pt(11)
p.font.bold = True
p.font.color.rgb = ACCENT_BLUE
p = tf.add_paragraph()
p.text = "Nguyễn Danh Học\nMSV: 23A1001D0158\nTrưởng nhóm • Phụ trách CSDL & API"
p.font.size = Pt(12)
p.font.color.rgb = TEXT_HEAD
p.space_before = Pt(4)

m2 = s1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.0) + col_w + gap, Inches(4.7), col_w, Inches(2.0))
m2.fill.solid()
m2.fill.fore_color.rgb = CARD_BG
m2.line.color.rgb = RED_BORDER
tf = m2.text_frame
tf.word_wrap = True
p = tf.paragraphs[0]
p.text = "CHỦ ĐỀ 2: STORED XSS"
p.font.size = Pt(11)
p.font.bold = True
p.font.color.rgb = ACCENT_RED
p = tf.add_paragraph()
p.text = "Nguyễn Thanh Bình\nMSV: 23A1001D0041\nThành viên • Phụ trách Cookie & DOM"
p.font.size = Pt(12)
p.font.color.rgb = TEXT_HEAD
p.space_before = Pt(4)

m3 = s1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.0) + (col_w + gap)*2, Inches(4.7), col_w, Inches(2.0))
m3.fill.solid()
m3.fill.fore_color.rgb = CARD_BG
m3.line.color.rgb = GREEN_BORDER
tf = m3.text_frame
tf.word_wrap = True
p = tf.paragraphs[0]
p.text = "CHỦ ĐỀ 3: CSRF ATTACK"
p.font.size = Pt(11)
p.font.bold = True
p.font.color.rgb = ACCENT_GREEN
p = tf.add_paragraph()
p.text = "Nguyễn Minh Cường\nMSV: 23A1001D0058\nThành viên • Phụ trách Giao dịch & Token"
p.font.size = Pt(12)
p.font.color.rgb = TEXT_HEAD
p.space_before = Pt(4)

# ==============================================================================
# SLIDE 2: MỤC TIÊU BÁO CÁO
# ==============================================================================
s2 = prs.slides.add_slide(blank_layout)
add_header(s2, "Mục Tiêu & Phương Pháp Thực Nghiệm")

c_t1 = s2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.4), Inches(5.65), Inches(5.3))
c_t1.fill.solid()
c_t1.fill.fore_color.rgb = CARD_BG
c_t1.line.color.rgb = CARD_BORDER
tf = c_t1.text_frame
tf.word_wrap = True
p = tf.paragraphs[0]
p.text = "Mục Tiêu Nghiên Cứu"
p.font.size = Pt(16)
p.font.bold = True
p.font.color.rgb = ACCENT_RED
p = tf.add_paragraph()
p.text = "• Bám sát yêu cầu thực tiễn: Tập trung vào thực nghiệm và phân tích mã nguồn thay vì lý thuyết thuần túy.\n• Nhận thức mức độ rủi ro: Đánh giá tác động an ninh thực tế trên ứng dụng web.\n• Phương pháp tiếp cận: Để phòng thủ hiệu quả, cần phân tích chính xác vector và kỹ thuật khai thác của kẻ tấn công.\n• Đối chứng rõ ràng: So sánh trực tiếp giữa mã nguồn dính lỗ hổng và mã nguồn đã được chuẩn hóa."
p.font.size = Pt(12.5)
p.font.color.rgb = TEXT_BODY
p.space_before = Pt(10)

c_t2 = s2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.85), Inches(1.4), Inches(5.65), Inches(5.3))
c_t2.fill.solid()
c_t2.fill.fore_color.rgb = BLUE_LIGHT
c_t2.line.color.rgb = BLUE_BORDER
tf = c_t2.text_frame
tf.word_wrap = True
p = tf.paragraphs[0]
p.text = "Môi Trường & Quy Trình Thực Nghiệm"
p.font.size = Pt(16)
p.font.bold = True
p.font.color.rgb = ACCENT_BLUE
p = tf.add_paragraph()
p.text = "• Nền tảng ứng dụng: Xây dựng ứng dụng mẫu trên ASP.NET Core MVC (.NET 10).\n• Hệ quản trị CSDL: Kết nối trực tiếp máy chủ Microsoft SQL Server 2025 Developer.\n• Thiết kế đối chứng: Chia 3 phân hệ tương tác riêng biệt cho 3 thành viên, nạp payload và theo dõi phản hồi thời gian thực.\n• Ứng dụng thực tế: Đưa các giải pháp phòng thủ vào thiết kế kiến trúc (Security View) của đề tài BTL."
p.font.size = Pt(12.5)
p.font.color.rgb = TEXT_BODY
p.space_before = Pt(10)

# ==============================================================================
# PHẦN 1: SQL INJECTION — NGUYỄN DANH HỌC (SLIDES 3 - 8)
# ==============================================================================
# Slide 3
s3 = prs.slides.add_slide(blank_layout)
add_header(s3, "1.1. SQL Injection: Bản Chất Kỹ Thuật & Lỗi Ghép Chuỗi", "CHỦ ĐỀ 1 • NGUYỄN DANH HỌC (23A1001D0158)")
c3_1 = s3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.4), Inches(5.65), Inches(5.3))
c3_1.fill.solid()
c3_1.fill.fore_color.rgb = CARD_BG
c3_1.line.color.rgb = CARD_BORDER
tf = c3_1.text_frame
tf.word_wrap = True
p = tf.paragraphs[0]
p.text = "Nguyên Nhân Kỹ Thuật Cốt Lõi"
p.font.size = Pt(15)
p.font.bold = True
p.font.color.rgb = ACCENT_RED
p = tf.add_paragraph()
p.text = "• Thiếu sự phân tách giữa Dữ liệu (Data) và Cú pháp lệnh (Code).\n• Ghép chuỗi đầu vào trực tiếp (String Concatenation / Interpolation) vào truy vấn SQL.\n• Trình thông dịch CSDL không thể phân định ranh giới giữa dữ liệu người dùng và cấu trúc truy vấn của hệ thống.\n• Ký tự điều khiển (', --, ;) làm thay đổi luồng thực thi của câu lệnh trên máy chủ SQL Server."
p.font.size = Pt(12.5)
p.font.color.rgb = TEXT_BODY
p.space_before = Pt(8)

c3_2 = s3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.85), Inches(1.4), Inches(5.65), Inches(5.3))
c3_2.fill.solid()
c3_2.fill.fore_color.rgb = RED_LIGHT
c3_2.line.color.rgb = RED_BORDER
tf = c3_2.text_frame
tf.word_wrap = True
p = tf.paragraphs[0]
p.text = "Minh Họa Đoạn Mã Nguồn Gây Lỗi"
p.font.size = Pt(15)
p.font.bold = True
p.font.color.rgb = ACCENT_RED
p = tf.add_paragraph()
p.text = "// Nối chuỗi trực tiếp từ tham số người dùng:\nstring rawSql = $\"SELECT * FROM Accounts \" +\n                $\"WHERE Username = '{username}' \" +\n                $\"AND Password = '{password}'\";\n\nvar accounts = _context.Accounts\n                       .FromSqlRaw(rawSql)\n                       .ToList();"
p.font.name = "Courier New"
p.font.size = Pt(11)
p.font.color.rgb = RGBColor(159, 18, 57)
p.space_before = Pt(8)
p = tf.add_paragraph()
p.text = "Lỗ hổng phát sinh khi lập trình viên giả định đầu vào luôn là văn bản chuẩn. Ký tự nháy đơn (') sẽ đóng sớm chuỗi giá trị và chuyển phần còn lại thành cấu trúc điều khiển."
p.font.size = Pt(12)
p.font.color.rgb = TEXT_HEAD
p.space_before = Pt(14)

# Slide 4
s4 = prs.slides.add_slide(blank_layout)
add_header(s4, "1.2. SQL Injection: Đánh Giá Tác Động An Ninh (CVSS 9.8)", "CHỦ ĐỀ 1 • NGUYỄN DANH HỌC (23A1001D0158)")
c4_1 = s4.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.4), col_w, Inches(5.3))
c4_1.fill.solid()
c4_1.fill.fore_color.rgb = RED_LIGHT
c4_1.line.color.rgb = RED_BORDER
tf = c4_1.text_frame
tf.word_wrap = True
p = tf.paragraphs[0]
p.text = "VƯỢT QUA XÁC THỰC\n(Authentication Bypass)"
p.font.size = Pt(14)
p.font.bold = True
p.font.color.rgb = ACCENT_RED
p = tf.add_paragraph()
p.text = "• Không cần thông tin xác thực hợp lệ.\n• Bỏ qua bước kiểm tra mật khẩu để đăng nhập tài khoản quản trị.\n• Chiếm đoạt phiên làm việc và thay đổi quyền hạn người dùng trên hệ thống."
p.font.size = Pt(12)
p.font.color.rgb = TEXT_BODY
p.space_before = Pt(10)

c4_2 = s4.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8) + col_w + Inches(0.39), Inches(1.4), col_w, Inches(5.3))
c4_2.fill.solid()
c4_2.fill.fore_color.rgb = AMBER_LIGHT
c4_2.line.color.rgb = AMBER_BORDER
tf = c4_2.text_frame
tf.word_wrap = True
p = tf.paragraphs[0]
p.text = "TRÍCH XUẤT CƠ SỞ DỮ LIỆU\n(Data Exfiltration)"
p.font.size = Pt(14)
p.font.bold = True
p.font.color.rgb = ACCENT_AMBER
p = tf.add_paragraph()
p.text = "• Trích xuất thông tin người dùng, hồ sơ chuyên gia và số dư ví điện tử.\n• Làm lộ mật khẩu và các thông tin cá nhân lưu trong bảng CSDL.\n• Vi phạm quy định về bảo mật dữ liệu người dùng."
p.font.size = Pt(12)
p.font.color.rgb = TEXT_BODY
p.space_before = Pt(10)

c4_3 = s4.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8) + (col_w + Inches(0.39))*2, Inches(1.4), col_w, Inches(5.3))
c4_3.fill.solid()
c4_3.fill.fore_color.rgb = BLUE_LIGHT
c4_3.line.color.rgb = BLUE_BORDER
tf = c4_3.text_frame
tf.word_wrap = True
p = tf.paragraphs[0]
p.text = "THAY ĐỔI / PHÁ HỦY DỮ LIỆU\n(Data Manipulation)"
p.font.size = Pt(14)
p.font.bold = True
p.font.color.rgb = ACCENT_BLUE
p = tf.add_paragraph()
p.text = "• Can thiệp, sửa đổi số dư tài khoản trái phép.\n• Thực thi các câu lệnh DDL nguy hiểm như DROP TABLE hoặc TRUNCATE.\n• Nguy cơ thực thi lệnh hệ điều hành nếu máy chủ CSDL cấp quyền quá mức."
p.font.size = Pt(12)
p.font.color.rgb = TEXT_BODY
p.space_before = Pt(10)

# Slide 5
s5 = prs.slides.add_slide(blank_layout)
add_header(s5, "1.3. SQL Injection: Phân Tích Cấu Trúc Khai Thác", "CHỦ ĐỀ 1 • NGUYỄN DANH HỌC (23A1001D0158)")
c5_main = s5.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.4), Inches(11.733), Inches(5.3))
c5_main.fill.solid()
c5_main.fill.fore_color.rgb = CARD_BG
c5_main.line.color.rgb = CARD_BORDER
tf = c5_main.text_frame
tf.word_wrap = True
p = tf.paragraphs[0]
p.text = "Kịch Bản: Khai Thác Vượt Qua Xác Thực Đăng Nhập"
p.font.size = Pt(15)
p.font.bold = True
p.font.color.rgb = ACCENT_RED

p = tf.add_paragraph()
p.text = "Giá trị Username đưa vào: ' OR '1'='1' --  | Mật khẩu: (tùy ý hoặc để trống)\nTruy vấn được máy chủ SQL Server phân tích & thông dịch:\nSELECT * FROM Accounts WHERE Username = '' OR '1'='1' --' AND Password = 'xyz'"
p.font.name = "Courier New"
p.font.size = Pt(12)
p.font.bold = True
p.font.color.rgb = ACCENT_BLUE
p.space_before = Pt(6)

p = tf.add_paragraph()
p.text = "Phân Tích Thành Phần Chuỗi Khai Thác:\n1. Dấu nháy đơn (') : Đóng sớm chuỗi ký tự của trường Username, chuyển ngữ cảnh về cú pháp SQL.\n2. Mệnh đề OR '1'='1' : Biểu thức logic chân lý (luôn TRUE với mọi bản ghi trong bảng Accounts).\n3. Ký tự chú thích (--) : Vô hiệu hóa mệnh đề kiểm tra mật khẩu phía sau (AND Password = '...')."
p.font.size = Pt(12.5)
p.font.color.rgb = TEXT_BODY
p.space_before = Pt(10)

p = tf.add_paragraph()
p.text = "Mở Rộng Kỹ Thuật Union-Based:\nPayload: ' UNION SELECT Id, Username, Password, SecretNote FROM Accounts --\nCho phép ghép nối kết quả từ bảng dữ liệu nhạy cảm vào màn hình tra cứu thông thường."
p.font.size = Pt(12.5)
p.font.bold = True
p.font.color.rgb = ACCENT_AMBER
p.space_before = Pt(10)

# Slide 6
s6 = prs.slides.add_slide(blank_layout)
add_header(s6, "1.4. SQL Injection: Minh Chứng Thực Nghiệm Trực Quan Trên Web UI", "CHỦ ĐỀ 1 • NGUYỄN DANH HỌC (23A1001D0158)")
p_sqli = "/media/hocjsoo/New Volume/OWASP_Demo/screenshots/only_cards.png"
if os.path.exists(p_sqli):
    s6.shapes.add_picture(p_sqli, Inches(1.8), Inches(1.35), width=Inches(9.7))

callout6 = s6.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(6.05), Inches(11.733), Inches(0.9))
callout6.fill.solid()
callout6.fill.fore_color.rgb = BLUE_LIGHT
callout6.line.color.rgb = BLUE_BORDER
tf = callout6.text_frame
p = tf.paragraphs[0]
p.text = "Minh chứng đối chứng trực tiếp trên ứng dụng (http://localhost:5076/SqlInjection)"
p.font.size = Pt(12)
p.font.bold = True
p.font.color.rgb = ACCENT_BLUE
p = tf.add_paragraph()
p.text = "• Nhánh dính lỗi (Trái): Chuỗi ' OR '1'='1' -- cho phép đăng nhập thành công và trích xuất danh sách tài khoản.\n• Nhánh phòng thủ (Phải): Ứng dụng xử lý an toàn qua EF Core LINQ, từ chối truy cập và bảo vệ dữ liệu."
p.font.size = Pt(11)
p.font.color.rgb = TEXT_BODY
p.space_before = Pt(2)

# Slide 7
s7 = prs.slides.add_slide(blank_layout)
add_header(s7, "1.5. SQL Injection: Minh Chứng Tầng API (cURL) & SQL Server 2025", "CHỦ ĐỀ 1 • NGUYỄN DANH HỌC (23A1001D0158)")
b7_api = s7.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.4), card_w, Inches(4.2))
b7_api.fill.solid()
b7_api.fill.fore_color.rgb = CARD_BG
b7_api.line.color.rgb = CARD_BORDER
tf = b7_api.text_frame
tf.word_wrap = True
p = tf.paragraphs[0]
p.text = "Minh Chứng Tầng API / Dòng Lệnh (cURL)"
p.font.size = Pt(14)
p.font.bold = True
p.font.color.rgb = ACCENT_BLUE
p = tf.add_paragraph()
p.text = "$ curl -X POST http://localhost:5076/SqlInjection/LoginVulnerable \\\n       -F \"username=' OR '1'='1' --\" -F \"password=any\"\n\nPhản hồi JSON từ Server:\n{\n  \"success\": true,\n  \"isExploited\": true,\n  \"count\": 4,\n  \"accounts\": [ ... danh sách tài khoản trích xuất ... ]\n}"
p.font.name = "Courier New"
p.font.size = Pt(10)
p.font.color.rgb = TEXT_HEAD
p.space_before = Pt(6)
p = tf.add_paragraph()
p.text = "Xác nhận lỗi xuất phát từ tầng xử lý C# Backend, kẻ tấn công có thể khai thác qua API mà không cần tương tác qua giao diện web."
p.font.size = Pt(11.5)
p.font.color.rgb = TEXT_BODY
p.space_before = Pt(8)

b7_db = s7.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.85), Inches(1.4), card_w, Inches(4.2))
b7_db.fill.solid()
b7_db.fill.fore_color.rgb = CARD_BG
b7_db.line.color.rgb = CARD_BORDER
tf = b7_db.text_frame
tf.word_wrap = True
p = tf.paragraphs[0]
p.text = "Minh Chứng Tầng CSDL (SQL Server 2025)"
p.font.size = Pt(14)
p.font.bold = True
p.font.color.rgb = ACCENT_GREEN
p = tf.add_paragraph()
p.text = "-- Database: OwaspDemoDB | Bảng: dbo.Accounts\n-- 1. Truy vấn dính lỗi:\nSELECT * FROM Accounts WHERE Username='' OR '1'='1' --'\n>> Kết quả: 4 rows returned (Bypass thành công)\n\n-- 2. Truy vấn tham số hóa:\nEXEC sp_executesql N'SELECT * FROM Accounts WHERE...',\n     N'@p0 nvarchar(50)', @p0 = N''' OR ''1''=''1'' --'\n>> Kết quả: 0 rows returned (An toàn)"
p.font.name = "Courier New"
p.font.size = Pt(10)
p.font.color.rgb = TEXT_HEAD
p.space_before = Pt(6)
p = tf.add_paragraph()
p.text = "Đối chiếu trực tiếp trên máy chủ CSDL chứng minh cơ chế phân tách rạch ròi giữa cú pháp lệnh và tham số giá trị."
p.font.size = Pt(11.5)
p.font.color.rgb = TEXT_BODY
p.space_before = Pt(8)

callout7 = s7.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(5.8), Inches(11.733), Inches(0.95))
callout7.fill.solid()
callout7.fill.fore_color.rgb = GREEN_LIGHT
callout7.line.color.rgb = GREEN_BORDER
tf = callout7.text_frame
p = tf.paragraphs[0]
p.text = "Thao tác trên công cụ: VS Code MSSQL Extension & File verify_sql_server.sql"
p.font.size = Pt(12)
p.font.bold = True
p.font.color.rgb = ACCENT_GREEN
p = tf.add_paragraph()
p.text = "Dữ liệu được truy vấn và kiểm thử trực tiếp trên máy chủ CSDL SQL Server 2025 (Database: OwaspDemoDB)."
p.font.size = Pt(11)
p.font.color.rgb = TEXT_BODY
p.space_before = Pt(2)

# Slide 8
s8 = prs.slides.add_slide(blank_layout)
add_header(s8, "1.6. SQL Injection: Giải Pháp Phòng Thủ Chuẩn Hóa", "CHỦ ĐỀ 1 • NGUYỄN DANH HỌC (23A1001D0158)")
c8_1 = s8.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.4), Inches(5.65), Inches(5.3))
c8_1.fill.solid()
c8_1.fill.fore_color.rgb = GREEN_LIGHT
c8_1.line.color.rgb = GREEN_BORDER
tf = c8_1.text_frame
tf.word_wrap = True
p = tf.paragraphs[0]
p.text = "Kỹ Thuật Parameterized Query Với EF Core"
p.font.size = Pt(15)
p.font.bold = True
p.font.color.rgb = ACCENT_GREEN
p = tf.add_paragraph()
p.text = "// SỬ DỤNG LINQ PARAMETERIZED (KHUYÊN DÙNG):\nvar account = _context.Accounts\n    .FirstOrDefault(a => a.Username == user \n                      && a.Password == pass);\n\n// TRUY VẤN THUẦN QUA FROMSQLINTERPOLATED:\nvar account = _context.Accounts\n    .FromSqlInterpolated($\"SELECT * FROM Accounts WHERE Username={user} AND Password={pass}\")\n    .FirstOrDefault();"
p.font.name = "Courier New"
p.font.size = Pt(10)
p.font.color.rgb = RGBColor(6, 95, 70)
p.space_before = Pt(8)
p = tf.add_paragraph()
p.text = "Cơ chế bảo vệ:\nSQL Server biên dịch kế hoạch thực thi (Execution Plan) trước khi nhận tham số. Đầu vào của người dùng được gán vào biến @p0 như một chuỗi văn bản (literal), vô hiệu hóa hoàn toàn khả năng chèn lệnh."
p.font.size = Pt(12)
p.font.color.rgb = TEXT_BODY
p.space_before = Pt(10)

c8_2 = s8.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.85), Inches(1.4), Inches(5.65), Inches(5.3))
c8_2.fill.solid()
c8_2.fill.fore_color.rgb = CARD_BG
c8_2.line.color.rgb = CARD_BORDER
tf = c8_2.text_frame
tf.word_wrap = True
p = tf.paragraphs[0]
p.text = "Bộ Quy Tắc Phòng Thủ Toàn Diện (Defense-in-Depth)"
p.font.size = Pt(15)
p.font.bold = True
p.font.color.rgb = ACCENT_BLUE
p = tf.add_paragraph()
p.text = "1. Tham số hóa truy vấn:\nTuyệt đối không ghép chuỗi SQL thủ công trong mã nguồn.\n\n2. Đặc quyền tối thiểu (Least Privilege):\nTài khoản kết nối CSDL chỉ cấp các quyền SELECT, INSERT, UPDATE cần thiết; không dùng tài khoản 'sa' hay quyền DDL trong môi trường chạy ứng dụng.\n\n3. Kiểm tra dữ liệu đầu vào (Input Validation):\nSử dụng Data Annotations để giới hạn định dạng và độ dài dữ liệu.\n\n4. Băm mật khẩu (Password Hashing):\nSử dụng ASP.NET Core Identity (PBKDF2) để bảo vệ thông tin ngay cả khi CSDL bị lộ."
p.font.size = Pt(12)
p.font.color.rgb = TEXT_BODY
p.space_before = Pt(8)

# ==============================================================================
# PHẦN 2: STORED XSS — NGUYỄN THANH BÌNH (SLIDES 9 - 14)
# ==============================================================================
# Slide 9
s9 = prs.slides.add_slide(blank_layout)
add_header(s9, "2.1. Stored XSS: Bản Chất Kỹ Thuật & Ngữ Cảnh Ứng Dụng", "CHỦ ĐỀ 2 • NGUYỄN THANH BÌNH (23A1001D0041)")
c9_1 = s9.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.4), Inches(5.65), Inches(5.3))
c9_1.fill.solid()
c9_1.fill.fore_color.rgb = CARD_BG
c9_1.line.color.rgb = CARD_BORDER
tf = c9_1.text_frame
tf.word_wrap = True
p = tf.paragraphs[0]
p.text = "Ngữ Cảnh Nghiệp Vụ Trong Hệ Thống"
p.font.size = Pt(15)
p.font.bold = True
p.font.color.rgb = ACCENT_RED
p = tf.add_paragraph()
p.text = "• Hệ thống cho phép người dùng đăng tải nhận xét, đánh giá chất lượng buổi tư vấn.\n• Bản chất Stored XSS (Persistent XSS): Mã kịch bản (JavaScript) độc hại được gửi lên và lưu trữ trực tiếp trong CSDL (bảng dbo.Comments).\n• Đặc điểm phát tán: Khi người dùng khác truy cập trang web, mã kịch bản lưu trong CSDL được tải về và thực thi tự động trên trình duyệt của nạn nhân mà không cần tương tác thêm."
p.font.size = Pt(12.5)
p.font.color.rgb = TEXT_BODY
p.space_before = Pt(8)

c9_2 = s9.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.85), Inches(1.4), Inches(5.65), Inches(5.3))
c9_2.fill.solid()
c9_2.fill.fore_color.rgb = RED_LIGHT
c9_2.line.color.rgb = RED_BORDER
tf = c9_2.text_frame
tf.word_wrap = True
p = tf.paragraphs[0]
p.text = "Mã Nguồn Minh Họa Lỗi & Sự Lạm Dụng @@Html.Raw"
p.font.size = Pt(15)
p.font.bold = True
p.font.color.rgb = ACCENT_RED
p = tf.add_paragraph()
p.text = "// 1. CONTROLLER LƯU TRỰC TIẾP KHÔNG QUA XỬ LÝ:\nvar comment = new Comment {\n    Author = author,\n    Content = content // Chứa mã <script> độc hại\n};\n_context.Comments.Add(comment);\n\n// 2. RAZOR VIEW HIỂN THỊ DỮ LIỆU THÔ:\n<div class=\"comment-body\">\n    @Html.Raw(comment.Content)\n</div>"
p.font.name = "Courier New"
p.font.size = Pt(10.5)
p.font.color.rgb = RGBColor(159, 18, 57)
p.space_before = Pt(8)
p = tf.add_paragraph()
p.text = "Việc sử dụng @Html.Raw() nhằm mục đích hiển thị định dạng văn bản vô tình làm vô hiệu hóa cơ chế phòng vệ mã hóa HTML tự động của Razor View."
p.font.size = Pt(12)
p.font.color.rgb = TEXT_HEAD
p.space_before = Pt(12)

# Slide 10
s10 = prs.slides.add_slide(blank_layout)
add_header(s10, "2.2. Stored XSS: Đánh Giá Rủi Ro & Tác Động Phía Trình Duyệt", "CHỦ ĐỀ 2 • NGUYỄN THANH BÌNH (23A1001D0041)")
c10_1 = s10.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.4), col_w, Inches(5.3))
c10_1.fill.solid()
c10_1.fill.fore_color.rgb = RED_LIGHT
c10_1.line.color.rgb = RED_BORDER
tf = c10_1.text_frame
tf.word_wrap = True
p = tf.paragraphs[0]
p.text = "ĐÁNH CẮP COOKIE & PHIÊN\n(Session Hijacking)"
p.font.size = Pt(14)
p.font.bold = True
p.font.color.rgb = ACCENT_RED
p = tf.add_paragraph()
p.text = "• Đoạn script đọc thuộc tính document.cookie để trích xuất token phiên làm việc.\n• Dữ liệu xác thực được gửi ngầm về máy chủ điều khiển của kẻ tấn công.\n• Kẻ tấn công sử dụng token để mạo danh người dùng hợp lệ."
p.font.size = Pt(12)
p.font.color.rgb = TEXT_BODY
p.space_before = Pt(10)

c10_2 = s10.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8) + col_w + Inches(0.39), Inches(1.4), col_w, Inches(5.3))
c10_2.fill.solid()
c10_2.fill.fore_color.rgb = AMBER_LIGHT
c10_2.line.color.rgb = AMBER_BORDER
tf = c10_2.text_frame
tf.word_wrap = True
p = tf.paragraphs[0]
p.text = "THAY ĐỔI CẤU TRÚC DOM\n(DOM Manipulation / Phishing)"
p.font.size = Pt(14)
p.font.bold = True
p.font.color.rgb = ACCENT_AMBER
p = tf.add_paragraph()
p.text = "• Thay đổi cấu trúc giao diện trang web, chèn các form đăng nhập giả mạo.\n• Lừa người dùng nhập thông tin nhạy cảm trên giao diện đã bị sửa đổi.\n• Tự động điều hướng người dùng sang các liên kết độc hại."
p.font.size = Pt(12)
p.font.color.rgb = TEXT_BODY
p.space_before = Pt(10)

c10_3 = s10.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8) + (col_w + Inches(0.39))*2, Inches(1.4), col_w, Inches(5.3))
c10_3.fill.solid()
c10_3.fill.fore_color.rgb = BLUE_LIGHT
c10_3.line.color.rgb = BLUE_BORDER
tf = c10_3.text_frame
tf.word_wrap = True
p = tf.paragraphs[0]
p.text = "THỰC THI THAO TÁC TRÁI PHÉP\n(Client-Side Exploitation)"
p.font.size = Pt(14)
p.font.bold = True
p.font.color.rgb = ACCENT_BLUE
p = tf.add_paragraph()
p.text = "• Ghi nhận thao tác bàn phím (Keylogging) khi người dùng nhập liệu.\n• Tự động thực hiện các hành động gửi tin nhắn hoặc đánh giá mạo danh người dùng.\n• Tận dụng quyền truy cập trình duyệt để mở rộng phạm vi tấn công."
p.font.size = Pt(12)
p.font.color.rgb = TEXT_BODY
p.space_before = Pt(10)

# Slide 11
s11 = prs.slides.add_slide(blank_layout)
add_header(s11, "2.3. Stored XSS: Phân Tích Kỹ Thuật Đưa Payload Vào CSDL", "CHỦ ĐỀ 2 • NGUYỄN THANH BÌNH (23A1001D0041)")
c11_main = s11.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.4), Inches(11.733), Inches(5.3))
c11_main.fill.solid()
c11_main.fill.fore_color.rgb = CARD_BG
c11_main.line.color.rgb = CARD_BORDER
tf = c11_main.text_frame
tf.word_wrap = True
p = tf.paragraphs[0]
p.text = "3 Biến Thể Payload XSS Thực Nghiệm Trong Hệ Thống"
p.font.size = Pt(15)
p.font.bold = True
p.font.color.rgb = ACCENT_RED

p = tf.add_paragraph()
p.text = "1. Payload 1 (Thẻ Script Cơ Bản):\n   <script>alert('Session Cookie: ' + document.cookie);</script>\n   Cơ chế: Trình duyệt phân tích cú pháp HTML, nhận diện khối mã và thực thi câu lệnh JavaScript."
p.font.size = Pt(12.5)
p.font.color.rgb = TEXT_BODY
p.space_before = Pt(8)

p = tf.add_paragraph()
p.text = "2. Payload 2 (Vượt Qua Bộ Lọc Bằng Trình Xử Lý Sự Kiện HTML):\n   <img src=x onerror=alert('XSS triggered via image onerror')>\n   Cơ chế: Khi ứng dụng chỉ lọc từ khóa '<script>', kẻ tấn công sử dụng thuộc tính 'onerror' của thẻ hình ảnh khi đường dẫn ảnh không tồn tại để kích hoạt JavaScript."
p.font.size = Pt(12.5)
p.font.color.rgb = TEXT_BODY
p.space_before = Pt(8)

p = tf.add_paragraph()
p.text = "3. Payload 3 (Thay Đổi Toàn Bộ Cấu Trúc Trang Web):\n   <script>document.body.innerHTML='<h2 style=color:red;text-align:center>Trang Web Đã Bị Thay Đổi</h2>';</script>\n   Cơ chế: Can thiệp trực tiếp vào thuộc tính body.innerHTML để thay đổi toàn bộ nội dung hiển thị của website."
p.font.size = Pt(12.5)
p.font.color.rgb = TEXT_BODY
p.space_before = Pt(8)

# Slide 12
s12 = prs.slides.add_slide(blank_layout)
add_header(s12, "2.4. Stored XSS: Minh Chứng Thực Nghiệm Trực Quan Trên Web UI", "CHỦ ĐỀ 2 • NGUYỄN THANH BÌNH (23A1001D0041)")
p_xss = "/media/hocjsoo/New Volume/OWASP_Demo/screenshots/crop_xss_ui.png"
if os.path.exists(p_xss):
    s12.shapes.add_picture(p_xss, Inches(1.8), Inches(1.35), width=Inches(9.7))

callout12 = s12.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(5.95), Inches(11.733), Inches(1.0))
callout12.fill.solid()
callout12.fill.fore_color.rgb = RED_LIGHT
callout12.line.color.rgb = RED_BORDER
tf = callout12.text_frame
p = tf.paragraphs[0]
p.text = "Minh chứng đối chứng chức năng Nhận xét đánh giá chuyên gia (/Xss)"
p.font.size = Pt(12)
p.font.bold = True
p.font.color.rgb = ACCENT_RED
p = tf.add_paragraph()
p.text = "• Nhánh dính lỗi (Trái): Chuỗi kịch bản được lưu và render thô, kích hoạt thông báo alert hiển thị Cookie.\n• Nhánh an toàn (Phải): Razor View tự động mã hóa chuỗi thành văn bản an toàn, triệt tiêu khả năng thực thi mã."
p.font.size = Pt(11)
p.font.color.rgb = TEXT_BODY
p.space_before = Pt(2)

# Slide 13
s13 = prs.slides.add_slide(blank_layout)
add_header(s13, "2.5. Stored XSS: Minh Chứng Dữ Liệu CSDL & Bảo Vệ Cookie", "CHỦ ĐỀ 2 • NGUYỄN THANH BÌNH (23A1001D0041)")
b13_db = s13.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.4), card_w, Inches(5.3))
b13_db.fill.solid()
b13_db.fill.fore_color.rgb = CARD_BG
b13_db.line.color.rgb = CARD_BORDER
tf = b13_db.text_frame
tf.word_wrap = True
p = tf.paragraphs[0]
p.text = "Dữ Liệu Lưu Trong Bảng dbo.Comments"
p.font.size = Pt(14)
p.font.bold = True
p.font.color.rgb = ACCENT_BLUE
p = tf.add_paragraph()
p.text = "SELECT Id, Author, Content, IsSecureStored FROM dbo.Comments;\n\nId | Author        | Content                           | IsSecure\n---+---------------+-----------------------------------+---------\n 1 | Le Thanh Binh | Dich vu tu van rat chuyen nghiep! | 1\n 2 | Hacker_XSS    | <script>alert(document.cookie);.. | 0"
p.font.name = "Courier New"
p.font.size = Pt(9.5)
p.font.color.rgb = TEXT_HEAD
p.space_before = Pt(6)
p = tf.add_paragraph()
p.text = "Phân tích đặc tính lưu trữ:\nMã kịch bản độc hại được lưu trữ trực tiếp trong CSDL. Do đó, lỗ hổng Stored XSS có tính chất nguy hại kéo dài và phát tán tới mọi người dùng xem nội dung đó."
p.font.size = Pt(11.5)
p.font.color.rgb = TEXT_BODY
p.space_before = Pt(8)

b13_c = s13.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.85), Inches(1.4), card_w, Inches(5.3))
b13_c.fill.solid()
b13_c.fill.fore_color.rgb = CARD_BG
b13_c.line.color.rgb = CARD_BORDER
tf = b13_c.text_frame
tf.word_wrap = True
p = tf.paragraphs[0]
p.text = "Cơ Chế Bảo Vệ Với Cờ HttpOnly Cookie"
p.font.size = Pt(14)
p.font.bold = True
p.font.color.rgb = ACCENT_GREEN
p = tf.add_paragraph()
p.text = "// CẤU HÌNH COOKIE AN TOÀN TRONG ASP.NET CORE:\nResponse.Cookies.Append(\"AuthSessionToken\", token, new CookieOptions\n{\n    HttpOnly = true, // CẤM JAVASCRIPT ĐỌC TRỘM\n    Secure = true,   // BẮT BUỘC TRUYỀN QUA HTTPS\n    SameSite = SameSiteMode.Strict\n});"
p.font.name = "Courier New"
p.font.size = Pt(9.5)
p.font.color.rgb = RGBColor(6, 95, 70)
p.space_before = Pt(6)
p = tf.add_paragraph()
p.text = "Vai trò của cờ HttpOnly:\nNgăn chặn mã JavaScript đọc thuộc tính document.cookie. Kể cả khi trang web xuất hiện lỗi XSS, kẻ tấn công cũng không thể trích xuất token phiên làm việc."
p.font.size = Pt(11.5)
p.font.color.rgb = TEXT_BODY
p.space_before = Pt(8)

# Slide 14
s14 = prs.slides.add_slide(blank_layout)
add_header(s14, "2.6. Stored XSS: Bộ Giải Pháp Phòng Thủ Chuẩn Trong ASP.NET Core", "CHỦ ĐỀ 2 • NGUYỄN THANH BÌNH (23A1001D0041)")
c14_main = s14.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.4), Inches(11.733), Inches(5.3))
c14_main.fill.solid()
c14_main.fill.fore_color.rgb = GREEN_LIGHT
c14_main.line.color.rgb = GREEN_BORDER
tf = c14_main.text_frame
tf.word_wrap = True
p = tf.paragraphs[0]
p.text = "CÁC NGUYÊN TẮC PHÒNG THỦ XSS THEO KHUYẾN NGHỊ OWASP:"
p.font.size = Pt(15)
p.font.bold = True
p.font.color.rgb = ACCENT_GREEN

p = tf.add_paragraph()
p.text = "1. Tận dụng cơ chế mã hóa tự động của Razor View:\n   Sử dụng cú pháp mặc định @model.Property thay vì @Html.Raw() đối với mọi dữ liệu bắt nguồn từ người dùng. Các ký tự nhạy cảm (<, >, \", ') tự động được chuyển đổi sang thực thể HTML an toàn."
p.font.size = Pt(13)
p.font.color.rgb = TEXT_BODY
p.space_before = Pt(8)

p = tf.add_paragraph()
p.text = "2. Mã hóa dữ liệu với System.Text.Encodings.Web:\n   Sử dụng HtmlEncoder.Default.Encode() để xử lý dữ liệu trước khi lưu trữ hoặc phản hồi."
p.font.size = Pt(13)
p.font.color.rgb = TEXT_BODY
p.space_before = Pt(8)

p = tf.add_paragraph()
p.text = "3. Thiết lập thuộc tính an toàn cho Cookie:\n   Kích hoạt HttpOnly = true và Secure = true để bảo vệ token phiên làm việc khỏi sự can thiệp của mã kịch bản phía client."
p.font.size = Pt(13)
p.font.color.rgb = TEXT_BODY
p.space_before = Pt(8)

p = tf.add_paragraph()
p.text = "4. Triển khai Content Security Policy (CSP Header):\n   Cấu hình chính sách bảo mật nội dung để hạn chế phạm vi thực thi kịch bản và chặn các nguồn tài nguyên không tin cậy."
p.font.size = Pt(13)
p.font.color.rgb = TEXT_BODY
p.space_before = Pt(8)

# ==============================================================================
# PHẦN 3: CSRF ATTACK — NGUYỄN MINH CƯỜNG (SLIDES 15 - 20)
# ==============================================================================
# Slide 15
s15 = prs.slides.add_slide(blank_layout)
add_header(s15, "3.1. CSRF Attack: Bản Chất Kỹ Thuật & Cơ Chế Giả Mạo Request", "CHỦ ĐỀ 3 • NGUYỄN MINH CƯỜNG (23A1001D0058)")
c15_1 = s15.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.4), Inches(5.65), Inches(5.3))
c15_1.fill.solid()
c15_1.fill.fore_color.rgb = CARD_BG
c15_1.line.color.rgb = CARD_BORDER
tf = c15_1.text_frame
tf.word_wrap = True
p = tf.paragraphs[0]
p.text = "Bản Chất Kỹ Thuật Của Tấn Công CSRF"
p.font.size = Pt(15)
p.font.bold = True
p.font.color.rgb = ACCENT_RED
p = tf.add_paragraph()
p.text = "• CSRF (A01:2021 - Broken Access Control) là kỹ thuật tấn công mượn quyền người dùng hợp lệ.\n• Điều kiện tiên quyết:\n  1. Nạn nhân có phiên đăng nhập hợp lệ tại ứng dụng mục tiêu (Cookie xác thực còn hiệu lực).\n  2. Trình duyệt tự động đính kèm Cookie xác thực theo mọi HTTP Request gửi tới tên miền của ứng dụng.\n  3. Kẻ tấn công lừa nạn nhân truy cập một trang web độc hại do kẻ tấn công kiểm soát.\n• Điểm mù kiểm soát: Máy chủ chỉ xác thực Cookie của người dùng mà không kiểm tra nguồn gốc phát sinh Request."
p.font.size = Pt(12.5)
p.font.color.rgb = TEXT_BODY
p.space_before = Pt(8)

c15_2 = s15.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.85), Inches(1.4), Inches(5.65), Inches(5.3))
c15_2.fill.solid()
c15_2.fill.fore_color.rgb = RED_LIGHT
c15_2.line.color.rgb = RED_BORDER
tf = c15_2.text_frame
tf.word_wrap = True
p = tf.paragraphs[0]
p.text = "Action Xử Lý Giao Dịch Thiếu Kiểm Soát"
p.font.size = Pt(15)
p.font.bold = True
p.font.color.rgb = ACCENT_RED
p = tf.add_paragraph()
p.text = "// ACTION NHẬN YÊU CẦU KHÔNG CÓ BẢO VỆ TOKEN:\n[HttpPost]\npublic IActionResult TransferVulnerable(decimal amount)\n{\n    var victim = _context.Wallets.Find(1);\n    var hacker = _context.Wallets.Find(2);\n    \n    victim.Balance -= amount;\n    hacker.Balance += amount;\n    _context.SaveChanges();\n    return RedirectToAction(\"Index\");\n}"
p.font.name = "Courier New"
p.font.size = Pt(10)
p.font.color.rgb = RGBColor(159, 18, 57)
p.space_before = Pt(8)
p = tf.add_paragraph()
p.text = "Khi Action không yêu cầu mã xác thực chống giả mạo, các yêu cầu POST bắt nguồn từ một website khác vẫn được máy chủ thực thi bình thường do Cookie xác thực được đính kèm tự động."
p.font.size = Pt(12)
p.font.color.rgb = TEXT_HEAD
p.space_before = Pt(10)

# Slide 16
s16 = prs.slides.add_slide(blank_layout)
add_header(s16, "3.2. CSRF Attack: Kịch Bản Khai Thác Bằng Form Ẩn", "CHỦ ĐỀ 3 • NGUYỄN MINH CƯỜNG (23A1001D0058)")
c16_main = s16.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.4), Inches(11.733), Inches(5.3))
c16_main.fill.solid()
c16_main.fill.fore_color.rgb = RED_LIGHT
c16_main.line.color.rgb = RED_BORDER
tf = c16_main.text_frame
tf.word_wrap = True
p = tf.paragraphs[0]
p.text = "Kịch Bản: Trang Web Giả Mạo Điều Khiển Yêu Cầu Chuyển Tiền"
p.font.size = Pt(16)
p.font.bold = True
p.font.color.rgb = ACCENT_RED

p = tf.add_paragraph()
p.text = "1. Nạn nhân có phiên làm việc đang hoạt động trên hệ thống tư vấn trực tuyến.\n2. Nạn nhân được dẫn dụ truy cập một trang web độc hại ở một tab khác (AttackerSite.cshtml).\n3. Trang web độc hại chứa một biểu mẫu ẩn trỏ trực tiếp tới địa chỉ nhận lệnh của ứng dụng mục tiêu:"
p.font.size = Pt(13)
p.font.color.rgb = TEXT_BODY
p.space_before = Pt(8)

p = tf.add_paragraph()
p.text = "<!-- BIỂU MẪU ẨN TRÊN TRANG WEB CỦA KẺ TẤN CÔNG -->\n<form id=\"exploit\" action=\"http://localhost:5076/Csrf/TransferVulnerable\" method=\"POST\">\n    <input type=\"hidden\" name=\"amount\" value=\"20000000\" />\n</form>\n<script>document.getElementById('exploit').submit();</script>"
p.font.name = "Courier New"
p.font.size = Pt(11)
p.font.bold = True
p.font.color.rgb = RGBColor(159, 18, 57)
p.space_before = Pt(6)

p = tf.add_paragraph()
p.text = "Hậu quả thực tế: Yêu cầu giao dịch được thực hiện mà không có sự đồng thuận chủ động của người dùng, dẫn đến thay đổi số dư hoặc cấu hình tài khoản."
p.font.size = Pt(13)
p.font.bold = True
p.font.color.rgb = ACCENT_RED
p.space_before = Pt(8)

# Slide 17
s17 = prs.slides.add_slide(blank_layout)
add_header(s17, "3.3. CSRF Attack: Minh Chứng Thực Nghiệm Trực Quan Trên Web UI", "CHỦ ĐỀ 3 • NGUYỄN MINH CƯỜNG (23A1001D0058)")
p_csrf = "/media/hocjsoo/New Volume/OWASP_Demo/screenshots/crop_csrf_ui.png"
if os.path.exists(p_csrf):
    s17.shapes.add_picture(p_csrf, Inches(1.8), Inches(1.35), width=Inches(9.7))

callout17 = s17.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(5.95), Inches(11.733), Inches(1.0))
callout17.fill.solid()
callout17.fill.fore_color.rgb = GREEN_LIGHT
callout17.line.color.rgb = GREEN_BORDER
tf = callout17.text_frame
p = tf.paragraphs[0]
p.text = "Minh chứng đối chứng phân hệ quản lý ví tiền & kịch bản CSRF (/Csrf)"
p.font.size = Pt(12)
p.font.bold = True
p.font.color.rgb = ACCENT_GREEN
p = tf.add_paragraph()
p.text = "• Mở Attacker Site: Biểu mẫu ẩn tự động gửi POST sang Action dính lỗi -> Số dư ví bị trừ 20.000.000 VNĐ.\n• Nhánh đã phòng thủ: Biểu mẫu chính chủ được tích hợp Anti-Forgery Token -> Từ chối 100% yêu cầu giả mạo từ bên ngoài."
p.font.size = Pt(11)
p.font.color.rgb = TEXT_BODY
p.space_before = Pt(2)

# Slide 18
s18 = prs.slides.add_slide(blank_layout)
add_header(s18, "3.4. CSRF Attack: Minh Chứng Biến Động CSDL dbo.UserWallets", "CHỦ ĐỀ 3 • NGUYỄN MINH CƯỜNG (23A1001D0058)")
b18_1 = s18.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.4), card_w, Inches(5.3))
b18_1.fill.solid()
b18_1.fill.fore_color.rgb = CARD_BG
b18_1.line.color.rgb = CARD_BORDER
tf = b18_1.text_frame
tf.word_wrap = True
p = tf.paragraphs[0]
p.text = "1. Trạng Thái CSDL Ban Đầu"
p.font.size = Pt(14)
p.font.bold = True
p.font.color.rgb = ACCENT_BLUE
p = tf.add_paragraph()
p.text = "SELECT Id, OwnerName, Balance FROM dbo.UserWallets;\n\nId | OwnerName                      | Balance (VND)\n---+--------------------------------+---------------\n 1 | Nạn nhân (Nguyễn Danh Học)     | 50,000,000.00\n 2 | Kẻ tấn công (Hacker BlackHat)  |          0.00\n(2 rows affected)"
p.font.name = "Courier New"
p.font.size = Pt(9.5)
p.font.color.rgb = TEXT_HEAD
p.space_before = Pt(6)
p = tf.add_paragraph()
p.text = "Số dư ví tài khoản nạn nhân ở trạng thái ban đầu là 50 triệu VNĐ, tài khoản đối tượng tấn công là 0 VNĐ."
p.font.size = Pt(11.5)
p.font.color.rgb = TEXT_BODY
p.space_before = Pt(8)

b18_2 = s18.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.85), Inches(1.4), card_w, Inches(5.3))
b18_2.fill.solid()
b18_2.fill.fore_color.rgb = RED_LIGHT
b18_2.line.color.rgb = RED_BORDER
tf = b18_2.text_frame
tf.word_wrap = True
p = tf.paragraphs[0]
p.text = "2. Trạng Thái CSDL Sau Khai Thác CSRF"
p.font.size = Pt(14)
p.font.bold = True
p.font.color.rgb = ACCENT_RED
p = tf.add_paragraph()
p.text = "SELECT Id, OwnerName, Balance FROM dbo.UserWallets;\n\nId | OwnerName                      | Balance (VND)\n---+--------------------------------+---------------\n 1 | Nạn nhân (Nguyễn Danh Học)     | 30,000,000.00\n 2 | Kẻ tấn công (Hacker BlackHat)  | 20,000,000.00\n(2 rows affected)"
p.font.name = "Courier New"
p.font.size = Pt(9.5)
p.font.color.rgb = RGBColor(159, 18, 57)
p.space_before = Pt(6)
p = tf.add_paragraph()
p.text = "Số dư thực tế trong CSDL bị thay đổi hoàn toàn tự động khi nạn nhân truy cập liên kết độc hại, chứng minh mức độ rủi ro nghiêm trọng của giao dịch thiếu token."
p.font.size = Pt(11.5)
p.font.color.rgb = TEXT_BODY
p.space_before = Pt(8)

# Slide 19
s19 = prs.slides.add_slide(blank_layout)
add_header(s19, "3.5. CSRF Attack: Cơ Chế Phòng Thủ Bằng Anti-Forgery Token", "CHỦ ĐỀ 3 • NGUYỄN MINH CƯỜNG (23A1001D0058)")
c19_1 = s19.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.4), Inches(5.65), Inches(5.3))
c19_1.fill.solid()
c19_1.fill.fore_color.rgb = GREEN_LIGHT
c19_1.line.color.rgb = GREEN_BORDER
tf = c19_1.text_frame
tf.word_wrap = True
p = tf.paragraphs[0]
p.text = "Mô Hình Synchronizer Token Pattern"
p.font.size = Pt(15)
p.font.bold = True
p.font.color.rgb = ACCENT_GREEN
p = tf.add_paragraph()
p.text = "// 1. TRONG RAZOR VIEW: NHÚNG TOKEN VÀO FORM\n<form asp-action=\"TransferSecure\" method=\"post\">\n    @Html.AntiForgeryToken()\n    <input type=\"number\" name=\"amount\" />\n    <button type=\"submit\">Chuyển tiền</button>\n</form>\n\n// 2. TRONG CONTROLLER: BẮT BUỘC KIỂM TRA TOKEN\n[HttpPost]\n[ValidateAntiForgeryToken]\npublic IActionResult TransferSecure(decimal amount)\n{\n    // Chỉ thực thi khi Token từ Form khớp với Token trong Cookie\n}"
p.font.name = "Courier New"
p.font.size = Pt(9.5)
p.font.color.rgb = RGBColor(6, 95, 70)
p.space_before = Pt(8)

c19_2 = s19.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.85), Inches(1.4), Inches(5.65), Inches(5.3))
c19_2.fill.solid()
c19_2.fill.fore_color.rgb = CARD_BG
c19_2.line.color.rgb = CARD_BORDER
tf = c19_2.text_frame
tf.word_wrap = True
p = tf.paragraphs[0]
p.text = "Cơ Chế Ngăn Chặn Yêu Cầu Giả Mạo"
p.font.size = Pt(15)
p.font.bold = True
p.font.color.rgb = ACCENT_BLUE
p = tf.add_paragraph()
p.text = "1. Nguyên tắc cặp Token:\nMáy chủ phát hành một cặp token: một giá trị lưu trong Cookie phiên và một giá trị nhúng vào trường ẩn của biểu mẫu.\n\n2. Bảo vệ bằng Chính sách Cùng nguồn gốc (SOP):\nWebsite độc hại ở tên miền khác bị trình duyệt ngăn cấm truy cập dữ liệu DOM của biểu mẫu hợp lệ, do đó không thể sao chép token.\n\n3. Thuộc tính SameSite cho Cookie:\nThiết lập SameSite = Strict hoặc Lax giúp kiểm soát việc trình duyệt tự động gửi kèm cookie trong các yêu cầu cross-site."
p.font.size = Pt(12)
p.font.color.rgb = TEXT_BODY
p.space_before = Pt(10)

# Slide 20
s20 = prs.slides.add_slide(blank_layout)
add_header(s20, "3.6. CSRF Attack: Tích Hợp Token Trong AJAX & Fetch API (Buổi 9)", "CHỦ ĐỀ 3 • NGUYỄN MINH CƯỜNG (23A1001D0058)")
c20_main = s20.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.4), Inches(11.733), Inches(5.3))
c20_main.fill.solid()
c20_main.fill.fore_color.rgb = CARD_BG
c20_main.line.color.rgb = CARD_BORDER
tf = c20_main.text_frame
tf.word_wrap = True
p = tf.paragraphs[0]
p.text = "Kỹ Thuật Đính Kèm Token Khi Sử Dụng Fetch API Không Tải Lại Trang"
p.font.size = Pt(15)
p.font.bold = True
p.font.color.rgb = ACCENT_BLUE

p = tf.add_paragraph()
p.text = "// TRÍCH XUẤT TOKEN TỪ DOM VÀ ĐÍNH KÈM VÀO HTTP HEADER:\nconst token = document.querySelector('input[name=\"__RequestVerificationToken\"]').value;\n\nfetch('/Csrf/TransferAjax', {\n    method: 'POST',\n    headers: {\n        'Content-Type': 'application/json',\n        'RequestVerificationToken': token // TRUYỀN TOKEN QUA HEADER\n    },\n    body: JSON.stringify({ amount: 5000000 })\n})\n.then(response => response.json());"
p.font.name = "Courier New"
p.font.size = Pt(10)
p.font.color.rgb = TEXT_HEAD
p.space_before = Pt(8)

p = tf.add_paragraph()
p.text = "Ứng dụng trong đồ án BTL:\nĐối với các tác vụ bất đồng bộ (như đặt lịch hẹn, cập nhật ca làm việc) sử dụng Fetch API, hệ thống đính kèm header RequestVerificationToken để đảm bảo an toàn giao dịch mà không làm suy giảm trải nghiệm người dùng."
p.font.size = Pt(12.5)
p.font.color.rgb = TEXT_BODY
p.space_before = Pt(12)

# ==============================================================================
# TỔNG KẾT & ĐỒ ÁN BTL — 2 SLIDES
# ==============================================================================
# Slide 21
s21 = prs.slides.add_slide(blank_layout)
add_header(s21, "Tổng Kết: So Sánh Đối Chứng 3 Lỗ Hổng An Ninh Web")
table_shape = s21.shapes.add_table(4, 4, Inches(0.8), Inches(1.5), Inches(11.733), Inches(5.2))
tbl = table_shape.table

headers = ["Tiêu chí phân tích", "SQL Injection (Học)", "Stored XSS (Bình)", "CSRF Attack (Cường)"]
for i, h in enumerate(headers):
    cell = tbl.cell(0, i)
    cell.fill.solid()
    cell.fill.fore_color.rgb = ACCENT_BLUE
    p = cell.text_frame.paragraphs[0]
    p.text = h
    p.font.bold = True
    p.font.size = Pt(12)
    p.font.color.rgb = WHITE

rows_data = [
    ("Mục tiêu tấn công", "Tầng Cơ sở dữ liệu (Database Engine)", "Tầng Trình duyệt người dùng (Client DOM)", "Quyền hạn phiên làm việc (Session/Cookie)"),
    ("Mức độ rủi ro", "CRITICAL (9.8 / 10.0)", "HIGH (7.5 - 8.5 / 10.0)", "MEDIUM / HIGH (6.5 - 8.0 / 10.0)"),
    ("Giải pháp phòng thủ", "Parameterized Query / EF Core LINQ", "Razor HTML Encoding + HttpOnly Cookie", "Anti-Forgery Token + SameSite Cookie")
]

for row_idx, rdata in enumerate(rows_data, start=1):
    for col_idx, text in enumerate(rdata):
        cell = tbl.cell(row_idx, col_idx)
        cell.fill.solid()
        cell.fill.fore_color.rgb = CARD_BG if row_idx % 2 == 1 else WHITE
        p = cell.text_frame.paragraphs[0]
        p.text = text
        p.font.size = Pt(11)
        p.font.color.rgb = TEXT_HEAD if col_idx == 0 else TEXT_BODY
        if col_idx == 1 and "CRITICAL" in text:
            p.font.bold = True
            p.font.color.rgb = ACCENT_RED

# Slide 22
s22 = prs.slides.add_slide(blank_layout)
add_header(s22, "Áp Dụng Kiến Trúc An Toàn Vào Đồ Án Website Tư Vấn Trực Tuyến")
main_c22 = s22.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.4), Inches(11.733), Inches(5.3))
main_c22.fill.solid()
main_c22.fill.fore_color.rgb = CARD_BG
main_c22.line.color.rgb = CARD_BORDER
tf = main_c22.text_frame
tf.word_wrap = True

p = tf.paragraphs[0]
p.text = "ÁP DỤNG KIẾN TRÚC BẢO MẬT ĐA TẦNG CHO BTL (WNC.G01):"
p.font.size = Pt(16)
p.font.bold = True
p.font.color.rgb = ACCENT_BLUE

p = tf.add_paragraph()
p.text = "1. Kiến trúc Security by Design (Buổi 8): Tích hợp đồng bộ các lớp bảo vệ ngay từ giai đoạn thiết kế hệ thống, phân quyền chặt chẽ theo 3 Role (Customer, Specialist, Admin).\n\n2. Tầng CSDL (Nguyễn Danh Học): Toàn bộ thao tác nghiệp vụ khai thác CSDL qua Entity Framework Core Code-First với LINQ Parameterized, loại bỏ hoàn toàn việc ghép chuỗi truy vấn.\n\n3. Tầng Giao diện (Nguyễn Thanh Bình): Dữ liệu phản hồi và đánh giá chuyên gia được Razor View mã hóa tự động, Cookie thiết lập đầy đủ thuộc tính HttpOnly và Secure.\n\n4. Tầng Giao dịch (Nguyễn Minh Cường): Bắt buộc kiểm tra Anti-Forgery Token cho toàn bộ các biểu mẫu POST và yêu cầu Fetch API đặt lịch hẹn, giao dịch ví."
p.font.size = Pt(12.5)
p.font.color.rgb = TEXT_BODY
p.space_before = Pt(10)

p = tf.add_paragraph()
p.text = "Kính chúc thầy sức khỏe. Nhóm WNC.G01 xin trân trọng cảm ơn!"
p.font.size = Pt(14)
p.font.bold = True
p.font.color.rgb = ACCENT_RED
p.space_before = Pt(18)

out_file = "/media/hocjsoo/New Volume/OWASP_Demo/Slide_OWASP_Top3_Nhom_WNC_G01.pptx"
prs.save(out_file)
print(f"Academic presentation saved successfully to: {out_file}")
