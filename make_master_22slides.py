import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)

# Bảng màu Gamma / Canva Light Theme (Chuẩn máy chiếu giảng đường)
BG_WHITE = RGBColor(255, 255, 255)
WHITE = RGBColor(255, 255, 255)
CARD_BG = RGBColor(248, 250, 252)          # Slate 50
CARD_BORDER = RGBColor(226, 232, 240)      # Slate 200

TEXT_HEAD = RGBColor(15, 23, 42)           # Slate 900
TEXT_BODY = RGBColor(51, 65, 85)           # Slate 700
TEXT_MUTED = RGBColor(100, 116, 139)       # Slate 500

ACCENT_BLUE = RGBColor(29, 78, 216)        # Blue 700 (HOU)
BLUE_LIGHT = RGBColor(239, 246, 255)       # Blue 50
BLUE_BORDER = RGBColor(191, 219, 254)

ACCENT_RED = RGBColor(190, 18, 60)         # Rose 700 (Danger)
RED_LIGHT = RGBColor(255, 241, 242)        # Rose 50
RED_BORDER = RGBColor(254, 205, 211)

ACCENT_GREEN = RGBColor(4, 120, 87)        # Emerald 700 (Secure)
GREEN_LIGHT = RGBColor(236, 253, 245)      # Emerald 50
GREEN_BORDER = RGBColor(167, 243, 208)

ACCENT_AMBER = RGBColor(180, 83, 9)        # Amber 700
AMBER_LIGHT = RGBColor(255, 251, 235)      # Amber 50
AMBER_BORDER = RGBColor(254, 240, 138)

blank_layout = prs.slide_layouts[6]

def add_header(slide, title_text, category_tag="OWASP TOP 10 • THỰC NGHIỆM AN TOÀN WEB"):
    top_line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(0.08))
    top_line.fill.solid()
    top_line.fill.fore_color.rgb = ACCENT_BLUE
    top_line.line.fill.background()

    badge = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(0.32), Inches(4.8), Inches(0.38))
    badge.fill.solid()
    badge.fill.fore_color.rgb = BLUE_LIGHT
    badge.line.color.rgb = BLUE_BORDER

    tf_b = badge.text_frame
    tf_b.word_wrap = False
    p_b = tf_b.paragraphs[0]
    p_b.text = category_tag.upper()
    p_b.font.size = Pt(10)
    p_b.font.bold = True
    p_b.font.color.rgb = ACCENT_BLUE

    tx = slide.shapes.add_textbox(Inches(0.8), Inches(0.72), Inches(11.7), Inches(0.6))
    tf = tx.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = title_text
    p.font.size = Pt(20)
    p.font.bold = True
    p.font.color.rgb = TEXT_HEAD

    footer = slide.shapes.add_textbox(Inches(0.8), Inches(7.1), Inches(11.7), Inches(0.3))
    ft = footer.text_frame
    p_ft = ft.paragraphs[0]
    p_ft.text = "Nhóm WNC.G01 (Học - Bình - Cường) • Khoa CNTT - Đại học Mở Hà Nội • GVHD: ThS. Lê Hữu Dũng"
    p_ft.font.size = Pt(9.5)
    p_ft.font.color.rgb = TEXT_MUTED

# ==============================================================================
# SLIDE 1: BÌA
# ==============================================================================
s1 = prs.slides.add_slide(blank_layout)
bar1 = s1.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(0.12))
bar1.fill.solid()
bar1.fill.fore_color.rgb = ACCENT_BLUE
bar1.line.fill.background()

u_badge = s1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.0), Inches(0.7), Inches(5.8), Inches(0.42))
u_badge.fill.solid()
u_badge.fill.fore_color.rgb = BLUE_LIGHT
u_badge.line.color.rgb = BLUE_BORDER
p = u_badge.text_frame.paragraphs[0]
p.text = "TRƯỜNG ĐẠI HỌC MỞ HÀ NỘI — KHOA CÔNG NGHỆ THÔNG TIN"
p.font.size = Pt(10.5)
p.font.bold = True
p.font.color.rgb = ACCENT_BLUE

tb1 = s1.shapes.add_textbox(Inches(1.0), Inches(1.25), Inches(11.333), Inches(3.2))
tf1 = tb1.text_frame
tf1.word_wrap = True
p = tf1.paragraphs[0]
p.text = "BÁO CÁO THỰC NGHIỆM AN TOÀN WEB (OWASP TOP 10)"
p.font.size = Pt(16)
p.font.bold = True
p.font.color.rgb = ACCENT_RED

p = tf1.add_paragraph()
p.text = "Bộ 3 Lỗ Hổng Kinh Điển & Cơ Chế Phòng Thủ\n(SQL Injection • Stored XSS • CSRF Attack)"
p.font.size = Pt(30)
p.font.bold = True
p.font.color.rgb = TEXT_HEAD
p.space_before = Pt(8)

p = tf1.add_paragraph()
p.text = "Môn học: Lập trình Web nâng cao  •  Giảng viên hướng dẫn: ThS. Lê Hữu Dũng"
p.font.size = Pt(14)
p.font.color.rgb = TEXT_BODY
p.space_before = Pt(8)

col_w = Inches(3.64)
gap = Inches(0.2)

m1 = s1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.0), Inches(4.7), col_w, Inches(2.0))
m1.fill.solid()
m1.fill.fore_color.rgb = CARD_BG
m1.line.color.rgb = BLUE_BORDER
tf = m1.text_frame
tf.word_wrap = True
p = tf.paragraphs[0]
p.text = "CHỦ ĐỀ 1: SQL INJECTION"
p.font.size = Pt(11)
p.font.bold = True
p.font.color.rgb = ACCENT_BLUE
p = tf.add_paragraph()
p.text = "Nguyễn Danh Học\nMSV: 23A1001D0158\nTrưởng nhóm • Demo CSDL & API"
p.font.size = Pt(12)
p.font.color.rgb = TEXT_HEAD
p.space_before = Pt(4)

m2 = s1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.0) + col_w + gap, Inches(4.7), col_w, Inches(2.0))
m2.fill.solid()
m2.fill.fore_color.rgb = CARD_BG
m2.line.color.rgb = RED_BORDER
tf = m2.text_frame
tf.word_wrap = True
p = tf.paragraphs[0]
p.text = "CHỦ ĐỀ 2: STORED XSS"
p.font.size = Pt(11)
p.font.bold = True
p.font.color.rgb = ACCENT_RED
p = tf.add_paragraph()
p.text = "Nguyễn Thanh Bình\nMSV: 23A1001D0041\nThành viên • Demo Cookie & DOM"
p.font.size = Pt(12)
p.font.color.rgb = TEXT_HEAD
p.space_before = Pt(4)

m3 = s1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.0) + (col_w + gap)*2, Inches(4.7), col_w, Inches(2.0))
m3.fill.solid()
m3.fill.fore_color.rgb = CARD_BG
m3.line.color.rgb = GREEN_BORDER
tf = m3.text_frame
tf.word_wrap = True
p = tf.paragraphs[0]
p.text = "CHỦ ĐỀ 3: CSRF ATTACK"
p.font.size = Pt(11)
p.font.bold = True
p.font.color.rgb = ACCENT_GREEN
p = tf.add_paragraph()
p.text = "Nguyễn Minh Cường\nMSV: 23A1001D0058\nThành viên • Demo Form ẩn giả mạo"
p.font.size = Pt(12)
p.font.color.rgb = TEXT_HEAD
p.space_before = Pt(4)

# ==============================================================================
# SLIDE 2: MỤC TIÊU & PHƯƠNG CHÂM
# ==============================================================================
s2 = prs.slides.add_slide(blank_layout)
add_header(s2, "Mục Tiêu & Phương Châm Báo Cáo Thực Nghiệm OWASP")

c_t1 = s2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.4), Inches(5.65), Inches(5.3))
c_t1.fill.solid()
c_t1.fill.fore_color.rgb = CARD_BG
c_t1.line.color.rgb = CARD_BORDER
tf = c_t1.text_frame
tf.word_wrap = True
p = tf.paragraphs[0]
p.text = "🎯 Yêu Cầu Của Giảng Viên (ThS. Lê Hữu Dũng)"
p.font.size = Pt(16)
p.font.bold = True
p.font.color.rgb = ACCENT_RED
p = tf.add_paragraph()
p.text = "• Không dịch lý thuyết tiếng Anh sách vở.\n• Thấy được lỗ hổng thực sự nguy hiểm đến mức nào trong môi trường thực tế.\n• Bắt buộc phải thực nghiệm demo: Muốn phòng thủ an toàn, phải hiểu hacker tấn công hệ thống theo cách nào.\n• Báo cáo đối chứng: Minh chứng rõ ràng giữa code dính lỗi và code đã vá chuẩn."
p.font.size = Pt(13)
p.font.color.rgb = TEXT_BODY
p.space_before = Pt(10)

c_t2 = s2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.85), Inches(1.4), Inches(5.65), Inches(5.3))
c_t2.fill.solid()
c_t2.fill.fore_color.rgb = BLUE_LIGHT
c_t2.line.color.rgb = BLUE_BORDER
tf = c_t2.text_frame
tf.word_wrap = True
p = tf.paragraphs[0]
p.text = "🚀 Phương Pháp Thực Hiện Của Nhóm WNC.G01"
p.font.size = Pt(16)
p.font.bold = True
p.font.color.rgb = ACCENT_BLUE
p = tf.add_paragraph()
p.text = "• Xây dựng ứng dụng thực nghiệm độc lập (OWASP_Demo) trên ASP.NET Core MVC .NET 10.\n• Kết nối trực tiếp máy chủ CSDL Microsoft SQL Server 2025 Developer trên Localhost.\n• Tạo sẵn 3 phân hệ tương tác riêng biệt cho 3 thành viên: Bấm 1 click nạp payload và quan sát kết quả tức thì.\n• Áp dụng kết quả thực nghiệm vào thiết kế kiến trúc bảo mật (Security View) của Đồ án BTL."
p.font.size = Pt(13)
p.font.color.rgb = TEXT_BODY
p.space_before = Pt(10)

# ==============================================================================
# PHẦN 1: SQL INJECTION — NGUYỄN DANH HỌC (SLIDES 3 - 8)
# ==============================================================================
# Slide 3: Bản chất kỹ thuật
s3 = prs.slides.add_slide(blank_layout)
add_header(s3, "1.1. SQL Injection: Bản Chất Kỹ Thuật & Điểm Mù Ghép Chuỗi", "CHỦ ĐỀ 1 • NGUYỄN DANH HỌC (23A1001D0158)")
c3_1 = s3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.4), Inches(5.65), Inches(5.3))
c3_1.fill.solid()
c3_1.fill.fore_color.rgb = CARD_BG
c3_1.line.color.rgb = CARD_BORDER
tf = c3_1.text_frame
tf.word_wrap = True
p = tf.paragraphs[0]
p.text = "Nguyên Nhân Cốt Lõi: Nhầm Lẫn Giữa Code & Data"
p.font.size = Pt(15)
p.font.bold = True
p.font.color.rgb = ACCENT_RED
p = tf.add_paragraph()
p.text = "• Không phân tách giữa DỮ LIỆU (Data) và CÂU LỆNH (Code).\n• Lập trình viên sử dụng kỹ thuật ghép chuỗi trực tiếp (String Concatenation hoặc Interpolation) để tạo câu SQL từ input của người dùng.\n• Hệ quản trị CSDL khi phân tích cú pháp (Parsing) không thể biết đâu là dữ liệu của user và đâu là cấu trúc lệnh do lập trình viên định nghĩa.\n• Ký tự điều khiển (', --, ;) giúp hacker 'thoát' khỏi phạm vi dữ liệu và ép SQL Server thực thi logic của hacker."
p.font.size = Pt(12.5)
p.font.color.rgb = TEXT_BODY
p.space_before = Pt(8)

c3_2 = s3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.85), Inches(1.4), Inches(5.65), Inches(5.3))
c3_2.fill.solid()
c3_2.fill.fore_color.rgb = RED_LIGHT
c3_2.line.color.rgb = RED_BORDER
tf = c3_2.text_frame
tf.word_wrap = True
p = tf.paragraphs[0]
p.text = "Đoạn Code Gây Lỗi Thực Tế Trong Controller"
p.font.size = Pt(15)
p.font.bold = True
p.font.color.rgb = ACCENT_RED
p = tf.add_paragraph()
p.text = "// LỖI NỐI CHUỖI TRỰC TIẾP:\nstring rawSql = $\"SELECT * FROM Accounts \" +\n                $\"WHERE Username = '{username}' \" +\n                $\"AND Password = '{password}'\";\n\nvar accounts = _context.Accounts\n                       .FromSqlRaw(rawSql)\n                       .ToList();"
p.font.name = "Courier New"
p.font.size = Pt(11)
p.font.color.rgb = RGBColor(159, 18, 57)
p.space_before = Pt(8)
p = tf.add_paragraph()
p.text = "⚠️ Điểm mù chết người:\nLập trình viên mặc định tin rằng input chỉ chứa chữ cái. Khi hacker truyền vào dấu nháy đơn ('), toàn bộ cây cú pháp của câu truy vấn bị bẻ gãy hoàn toàn."
p.font.size = Pt(12)
p.font.color.rgb = TEXT_HEAD
p.space_before = Pt(14)

# Slide 4: Độ nguy hiểm & Thang điểm CVSS
s4 = prs.slides.add_slide(blank_layout)
add_header(s4, "1.2. SQL Injection: Thang Đo Nguy Hiểm & Hậu Quả Thực Tế (CVSS 9.8)", "CHỦ ĐỀ 1 • NGUYỄN DANH HỌC (23A1001D0158)")
col_w = Inches(3.64)
gap = Inches(0.39)
c4_1 = s4.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.4), col_w, Inches(5.3))
c4_1.fill.solid()
c4_1.fill.fore_color.rgb = RED_LIGHT
c4_1.line.color.rgb = RED_BORDER
tf = c4_1.text_frame
tf.word_wrap = True
p = tf.paragraphs[0]
p.text = "🚨 CHIẾM QUYỀN HỆ THỐNG\n(Authentication Bypass)"
p.font.size = Pt(14)
p.font.bold = True
p.font.color.rgb = ACCENT_RED
p = tf.add_paragraph()
p.text = "• Không cần biết mật khẩu của bất kỳ ai.\n• Chỉ với 1 chuỗi payload ngắn, kẻ tấn công đăng nhập thẳng vào tài khoản Quản trị viên (Admin).\n• Chiếm đoạt phiên làm việc, đổi quyền hạn và kiểm soát toàn bộ hệ thống."
p.font.size = Pt(12)
p.font.color.rgb = TEXT_BODY
p.space_before = Pt(10)

c4_2 = s4.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8) + col_w + gap, Inches(1.4), col_w, Inches(5.3))
c4_2.fill.solid()
c4_2.fill.fore_color.rgb = AMBER_LIGHT
c4_2.line.color.rgb = AMBER_BORDER
tf = c4_2.text_frame
tf.word_wrap = True
p = tf.paragraphs[0]
p.text = "📂 ĐÁNH CẮP TOÀN BỘ CSDL\n(Data Exfiltration)"
p.font.size = Pt(14)
p.font.bold = True
p.font.color.rgb = ACCENT_AMBER
p = tf.add_paragraph()
p.text = "• Trích xuất thông tin khách hàng, hồ sơ chuyên gia, số dư tài khoản ngân hàng.\n• Rò rỉ mật khẩu và bệnh án cá nhân (Secret Notes).\n• Vi phạm nghiêm trọng an ninh dữ liệu, phá vỡ uy tín của nền tảng tư vấn trực tuyến."
p.font.size = Pt(12)
p.font.color.rgb = TEXT_BODY
p.space_before = Pt(10)

c4_3 = s4.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8) + (col_w + gap)*2, Inches(1.4), col_w, Inches(5.3))
c4_3.fill.solid()
c4_3.fill.fore_color.rgb = BLUE_LIGHT
c4_3.line.color.rgb = BLUE_BORDER
tf = c4_3.text_frame
tf.word_wrap = True
p = tf.paragraphs[0]
p.text = "💣 PHÁ HỦY / THỰC THI LỆNH\n(System Takeover)"
p.font.size = Pt(14)
p.font.bold = True
p.font.color.rgb = ACCENT_BLUE
p = tf.add_paragraph()
p.text = "• Xóa sạch CSDL với lệnh DROP TABLE hoặc TRUNCATE.\n• Giả mạo số dư tiền trong tài khoản.\n• Trong các CSDL cấu hình lỏng lẻo, hacker có thể bật xp_cmdshell để chiếm quyền điều khiển cả máy chủ hệ điều hành."
p.font.size = Pt(12)
p.font.color.rgb = TEXT_BODY
p.space_before = Pt(10)

# Slide 5: Kịch bản giải phẫu Payload
s5 = prs.slides.add_slide(blank_layout)
add_header(s5, "1.3. SQL Injection: Giải Phẫu Chi Tiết Cơ Chế Bẻ Gãy Cú Pháp", "CHỦ ĐỀ 1 • NGUYỄN DANH HỌC (23A1001D0158)")
c5_main = s5.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.4), Inches(11.733), Inches(5.3))
c5_main.fill.solid()
c5_main.fill.fore_color.rgb = CARD_BG
c5_main.line.color.rgb = CARD_BORDER
tf = c5_main.text_frame
tf.word_wrap = True
p = tf.paragraphs[0]
p.text = "Kịch bản: Bắn Payload Bypass Xác Thực Đăng Nhập Quản Trị Viên"
p.font.size = Pt(15)
p.font.bold = True
p.font.color.rgb = ACCENT_RED

p = tf.add_paragraph()
p.text = "Payload đưa vào ô Username: ' OR '1'='1' --  | Mật khẩu: (nhập bất kỳ hoặc bỏ trống)\nCâu truy vấn thực tế được SQL Server phân tích & thông dịch:\nSELECT * FROM Accounts WHERE Username = '' OR '1'='1' --' AND Password = 'xyz'"
p.font.name = "Courier New"
p.font.size = Pt(12)
p.font.bold = True
p.font.color.rgb = ACCENT_BLUE
p.space_before = Pt(6)

p = tf.add_paragraph()
p.text = "3 Thành Phần Khai Thác Cốt Lõi:\n1. Dấu nháy đơn (') : Đóng sớm chuỗi Username, đưa ngữ cảnh trở lại câu lệnh SQL.\n2. Mệnh đề OR '1'='1' : Biểu thức logic chân lý (luôn TRUE cho mọi bản ghi trong bảng Accounts).\n3. Ký tự chú thích (--) : Vô hiệu hóa hoàn toàn bước kiểm tra mật khẩu phía sau (AND Password = '...')."
p.font.size = Pt(12.5)
p.font.color.rgb = TEXT_BODY
p.space_before = Pt(10)

p = tf.add_paragraph()
p.text = "📌 Cơ chế tấn công Union-Based (Rút dữ liệu từ bảng khác):\nPayload: ' UNION SELECT Id, Username, Password, SecretNote FROM Accounts --\nCho phép kẻ tấn công ghép nối bảng nhạy cảm vào kết quả hiển thị của màn hình tìm kiếm thông thường!"
p.font.size = Pt(12.5)
p.font.bold = True
p.font.color.rgb = ACCENT_AMBER
p.space_before = Pt(10)

# Slide 6: Đối chứng thực nghiệm Web UI & Live Demo
s6 = prs.slides.add_slide(blank_layout)
add_header(s6, "1.4. SQL Injection: Đối Chứng Thực Nghiệm Trực Quan Trên Web UI", "CHỦ ĐỀ 1 • NGUYỄN DANH HỌC (23A1001D0158)")
card_w = Inches(5.65)
b6_v = s6.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.4), card_w, Inches(4.2))
b6_v.fill.solid()
b6_v.fill.fore_color.rgb = RED_LIGHT
b6_v.line.color.rgb = RED_BORDER
tf = b6_v.text_frame
tf.word_wrap = True
p = tf.paragraphs[0]
p.text = "❌ 1. NHÁNH BỊ LỖI (Nối Chuỗi SQL Thuần)"
p.font.size = Pt(14)
p.font.bold = True
p.font.color.rgb = ACCENT_RED
p = tf.add_paragraph()
p.text = "• Input khai thác: ' OR '1'='1' --\n• Câu lệnh thực thi:\n  SELECT * FROM Accounts WHERE Username = '' OR '1'='1' --'\n• Cơ chế thông dịch: Mệnh đề WHERE luôn TRUE.\n• Hậu quả thực tế:\n  - Đăng nhập thành công mà KHÔNG CẦN mật khẩu.\n  - Rò rỉ toàn bộ 4 tài khoản CSDL (Admin, Chuyên gia, Khách hàng).\n  - Lộ mật khẩu gốc và số dư tài khoản ngân hàng."
p.font.size = Pt(12)
p.font.color.rgb = TEXT_BODY
p.space_before = Pt(8)

b6_s = s6.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.85), Inches(1.4), card_w, Inches(4.2))
b6_s.fill.solid()
b6_s.fill.fore_color.rgb = GREEN_LIGHT
b6_s.line.color.rgb = GREEN_BORDER
tf = b6_s.text_frame
tf.word_wrap = True
p = tf.paragraphs[0]
p.text = "✅ 2. NHÁNH AN TOÀN (EF Core Parameterized)"
p.font.size = Pt(14)
p.font.bold = True
p.font.color.rgb = ACCENT_GREEN
p = tf.add_paragraph()
p.text = "• Cùng Input khai thác: ' OR '1'='1' --\n• Câu lệnh thực thi:\n  SELECT * FROM Accounts WHERE Username = @p0 AND Pass = @p1\n• Cơ chế thông dịch:\n  - CSDL biên dịch cây truy vấn (Execution Plan) TRƯỚC.\n  - Toàn bộ chuỗi payload được xem là văn bản (String Literal) thuần túy.\n• Kết quả phòng thủ:\n  - Hệ thống chặn đứng, trả về 0 bản ghi.\n  - Báo 'Sai tài khoản/mật khẩu', bảo vệ an toàn 100%."
p.font.size = Pt(12)
p.font.color.rgb = TEXT_BODY
p.space_before = Pt(8)

callout6 = s6.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(5.8), Inches(11.733), Inches(0.95))
callout6.fill.solid()
callout6.fill.fore_color.rgb = BLUE_LIGHT
callout6.line.color.rgb = BLUE_BORDER
tf = callout6.text_frame
p = tf.paragraphs[0]
p.text = "💻 [HỌC LIVE DEMO]: Chuyển sang Web UI Tab 1 (http://localhost:5076/SqlInjection)"
p.font.size = Pt(12.5)
p.font.bold = True
p.font.color.rgb = ACCENT_BLUE
p = tf.add_paragraph()
p.text = "Thao tác: Bấm nạp payload ' OR '1'='1' -- -> Submit bên Đỏ (rò rỉ toàn bộ CSDL) -> Submit bên Xanh (chặn đứng an toàn)."
p.font.size = Pt(11)
p.font.color.rgb = TEXT_BODY
p.space_before = Pt(2)

# Slide 7: Minh chứng tầng sâu CLI & SQL Server 2025
s7 = prs.slides.add_slide(blank_layout)
add_header(s7, "1.5. SQL Injection: Minh Chứng Tầng API (cURL) & SQL Server 2025", "CHỦ ĐỀ 1 • NGUYỄN DANH HỌC (23A1001D0158)")
b7_api = s7.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.4), card_w, Inches(4.2))
b7_api.fill.solid()
b7_api.fill.fore_color.rgb = CARD_BG
b7_api.line.color.rgb = CARD_BORDER
tf = b7_api.text_frame
tf.word_wrap = True
p = tf.paragraphs[0]
p.text = "Minh Chứng Tầng API / Dòng Lệnh (cURL)"
p.font.size = Pt(14)
p.font.bold = True
p.font.color.rgb = ACCENT_BLUE
p = tf.add_paragraph()
p.text = "$ curl -X POST http://localhost:5076/SqlInjection/LoginVulnerable \\\n       -F \"username=' OR '1'='1' --\" -F \"password=any\"\n\nPhản hồi JSON từ Server:\n{\n  \"success\": true,\n  \"isExploited\": true,\n  \"count\": 4,\n  \"accounts\": [ ... 4 tài khoản bị rò rỉ ... ]\n}"
p.font.name = "Courier New"
p.font.size = Pt(10)
p.font.color.rgb = TEXT_HEAD
p.space_before = Pt(6)
p = tf.add_paragraph()
p.text = "📌 Ý nghĩa: Chứng minh lỗ hổng nằm ở logic C# Backend, kẻ tấn công dùng script tự động quét và khai thác mà không cần mở web."
p.font.size = Pt(11.5)
p.font.color.rgb = TEXT_BODY
p.space_before = Pt(8)

b7_db = s7.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.85), Inches(1.4), card_w, Inches(4.2))
b7_db.fill.solid()
b7_db.fill.fore_color.rgb = CARD_BG
b7_db.line.color.rgb = CARD_BORDER
tf = b7_db.text_frame
tf.word_wrap = True
p = tf.paragraphs[0]
p.text = "Minh Chứng Tầng CSDL (SQL Server 2025 Thật)"
p.font.size = Pt(14)
p.font.bold = True
p.font.color.rgb = ACCENT_GREEN
p = tf.add_paragraph()
p.text = "-- Database: OwaspDemoDB | Bảng: dbo.Accounts\n-- 1. Câu lệnh dính lỗi:\nSELECT * FROM Accounts WHERE Username='' OR '1'='1' --'\n>> Kết quả: 4 rows returned (Bypass thành công)\n\n-- 2. Câu lệnh tham số hóa chuẩn:\nEXEC sp_executesql N'SELECT * FROM Accounts WHERE...',\n     N'@p0 nvarchar(50)', @p0 = N''' OR ''1''=''1'' --'\n>> Kết quả: 0 rows returned (An toàn 100%)"
p.font.name = "Courier New"
p.font.size = Pt(10)
p.font.color.rgb = TEXT_HEAD
p.space_before = Pt(6)
p = tf.add_paragraph()
p.text = "📌 Ý nghĩa: Đối chiếu trực tiếp trên máy chủ SQL Server 2025 chứng minh cơ chế phân tách rạch ròi giữa cú pháp lệnh và tham số giá trị."
p.font.size = Pt(11.5)
p.font.color.rgb = TEXT_BODY
p.space_before = Pt(8)

callout7 = s7.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(5.8), Inches(11.733), Inches(0.95))
callout7.fill.solid()
callout7.fill.fore_color.rgb = GREEN_LIGHT
callout7.line.color.rgb = GREEN_BORDER
tf = callout7.text_frame
p = tf.paragraphs[0]
p.text = "🗄️ [HỌC LIVE DEMO]: Chuyển sang VS Code mở bảng dbo.Accounts & file verify_sql_server.sql"
p.font.size = Pt(12.5)
p.font.bold = True
p.font.color.rgb = ACCENT_GREEN
p = tf.add_paragraph()
p.text = "Thao tác trực tiếp: Mở extension MSSQL trên VS Code, đối chiếu bảng dữ liệu gốc và chạy câu query đối chứng."
p.font.size = Pt(11)
p.font.color.rgb = TEXT_BODY
p.space_before = Pt(2)

# Slide 8: Phòng thủ chuẩn hóa SQLi
s8 = prs.slides.add_slide(blank_layout)
add_header(s8, "1.6. SQL Injection: Giải Pháp Phòng Thủ Toàn Diện (Defense-in-Depth)", "CHỦ ĐỀ 1 • NGUYỄN DANH HỌC (23A1001D0158)")
c8_1 = s8.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.4), Inches(5.65), Inches(5.3))
c8_1.fill.solid()
c8_1.fill.fore_color.rgb = GREEN_LIGHT
c8_1.line.color.rgb = GREEN_BORDER
tf = c8_1.text_frame
tf.word_wrap = True
p = tf.paragraphs[0]
p.text = "Cách Viết Code Chuẩn Hóa Với EF Core"
p.font.size = Pt(15)
p.font.bold = True
p.font.color.rgb = ACCENT_GREEN
p = tf.add_paragraph()
p.text = "// Dùng LINQ Parameterized (KHUYÊN DÙNG):\nvar account = _context.Accounts\n    .FirstOrDefault(a => a.Username == user \n                      && a.Password == pass);\n\n// Hoặc dùng FromSqlInterpolated nếu viết SQL thuần:\nvar account = _context.Accounts\n    .FromSqlInterpolated($\"SELECT * FROM Accounts WHERE Username={user} AND Password={pass}\")\n    .FirstOrDefault();"
p.font.name = "Courier New"
p.font.size = Pt(10)
p.font.color.rgb = RGBColor(6, 95, 70)
p.space_before = Pt(8)
p = tf.add_paragraph()
p.text = "Tại sao cách này an toàn tuyệt đối?\nSQL Server biên dịch cấu trúc lệnh TRƯỚC khi gán dữ liệu. Biến @p0 nhận toàn bộ payload như một chuỗi chữ bình thường, vô hiệu hóa hoàn toàn ý đồ chèn lệnh."
p.font.size = Pt(12)
p.font.color.rgb = TEXT_BODY
p.space_before = Pt(10)

c8_2 = s8.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.85), Inches(1.4), Inches(5.65), Inches(5.3))
c8_2.fill.solid()
c8_2.fill.fore_color.rgb = CARD_BG
c8_2.line.color.rgb = CARD_BORDER
tf = c8_2.text_frame
tf.word_wrap = True
p = tf.paragraphs[0]
p.text = "Bộ Quy Tắc Phòng Thủ Toàn Diện (Defense in Depth)"
p.font.size = Pt(15)
p.font.bold = True
p.font.color.rgb = ACCENT_BLUE
p = tf.add_paragraph()
p.text = "1. Luôn sử dụng Parameterized Query / ORM:\nTuyệt đối không ghép chuỗi SQL thủ công dưới bất kỳ hình thức nào.\n\n2. Nguyên tắc đặc quyền tối thiểu (Least Privilege):\nTài khoản kết nối CSDL của ứng dụng chỉ có quyền SELECT/INSERT/UPDATE trên các bảng cần thiết, không bao giờ dùng tài khoản 'sa' hay quyền DDL (DROP, ALTER).\n\n3. Xác thực dữ liệu đầu vào (Input Validation):\nSử dụng Data Annotations ([RegularExpression], [StringLength]) kiểm tra chặt chẽ khuôn dạng dữ liệu.\n\n4. Mã hóa mật khẩu một chiều (Password Hashing):\nSử dụng ASP.NET Core Identity (PBKDF2/BCrypt) để nếu CSDL có bị lộ, mật khẩu vẫn được bảo vệ an toàn."
p.font.size = Pt(12)
p.font.color.rgb = TEXT_BODY
p.space_before = Pt(8)


# ==============================================================================
# PHẦN 2: STORED XSS — NGUYỄN THANH BÌNH (SLIDES 9 - 14)
# ==============================================================================
# Slide 9: Ngữ cảnh & Bản chất XSS
s9 = prs.slides.add_slide(blank_layout)
add_header(s9, "2.1. Stored XSS: Bản Chất Kỹ Thuật & Ngữ Cảnh Bài Toán", "CHỦ ĐỀ 2 • NGUYỄN THANH BÌNH (23A1001D0041)")
c9_1 = s9.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.4), Inches(5.65), Inches(5.3))
c9_1.fill.solid()
c9_1.fill.fore_color.rgb = CARD_BG
c9_1.line.color.rgb = CARD_BORDER
tf = c9_1.text_frame
tf.word_wrap = True
p = tf.paragraphs[0]
p.text = "Ngữ Cảnh Nghiệp Vụ Trong Website Tư Vấn"
p.font.size = Pt(15)
p.font.bold = True
p.font.color.rgb = ACCENT_RED
p = tf.add_paragraph()
p.text = "• Trong website tư vấn trực tuyến, khách hàng có chức năng đăng đánh giá, nhận xét và gửi câu hỏi tư vấn.\n• Bản chất Stored XSS (Persistent XSS): Kẻ tấn công nhúng mã JavaScript độc hại vào nội dung nhận xét.\n• Máy chủ không lọc mà lưu thẳng mã độc vào CSDL (bảng dbo.Comments).\n• Nguy hiểm: Mỗi khi bất kỳ người dùng khác, chuyên gia hoặc quản trị viên mở trang đánh giá, trình duyệt của họ sẽ tự động biên dịch và thực thi đoạn script độc hại đó."
p.font.size = Pt(12.5)
p.font.color.rgb = TEXT_BODY
p.space_before = Pt(8)

c9_2 = s9.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.85), Inches(1.4), Inches(5.65), Inches(5.3))
c9_2.fill.solid()
c9_2.fill.fore_color.rgb = RED_LIGHT
c9_2.line.color.rgb = RED_BORDER
tf = c9_2.text_frame
tf.word_wrap = True
p = tf.paragraphs[0]
p.text = "Mã Nguồn C# Gây Lỗi & Sự Lạm Dụng @@Html.Raw"
p.font.size = Pt(15)
p.font.bold = True
p.font.color.rgb = ACCENT_RED
p = tf.add_paragraph()
p.text = "// 1. CONTROLLER LƯU NGUYÊN VĂN KHÔNG LÀM SẠCH:\nvar comment = new Comment {\n    Author = author,\n    Content = content // Chứa thẻ <script> độc hại\n};\n_context.Comments.Add(comment);\n\n// 2. RAZOR VIEW LẠM DỤNG HTML.RAW:\n<div class=\"comment-body\">\n    @Html.Raw(comment.Content) \n</div>"
p.font.name = "Courier New"
p.font.size = Pt(10.5)
p.font.color.rgb = RGBColor(159, 18, 57)
p.space_before = Pt(8)
p = tf.add_paragraph()
p.text = "⚠️ Điểm mù lập trình viên:\nSử dụng @Html.Raw() với mục đích cho phép hiển thị format HTML nhưng vô tình mở toang cánh cửa cho hacker thực thi JavaScript trên máy nạn nhân."
p.font.size = Pt(12)
p.font.color.rgb = TEXT_HEAD
p.space_before = Pt(12)

# Slide 10: Phân loại XSS & Thang đo nguy hiểm
s10 = prs.slides.add_slide(blank_layout)
add_header(s10, "2.2. Stored XSS: Phân Loại Chuẩn & Mối Đe Dọa Nghiêm Trọng", "CHỦ ĐỀ 2 • NGUYỄN THANH BÌNH (23A1001D0041)")
c10_1 = s10.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.4), col_w, Inches(5.3))
c10_1.fill.solid()
c10_1.fill.fore_color.rgb = RED_LIGHT
c10_1.line.color.rgb = RED_BORDER
tf = c10_1.text_frame
tf.word_wrap = True
p = tf.paragraphs[0]
p.text = "🍪 ĐÁNH CẮP SESSION COOKIE\n(Cookie Stealing)"
p.font.size = Pt(14)
p.font.bold = True
p.font.color.rgb = ACCENT_RED
p = tf.add_paragraph()
p.text = "• Mã độc truy cập document.cookie để trích xuất Session Token của nạn nhân.\n• Tự động gửi token về máy chủ của hacker.\n• Hacker giả mạo danh tính nạn nhân (Session Hijacking) đăng nhập vào tài khoản mà không cần mật khẩu."
p.font.size = Pt(12)
p.font.color.rgb = TEXT_BODY
p.space_before = Pt(10)

c10_2 = s10.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8) + col_w + gap, Inches(1.4), col_w, Inches(5.3))
c10_2.fill.solid()
c10_2.fill.fore_color.rgb = AMBER_LIGHT
c10_2.line.color.rgb = AMBER_BORDER
tf = c10_2.text_frame
tf.word_wrap = True
p = tf.paragraphs[0]
p.text = "🎣 THAY ĐỔI GIAO DIỆN / LỪA ĐẢO\n(DOM Defacement & Phishing)"
p.font.size = Pt(14)
p.font.bold = True
p.font.color.rgb = ACCENT_AMBER
p = tf.add_paragraph()
p.text = "• Hacker dùng JavaScript chèn form đăng nhập giả mạo đè lên giao diện website thật.\n• Lừa người dùng nhập lại tài khoản ngân hàng hoặc mật khẩu.\n• Chuyển hướng người dùng sang trang web lừa đảo (Open Redirect)."
p.font.size = Pt(12)
p.font.color.rgb = TEXT_BODY
p.space_before = Pt(10)

c10_3 = s10.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8) + (col_w + gap)*2, Inches(1.4), col_w, Inches(5.3))
c10_3.fill.solid()
c10_3.fill.fore_color.rgb = BLUE_LIGHT
c10_3.line.color.rgb = BLUE_BORDER
tf = c10_3.text_frame
tf.word_wrap = True
p = tf.paragraphs[0]
p.text = "🤖 THAO TÁC NGẦM & KEYLOGGER\n(Client-Side Hijacking)"
p.font.size = Pt(14)
p.font.bold = True
p.font.color.rgb = ACCENT_BLUE
p = tf.add_paragraph()
p.text = "• Lắng nghe sự kiện bàn phím của nạn nhân (Keylogging) khi họ chat với chuyên gia.\n• Tự động gửi tin nhắn hoặc đánh giá 5 sao giả mạo danh nghĩa của người dùng.\n• Biến trình duyệt nạn nhân thành botnet tấn công từ chối dịch vụ."
p.font.size = Pt(12)
p.font.color.rgb = TEXT_BODY
p.space_before = Pt(10)

# Slide 11: Kịch bản XSS & Giải phẫu Payload
s11 = prs.slides.add_slide(blank_layout)
add_header(s11, "2.3. Stored XSS: Phân Tích Các Kỹ Thuật Đưa Payload Vào CSDL", "CHỦ ĐỀ 2 • NGUYỄN THANH BÌNH (23A1001D0041)")
c11_main = s11.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.4), Inches(11.733), Inches(5.3))
c11_main.fill.solid()
c11_main.fill.fore_color.rgb = CARD_BG
c11_main.line.color.rgb = CARD_BORDER
tf = c11_main.text_frame
tf.word_wrap = True
p = tf.paragraphs[0]
p.text = "3 Biến Thể Payload XSS Thực Nghiệm Trong Ứng Dụng"
p.font.size = Pt(15)
p.font.bold = True
p.font.color.rgb = ACCENT_RED

p = tf.add_paragraph()
p.text = "1. Payload 1 (Thẻ Script Truyền Thống):\n   <script>alert('🚨 BỊ HACK XSS! Lộ Session Cookie: ' + document.cookie);</script>\n   Cơ chế: Trình duyệt gặp thẻ <script> sẽ tạm dừng render HTML và thực thi mã JavaScript ngay lập tức."
p.font.size = Pt(12.5)
p.font.color.rgb = TEXT_BODY
p.space_before = Pt(8)

p = tf.add_paragraph()
p.text = "2. Payload 2 (Bypass Bộ Lọc Qua Thẻ Ảnh - Event Handler):\n   <img src=x onerror=alert('🚨 XSS kích hoạt qua thẻ IMG onerror!')>\n   Cơ chế: Hacker qua mặt các bộ lọc từ khóa '<script>' bằng cách dùng thuộc tính sự kiện 'onerror' của thẻ <img> khi đường dẫn ảnh bị lỗi (src=x)."
p.font.size = Pt(12.5)
p.font.color.rgb = TEXT_BODY
p.space_before = Pt(8)

p = tf.add_paragraph()
p.text = "3. Payload 3 (Thay Đổi Toàn Bộ Cây DOM - Web Defacement):\n   <script>document.body.innerHTML='<h1 style=color:red;text-align:center>WEB ĐÃ BỊ CHIẾM QUYỀN</h1>';</script>\n   Cơ chế: Xóa sạch toàn bộ nội dung trang web chính thống và hiển thị thông điệp của hacker!"
p.font.size = Pt(12.5)
p.font.color.rgb = TEXT_BODY
p.space_before = Pt(8)

# Slide 12: Đối chứng thực nghiệm XSS & Live Demo
s12 = prs.slides.add_slide(blank_layout)
add_header(s12, "2.4. Stored XSS: Đối Chứng Thực Nghiệm & Live Demo Trực Tiếp", "CHỦ ĐỀ 2 • NGUYỄN THANH BÌNH (23A1001D0041)")
b12_v = s12.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.4), card_w, Inches(4.2))
b12_v.fill.solid()
b12_v.fill.fore_color.rgb = RED_LIGHT
b12_v.line.color.rgb = RED_BORDER
tf = b12_v.text_frame
tf.word_wrap = True
p = tf.paragraphs[0]
p.text = "❌ 1. NHÁNH BỊ LỖI (Dùng @@Html.Raw)"
p.font.size = Pt(14)
p.font.bold = True
p.font.color.rgb = ACCENT_RED
p = tf.add_paragraph()
p.text = "• Input khai thác: <script>alert(document.cookie);</script>\n• Cách render:\n  @Html.Raw(comment.Content)\n• Cơ chế:\n  - Trình duyệt coi đoạn chuỗi là mã lệnh JavaScript thực thi được.\n• Hậu quả thực tế:\n  - Bật popup alert hiển thị Cookie nhạy cảm AuthSessionToken.\n  - Bất kỳ ai vào xem đánh giá chuyên gia đều bị mã độc tấn công."
p.font.size = Pt(12)
p.font.color.rgb = TEXT_BODY
p.space_before = Pt(8)

b12_s = s12.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.85), Inches(1.4), card_w, Inches(4.2))
b12_s.fill.solid()
b12_s.fill.fore_color.rgb = GREEN_LIGHT
b12_s.line.color.rgb = GREEN_BORDER
tf = b12_s.text_frame
tf.word_wrap = True
p = tf.paragraphs[0]
p.text = "✅ 2. NHÁNH AN TOÀN (Tự Động HTML Encoding)"
p.font.size = Pt(14)
p.font.bold = True
p.font.color.rgb = ACCENT_GREEN
p = tf.add_paragraph()
p.text = "• Cùng Input khai thác: <script>alert(document.cookie);</script>\n• Cách render chuẩn:\n  @comment.Content (hoặc HtmlEncoder.Default.Encode)\n• Cơ chế:\n  - Các ký tự <, > được chuyển đổi thành &lt;, &gt;.\n• Kết quả phòng thủ:\n  - Trình duyệt chỉ in chuỗi text ra màn hình như chữ bình thường.\n  - Tuyệt đối không có bất kỳ popup hay script nào được kích hoạt."
p.font.size = Pt(12)
p.font.color.rgb = TEXT_BODY
p.space_before = Pt(8)

callout12 = s12.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(5.8), Inches(11.733), Inches(0.95))
callout12.fill.solid()
callout12.fill.fore_color.rgb = RED_LIGHT
callout12.line.color.rgb = RED_BORDER
tf = callout12.text_frame
p = tf.paragraphs[0]
p.text = "💬 [BÌNH LIVE DEMO]: Chuyển sang Web UI Tab 2 (Stored XSS - /Xss)"
p.font.size = Pt(12.5)
p.font.bold = True
p.font.color.rgb = ACCENT_RED
p = tf.add_paragraph()
p.text = "Thao tác: Bấm nạp payload đánh cắp Cookie -> Đăng bên Lỗi (bật popup alert lộ Session) -> Đăng bên An toàn (chữ hiển thị an toàn)."
p.font.size = Pt(11)
p.font.color.rgb = TEXT_BODY
p.space_before = Pt(2)

# Slide 13: Minh chứng CSDL & Cookie
s13 = prs.slides.add_slide(blank_layout)
add_header(s13, "2.5. Stored XSS: Minh Chứng CSDL dbo.Comments & Cờ Bảo Vệ Cookie", "CHỦ ĐỀ 2 • NGUYỄN THANH BÌNH (23A1001D0041)")
b13_db = s13.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.4), card_w, Inches(5.3))
b13_db.fill.solid()
b13_db.fill.fore_color.rgb = CARD_BG
b13_db.line.color.rgb = CARD_BORDER
tf = b13_db.text_frame
tf.word_wrap = True
p = tf.paragraphs[0]
p.text = "Dữ Liệu Thật Lưu Trong Bảng dbo.Comments"
p.font.size = Pt(14)
p.font.bold = True
p.font.color.rgb = ACCENT_BLUE
p = tf.add_paragraph()
p.text = "SELECT Id, Author, Content, IsSecureStored FROM dbo.Comments;\n\nId | Author        | Content                           | IsSecure\n---+---------------+-----------------------------------+---------\n 1 | Le Thanh Binh | Dich vu tu van rat chuyen nghiep! | 1\n 2 | Hacker_XSS    | <script>alert(document.cookie);.. | 0"
p.font.name = "Courier New"
p.font.size = Pt(9.5)
p.font.color.rgb = TEXT_HEAD
p.space_before = Pt(6)
p = tf.add_paragraph()
p.text = "📌 Bản chất lưu trữ:\nĐoạn script độc hại được lưu vĩnh viễn trong CSDL. Do đó gọi là 'Stored XSS' - nguy hiểm hơn rất nhiều so với 'Reflected XSS' vì nó nằm sẵn trong hệ thống và gây hại cho mọi lượt truy cập sau đó."
p.font.size = Pt(11.5)
p.font.color.rgb = TEXT_BODY
p.space_before = Pt(8)

b13_c = s13.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.85), Inches(1.4), card_w, Inches(5.3))
b13_c.fill.solid()
b13_c.fill.fore_color.rgb = CARD_BG
b13_c.line.color.rgb = CARD_BORDER
tf = b13_c.text_frame
tf.word_wrap = True
p = tf.paragraphs[0]
p.text = "Phòng Vệ Chiều Sâu: Cấu Hình HttpOnly Cookie"
p.font.size = Pt(14)
p.font.bold = True
p.font.color.rgb = ACCENT_GREEN
p = tf.add_paragraph()
p.text = "// CẤU HÌNH COOKIE AN TOÀN TRONG ASP.NET CORE:\nResponse.Cookies.Append(\"AuthSessionToken\", token, new CookieOptions\n{\n    HttpOnly = true, // CẤM JAVASCRIPT TRUY CẬP\n    Secure = true,   // CHỈ TRUYỀN QUA HTTPS\n    SameSite = SameSiteMode.Strict\n});"
p.font.name = "Courier New"
p.font.size = Pt(9.5)
p.font.color.rgb = RGBColor(6, 95, 70)
p.space_before = Pt(6)
p = tf.add_paragraph()
p.text = "📌 Tác dụng của HttpOnly:\nDù ứng dụng có vô tình dính lỗi XSS thì lệnh 'document.cookie' của hacker vẫn trả về rỗng! Hacker không thể nào đánh cắp được token phiên làm việc."
p.font.size = Pt(11.5)
p.font.color.rgb = TEXT_BODY
p.space_before = Pt(8)

# Slide 14: Phòng thủ chuẩn hóa XSS
s14 = prs.slides.add_slide(blank_layout)
add_header(s14, "2.6. Stored XSS: Bộ 4 Nguyên Tắc Phòng Thủ Toàn Diện", "CHỦ ĐỀ 2 • NGUYỄN THANH BÌNH (23A1001D0041)")
c14_main = s14.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.4), Inches(11.733), Inches(5.3))
c14_main.fill.solid()
c14_main.fill.fore_color.rgb = GREEN_LIGHT
c14_main.line.color.rgb = GREEN_BORDER
tf = c14_main.text_frame
tf.word_wrap = True
p = tf.paragraphs[0]
p.text = "BỘ NGUYÊN TẮC PHÒNG THỦ XSS CHUẨN OWASP TRONG ASP.NET CORE:"
p.font.size = Pt(15)
p.font.bold = True
p.font.color.rgb = ACCENT_GREEN

p = tf.add_paragraph()
p.text = "1. Luôn tận dụng Context-Aware HTML Encoding của Razor View:\n   Tuyệt đối không dùng @Html.Raw() cho dữ liệu nhập từ người dùng. Sử dụng cú pháp mặc định @model.Property để framework tự động mã hóa ký tự nhạy cảm."
p.font.size = Pt(13)
p.font.color.rgb = TEXT_BODY
p.space_before = Pt(8)

p = tf.add_paragraph()
p.text = "2. Mã hóa tại Controller với System.Text.Encodings.Web:\n   Dùng HtmlEncoder.Default.Encode() hoặc JavaScriptEncoder.Default.Encode() để làm sạch chuỗi trước khi lưu vào CSDL."
p.font.size = Pt(13)
p.font.color.rgb = TEXT_BODY
p.space_before = Pt(8)

p = tf.add_paragraph()
p.text = "3. Thiết lập cờ HttpOnly = true và Secure = true cho toàn bộ Cookie nhạy cảm:\n   Bảo vệ token phiên làm việc (Session) không bao giờ bị lộ qua thuộc tính document.cookie."
p.font.size = Pt(13)
p.font.color.rgb = TEXT_BODY
p.space_before = Pt(8)

p = tf.add_paragraph()
p.text = "4. Triển khai Content Security Policy (CSP Header):\n   Thiết lập HTTP Response Header (script-src 'self') để trình duyệt chặn đứng hoàn toàn việc chạy inline script và tải mã độc từ các tên miền lạ."
p.font.size = Pt(13)
p.font.color.rgb = TEXT_BODY
p.space_before = Pt(8)


# ==============================================================================
# PHẦN 3: CSRF ATTACK — NGUYỄN MINH CƯỜNG (SLIDES 15 - 20)
# ==============================================================================
# Slide 15: Bản chất CSRF
s15 = prs.slides.add_slide(blank_layout)
add_header(s15, "3.1. CSRF Attack: Bản Chất Kỹ Thuật & Cơ Chế Giả Mạo Request", "CHỦ ĐỀ 3 • NGUYỄN MINH CƯỜNG (23A1001D0058)")
c15_1 = s15.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.4), Inches(5.65), Inches(5.3))
c15_1.fill.solid()
c15_1.fill.fore_color.rgb = CARD_BG
c15_1.line.color.rgb = CARD_BORDER
tf = c15_1.text_frame
tf.word_wrap = True
p = tf.paragraphs[0]
p.text = "Bản Chất Của Lỗ Hổng Cross-Site Request Forgery"
p.font.size = Pt(15)
p.font.bold = True
p.font.color.rgb = ACCENT_RED
p = tf.add_paragraph()
p.text = "• CSRF (A01:2021 - Broken Access Control) là kỹ thuật tấn công mượn quyền của người dùng hợp lệ.\n• Điều kiện xảy ra:\n  1. Nạn nhân đã đăng nhập vào hệ thống tư vấn trực tuyến (phiên làm việc Session/Cookie vẫn đang hợp lệ).\n  2. Trình duyệt tự động đính kèm Cookie xác thực vào mọi HTTP Request gửi tới hệ thống đó.\n  3. Kẻ tấn công lừa nạn nhân truy cập một trang web độc hại ở một tab khác (Attacker Site).\n• Điểm mù: Máy chủ chỉ kiểm tra Cookie mà không kiểm tra xem Request đó bắt nguồn từ đâu!"
p.font.size = Pt(12.5)
p.font.color.rgb = TEXT_BODY
p.space_before = Pt(8)

c15_2 = s15.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.85), Inches(1.4), Inches(5.65), Inches(5.3))
c15_2.fill.solid()
c15_2.fill.fore_color.rgb = RED_LIGHT
c15_2.line.color.rgb = RED_BORDER
tf = c15_2.text_frame
tf.word_wrap = True
p = tf.paragraphs[0]
p.text = "Action Bị Lỗi: Thiếu Thuộc Tính Phòng Vệ"
p.font.size = Pt(15)
p.font.bold = True
p.font.color.rgb = ACCENT_RED
p = tf.add_paragraph()
p.text = "// ACTION NHẬN TIỀN KHÔNG CÓ BẢO VỆ TOKEN:\n[HttpPost]\npublic IActionResult TransferVulnerable(decimal amount)\n{\n    var victim = _context.Wallets.Find(1);\n    var hacker = _context.Wallets.Find(2);\n    \n    // Trừ tiền nạn nhân, cộng tiền hacker:\n    victim.Balance -= amount;\n    hacker.Balance += amount;\n    _context.SaveChanges();\n    return RedirectToAction(\"Index\");\n}"
p.font.name = "Courier New"
p.font.size = Pt(10)
p.font.color.rgb = RGBColor(159, 18, 57)
p.space_before = Pt(8)
p = tf.add_paragraph()
p.text = "⚠️ Nguy hiểm:\nBất kỳ trang web nào trên Internet cũng có thể gửi form POST tới URL này. Máy chủ thấy Cookie hợp lệ sẽ trừ tiền ngay lập tức!"
p.font.size = Pt(12)
p.font.color.rgb = TEXT_HEAD
p.space_before = Pt(10)

# Slide 16: Kịch bản lừa đảo bẫy trúng thưởng
s16 = prs.slides.add_slide(blank_layout)
add_header(s16, "3.2. CSRF Attack: Kịch Bản Lừa Đảo Chuyển Tiền Ví Ngầm", "CHỦ ĐỀ 3 • NGUYỄN MINH CƯỜNG (23A1001D0058)")
c16_main = s16.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.4), Inches(11.733), Inches(5.3))
c16_main.fill.solid()
c16_main.fill.fore_color.rgb = RED_LIGHT
c16_main.line.color.rgb = RED_BORDER
tf = c16_main.text_frame
tf.word_wrap = True
p = tf.paragraphs[0]
p.text = "Kịch bản: Trang Web Giả Mạo Trúng Thưởng (Attacker Site)"
p.font.size = Pt(16)
p.font.bold = True
p.font.color.rgb = ACCENT_RED

p = tf.add_paragraph()
p.text = "1. Nạn nhân có 50.000.000 VNĐ trong ví đang mở tab làm việc tại nền tảng tư vấn trực tuyến.\n2. Kẻ tấn công gửi link: 'Chúc mừng bạn trúng thưởng iPhone 16 Pro Max! Click để nhận quà!'.\n3. Khi nạn nhân click vào link, một trang web bẫy (AttackerSite.cshtml) được mở ra, bên trong chứa mã HTML ẩn:"
p.font.size = Pt(13)
p.font.color.rgb = TEXT_BODY
p.space_before = Pt(8)

p = tf.add_paragraph()
p.text = "<!-- FORM ẨN TỰ ĐỘNG GỬI SANG HỆ THỐNG MỤC TIÊU -->\n<form id=\"exploit\" action=\"http://localhost:5076/Csrf/TransferVulnerable\" method=\"POST\">\n    <input type=\"hidden\" name=\"amount\" value=\"20000000\" />\n</form>\n<script>document.getElementById('exploit').submit();</script>"
p.font.name = "Courier New"
p.font.size = Pt(11)
p.font.bold = True
p.font.color.rgb = RGBColor(159, 18, 57)
p.space_before = Pt(6)

p = tf.add_paragraph()
p.text = "Hậu quả thực tế: Nạn nhân không hề bấm chuyển tiền, nhưng ví tiền bị trừ ngay 20.000.000 VNĐ sang ví hacker!"
p.font.size = Pt(13)
p.font.bold = True
p.font.color.rgb = ACCENT_RED
p.space_before = Pt(8)

# Slide 17: Live demo CSRF
s17 = prs.slides.add_slide(blank_layout)
add_header(s17, "3.3. CSRF Attack: Thực Nghiệm Trực Tiếp & Kiểm Thử Rút Tiền", "CHỦ ĐỀ 3 • NGUYỄN MINH CƯỜNG (23A1001D0058)")
c17_main = s17.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.4), Inches(11.733), Inches(4.2))
c17_main.fill.solid()
c17_main.fill.fore_color.rgb = CARD_BG
c17_main.line.color.rgb = CARD_BORDER
tf = c17_main.text_frame
tf.word_wrap = True
p = tf.paragraphs[0]
p.text = "Quy Trình 3 Bước Thực Nghiệm Trên Giao Diện Web"
p.font.size = Pt(15)
p.font.bold = True
p.font.color.rgb = ACCENT_GREEN

p = tf.add_paragraph()
p.text = "• Bước 1: Kiểm tra số dư ví ban đầu tại Tab 3:\n  - Ví Nạn nhân: 50.000.000 VNĐ | Ví Hacker: 0 VNĐ.\n• Bước 2: Bấm nút đỏ 'Mở trang web bẫy của Hacker (Attacker Site)' ở tab mới -> Bấm nút 'Nhận thưởng ngay':\n  - Ngay lập tức, số dư của nạn nhân bị trừ 20.000.000 VNĐ và chuyển sang ví hacker!\n• Bước 3: Thử nghiệm chuyển tiền tại Form bên phải (Đã phòng thủ):\n  - Giao dịch chỉ hợp lệ khi được gửi từ chính form trên hệ thống với Anti-Forgery Token hợp lệ."
p.font.size = Pt(13)
p.font.color.rgb = TEXT_BODY
p.space_before = Pt(8)

callout17 = s17.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(5.8), Inches(11.733), Inches(0.95))
callout17.fill.solid()
callout17.fill.fore_color.rgb = GREEN_LIGHT
callout17.line.color.rgb = GREEN_BORDER
tf = callout17.text_frame
p = tf.paragraphs[0]
p.text = "💳 [CƯỜNG LIVE DEMO]: Chuyển sang Web UI Tab 3 (CSRF - /Csrf)"
p.font.size = Pt(12.5)
p.font.bold = True
p.font.color.rgb = ACCENT_GREEN
p = tf.add_paragraph()
p.text = "Thao tác: Mở Attacker Site -> Bấm nhận thưởng -> Quay lại thấy số dư ví nạn nhân bị trừ 20 triệu -> Thử lại form có Token."
p.font.size = Pt(11)
p.font.color.rgb = TEXT_BODY
p.space_before = Pt(2)

# Slide 18: Minh chứng CSDL ví tiền
s18 = prs.slides.add_slide(blank_layout)
add_header(s18, "3.4. CSRF Attack: Minh Chứng CSDL dbo.UserWallets Trước & Sau Tấn Công", "CHỦ ĐỀ 3 • NGUYỄN MINH CƯỜNG (23A1001D0058)")
b18_1 = s18.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.4), card_w, Inches(5.3))
b18_1.fill.solid()
b18_1.fill.fore_color.rgb = CARD_BG
b18_1.line.color.rgb = CARD_BORDER
tf = b18_1.text_frame
tf.word_wrap = True
p = tf.paragraphs[0]
p.text = "1. Trạng Thái CSDL Ban Đầu (Chưa Bị Tấn Công)"
p.font.size = Pt(14)
p.font.bold = True
p.font.color.rgb = ACCENT_BLUE
p = tf.add_paragraph()
p.text = "SELECT Id, OwnerName, Balance FROM dbo.UserWallets;\n\nId | OwnerName                      | Balance (VND)\n---+--------------------------------+---------------\n 1 | Nan nhan (Nguyen Danh Hoc)     | 50,000,000.00\n 2 | Ke tan cong (Hacker BlackHat)  |          0.00\n(2 rows affected)"
p.font.name = "Courier New"
p.font.size = Pt(9.5)
p.font.color.rgb = TEXT_HEAD
p.space_before = Pt(6)
p = tf.add_paragraph()
p.text = "📌 Ví của nạn nhân có 50 triệu VNĐ, ví của hacker hoàn toàn không có tiền."
p.font.size = Pt(11.5)
p.font.color.rgb = TEXT_BODY
p.space_before = Pt(8)

b18_2 = s18.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.85), Inches(1.4), card_w, Inches(5.3))
b18_2.fill.solid()
b18_2.fill.fore_color.rgb = RED_LIGHT
b18_2.line.color.rgb = RED_BORDER
tf = b18_2.text_frame
tf.word_wrap = True
p = tf.paragraphs[0]
p.text = "2. Trạng Thái CSDL Sau Khi Bị Khai Thác CSRF"
p.font.size = Pt(14)
p.font.bold = True
p.font.color.rgb = ACCENT_RED
p = tf.add_paragraph()
p.text = "SELECT Id, OwnerName, Balance FROM dbo.UserWallets;\n\nId | OwnerName                      | Balance (VND)\n---+--------------------------------+---------------\n 1 | Nan nhan (Nguyen Danh Hoc)     | 30,000,000.00\n 2 | Ke tan cong (Hacker BlackHat)  | 20,000,000.00\n(2 rows affected)"
p.font.name = "Courier New"
p.font.size = Pt(9.5)
p.font.color.rgb = RGBColor(159, 18, 57)
p.space_before = Pt(6)
p = tf.add_paragraph()
p.text = "📌 Hậu quả kiểm chứng trên CSDL thật:\n20.000.000 VNĐ đã biến mất khỏi tài khoản nạn nhân và chuyển sang số dư của kẻ tấn công mà nạn nhân không hề hay biết."
p.font.size = Pt(11.5)
p.font.color.rgb = TEXT_BODY
p.space_before = Pt(8)

# Slide 19: Cơ chế phòng thủ CSRF Token
s19 = prs.slides.add_slide(blank_layout)
add_header(s19, "3.5. CSRF Attack: Cơ Chế Phòng Thủ Bằng Anti-Forgery Token", "CHỦ ĐỀ 3 • NGUYỄN MINH CƯỜNG (23A1001D0058)")
c19_1 = s19.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.4), Inches(5.65), Inches(5.3))
c19_1.fill.solid()
c19_1.fill.fore_color.rgb = GREEN_LIGHT
c19_1.line.color.rgb = GREEN_BORDER
tf = c19_1.text_frame
tf.word_wrap = True
p = tf.paragraphs[0]
p.text = "Cơ Chế Synchronizer Token Pattern Trong ASP.NET Core"
p.font.size = Pt(15)
p.font.bold = True
p.font.color.rgb = ACCENT_GREEN
p = tf.add_paragraph()
p.text = "// 1. TRONG RAZOR VIEW: NHÚNG TOKEN VÀO FORM\n<form asp-action=\"TransferSecure\" method=\"post\">\n    @Html.AntiForgeryToken()\n    <input type=\"number\" name=\"amount\" />\n    <button type=\"submit\">Chuyển tiền</button>\n</form>\n\n// 2. TRONG CONTROLLER: BẮT BUỘC KIỂM TRA TOKEN\n[HttpPost]\n[ValidateAntiForgeryToken]\npublic IActionResult TransferSecure(decimal amount)\n{\n    // Chỉ thực thi khi Token từ Form khớp với Token trong Cookie\n}"
p.font.name = "Courier New"
p.font.size = Pt(9.5)
p.font.color.rgb = RGBColor(6, 95, 70)
p.space_before = Pt(8)

c19_2 = s19.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.85), Inches(1.4), Inches(5.65), Inches(5.3))
c19_2.fill.solid()
c19_2.fill.fore_color.rgb = CARD_BG
c19_2.line.color.rgb = CARD_BORDER
tf = c19_2.text_frame
tf.word_wrap = True
p = tf.paragraphs[0]
p.text = "Tại Sao Kẻ Tấn Công Không Thể Vượt Qua?"
p.font.size = Pt(15)
p.font.bold = True
p.font.color.rgb = ACCENT_BLUE
p = tf.add_paragraph()
p.text = "1. Nguyên tắc cặp Token:\nMáy chủ phát hành 1 token mã hóa trong Cookie và 1 token bí mật nhúng vào trường input ẩn của Form.\n\n2. Bảo vệ bằng Same-Origin Policy (SOP):\nTrang web độc hại của hacker ở tên miền khác (Cross-Origin) TUYỆT ĐỐI KHÔNG THỂ đọc trộm token bí mật trong form của nạn nhân.\n\n3. Cấu hình Cookie SameSite = Strict:\nNgăn chặn trình duyệt tự động gửi cookie trong các cross-site request, triệt tiêu tận gốc tiền đề tấn công của CSRF."
p.font.size = Pt(12)
p.font.color.rgb = TEXT_BODY
p.space_before = Pt(10)

# Slide 20: CSRF trong AJAX & Header
s20 = prs.slides.add_slide(blank_layout)
add_header(s20, "3.6. CSRF Attack: Phòng Thủ Trong AJAX & Fetch API (Buổi 9)", "CHỦ ĐỀ 3 • NGUYỄN MINH CƯỜNG (23A1001D0058)")
c20_main = s20.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.4), Inches(11.733), Inches(5.3))
c20_main.fill.solid()
c20_main.fill.fore_color.rgb = CARD_BG
c20_main.line.color.rgb = CARD_BORDER
tf = c20_main.text_frame
tf.word_wrap = True
p = tf.paragraphs[0]
p.text = "Kỹ Thuật Truyền Token Khi Sử Dụng AJAX Fetch API (Slide Buổi 9)"
p.font.size = Pt(15)
p.font.bold = True
p.font.color.rgb = ACCENT_BLUE

p = tf.add_paragraph()
p.text = "// TRONG JAVASCRIPT FETCH API:\nconst token = document.querySelector('input[name=\"__RequestVerificationToken\"]').value;\n\nfetch('/Csrf/TransferAjax', {\n    method: 'POST',\n    headers: {\n        'Content-Type': 'application/json',\n        'RequestVerificationToken': token // TRUYỀN TOKEN QUA HEADER\n    },\n    body: JSON.stringify({ amount: 5000000 })\n})\n.then(response => response.json());"
p.font.name = "Courier New"
p.font.size = Pt(10)
p.font.color.rgb = TEXT_HEAD
p.space_before = Pt(8)

p = tf.add_paragraph()
p.text = "📌 Tích hợp vào bài tập lớn:\nĐối với các thao tác đặt lịch không reload trang (AJAX Fetch API), nhóm đính kèm header RequestVerificationToken để đảm bảo vừa có trải nghiệm người dùng mượt mà, vừa đảm bảo an toàn tuyệt đối trước các đợt tấn công giả mạo yêu cầu."
p.font.size = Pt(12.5)
p.font.color.rgb = TEXT_BODY
p.space_before = Pt(12)


# ==============================================================================
# TỔNG KẾT & ĐỒ ÁN BTL — 2 SLIDES (SLIDES 21 - 22)
# ==============================================================================
# Slide 21: Bảng so sánh tổng kết 3 lỗ hổng
s21 = prs.slides.add_slide(blank_layout)
add_header(s21, "Tổng Kết: So Sánh Đối Chứng 3 Lỗ Hổng Web Kinh Điển")
table_shape = s21.shapes.add_table(4, 4, Inches(0.8), Inches(1.5), Inches(11.733), Inches(5.2))
tbl = table_shape.table

headers = ["Tiêu chí so sánh", "SQL Injection (Học)", "Stored XSS (Bình)", "CSRF Attack (Cường)"]
for i, h in enumerate(headers):
    cell = tbl.cell(0, i)
    cell.fill.solid()
    cell.fill.fore_color.rgb = ACCENT_BLUE
    p = cell.text_frame.paragraphs[0]
    p.text = h
    p.font.bold = True
    p.font.size = Pt(12)
    p.font.color.rgb = WHITE

rows_data = [
    ("Mục tiêu tấn công", "Tầng Cơ sở dữ liệu (SQL Server Engine)", "Tầng Trình duyệt người dùng (Client DOM)", "Quyền hạn của phiên làm việc (Session/Cookie)"),
    ("Mức độ nghiêm trọng", "CRITICAL (9.8 / 10.0)", "HIGH (7.5 - 8.5 / 10.0)", "MEDIUM / HIGH (6.5 - 8.0 / 10.0)"),
    ("Giải pháp cốt lõi", "Parameterized Query / EF Core LINQ", "Razor HTML Encoding + HttpOnly Cookie", "Anti-Forgery Token + SameSite Cookie")
]

for row_idx, rdata in enumerate(rows_data, start=1):
    for col_idx, text in enumerate(rdata):
        cell = tbl.cell(row_idx, col_idx)
        cell.fill.solid()
        cell.fill.fore_color.rgb = CARD_BG if row_idx % 2 == 1 else WHITE
        p = cell.text_frame.paragraphs[0]
        p.text = text
        p.font.size = Pt(11)
        p.font.color.rgb = TEXT_HEAD if col_idx == 0 else TEXT_BODY
        if col_idx == 1 and "CRITICAL" in text:
            p.font.bold = True
            p.font.color.rgb = ACCENT_RED

# Slide 22: Cam kết đồ án BTL & Lời cảm ơn
s22 = prs.slides.add_slide(blank_layout)
add_header(s22, "Áp Dụng Kiến Trúc An Toàn Vào Đồ Án Website Tư Vấn Trực Tuyến")

main_c22 = s22.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.4), Inches(11.733), Inches(5.3))
main_c22.fill.solid()
main_c22.fill.fore_color.rgb = CARD_BG
main_c22.line.color.rgb = CARD_BORDER
tf = main_c22.text_frame
tf.word_wrap = True

p = tf.paragraphs[0]
p.text = "CAM KẾT CHUẨN AN TOÀN ĐA TẦNG CHO BTL (NHÓM WNC.G01):"
p.font.size = Pt(16)
p.font.bold = True
p.font.color.rgb = ACCENT_BLUE

p = tf.add_paragraph()
p.text = "1. Kiến trúc bảo mật (Security View - Buổi 8): Toàn bộ hệ thống được xây dựng theo nguyên tắc 'Security by Design'. Không xem an ninh là tính năng chắp vá mà tích hợp sâu vào kiến trúc 3 tầng.\n\n2. Bảo vệ CSDL tuyệt đối (Nguyễn Danh Học): 100% các câu truy vấn đặt lịch, hồ sơ chuyên gia và đánh giá chất lượng được quản lý bằng Entity Framework Core Code-First với LINQ Parameterized.\n\n3. Bảo vệ dữ liệu người dùng (Nguyễn Thanh Bình): 100% nội dung phản hồi tư vấn và bình luận được Razor View mã hóa tự động, Cookie cấu hình HttpOnly = true và Secure = true.\n\n4. Bảo vệ giao dịch (Nguyễn Minh Cường): Bắt buộc kiểm tra Anti-Forgery Token trên toàn bộ các form đặt lịch hẹn, nạp tiền ví và đổi mật khẩu."
p.font.size = Pt(12.5)
p.font.color.rgb = TEXT_BODY
p.space_before = Pt(10)

p = tf.add_paragraph()
p.text = "XIN TRÂN TRỌNG CẢM ƠN THẦY VÀ CÁC BẠN ĐÃ THEO DÕI!\nNhóm WNC.G01 sẵn sàng lắng nghe câu hỏi và nhận xét từ giảng viên."
p.font.size = Pt(15)
p.font.bold = True
p.font.color.rgb = ACCENT_RED
p.space_before = Pt(16)

out_file = "/media/hocjsoo/New Volume/OWASP_Demo/Slide_OWASP_Top3_Nhom_WNC_G01.pptx"
prs.save(out_file)
print(f"Master 22-slides presentation saved successfully to: {out_file}")
