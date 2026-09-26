import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)

# ==============================================================================
# BẢNG MÀU GAMMA / CANVA LIGHT THEME CHO MÁY CHIẾU (HIGH CONTRAST)
# ==============================================================================
BG_WHITE = RGBColor(255, 255, 255)         # Nền trắng tinh khiết phản xạ ánh sáng tốt nhất
CARD_BG = RGBColor(248, 250, 252)          # Slate 50 (Thẻ bo góc mềm)
CARD_BORDER = RGBColor(226, 232, 240)      # Slate 200 (Viền tinh tế)

TEXT_HEAD = RGBColor(15, 23, 42)           # Slate 900 (Đen sẫm tương phản cao)
TEXT_BODY = RGBColor(51, 65, 85)           # Slate 700 (Dễ đọc từ xa)
TEXT_MUTED = RGBColor(100, 116, 139)       # Slate 500

ACCENT_BLUE = RGBColor(29, 78, 216)        # Blue 700 (Chuẩn ĐH Mở HOU)
BLUE_LIGHT = RGBColor(239, 246, 255)       # Blue 50
BLUE_BORDER = RGBColor(191, 219, 254)      # Blue 200

ACCENT_RED = RGBColor(190, 18, 60)         # Rose 700 (Cảnh báo lỗi / Tấn công)
RED_LIGHT = RGBColor(255, 241, 242)        # Rose 50
RED_BORDER = RGBColor(254, 205, 211)       # Rose 200

ACCENT_GREEN = RGBColor(4, 120, 87)        # Emerald 700 (An toàn / Phòng thủ)
GREEN_LIGHT = RGBColor(236, 253, 245)      # Emerald 50
GREEN_BORDER = RGBColor(167, 243, 208)     # Emerald 200

ACCENT_AMBER = RGBColor(180, 83, 9)        # Amber 700
AMBER_LIGHT = RGBColor(255, 251, 235)      # Amber 50
AMBER_BORDER = RGBColor(254, 240, 138)     # Amber 200

blank_layout = prs.slide_layouts[6]

def add_header(slide, title_text, category_tag="OWASP TOP 10 • A03:2021 (INJECTION)"):
    # Header Accent Top Line
    top_line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(0.08))
    top_line.fill.solid()
    top_line.fill.fore_color.rgb = ACCENT_BLUE
    top_line.line.fill.background()

    # Category Badge (Pill Style)
    badge = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(0.35), Inches(3.8), Inches(0.38))
    badge.fill.solid()
    badge.fill.fore_color.rgb = BLUE_LIGHT
    badge.line.color.rgb = BLUE_BORDER

    tf_b = badge.text_frame
    tf_b.word_wrap = False
    p_b = tf_b.paragraphs[0]
    p_b.text = category_tag.upper()
    p_b.font.size = Pt(10)
    p_b.font.bold = True
    p_b.font.color.rgb = ACCENT_BLUE

    # Title
    tx = slide.shapes.add_textbox(Inches(0.8), Inches(0.75), Inches(11.7), Inches(0.6))
    tf = tx.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = title_text
    p.font.size = Pt(21)
    p.font.bold = True
    p.font.color.rgb = TEXT_HEAD

    # Footer
    footer = slide.shapes.add_textbox(Inches(0.8), Inches(7.1), Inches(11.7), Inches(0.3))
    ft = footer.text_frame
    p_ft = ft.paragraphs[0]
    p_ft.text = "Nhóm WNC.G01 • Trường Đại học Mở Hà Nội (HOU) • GVHD: ThS. Lê Hữu Dũng"
    p_ft.font.size = Pt(9.5)
    p_ft.font.color.rgb = TEXT_MUTED

# ==============================================================================
# SLIDE 1: BÌA THEO PHONG CÁCH GAMMA / CANVA
# ==============================================================================
s1 = prs.slides.add_slide(blank_layout)

# Top Bar
bar1 = s1.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(0.12))
bar1.fill.solid()
bar1.fill.fore_color.rgb = ACCENT_BLUE
bar1.line.fill.background()

# University Tag
u_badge = s1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.0), Inches(0.8), Inches(5.8), Inches(0.42))
u_badge.fill.solid()
u_badge.fill.fore_color.rgb = BLUE_LIGHT
u_badge.line.color.rgb = BLUE_BORDER
tf = u_badge.text_frame
p = tf.paragraphs[0]
p.text = "TRƯỜNG ĐẠI HỌC MỞ HÀ NỘI — KHOA CÔNG NGHỆ THÔNG TIN"
p.font.size = Pt(10.5)
p.font.bold = True
p.font.color.rgb = ACCENT_BLUE

# Title & Subtitle Box
tb1 = s1.shapes.add_textbox(Inches(1.0), Inches(1.35), Inches(11.333), Inches(3.2))
tf1 = tb1.text_frame
tf1.word_wrap = True

p = tf1.paragraphs[0]
p.text = "BÁO CÁO THỰC NGHIỆM AN TOÀN WEB (OWASP TOP 10)"
p.font.size = Pt(16)
p.font.bold = True
p.font.color.rgb = ACCENT_RED

p = tf1.add_paragraph()
p.text = "Lỗ Hổng SQL Injection (A03:2021)\nKịch Bản Tấn Công & Phòng Thủ Đa Tầng"
p.font.size = Pt(32)
p.font.bold = True
p.font.color.rgb = TEXT_HEAD
p.space_before = Pt(8)

p = tf1.add_paragraph()
p.text = "Môn học: Lập trình Web nâng cao  •  Giảng viên hướng dẫn: ThS. Lê Hữu Dũng"
p.font.size = Pt(14)
p.font.color.rgb = TEXT_BODY
p.space_before = Pt(10)

# Bottom Group Card (3 Students)
g_card = s1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.0), Inches(4.8), Inches(11.333), Inches(1.9))
g_card.fill.solid()
g_card.fill.fore_color.rgb = CARD_BG
g_card.line.color.rgb = CARD_BORDER

tb_g = s1.shapes.add_textbox(Inches(1.3), Inches(4.95), Inches(10.7), Inches(1.6))
tf_g = tb_g.text_frame
tf_g.word_wrap = True
p = tf_g.paragraphs[0]
p.text = "NHÓM THỰC HIỆN: WNC.G01"
p.font.size = Pt(13)
p.font.bold = True
p.font.color.rgb = ACCENT_BLUE

p = tf_g.add_paragraph()
p.text = "1. Nguyễn Danh Học  — MSV: 23A1001D0158 (Phụ trách báo cáo & demo thực nghiệm chính)\n2. Nguyễn Thanh Bình — MSV: 23A1001D0041\n3. Nguyễn Minh Cường — MSV: 23A1001D0058"
p.font.size = Pt(12.5)
p.font.color.rgb = TEXT_HEAD
p.space_before = Pt(6)


# ==============================================================================
# SLIDE 2: BẢN CHẤT LỖ HỔNG
# ==============================================================================
s2 = prs.slides.add_slide(blank_layout)
add_header(s2, "1. Bản Chất Kỹ Thuật Của Lỗ Hổng SQL Injection")

# Left Column (Theoretical Root Cause)
c_left = s2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.45), Inches(5.6), Inches(5.3))
c_left.fill.solid()
c_left.fill.fore_color.rgb = CARD_BG
c_left.line.color.rgb = CARD_BORDER

tb2_1 = s2.shapes.add_textbox(Inches(1.0), Inches(1.6), Inches(5.2), Inches(5.0))
tf2_1 = tb2_1.text_frame
tf2_1.word_wrap = True

p = tf2_1.paragraphs[0]
p.text = "Nguyên Nhân Cốt Lõi"
p.font.size = Pt(16)
p.font.bold = True
p.font.color.rgb = ACCENT_RED

p = tf2_1.add_paragraph()
p.text = "• Không phân tách giữa DỮ LIỆU (Data) và CÂU LỆNH (Code).\n• Ứng dụng nối chuỗi trực tiếp từ người dùng vào câu SQL khiến CSDL không phân biệt được đâu là dữ liệu, đâu là lệnh điều khiển.\n• Hacker dùng các ký tự đặc biệt (', --, ;) để bẻ gãy cấu trúc logic của câu lệnh."
p.font.size = Pt(12.5)
p.font.color.rgb = TEXT_BODY
p.space_before = Pt(8)

p = tf2_1.add_paragraph()
p.text = "Phân Loại Theo Chuẩn OWASP WSTG"
p.font.size = Pt(15)
p.font.bold = True
p.font.color.rgb = ACCENT_BLUE
p.space_before = Pt(16)

p = tf2_1.add_paragraph()
p.text = "1. In-band SQLi: Khai thác trực diện qua Error hoặc UNION.\n2. Inferential SQLi (Blind): Suy luận mù qua Boolean hoặc Time delay.\n3. Out-of-band SQLi: Kích hoạt DNS/HTTP request từ CSDL."
p.font.size = Pt(12)
p.font.color.rgb = TEXT_BODY
p.space_before = Pt(6)

# Right Column (Vulnerable Code Sample)
c_right = s2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.8), Inches(1.45), Inches(5.7), Inches(5.3))
c_right.fill.solid()
c_right.fill.fore_color.rgb = RED_LIGHT
c_right.line.color.rgb = RED_BORDER

tb2_2 = s2.shapes.add_textbox(Inches(7.0), Inches(1.6), Inches(5.3), Inches(5.0))
tf2_2 = tb2_2.text_frame
tf2_2.word_wrap = True

p = tf2_2.paragraphs[0]
p.text = "Mã Nguồn C# Gây Lỗi (Cố Tình Ghép Chuỗi)"
p.font.size = Pt(15)
p.font.bold = True
p.font.color.rgb = ACCENT_RED

p = tf2_2.add_paragraph()
p.text = "// Nối chuỗi trực tiếp từ Request:\nstring query = $\"SELECT * FROM Accounts \" +\n               $\"WHERE Username = '{user}' \" +\n               $\"AND Password = '{pass}'\";\n\nvar users = _context.Accounts\n                    .FromSqlRaw(query)\n                    .ToList();"
p.font.name = "Courier New"
p.font.size = Pt(11)
p.font.color.rgb = RGBColor(159, 18, 57)
p.space_before = Pt(10)

p = tf2_2.add_paragraph()
p.text = "⚠️ Điểm mù chết người:"
p.font.size = Pt(13)
p.font.bold = True
p.font.color.rgb = TEXT_HEAD
p.space_before = Pt(16)

p = tf2_2.add_paragraph()
p.text = "Lập trình viên tin tưởng input '{user}' là vô hại. Dấu nháy đơn (') của hacker lập tức đóng sớm chuỗi dữ liệu, biến phần còn lại thành câu lệnh SQL có chủ đích."
p.font.size = Pt(12)
p.font.color.rgb = TEXT_BODY
p.space_before = Pt(4)


# ==============================================================================
# SLIDE 3: 3 CẤP ĐỘ NGUY HIỂM
# ==============================================================================
s3 = prs.slides.add_slide(blank_layout)
add_header(s3, "2. Độ Nguy Hiểm Thực Tế Của SQL Injection (Tác Động Nghiêm Trọng)")

col_w = Inches(3.64)
gap = Inches(0.39)

# Card 1
c1 = s3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.45), col_w, Inches(5.3))
c1.fill.solid()
c1.fill.fore_color.rgb = RED_LIGHT
c1.line.color.rgb = RED_BORDER

t1 = s3.shapes.add_textbox(Inches(0.95), Inches(1.65), col_w - Inches(0.3), Inches(4.9))
tf = t1.text_frame
tf.word_wrap = True
p = tf.paragraphs[0]
p.text = "🚨 CHIẾM QUYỀN HỆ THỐNG\n(Authentication Bypass)"
p.font.size = Pt(14)
p.font.bold = True
p.font.color.rgb = ACCENT_RED
p = tf.add_paragraph()
p.text = "• Không cần biết mật khẩu của bất kỳ ai.\n• Chỉ với 1 chuỗi payload ngắn, kẻ tấn công đăng nhập thẳng vào tài khoản Quản trị viên (Admin).\n• Chiếm đoạt phiên làm việc, đổi quyền hạn và kiểm soát toàn bộ hệ thống."
p.font.size = Pt(12)
p.font.color.rgb = TEXT_BODY
p.space_before = Pt(10)

# Card 2
c2 = s3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8) + col_w + gap, Inches(1.45), col_w, Inches(5.3))
c2.fill.solid()
c2.fill.fore_color.rgb = AMBER_LIGHT
c2.line.color.rgb = AMBER_BORDER

t2 = s3.shapes.add_textbox(Inches(0.95) + col_w + gap, Inches(1.65), col_w - Inches(0.3), Inches(4.9))
tf = t2.text_frame
tf.word_wrap = True
p = tf.paragraphs[0]
p.text = "📂 ĐÁNH CẮP TOÀN BỘ CSDL\n(Data Exfiltration)"
p.font.size = Pt(14)
p.font.bold = True
p.font.color.rgb = ACCENT_AMBER
p = tf.add_paragraph()
p.text = "• Trích xuất thông tin khách hàng, hồ sơ chuyên gia, số dư tài khoản ngân hàng.\n• Rò rỉ mật khẩu và bệnh án cá nhân (Secret Notes).\n• Vi phạm nghiêm trọng an ninh dữ liệu, phá vỡ uy tín của nền tảng tư vấn trực tuyến."
p.font.size = Pt(12)
p.font.color.rgb = TEXT_BODY
p.space_before = Pt(10)

# Card 3
c3 = s3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8) + (col_w + gap)*2, Inches(1.45), col_w, Inches(5.3))
c3.fill.solid()
c3.fill.fore_color.rgb = BLUE_LIGHT
c3.line.color.rgb = BLUE_BORDER

t3 = s3.shapes.add_textbox(Inches(0.95) + (col_w + gap)*2, Inches(1.65), col_w - Inches(0.3), Inches(4.9))
tf = t3.text_frame
tf.word_wrap = True
p = tf.paragraphs[0]
p.text = "💣 PHÁ HỦY / THỰC THI LỆNH\n(System Takeover)"
p.font.size = Pt(14)
p.font.bold = True
p.font.color.rgb = ACCENT_BLUE
p = tf.add_paragraph()
p.text = "• Xóa sạch CSDL với lệnh DROP TABLE hoặc TRUNCATE.\n• Giả mạo số dư tiền trong tài khoản.\n• Trong các CSDL cấu hình lỏng lẻo, hacker có thể bật xp_cmdshell để chiếm quyền điều khiển cả máy chủ hệ điều hành."
p.font.size = Pt(12)
p.font.color.rgb = TEXT_BODY
p.space_before = Pt(10)


# ==============================================================================
# SLIDE 4: PHÂN TÍCH CƠ CHẾ BẺ GÃY CÚ PHÁP
# ==============================================================================
s4 = prs.slides.add_slide(blank_layout)
add_header(s4, "3. Kịch Bản Tấn Công: Phân Tích Cơ Chế Bẻ Gãy Cú Pháp SQL")

# Main Container Card
main_box = s4.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.45), Inches(11.733), Inches(5.3))
main_box.fill.solid()
main_box.fill.fore_color.rgb = CARD_BG
main_box.line.color.rgb = CARD_BORDER

tb4 = s4.shapes.add_textbox(Inches(1.1), Inches(1.6), Inches(11.1), Inches(5.0))
tf4 = tb4.text_frame
tf4.word_wrap = True

p = tf4.paragraphs[0]
p.text = "Mục tiêu khai thác: Vượt qua xác thực đăng nhập (Authentication Bypass)"
p.font.size = Pt(15)
p.font.bold = True
p.font.color.rgb = ACCENT_RED

p = tf4.add_paragraph()
p.text = "Payload đưa vào Username:   ' OR '1'='1' --\nMật khẩu đưa vào Password: (Nhập bất kỳ hoặc bỏ trống)"
p.font.name = "Courier New"
p.font.size = Pt(13)
p.font.bold = True
p.font.color.rgb = ACCENT_BLUE
p.space_before = Pt(8)

p = tf4.add_paragraph()
p.text = "Câu truy vấn thực tế được SQL Server phân tích & thông dịch:"
p.font.size = Pt(14)
p.font.bold = True
p.font.color.rgb = TEXT_HEAD
p.space_before = Pt(12)

p = tf4.add_paragraph()
p.text = "SELECT * FROM Accounts WHERE Username = '' OR '1'='1' --' AND Password = 'xyz'"
p.font.name = "Courier New"
p.font.size = Pt(13.5)
p.font.bold = True
p.font.color.rgb = ACCENT_RED
p.space_before = Pt(4)

p = tf4.add_paragraph()
p.text = "Giải Phẫu 3 Thành Phần Của Payload:"
p.font.size = Pt(14)
p.font.bold = True
p.font.color.rgb = TEXT_HEAD
p.space_before = Pt(14)

p = tf4.add_paragraph()
p.text = "1. Dấu nháy đơn (') : Đóng sớm chuỗi ký tự hợp lệ của Username, đưa ngữ cảnh trở lại câu lệnh SQL.\n2. Mệnh đề OR '1'='1' : Biểu thức logic chân lý (luôn luôn TRUE cho mọi dòng trong bảng Accounts).\n3. Ký tự chú thích (--) : Vô hiệu hóa hoàn toàn phần kiểm tra mật khẩu phía sau (AND Password = '...')."
p.font.size = Pt(12.5)
p.font.color.rgb = TEXT_BODY
p.space_before = Pt(6)


# ==============================================================================
# SLIDE 5: ĐỐI CHỨNG THỰC NGHIỆM (CANVA / GAMMA STYLE)
# ==============================================================================
s5 = prs.slides.add_slide(blank_layout)
add_header(s5, "4. Đối Chứng Thực Nghiệm: Nhánh Dính Lỗi vs Nhánh Đã Phòng Thủ")

card_w = Inches(5.65)

# Left Card: Vulnerable
box_v = s5.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.45), card_w, Inches(4.2))
box_v.fill.solid()
box_v.fill.fore_color.rgb = RED_LIGHT
box_v.line.color.rgb = RED_BORDER

tb_v = s5.shapes.add_textbox(Inches(1.0), Inches(1.6), card_w - Inches(0.4), Inches(3.9))
tf_v = tb_v.text_frame
tf_v.word_wrap = True

p = tf_v.paragraphs[0]
p.text = "❌ 1. NHÁNH BỊ LỖI (Nối Chuỗi SQL Thuần)"
p.font.size = Pt(14)
p.font.bold = True
p.font.color.rgb = ACCENT_RED

p = tf_v.add_paragraph()
p.text = "• Input khai thác: ' OR '1'='1' --\n• Câu lệnh thực thi:\n  SELECT * FROM Accounts WHERE Username = '' OR '1'='1' --'\n• Cơ chế thông dịch: Mệnh đề WHERE luôn TRUE.\n• Hậu quả thực tế:\n  - Đăng nhập thành công mà KHÔNG CẦN mật khẩu.\n  - Rò rỉ toàn bộ 4 tài khoản CSDL (Admin, Chuyên gia, Khách hàng).\n  - Lộ mật khẩu gốc và số dư tài khoản ngân hàng."
p.font.size = Pt(12)
p.font.color.rgb = TEXT_BODY
p.space_before = Pt(8)

# Right Card: Secure
box_s = s5.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.85), Inches(1.45), card_w, Inches(4.2))
box_s.fill.solid()
box_s.fill.fore_color.rgb = GREEN_LIGHT
box_s.line.color.rgb = GREEN_BORDER

tb_s = s5.shapes.add_textbox(Inches(7.05), Inches(1.6), card_w - Inches(0.4), Inches(3.9))
tf_s = tb_s.text_frame
tf_s.word_wrap = True

p = tf_s.paragraphs[0]
p.text = "✅ 2. NHÁNH AN TOÀN (EF Core Parameterized)"
p.font.size = Pt(14)
p.font.bold = True
p.font.color.rgb = ACCENT_GREEN

p = tf_s.add_paragraph()
p.text = "• Cùng Input khai thác: ' OR '1'='1' --\n• Câu lệnh thực thi:\n  SELECT * FROM Accounts WHERE Username = @p0 AND Pass = @p1\n• Cơ chế thông dịch:\n  - CSDL biên dịch cây truy vấn (Execution Plan) TRƯỚC.\n  - Toàn bộ chuỗi payload được xem là văn bản (String Literal) thuần túy.\n• Kết quả phòng thủ:\n  - Hệ thống chặn đứng, trả về 0 bản ghi.\n  - Báo 'Sai tài khoản/mật khẩu', bảo vệ an toàn 100%."
p.font.size = Pt(12)
p.font.color.rgb = TEXT_BODY
p.space_before = Pt(8)

# Live Demo Banner
callout1 = s5.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(5.8), Inches(11.733), Inches(0.95))
callout1.fill.solid()
callout1.fill.fore_color.rgb = BLUE_LIGHT
callout1.line.color.rgb = BLUE_BORDER

tb_c1 = s5.shapes.add_textbox(Inches(1.0), Inches(5.85), Inches(11.3), Inches(0.85))
tf_c1 = tb_c1.text_frame
tf_c1.word_wrap = True
p = tf_c1.paragraphs[0]
p.text = "💻 [LIVE DEMO TRÊN MÁY]: Chuyển sang Trình duyệt Web (http://localhost:5076)"
p.font.size = Pt(12.5)
p.font.bold = True
p.font.color.rgb = ACCENT_BLUE
p = tf_c1.add_paragraph()
p.text = "Thao tác trực tiếp: Bấm nạp payload → Nhấn submit bên Đỏ (ra toàn bộ CSDL) → Nhấn submit bên Xanh (bị chặn an toàn)."
p.font.size = Pt(11)
p.font.color.rgb = TEXT_BODY
p.space_before = Pt(2)


# ==============================================================================
# SLIDE 6: ĐỐI CHỨNG TẦNG SÂU (API CLI & SQL SERVER 2025)
# ==============================================================================
s6 = prs.slides.add_slide(blank_layout)
add_header(s6, "5. Đối Chứng Tầng Sâu: Tầng API (cURL) & Microsoft SQL Server 2025")

# Box 1: API Proof
b_api = s6.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.45), card_w, Inches(4.2))
b_api.fill.solid()
b_api.fill.fore_color.rgb = CARD_BG
b_api.line.color.rgb = CARD_BORDER

tb_api = s6.shapes.add_textbox(Inches(1.0), Inches(1.6), card_w - Inches(0.4), Inches(3.9))
tf_api = tb_api.text_frame
tf_api.word_wrap = True
p = tf_api.paragraphs[0]
p.text = "Minh Chứng Tầng API / Dòng Lệnh (cURL)"
p.font.size = Pt(14)
p.font.bold = True
p.font.color.rgb = ACCENT_BLUE

p = tf_api.add_paragraph()
p.text = "$ curl -X POST http://localhost:5076/SqlInjection/LoginVulnerable \\\n       -F \"username=' OR '1'='1' --\" -F \"password=any\"\n\nPhản hồi JSON từ Server:\n{\n  \"success\": true,\n  \"isExploited\": true,\n  \"count\": 4,\n  \"accounts\": [ ... 4 tài khoản bị rò rỉ ... ]\n}"
p.font.name = "Courier New"
p.font.size = Pt(10)
p.font.color.rgb = TEXT_HEAD
p.space_before = Pt(6)

p = tf_api.add_paragraph()
p.text = "📌 Ý nghĩa: Chứng minh lỗ hổng nằm ở logic C# Backend, kẻ tấn công dùng script tự động quét và khai thác mà không cần mở web."
p.font.size = Pt(11.5)
p.font.color.rgb = TEXT_BODY
p.space_before = Pt(8)

# Box 2: Database Proof
b_db = s6.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.85), Inches(1.45), card_w, Inches(4.2))
b_db.fill.solid()
b_db.fill.fore_color.rgb = CARD_BG
b_db.line.color.rgb = CARD_BORDER

tb_db = s6.shapes.add_textbox(Inches(7.05), Inches(1.6), card_w - Inches(0.4), Inches(3.9))
tf_db = tb_db.text_frame
tf_db.word_wrap = True
p = tf_db.paragraphs[0]
p.text = "Minh Chứng Tầng CSDL (SQL Server 2025 Thật)"
p.font.size = Pt(14)
p.font.bold = True
p.font.color.rgb = ACCENT_GREEN

p = tf_db.add_paragraph()
p.text = "-- Database: OwaspDemoDB | Bảng: dbo.Accounts\n-- 1. Câu lệnh dính lỗi:\nSELECT * FROM Accounts WHERE Username='' OR '1'='1' --'\n>> Kết quả: 4 rows returned (Bypass thành công)\n\n-- 2. Câu lệnh tham số hóa chuẩn:\nEXEC sp_executesql N'SELECT * FROM Accounts WHERE...',\n     N'@p0 nvarchar(50)', @p0 = N''' OR ''1''=''1'' --'\n>> Kết quả: 0 rows returned (An toàn 100%)"
p.font.name = "Courier New"
p.font.size = Pt(10)
p.font.color.rgb = TEXT_HEAD
p.space_before = Pt(6)

p = tf_db.add_paragraph()
p.text = "📌 Ý nghĩa: Đối chiếu trực tiếp trên máy chủ SQL Server 2025 chứng minh cơ chế phân tách rạch ròi giữa cú pháp lệnh và tham số giá trị."
p.font.size = Pt(11.5)
p.font.color.rgb = TEXT_BODY
p.space_before = Pt(8)

# Live Demo Banner 2
callout2 = s6.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(5.8), Inches(11.733), Inches(0.95))
callout2.fill.solid()
callout2.fill.fore_color.rgb = GREEN_LIGHT
callout2.line.color.rgb = GREEN_BORDER

tb_c2 = s6.shapes.add_textbox(Inches(1.0), Inches(5.85), Inches(11.3), Inches(0.85))
tf_c2 = tb_c2.text_frame
tf_c2.word_wrap = True
p = tf_c2.paragraphs[0]
p.text = "🗄️ [LIVE DEMO TRÊN MÁY]: Chuyển sang VS Code mở bảng dbo.Accounts & file verify_sql_server.sql"
p.font.size = Pt(12.5)
p.font.bold = True
p.font.color.rgb = ACCENT_GREEN
p = tf_c2.add_paragraph()
p.text = "Thao tác trực tiếp: Mở extension MSSQL trên VS Code, đối chiếu bảng dữ liệu gốc và chạy câu query đối chứng."
p.font.size = Pt(11)
p.font.color.rgb = TEXT_BODY
p.space_before = Pt(2)


# ==============================================================================
# SLIDE 7: PHÒNG THỦ TRONG ASP.NET CORE
# ==============================================================================
s7 = prs.slides.add_slide(blank_layout)
add_header(s7, "6. Biện Pháp Phòng Thủ Chuẩn Trong ASP.NET Core MVC")

# Left Column: Code
box7 = s7.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.45), Inches(5.6), Inches(5.3))
box7.fill.solid()
box7.fill.fore_color.rgb = GREEN_LIGHT
box7.line.color.rgb = GREEN_BORDER

tb7_1 = s7.shapes.add_textbox(Inches(1.0), Inches(1.6), Inches(5.2), Inches(5.0))
tf7_1 = tb7_1.text_frame
tf7_1.word_wrap = True

p = tf7_1.paragraphs[0]
p.text = "Cách Viết Code Chuẩn Hóa Với EF Core"
p.font.size = Pt(15)
p.font.bold = True
p.font.color.rgb = ACCENT_GREEN

p = tf7_1.add_paragraph()
p.text = "// Dùng LINQ Parameterized (KHUYÊN DÙNG):\nvar account = _context.Accounts\n    .FirstOrDefault(a => a.Username == user \n                      && a.Password == pass);\n\n// Hoặc dùng FromSqlInterpolated nếu viết SQL thuần:\nvar account = _context.Accounts\n    .FromSqlInterpolated($\"SELECT * FROM Accounts WHERE Username={user} AND Password={pass}\")\n    .FirstOrDefault();"
p.font.name = "Courier New"
p.font.size = Pt(10)
p.font.color.rgb = RGBColor(6, 95, 70)
p.space_before = Pt(8)

p = tf7_1.add_paragraph()
p.text = "Tại sao cách này an toàn tuyệt đối?"
p.font.size = Pt(13)
p.font.bold = True
p.font.color.rgb = TEXT_HEAD
p.space_before = Pt(12)

p = tf7_1.add_paragraph()
p.text = "SQL Server biên dịch cấu trúc lệnh TRƯỚC khi gán dữ liệu. Biến @p0 nhận toàn bộ payload như một chuỗi chữ bình thường, vô hiệu hóa hoàn toàn ý đồ chèn lệnh."
p.font.size = Pt(12)
p.font.color.rgb = TEXT_BODY
p.space_before = Pt(4)

# Right Column: Rules
c_rules = s7.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.8), Inches(1.45), Inches(5.7), Inches(5.3))
c_rules.fill.solid()
c_rules.fill.fore_color.rgb = CARD_BG
c_rules.line.color.rgb = CARD_BORDER

tb7_2 = s7.shapes.add_textbox(Inches(7.0), Inches(1.6), Inches(5.3), Inches(5.0))
tf7_2 = tb7_2.text_frame
tf7_2.word_wrap = True

p = tf7_2.paragraphs[0]
p.text = "Bộ Quy Tắc Phòng Thủ Toàn Diện (Defense in Depth)"
p.font.size = Pt(15)
p.font.bold = True
p.font.color.rgb = ACCENT_BLUE

p = tf7_2.add_paragraph()
p.text = "1. Luôn sử dụng Parameterized Query / ORM:\nTuyệt đối không ghép chuỗi SQL thủ công dưới bất kỳ hình thức nào."
p.font.size = Pt(12)
p.font.color.rgb = TEXT_BODY
p.space_before = Pt(10)

p = tf7_2.add_paragraph()
p.text = "2. Nguyên tắc đặc quyền tối thiểu (Least Privilege):\nTài khoản kết nối CSDL của ứng dụng chỉ có quyền SELECT/INSERT/UPDATE trên các bảng cần thiết, không bao giờ dùng tài khoản 'sa' hay quyền DDL (DROP, ALTER)."
p.font.size = Pt(12)
p.font.color.rgb = TEXT_BODY
p.space_before = Pt(8)

p = tf7_2.add_paragraph()
p.text = "3. Xác thực dữ liệu đầu vào (Input Validation):\nSử dụng Data Annotations ([RegularExpression], [StringLength]) kiểm tra chặt chẽ khuôn dạng dữ liệu."
p.font.size = Pt(12)
p.font.color.rgb = TEXT_BODY
p.space_before = Pt(8)

p = tf7_2.add_paragraph()
p.text = "4. Mã hóa mật khẩu một chiều (Password Hashing):\nSử dụng ASP.NET Core Identity (PBKDF2/BCrypt) để nếu CSDL có bị lộ, mật khẩu vẫn được bảo vệ an toàn."
p.font.size = Pt(12)
p.font.color.rgb = TEXT_BODY
p.space_before = Pt(8)


# ==============================================================================
# SLIDE 8: TỔNG KẾT & CAM KẾT ĐỒ ÁN
# ==============================================================================
s8 = prs.slides.add_slide(blank_layout)
add_header(s8, "7. Tổng Kết & Bài Học Rút Ra Cho Dự Án Website Tư Vấn Trực Tuyến")

main_c8 = s8.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.45), Inches(11.733), Inches(5.3))
main_c8.fill.solid()
main_c8.fill.fore_color.rgb = CARD_BG
main_c8.line.color.rgb = CARD_BORDER

tb8 = s8.shapes.add_textbox(Inches(1.1), Inches(1.65), Inches(11.1), Inches(4.8))
tf8 = tb8.text_frame
tf8.word_wrap = True

p = tf8.paragraphs[0]
p.text = "CAM KẾT CHUẨN AN TOÀN CHO ĐỀ TÀI BTL (WNC.G01):"
p.font.size = Pt(16)
p.font.bold = True
p.font.color.rgb = ACCENT_RED

p = tf8.add_paragraph()
p.text = "1. Nhận thức an toàn: SQL Injection là lỗ hổng nghiêm trọng ở mức độ CRITICAL (9.8/10 theo CVSS v3.1). Lập trình viên phải luôn coi mọi dữ liệu đầu vào từ người dùng là 'không an toàn'."
p.font.size = Pt(13)
p.font.color.rgb = TEXT_BODY
p.space_before = Pt(12)

p = tf8.add_paragraph()
p.text = "2. Phòng thủ triệt để: Trong đồ án Website cung cấp dịch vụ tư vấn trực tuyến, nhóm cam kết áp dụng 100% Entity Framework Core Code-First với Parameterized LINQ cho toàn bộ các thao tác đặt lịch, quản lý chuyên gia và đánh giá chất lượng."
p.font.size = Pt(13)
p.font.color.rgb = TEXT_BODY
p.space_before = Pt(10)

p = tf8.add_paragraph()
p.text = "3. Bảo mật đa tầng (Security View - Buổi 8): Tích hợp chặt chẽ ASP.NET Core Identity để băm mật khẩu, AntiForgeryToken chống CSRF trên toàn bộ form và kiểm soát phân quyền 3 Role (Customer, Specialist, Admin)."
p.font.size = Pt(13)
p.font.color.rgb = TEXT_BODY
p.space_before = Pt(10)

p = tf8.add_paragraph()
p.text = "XIN TRÂN TRỌNG CẢM ƠN THẦY VÀ CÁC BẠN ĐÃ THEO DÕI!\nNhóm WNC.G01 sẵn sàng lắng nghe câu hỏi và nhận xét từ giảng viên."
p.font.size = Pt(15)
p.font.bold = True
p.font.color.rgb = ACCENT_BLUE
p.space_before = Pt(22)

out_file = "/media/hocjsoo/New Volume/OWASP_Demo/Slide_OWASP_SQL_Injection.pptx"
prs.save(out_file)
print(f"Gamma/Canva light theme presentation saved to: {out_file}")
