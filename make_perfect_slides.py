import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)

# Bảng màu chuẩn học thuật (Academic Light Theme)
BG_WHITE = RGBColor(255, 255, 255)
WHITE = RGBColor(255, 255, 255)
CARD_BG = RGBColor(248, 250, 252)          # Slate 50
CARD_BORDER = RGBColor(226, 232, 240)      # Slate 200

TEXT_HEAD = RGBColor(15, 23, 42)           # Slate 900 (Đen kỹ thuật)
TEXT_BODY = RGBColor(51, 65, 85)           # Slate 700 (Xám đậm dễ đọc)
TEXT_MUTED = RGBColor(100, 116, 139)       # Slate 500

ACCENT_BLUE = RGBColor(29, 78, 216)        # Blue 700 (HOU Chuẩn)
BLUE_LIGHT = RGBColor(239, 246, 255)       # Blue 50
BLUE_BORDER = RGBColor(191, 219, 254)

ACCENT_RED = RGBColor(190, 18, 60)         # Rose 700 (Cảnh báo lỗi)
RED_LIGHT = RGBColor(255, 241, 242)        # Rose 50
RED_BORDER = RGBColor(254, 205, 211)

ACCENT_GREEN = RGBColor(4, 120, 87)        # Emerald 700 (An toàn)
GREEN_LIGHT = RGBColor(236, 253, 245)      # Emerald 50
GREEN_BORDER = RGBColor(167, 243, 208)

ACCENT_AMBER = RGBColor(180, 83, 9)        # Amber 700
AMBER_LIGHT = RGBColor(255, 251, 235)      # Amber 50
AMBER_BORDER = RGBColor(254, 240, 138)

blank_layout = prs.slide_layouts[6]

def add_header(slide, title_text, category_tag="OWASP TOP 10 • THỰC NGHIỆM AN TOÀN WEB"):
    top_line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(0.08))
    top_line.fill.solid()
    top_line.fill.fore_color.rgb = ACCENT_BLUE
    top_line.line.fill.background()

    badge = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(0.28), Inches(4.8), Inches(0.36))
    badge.fill.solid()
    badge.fill.fore_color.rgb = BLUE_LIGHT
    badge.line.color.rgb = BLUE_BORDER

    tf_b = badge.text_frame
    tf_b.word_wrap = False
    p_b = tf_b.paragraphs[0]
    p_b.text = category_tag.upper()
    p_b.font.name = "Arial"
    p_b.font.size = Pt(10)
    p_b.font.bold = True
    p_b.font.color.rgb = ACCENT_BLUE

    tx = slide.shapes.add_textbox(Inches(0.8), Inches(0.68), Inches(11.7), Inches(0.6))
    tf = tx.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = title_text
    p.font.name = "Arial"
    p.font.size = Pt(20)
    p.font.bold = True
    p.font.color.rgb = TEXT_HEAD

    footer = slide.shapes.add_textbox(Inches(0.8), Inches(7.15), Inches(11.7), Inches(0.3))
    ft = footer.text_frame
    ft.margin_left = ft.margin_right = ft.margin_top = ft.margin_bottom = 0
    p_ft = ft.paragraphs[0]
    p_ft.text = "Nhóm WNC.G01 • Trường Đại học Mở Hà Nội (HOU) • Giảng viên: ThS. Lê Hữu Dũng"
    p_ft.font.name = "Arial"
    p_ft.font.size = Pt(9.5)
    p_ft.font.color.rgb = TEXT_MUTED

def add_p(tf, text, size=12, bold=False, color=TEXT_BODY, space_before=0, font_name="Arial"):
    p = tf.add_paragraph()
    p.text = text
    p.font.name = font_name
    p.font.size = Pt(size)
    p.font.bold = bold
    p.font.color.rgb = color
    if space_before:
        p.space_before = Pt(space_before)
    return p

card_w = Inches(5.65)
col_w = Inches(3.64)
gap = Inches(0.2)

# ==============================================================================
# SLIDE 1: BÌA BÁO CÁO
# ==============================================================================
s1 = prs.slides.add_slide(blank_layout)
bar1 = s1.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(0.12))
bar1.fill.solid()
bar1.fill.fore_color.rgb = ACCENT_BLUE
bar1.line.fill.background()

u_badge = s1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.0), Inches(0.7), Inches(5.8), Inches(0.42))
u_badge.fill.solid()
u_badge.fill.fore_color.rgb = BLUE_LIGHT
u_badge.line.color.rgb = BLUE_BORDER
p = u_badge.text_frame.paragraphs[0]
p.text = "TRƯỜNG ĐẠI HỌC MỞ HÀ NỘI — KHOA CÔNG NGHỆ THÔNG TIN"
p.font.name = "Arial"
p.font.size = Pt(10.5)
p.font.bold = True
p.font.color.rgb = ACCENT_BLUE

tb1 = s1.shapes.add_textbox(Inches(1.0), Inches(1.25), Inches(11.333), Inches(3.2))
tf1 = tb1.text_frame
tf1.word_wrap = True
p = tf1.paragraphs[0]
p.text = "BÁO CÁO THỰC NGHIỆM AN TOÀN WEB (OWASP TOP 10)"
p.font.name = "Arial"
p.font.size = Pt(15)
p.font.bold = True
p.font.color.rgb = ACCENT_RED

add_p(tf1, "Nghiên Cứu Thực Nghiệm Lỗ Hổng Web & Cơ Chế Phòng Thủ\n(SQL Injection • Stored XSS • CSRF Attack)", size=28, bold=True, color=TEXT_HEAD, space_before=8)
add_p(tf1, "Học phần: Lập trình Web nâng cao  •  Giảng viên hướng dẫn: ThS. Lê Hữu Dũng", size=13.5, color=TEXT_BODY, space_before=8)

# 3 Thẻ thành viên
m1 = s1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.0), Inches(4.7), col_w, Inches(2.0))
m1.fill.solid()
m1.fill.fore_color.rgb = CARD_BG
m1.line.color.rgb = BLUE_BORDER
tf = m1.text_frame
p = tf.paragraphs[0]
p.text = "CHỦ ĐỀ 1: SQL INJECTION"
p.font.name = "Arial"
p.font.size = Pt(11)
p.font.bold = True
p.font.color.rgb = ACCENT_BLUE
add_p(tf, "Nguyễn Danh Học\nMSV: 23A1001D0158\nTrưởng nhóm • Phụ trách CSDL & API", size=12, color=TEXT_HEAD, space_before=4)

m2 = s1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.0) + col_w + gap, Inches(4.7), col_w, Inches(2.0))
m2.fill.solid()
m2.fill.fore_color.rgb = CARD_BG
m2.line.color.rgb = RED_BORDER
tf = m2.text_frame
p = tf.paragraphs[0]
p.text = "CHỦ ĐỀ 2: STORED XSS"
p.font.name = "Arial"
p.font.size = Pt(11)
p.font.bold = True
p.font.color.rgb = ACCENT_RED
add_p(tf, "Nguyễn Thanh Bình\nMSV: 23A1001D0041\nThành viên • Phụ trách Cookie & DOM", size=12, color=TEXT_HEAD, space_before=4)

m3 = s1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.0) + (col_w + gap)*2, Inches(4.7), col_w, Inches(2.0))
m3.fill.solid()
m3.fill.fore_color.rgb = CARD_BG
m3.line.color.rgb = GREEN_BORDER
tf = m3.text_frame
p = tf.paragraphs[0]
p.text = "CHỦ ĐỀ 3: CSRF ATTACK"
p.font.name = "Arial"
p.font.size = Pt(11)
p.font.bold = True
p.font.color.rgb = ACCENT_GREEN
add_p(tf, "Nguyễn Minh Cường\nMSV: 23A1001D0058\nThành viên • Phụ trách Giao dịch & Token", size=12, color=TEXT_HEAD, space_before=4)


# ==============================================================================
# CHỦ ĐỀ 1: SQL INJECTION — NGUYỄN DANH HỌC
# ==============================================================================
# Slide 2: Lý thuyết bản chất
s2 = prs.slides.add_slide(blank_layout)
add_header(s2, "1.1. SQL Injection: Vị Trí Trong Kiến Trúc 3 Tầng & Nguyên Nhân Gốc", "CHỦ ĐỀ 1 • NGUYỄN DANH HỌC (23A1001D0158)")
c2_1 = s2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.35), card_w, Inches(5.4))
c2_1.fill.solid()
c2_1.fill.fore_color.rgb = CARD_BG
c2_1.line.color.rgb = CARD_BORDER
tf = c2_1.text_frame
p = tf.paragraphs[0]
p.text = "Vị Trí Xảy Ra Trong Kiến Trúc 3 Tầng"
p.font.name = "Arial"
p.font.size = Pt(15)
p.font.bold = True
p.font.color.rgb = ACCENT_RED
add_p(tf, "• Mô hình 3 tầng: Client (Browser) ➔ Web Server (ASP.NET Core) ➔ Database Server (SQL Server 2025).", size=12.5, bold=True, color=TEXT_HEAD, space_before=6)
add_p(tf, "• Luồng dữ liệu bình thường: Người dùng gửi dữ liệu xác thực (username/password) qua HTTP POST. Controller nhận tham số và chuyển cho SQL Server thực thi.", size=12, space_before=6)
add_p(tf, "• Nguyên nhân gốc rễ (Root Cause): Lập trình viên sử dụng kỹ thuật ghép chuỗi trực tiếp (String Concatenation) để tạo câu lệnh SQL từ dữ liệu người dùng mà không qua cơ chế tham số hóa.", size=12, space_before=6)
add_p(tf, "• Hậu quả phân tích cú pháp: SQL Server biên dịch chuỗi truy vấn đã bị tiêm nhiễm, coi các ký tự điều khiển của hacker là một phần cấu trúc ngữ pháp của câu lệnh.", size=12, space_before=6)

c2_2 = s2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.85), Inches(1.35), card_w, Inches(5.4))
c2_2.fill.solid()
c2_2.fill.fore_color.rgb = RED_LIGHT
c2_2.line.color.rgb = RED_BORDER
tf = c2_2.text_frame
p = tf.paragraphs[0]
p.text = "Đoạn Mã Nguồn Gây Lỗi Cụ Thể (C# / ASP.NET Core)"
p.font.name = "Arial"
p.font.size = Pt(15)
p.font.bold = True
p.font.color.rgb = ACCENT_RED
add_p(tf, "// GHÉP CHUỖI TRỰC TIẾP TRONG CONTROLLER:\nstring rawSql = $\"SELECT * FROM Accounts \" +\n                $\"WHERE Username = '{username}' \" +\n                $\"AND Password = '{password}'\";\n\nvar accounts = _context.Accounts\n                       .FromSqlRaw(rawSql)\n                       .ToList();", size=10.5, color=RGBColor(159, 18, 57), space_before=8, font_name="Courier New")
add_p(tf, "Điểm mù kỹ thuật:\nLập trình viên giả định đầu vào '{username}' luôn là văn bản thông thường. Khi kẻ tấn công chèn ký tự nháy đơn ('), ranh giới giữa Dữ liệu (Data) và Cú pháp lệnh (Code) bị xóa nhòa hoàn toàn.", size=12, bold=True, color=TEXT_HEAD, space_before=14)

# Slide 3: Phân loại & CVSS 9.8
s3 = prs.slides.add_slide(blank_layout)
add_header(s3, "1.2. SQL Injection: Phân Loại Kỹ Thuật OWASP & Đánh Giá Rủi Ro", "CHỦ ĐỀ 1 • NGUYỄN DANH HỌC (23A1001D0158)")
c3_1 = s3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.35), col_w, Inches(5.4))
c3_1.fill.solid()
c3_1.fill.fore_color.rgb = RED_LIGHT
c3_1.line.color.rgb = RED_BORDER
tf = c3_1.text_frame
p = tf.paragraphs[0]
p.text = "VƯỢT QUA XÁC THỰC\n(Authentication Bypass)"
p.font.name = "Arial"
p.font.size = Pt(14)
p.font.bold = True
p.font.color.rgb = ACCENT_RED
add_p(tf, "• Cơ chế: Thay đổi cấu trúc logic của mệnh đề WHERE thành luôn đúng (TRUE).\n• Kẻ tấn công đăng nhập thẳng vào tài khoản Quản trị viên (Admin) mà không cần cung cấp mật khẩu.\n• Chiếm đoạt phiên làm việc hợp lệ và thay đổi quyền hạn người dùng trên hệ thống.", size=12, space_before=10)

c3_2 = s3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8) + col_w + Inches(0.39), Inches(1.35), col_w, Inches(5.4))
c3_2.fill.solid()
c3_2.fill.fore_color.rgb = AMBER_LIGHT
c3_2.line.color.rgb = AMBER_BORDER
tf = c3_2.text_frame
p = tf.paragraphs[0]
p.text = "TRÍCH XUẤT CƠ SỞ DỮ LIỆU\n(Data Exfiltration)"
p.font.name = "Arial"
p.font.size = Pt(14)
p.font.bold = True
p.font.color.rgb = ACCENT_AMBER
add_p(tf, "• Kỹ thuật Union-Based: Ghép nối bảng dữ liệu nhạy cảm vào kết quả truy vấn thông thường.\n• Kỹ thuật Error-Based / Blind: Suy luận dữ liệu qua thông báo lỗi hoặc độ trễ thời gian (Time Delay).\n• Làm rò rỉ toàn bộ hồ sơ khách hàng, mật khẩu và dữ liệu tài chính.", size=12, space_before=10)

c3_3 = s3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8) + (col_w + Inches(0.39))*2, Inches(1.35), col_w, Inches(5.4))
c3_3.fill.solid()
c3_3.fill.fore_color.rgb = BLUE_LIGHT
c3_3.line.color.rgb = BLUE_BORDER
tf = c3_3.text_frame
p = tf.paragraphs[0]
p.text = "PHÁ HỦY / THỰC THI LỆNH\n(System Takeover)"
p.font.name = "Arial"
p.font.size = Pt(14)
p.font.bold = True
p.font.color.rgb = ACCENT_BLUE
add_p(tf, "• Can thiệp DDL: Thực thi câu lệnh DROP TABLE, TRUNCATE làm sập dịch vụ.\n• Can thiệp DML: Thay đổi số dư ví, cấp đặc quyền trái phép.\n• Thực thi lệnh hệ điều hành: Kích hoạt thủ tục xp_cmdshell trên máy chủ SQL Server nếu cấu hình thiếu an toàn.", size=12, space_before=10)

# Slide 4: Giải phẫu kịch bản khai thác
s4 = prs.slides.add_slide(blank_layout)
add_header(s4, "1.3. SQL Injection: Phân Tích Cấu Trúc Khai Thác Của Payload", "CHỦ ĐỀ 1 • NGUYỄN DANH HỌC (23A1001D0158)")
c4_main = s4.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.35), Inches(11.733), Inches(5.4))
c4_main.fill.solid()
c4_main.fill.fore_color.rgb = CARD_BG
c4_main.line.color.rgb = CARD_BORDER
tf = c4_main.text_frame
p = tf.paragraphs[0]
p.text = "Kịch Bản: Bypass Đăng Nhập Với Chuỗi Khai Thác ' OR '1'='1' --"
p.font.name = "Arial"
p.font.size = Pt(15)
p.font.bold = True
p.font.color.rgb = ACCENT_RED
add_p(tf, "Giá trị Username đưa vào: ' OR '1'='1' --  | Mật khẩu: (tùy ý hoặc để trống)\nTruy vấn được máy chủ SQL Server phân tích & thông dịch:\nSELECT * FROM Accounts WHERE Username = '' OR '1'='1' --' AND Password = 'xyz'", size=12, bold=True, color=ACCENT_BLUE, space_before=6, font_name="Courier New")
add_p(tf, "Giải Phẫu 3 Thành Phần Khai Thác Cốt Lõi:\n1. Dấu nháy đơn (') : Đóng sớm chuỗi ký tự của trường Username, chuyển ngữ cảnh về cú pháp lệnh SQL.\n2. Mệnh đề OR '1'='1' : Biểu thức logic chân lý (luôn TRUE với mọi bản ghi trong bảng Accounts, khiến toàn bộ điều kiện lọc bị vô hiệu).\n3. Ký tự chú thích (--) : Báo cho máy chủ SQL bỏ qua toàn bộ phần kiểm tra mật khẩu phía sau (AND Password = '...').", size=12.5, space_before=10)
add_p(tf, "Mở Rộng Kỹ Thuật Union-Based:\nPayload: ' UNION SELECT Id, Username, Password, SecretNote FROM Accounts --\nCho phép kẻ tấn công trích xuất danh sách tài khoản quản trị và bí mật hệ thống trên các màn hình tra cứu thông thường.", size=12.5, bold=True, color=ACCENT_AMBER, space_before=10)

# Slide 5: Minh chứng Web UI (Nhúng ảnh thật - CĂN CHUẨN KHÔNG ĐÈ)
s5 = prs.slides.add_slide(blank_layout)
add_header(s5, "1.4. SQL Injection: Minh Chứng Kiểm Thử Qua Chrome DevTools (Network)", "CHỦ ĐỀ 1 • NGUYỄN DANH HỌC (23A1001D0158)")
p_sqli = "/media/hocjsoo/New Volume/OWASP_Demo/screenshots/01_sqli_devtools_proof.png"
if os.path.exists(p_sqli):
    s5.shapes.add_picture(p_sqli, Inches(2.06), Inches(1.35), width=Inches(9.2))

cap5 = s5.shapes.add_textbox(Inches(0.8), Inches(6.62), Inches(11.733), Inches(0.4))
tf5 = cap5.text_frame
# = tf5.margin_bottom = tf5.margin_left = tf5.margin_right = 0
p = tf5.paragraphs[0]
p.text = "Minh chứng đối chứng: Kiểm tra gói tin mạng (Network Fetch/XHR): Payload gửi lên và JSON trích xuất CSDL ' OR '1'='1' -- (Trái) vs Chặn đứng an toàn với EF Core (Phải)"
p.font.name = "Arial"
p.font.size = Pt(11)
p.font.bold = True
p.font.color.rgb = ACCENT_BLUE
p.alignment = PP_ALIGN.CENTER

# Slide 6: Minh chứng Tầng Sâu (cURL & SQL Server)
s6 = prs.slides.add_slide(blank_layout)
add_header(s6, "1.5. SQL Injection: Minh Chứng Tầng API (cURL) & SQL Server 2025", "CHỦ ĐỀ 1 • NGUYỄN DANH HỌC (23A1001D0158)")
b6_api = s6.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.35), card_w, Inches(4.7))
b6_api.fill.solid()
b6_api.fill.fore_color.rgb = CARD_BG
b6_api.line.color.rgb = CARD_BORDER
tf = b6_api.text_frame
p = tf.paragraphs[0]
p.text = "Minh Chứng Tầng API / Dòng Lệnh (cURL)"
p.font.name = "Arial"
p.font.size = Pt(14)
p.font.bold = True
p.font.color.rgb = ACCENT_BLUE
add_p(tf, "$ curl -X POST http://localhost:5076/SqlInjection/LoginVulnerable \\\n       -F \"username=' OR '1'='1' --\" -F \"password=any\"\n\nPhản hồi JSON từ Server:\n{\n  \"success\": true,\n  \"isExploited\": true,\n  \"count\": 4,\n  \"accounts\": [ ... danh sách tài khoản trích xuất ... ]\n}", size=10, color=TEXT_HEAD, space_before=6, font_name="Courier New")
add_p(tf, "Xác nhận lỗi xuất phát từ tầng xử lý C# Backend, kẻ tấn công có thể khai thác qua API mà không cần tương tác qua giao diện web.", size=11.5, space_before=8)

b6_db = s6.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.85), Inches(1.35), card_w, Inches(4.7))
b6_db.fill.solid()
b6_db.fill.fore_color.rgb = CARD_BG
b6_db.line.color.rgb = CARD_BORDER
tf = b6_db.text_frame
p = tf.paragraphs[0]
p.text = "Minh Chứng Tầng CSDL (SQL Server 2025)"
p.font.name = "Arial"
p.font.size = Pt(14)
p.font.bold = True
p.font.color.rgb = ACCENT_GREEN
add_p(tf, "-- Database: OwaspDemoDB | Bảng: dbo.Accounts\n-- 1. Truy vấn dính lỗi:\nSELECT * FROM Accounts WHERE Username='' OR '1'='1' --'\n>> Kết quả: 4 rows returned (Bypass thành công)\n\n-- 2. Truy vấn tham số hóa:\nEXEC sp_executesql N'SELECT * FROM Accounts WHERE...',\n     N'@p0 nvarchar(50)', @p0 = N''' OR ''1''=''1'' --'\n>> Kết quả: 0 rows returned (An toàn)", size=10, color=TEXT_HEAD, space_before=6, font_name="Courier New")
add_p(tf, "Đối chiếu trực tiếp trên máy chủ CSDL chứng minh cơ chế phân tách rạch ròi giữa cú pháp lệnh và tham số giá trị.", size=11.5, space_before=8)

callout6 = s6.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(6.15), Inches(11.733), Inches(0.85))
callout6.fill.solid()
callout6.fill.fore_color.rgb = GREEN_LIGHT
callout6.line.color.rgb = GREEN_BORDER
tf = callout6.text_frame
tf.margin_top = Inches(0.08)
p = tf.paragraphs[0]
p.text = "Thao tác trên công cụ: VS Code MSSQL Extension & File verify_sql_server.sql"
p.font.name = "Arial"
p.font.size = Pt(11)
p.font.bold = True
p.font.color.rgb = ACCENT_GREEN
add_p(tf, "Dữ liệu được truy vấn và kiểm thử trực tiếp trên máy chủ CSDL SQL Server 2025 (Database: OwaspDemoDB).", size=10.5, space_before=1)

# Slide 7: Cơ chế phòng thủ chuẩn hóa
s7 = prs.slides.add_slide(blank_layout)
add_header(s7, "1.6. SQL Injection: Giải Pháp Phòng Thủ Chuẩn Hóa", "CHỦ ĐỀ 1 • NGUYỄN DANH HỌC (23A1001D0158)")
c7_1 = s7.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.35), card_w, Inches(5.4))
c7_1.fill.solid()
c7_1.fill.fore_color.rgb = GREEN_LIGHT
c7_1.line.color.rgb = GREEN_BORDER
tf = c7_1.text_frame
p = tf.paragraphs[0]
p.text = "Kỹ Thuật Parameterized Query Với EF Core"
p.font.name = "Arial"
p.font.size = Pt(15)
p.font.bold = True
p.font.color.rgb = ACCENT_GREEN
add_p(tf, "// SỬ DỤNG LINQ PARAMETERIZED (KHUYÊN DÙNG):\nvar account = _context.Accounts\n    .FirstOrDefault(a => a.Username == user \n                      && a.Password == pass);\n\n// TRUY VẤN THUẦN QUA FROMSQLINTERPOLATED:\nvar account = _context.Accounts\n    .FromSqlInterpolated($\"SELECT * FROM Accounts WHERE Username={user} AND Password={pass}\")\n    .FirstOrDefault();", size=10, color=RGBColor(6, 95, 70), space_before=8, font_name="Courier New")
add_p(tf, "Cơ chế bảo vệ:\nSQL Server biên dịch kế hoạch thực thi (Execution Plan) trước khi nhận tham số. Đầu vào của người dùng được gán vào biến @p0 như một chuỗi văn bản (literal), vô hiệu hóa hoàn toàn khả năng chèn lệnh.", size=12, space_before=10)

c7_2 = s7.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.85), Inches(1.35), card_w, Inches(5.4))
c7_2.fill.solid()
c7_2.fill.fore_color.rgb = CARD_BG
c7_2.line.color.rgb = CARD_BORDER
tf = c7_2.text_frame
p = tf.paragraphs[0]
p.text = "Bộ Quy Tắc Phòng Thủ Toàn Diện (Defense-in-Depth)"
p.font.name = "Arial"
p.font.size = Pt(15)
p.font.bold = True
p.font.color.rgb = ACCENT_BLUE
add_p(tf, "1. Tham số hóa truy vấn:\nTuyệt đối không ghép chuỗi SQL thủ công trong mã nguồn.\n\n2. Đặc quyền tối thiểu (Least Privilege):\nTài khoản kết nối CSDL chỉ cấp các quyền SELECT, INSERT, UPDATE cần thiết; không dùng tài khoản 'sa' hay quyền DDL trong môi trường chạy ứng dụng.\n\n3. Kiểm tra dữ liệu đầu vào (Input Validation):\nSử dụng Data Annotations để giới hạn định dạng và độ dài dữ liệu.\n\n4. Băm mật khẩu (Password Hashing):\nSử dụng ASP.NET Core Identity (PBKDF2) để bảo vệ thông tin ngay cả khi CSDL bị lộ.", size=12, space_before=8)


# ==============================================================================
# CHỦ ĐỀ 2: STORED XSS — NGUYỄN THANH BÌNH
# ==============================================================================
# Slide 8: Bản chất XSS
s8 = prs.slides.add_slide(blank_layout)
add_header(s8, "2.1. Stored XSS: Bản Chất Kỹ Thuật & Phân Loại Trong Web", "CHỦ ĐỀ 2 • NGUYỄN THANH BÌNH (23A1001D0041)")
c8_1 = s8.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.35), card_w, Inches(5.4))
c8_1.fill.solid()
c8_1.fill.fore_color.rgb = CARD_BG
c8_1.line.color.rgb = CARD_BORDER
tf = c8_1.text_frame
p = tf.paragraphs[0]
p.text = "Phân Biệt 3 Loại Cross-Site Scripting (XSS)"
p.font.name = "Arial"
p.font.size = Pt(15)
p.font.bold = True
p.font.color.rgb = ACCENT_RED
add_p(tf, "• 1. Reflected XSS (Phản xạ): Mã độc nằm trong URL (query string), kích hoạt khi nạn nhân click vào liên kết độc hại.", size=12, space_before=6)
add_p(tf, "• 2. DOM-based XSS: Lỗ hổng nằm hoàn toàn ở mã kịch bản phía Client (JavaScript đọc và ghi trực tiếp vào DOM mà không qua máy chủ).", size=12, space_before=6)
add_p(tf, "• 3. Stored XSS (Lưu trữ - Nghiêm trọng nhất): Mã độc được lưu trữ vĩnh viễn trên máy chủ CSDL. Mọi người dùng truy cập trang web đều bị tấn công tự động mà không cần bấm liên kết lạ.", size=12, bold=True, color=ACCENT_RED, space_before=6)
add_p(tf, "• Vị trí luồng dữ liệu: Client (Inject) ➔ Web Server (Lưu) ➔ Database (Lưu trữ) ➔ Web Server (Đọc) ➔ Nạn nhân khác (Trình duyệt thực thi).", size=12, space_before=6)

c8_2 = s8.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.85), Inches(1.35), card_w, Inches(5.4))
c8_2.fill.solid()
c8_2.fill.fore_color.rgb = RED_LIGHT
c8_2.line.color.rgb = RED_BORDER
tf = c8_2.text_frame
p = tf.paragraphs[0]
p.text = "Ngữ Cảnh Thực Tế & Lạm Dụng @@Html.Raw"
p.font.name = "Arial"
p.font.size = Pt(15)
p.font.bold = True
p.font.color.rgb = ACCENT_RED
add_p(tf, "// CONTROLLER LƯU NGUYÊN VĂN NỘI DUNG VÀO CSDL:\nvar comment = new Comment {\n    Author = author,\n    Content = content // Chứa mã <script> độc hại\n};\n_context.Comments.Add(comment);\n\n// RAZOR VIEW LẠM DỤNG HTML.RAW ĐỂ HIỂN THỊ:\n<div class=\"comment-body\">\n    @Html.Raw(comment.Content)\n</div>", size=10, color=RGBColor(159, 18, 57), space_before=6, font_name="Courier New")
add_p(tf, "Nguyên nhân phát sinh:\nLập trình viên muốn hỗ trợ định dạng chữ in đậm, xuống dòng cho bài đánh giá nên dùng @Html.Raw(), vô tình cho phép trình duyệt của người dùng thực thi các thẻ HTML/JavaScript nguy hiểm.", size=12, bold=True, color=TEXT_HEAD, space_before=10)

# Slide 9: Mối đe dọa XSS
s9 = prs.slides.add_slide(blank_layout)
add_header(s9, "2.2. Stored XSS: Đánh Giá Rủi Ro & Cơ Chế Chiếm Đoạt Phiên", "CHỦ ĐỀ 2 • NGUYỄN THANH BÌNH (23A1001D0041)")
c9_1 = s9.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.35), col_w, Inches(5.4))
c9_1.fill.solid()
c9_1.fill.fore_color.rgb = RED_LIGHT
c9_1.line.color.rgb = RED_BORDER
tf = c9_1.text_frame
p = tf.paragraphs[0]
p.text = "ĐÁNH CẮP COOKIE & PHIÊN\n(Session Hijacking)"
p.font.name = "Arial"
p.font.size = Pt(14)
p.font.bold = True
p.font.color.rgb = ACCENT_RED
add_p(tf, "• Đoạn script đọc thuộc tính document.cookie để trích xuất token phiên làm việc.\n• Dữ liệu xác thực được gửi ngầm về máy chủ điều khiển của kẻ tấn công.\n• Kẻ tấn công sử dụng token để mạo danh người dùng hợp lệ.", size=12, space_before=10)

c9_2 = s9.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8) + col_w + Inches(0.39), Inches(1.35), col_w, Inches(5.4))
c9_2.fill.solid()
c9_2.fill.fore_color.rgb = AMBER_LIGHT
c9_2.line.color.rgb = AMBER_BORDER
tf = c9_2.text_frame
p = tf.paragraphs[0]
p.text = "THAY ĐỔI CẤU TRÚC DOM\n(DOM Manipulation / Phishing)"
p.font.name = "Arial"
p.font.size = Pt(14)
p.font.bold = True
p.font.color.rgb = ACCENT_AMBER
add_p(tf, "• Thay đổi cấu trúc giao diện trang web, chèn các form đăng nhập giả mạo.\n• Lừa người dùng nhập thông tin nhạy cảm trên giao diện đã bị sửa đổi.\n• Tự động điều hướng người dùng sang các liên kết độc hại.", size=12, space_before=10)

c9_3 = s9.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8) + (col_w + Inches(0.39))*2, Inches(1.35), col_w, Inches(5.4))
c9_3.fill.solid()
c9_3.fill.fore_color.rgb = BLUE_LIGHT
c9_3.line.color.rgb = BLUE_BORDER
tf = c9_3.text_frame
p = tf.paragraphs[0]
p.text = "THỰC THI THAO TÁC TRÁI PHÉP\n(Client-Side Exploitation)"
p.font.name = "Arial"
p.font.size = Pt(14)
p.font.bold = True
p.font.color.rgb = ACCENT_BLUE
add_p(tf, "• Ghi nhận thao tác bàn phím (Keylogging) khi người dùng nhập liệu.\n• Tự động thực hiện các hành động gửi tin nhắn hoặc đánh giá mạo danh người dùng.\n• Tận dụng quyền truy cập trình duyệt để mở rộng phạm vi tấn công.", size=12, space_before=10)

# Slide 10: Phân tích kỹ thuật payload XSS
s10 = prs.slides.add_slide(blank_layout)
add_header(s10, "2.3. Stored XSS: Phân Tích Các Kỹ Thuật Tiêm Payload", "CHỦ ĐỀ 2 • NGUYỄN THANH BÌNH (23A1001D0041)")
c10_main = s10.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.35), Inches(11.733), Inches(5.4))
c10_main.fill.solid()
c10_main.fill.fore_color.rgb = CARD_BG
c10_main.line.color.rgb = CARD_BORDER
tf = c10_main.text_frame
p = tf.paragraphs[0]
p.text = "3 Biến Thể Payload XSS Thực Nghiệm Trong Hệ Thống"
p.font.name = "Arial"
p.font.size = Pt(15)
p.font.bold = True
p.font.color.rgb = ACCENT_RED
add_p(tf, "1. Payload 1 (Thẻ Script Cơ Bản):\n   <script>alert('Session Cookie: ' + document.cookie);</script>\n   Cơ chế: Trình duyệt phân tích cú pháp HTML, nhận diện khối mã và thực thi câu lệnh JavaScript.", size=12.5, space_before=8)
add_p(tf, "2. Payload 2 (Vượt Qua Bộ Lọc Bằng Trình Xử Lý Sự Kiện HTML):\n   <img src=x onerror=alert('XSS triggered via image onerror')>\n   Cơ chế: Khi ứng dụng chỉ lọc từ khóa '<script>', kẻ tấn công sử dụng thuộc tính 'onerror' của thẻ hình ảnh khi đường dẫn ảnh không tồn tại để kích hoạt JavaScript.", size=12.5, space_before=8)
add_p(tf, "3. Payload 3 (Thay Đổi Toàn Bộ Cấu Trúc Trang Web):\n   <script>document.body.innerHTML='<h2 style=color:red;text-align:center>Trang Web Đã Bị Thay Đổi</h2>';</script>\n   Cơ chế: Can thiệp trực tiếp vào thuộc tính body.innerHTML để thay đổi toàn bộ nội dung hiển thị của website.", size=12.5, space_before=8)

# Slide 11: Minh chứng Web UI XSS (Nhúng ảnh thật - CĂN CHUẨN KHÔNG ĐÈ)
s11 = prs.slides.add_slide(blank_layout)
add_header(s11, "2.4. Stored XSS: Minh Chứng Kiểm Thử Qua Chrome DevTools (Application & Console)", "CHỦ ĐỀ 2 • NGUYỄN THANH BÌNH (23A1001D0041)")
p_xss = "/media/hocjsoo/New Volume/OWASP_Demo/screenshots/02_xss_devtools_proof.png"
if os.path.exists(p_xss):
    s11.shapes.add_picture(p_xss, Inches(2.06), Inches(1.35), width=Inches(9.2))

cap11 = s11.shapes.add_textbox(Inches(0.8), Inches(6.58), Inches(11.733), Inches(0.4))
tf11 = cap11.text_frame
tf11.margin_top = tf11.margin_bottom = tf11.margin_left = tf11.margin_right = 0
p = tf11.paragraphs[0]
p.text = "Minh chứng đối chứng: Soi Cookie lưu trữ (HttpOnly = false) & Lệnh JavaScript document.cookie đọc trộm Session Token (Trái) vs Cơ chế tự động mã hóa HTML của Razor View (Phải)"
p.font.name = "Arial"
p.font.size = Pt(11)
p.font.bold = True
p.font.color.rgb = ACCENT_RED
p.alignment = PP_ALIGN.CENTER

# Slide 12: Minh chứng CSDL & Cookie HttpOnly
s12 = prs.slides.add_slide(blank_layout)
add_header(s12, "2.5. Stored XSS: Minh Chứng Dữ Liệu CSDL & Bảo Vệ Cookie", "CHỦ ĐỀ 2 • NGUYỄN THANH BÌNH (23A1001D0041)")
b12_db = s12.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.35), card_w, Inches(5.4))
b12_db.fill.solid()
b12_db.fill.fore_color.rgb = CARD_BG
b12_db.line.color.rgb = CARD_BORDER
tf = b12_db.text_frame
p = tf.paragraphs[0]
p.text = "Dữ Liệu Lưu Trong Bảng dbo.Comments"
p.font.name = "Arial"
p.font.size = Pt(14)
p.font.bold = True
p.font.color.rgb = ACCENT_BLUE
add_p(tf, "SELECT Id, Author, Content, IsSecureStored FROM dbo.Comments;\n\nId | Author        | Content                           | IsSecure\n---+---------------+-----------------------------------+---------\n 1 | Le Thanh Binh | Dich vu tu van rat chuyen nghiep! | 1\n 2 | Hacker_XSS    | <script>alert(document.cookie);.. | 0", size=9.5, color=TEXT_HEAD, space_before=6, font_name="Courier New")
add_p(tf, "Phân tích đặc tính lưu trữ:\nMã kịch bản độc hại được lưu trữ trực tiếp trong CSDL. Do đó, lỗ hổng Stored XSS có tính chất nguy hại kéo dài và phát tán tới mọi người dùng xem nội dung đó.", size=11.5, space_before=8)

b12_c = s12.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.85), Inches(1.35), card_w, Inches(5.4))
b12_c.fill.solid()
b12_c.fill.fore_color.rgb = CARD_BG
b12_c.line.color.rgb = CARD_BORDER
tf = b12_c.text_frame
p = tf.paragraphs[0]
p.text = "Cơ Chế Bảo Vệ Với Cờ HttpOnly Cookie"
p.font.name = "Arial"
p.font.size = Pt(14)
p.font.bold = True
p.font.color.rgb = ACCENT_GREEN
add_p(tf, "// CẤU HÌNH COOKIE AN TOÀN TRONG ASP.NET CORE:\nResponse.Cookies.Append(\"AuthSessionToken\", token, new CookieOptions\n{\n    HttpOnly = true, // CẤM JAVASCRIPT ĐỌC TRỘM\n    Secure = true,   // BẮT BUỘC TRUYỀN QUA HTTPS\n    SameSite = SameSiteMode.Strict\n});", size=9.5, color=RGBColor(6, 95, 70), space_before=6, font_name="Courier New")
add_p(tf, "Vai trò của cờ HttpOnly:\nNgăn chặn mã JavaScript đọc thuộc tính document.cookie. Kể cả khi trang web xuất hiện lỗi XSS, kẻ tấn công cũng không thể trích xuất token phiên làm việc.", size=11.5, space_before=8)

# Slide 13: Phòng thủ chuẩn XSS
s13 = prs.slides.add_slide(blank_layout)
add_header(s13, "2.6. Stored XSS: Bộ Giải Pháp Phòng Thủ Chuẩn Trong ASP.NET Core", "CHỦ ĐỀ 2 • NGUYỄN THANH BÌNH (23A1001D0041)")
c13_main = s13.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.35), Inches(11.733), Inches(5.4))
c13_main.fill.solid()
c13_main.fill.fore_color.rgb = GREEN_LIGHT
c13_main.line.color.rgb = GREEN_BORDER
tf = c13_main.text_frame
p = tf.paragraphs[0]
p.text = "CÁC NGUYÊN TẮC PHÒNG THỦ XSS THEO KHUYẾN NGHỊ OWASP:"
p.font.name = "Arial"
p.font.size = Pt(15)
p.font.bold = True
p.font.color.rgb = ACCENT_GREEN
add_p(tf, "1. Tận dụng cơ chế mã hóa tự động của Razor View:\n   Sử dụng cú pháp mặc định @model.Property thay vì @Html.Raw() đối với mọi dữ liệu bắt nguồn từ người dùng. Các ký tự nhạy cảm (<, >, \", ') tự động được chuyển đổi sang thực thể HTML an toàn.", size=13, space_before=8)
add_p(tf, "2. Mã hóa dữ liệu với System.Text.Encodings.Web:\n   Sử dụng HtmlEncoder.Default.Encode() để xử lý dữ liệu trước khi lưu trữ hoặc phản hồi.", size=13, space_before=8)
add_p(tf, "3. Thiết lập thuộc tính an toàn cho Cookie:\n   Kích hoạt HttpOnly = true và Secure = true để bảo vệ token phiên làm việc khỏi sự can thiệp của mã kịch bản phía client.", size=13, space_before=8)
add_p(tf, "4. Triển khai Content Security Policy (CSP Header):\n   Cấu hình chính sách bảo mật nội dung để hạn chế phạm vi thực thi kịch bản và chặn các nguồn tài nguyên không tin cậy.", size=13, space_before=8)


# ==============================================================================
# CHỦ ĐỀ 3: CSRF ATTACK — NGUYỄN MINH CƯỜNG
# ==============================================================================
# Slide 14: Bản chất CSRF
s14 = prs.slides.add_slide(blank_layout)
add_header(s14, "3.1. CSRF Attack: Bản Chất Kỹ Thuật & Điều Kiện Tấn Công", "CHỦ ĐỀ 3 • NGUYỄN MINH CƯỜNG (23A1001D0058)")
c14_1 = s14.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.35), card_w, Inches(5.4))
c14_1.fill.solid()
c14_1.fill.fore_color.rgb = CARD_BG
c14_1.line.color.rgb = CARD_BORDER
tf = c14_1.text_frame
p = tf.paragraphs[0]
p.text = "Bản Chất Kỹ Thuật Của Tấn Công CSRF"
p.font.name = "Arial"
p.font.size = Pt(15)
p.font.bold = True
p.font.color.rgb = ACCENT_RED
add_p(tf, "• CSRF (A01:2021 - Broken Access Control) là kỹ thuật tấn công mượn quyền người dùng hợp lệ.\n• Điều kiện tiên quyết:\n  1. Nạn nhân có phiên đăng nhập hợp lệ tại ứng dụng mục tiêu (Cookie xác thực còn hiệu lực).\n  2. Trình duyệt tự động đính kèm Cookie xác thực theo mọi HTTP Request gửi tới tên miền của ứng dụng.\n  3. Kẻ tấn công lừa nạn nhân truy cập một trang web độc hại do kẻ tấn công kiểm soát.\n• Điểm mù kiểm soát: Máy chủ chỉ xác thực Cookie của người dùng mà không kiểm tra nguồn gốc phát sinh Request.", size=12.5, space_before=8)

c14_2 = s14.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.85), Inches(1.35), card_w, Inches(5.4))
c14_2.fill.solid()
c14_2.fill.fore_color.rgb = RED_LIGHT
c14_2.line.color.rgb = RED_BORDER
tf = c14_2.text_frame
p = tf.paragraphs[0]
p.text = "Action Xử Lý Giao Dịch Thiếu Kiểm Soát"
p.font.name = "Arial"
p.font.size = Pt(15)
p.font.bold = True
p.font.color.rgb = ACCENT_RED
add_p(tf, "// ACTION NHẬN YÊU CẦU KHÔNG CÓ BẢO VỆ TOKEN:\n[HttpPost]\npublic IActionResult TransferVulnerable(decimal amount)\n{\n    var victim = _context.Wallets.Find(1);\n    var hacker = _context.Wallets.Find(2);\n    \n    victim.Balance -= amount;\n    hacker.Balance += amount;\n    _context.SaveChanges();\n    return RedirectToAction(\"Index\");\n}", size=10, color=RGBColor(159, 18, 57), space_before=8, font_name="Courier New")
add_p(tf, "Khi Action không yêu cầu mã xác thực chống giả mạo, các yêu cầu POST bắt nguồn từ một website khác vẫn được máy chủ thực thi bình thường do Cookie xác thực được đính kèm tự động.", size=12, bold=True, color=TEXT_HEAD, space_before=10)

# Slide 15: Kịch bản lừa đảo bẫy trúng thưởng
s15 = prs.slides.add_slide(blank_layout)
add_header(s15, "3.2. CSRF Attack: Kịch Bản Khai Thác Bằng Biểu Mẫu Ẩn", "CHỦ ĐỀ 3 • NGUYỄN MINH CƯỜNG (23A1001D0058)")
c15_main = s15.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.35), Inches(11.733), Inches(5.4))
c15_main.fill.solid()
c15_main.fill.fore_color.rgb = RED_LIGHT
c15_main.line.color.rgb = RED_BORDER
tf = c15_main.text_frame
p = tf.paragraphs[0]
p.text = "Kịch Bản: Trang Web Giả Mạo Điều Khiển Yêu Cầu Chuyển Tiền"
p.font.name = "Arial"
p.font.size = Pt(16)
p.font.bold = True
p.font.color.rgb = ACCENT_RED
add_p(tf, "1. Nạn nhân có phiên làm việc đang hoạt động trên hệ thống tư vấn trực tuyến.\n2. Nạn nhân được dẫn dụ truy cập một trang web độc hại ở một tab khác (AttackerSite.cshtml).\n3. Trang web độc hại chứa một biểu mẫu ẩn trỏ trực tiếp tới địa chỉ nhận lệnh của ứng dụng mục tiêu:", size=13, space_before=8)
add_p(tf, "<!-- BIỂU MẪU ẨN TRÊN TRANG WEB CỦA KẺ TẤN CÔNG -->\n<form id=\"exploit\" action=\"http://localhost:5076/Csrf/TransferVulnerable\" method=\"POST\">\n    <input type=\"hidden\" name=\"amount\" value=\"20000000\" />\n</form>\n<script>document.getElementById('exploit').submit();</script>", size=11, bold=True, color=RGBColor(159, 18, 57), space_before=6, font_name="Courier New")
add_p(tf, "Hậu quả thực tế: Yêu cầu giao dịch được thực hiện mà không có sự đồng thuận chủ động của người dùng, dẫn đến thay đổi số dư hoặc cấu hình tài khoản.", size=13, bold=True, color=ACCENT_RED, space_before=8)

# Slide 16: Minh chứng Web UI CSRF (Nhúng ảnh thật - CĂN CHUẨN KHÔNG ĐÈ)
s16 = prs.slides.add_slide(blank_layout)
add_header(s16, "3.3. CSRF Attack: Minh Chứng Kiểm Thử Qua Chrome DevTools (Request Headers)", "CHỦ ĐỀ 3 • NGUYỄN MINH CƯỜNG (23A1001D0058)")
p_csrf = "/media/hocjsoo/New Volume/OWASP_Demo/screenshots/03_csrf_devtools_proof.png"
if os.path.exists(p_csrf):
    s16.shapes.add_picture(p_csrf, Inches(2.06), Inches(1.35), width=Inches(9.2))

cap16 = s16.shapes.add_textbox(Inches(0.8), Inches(6.58), Inches(11.733), Inches(0.4))
tf16 = cap16.text_frame
tf16.margin_top = tf16.margin_bottom = tf16.margin_left = tf16.margin_right = 0
p = tf16.paragraphs[0]
p.text = "Minh chứng đối chứng: Soi gói tin HTTP POST gửi ngầm từ Attacker Site: Đính kèm Cookie tự động nhưng thiếu Anti-Forgery Token của Hacker (Trái) vs Chặn đứng yêu cầu thiếu Token (Phải)"
p.font.name = "Arial"
p.font.size = Pt(11)
p.font.bold = True
p.font.color.rgb = ACCENT_GREEN
p.alignment = PP_ALIGN.CENTER

# Slide 17: Minh chứng CSDL ví tiền
s17 = prs.slides.add_slide(blank_layout)
add_header(s17, "3.4. CSRF Attack: Minh Chứng Biến Động CSDL dbo.UserWallets", "CHỦ ĐỀ 3 • NGUYỄN MINH CƯỜNG (23A1001D0058)")
b17_1 = s17.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.35), card_w, Inches(5.4))
b17_1.fill.solid()
b17_1.fill.fore_color.rgb = CARD_BG
b17_1.line.color.rgb = CARD_BORDER
tf = b17_1.text_frame
p = tf.paragraphs[0]
p.text = "1. Trạng Thái CSDL Ban Đầu"
p.font.name = "Arial"
p.font.size = Pt(14)
p.font.bold = True
p.font.color.rgb = ACCENT_BLUE
add_p(tf, "SELECT Id, OwnerName, Balance FROM dbo.UserWallets;\n\nId | OwnerName                      | Balance (VND)\n---+--------------------------------+---------------\n 1 | Nạn nhân (Nguyễn Danh Học)     | 50,000,000.00\n 2 | Kẻ tấn công (Hacker BlackHat)  |          0.00\n(2 rows affected)", size=9.5, color=TEXT_HEAD, space_before=6, font_name="Courier New")
add_p(tf, "Số dư ví tài khoản nạn nhân ở trạng thái ban đầu là 50 triệu VNĐ, tài khoản đối tượng tấn công là 0 VNĐ.", size=11.5, space_before=8)

b17_2 = s17.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.85), Inches(1.35), card_w, Inches(5.4))
b17_2.fill.solid()
b17_2.fill.fore_color.rgb = RED_LIGHT
b17_2.line.color.rgb = RED_BORDER
tf = b17_2.text_frame
p = tf.paragraphs[0]
p.text = "2. Trạng Thái CSDL Sau Khai Thác CSRF"
p.font.name = "Arial"
p.font.size = Pt(14)
p.font.bold = True
p.font.color.rgb = ACCENT_RED
add_p(tf, "SELECT Id, OwnerName, Balance FROM dbo.UserWallets;\n\nId | OwnerName                      | Balance (VND)\n---+--------------------------------+---------------\n 1 | Nạn nhân (Nguyễn Danh Học)     | 30,000,000.00\n 2 | Kẻ tấn công (Hacker BlackHat)  | 20,000,000.00\n(2 rows affected)", size=9.5, color=RGBColor(159, 18, 57), space_before=6, font_name="Courier New")
add_p(tf, "Số dư thực tế trong CSDL bị thay đổi hoàn toàn tự động khi nạn nhân truy cập liên kết độc hại, chứng minh mức độ rủi ro nghiêm trọng của giao dịch thiếu token.", size=11.5, space_before=8)

# Slide 18: Phòng thủ CSRF Token & Fetch API
s18 = prs.slides.add_slide(blank_layout)
add_header(s18, "3.5. CSRF Attack: Cơ Chế Synchronizer Token & Fetch API", "CHỦ ĐỀ 3 • NGUYỄN MINH CƯỜNG (23A1001D0058)")
c18_1 = s18.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.35), card_w, Inches(5.4))
c18_1.fill.solid()
c18_1.fill.fore_color.rgb = GREEN_LIGHT
c18_1.line.color.rgb = GREEN_BORDER
tf = c18_1.text_frame
p = tf.paragraphs[0]
p.text = "Mô Hình Synchronizer Token Pattern"
p.font.name = "Arial"
p.font.size = Pt(15)
p.font.bold = True
p.font.color.rgb = ACCENT_GREEN
add_p(tf, "// 1. TRONG RAZOR VIEW: NHÚNG TOKEN VÀO FORM\n<form asp-action=\"TransferSecure\" method=\"post\">\n    @Html.AntiForgeryToken()\n    <input type=\"number\" name=\"amount\" />\n    <button type=\"submit\">Chuyển tiền</button>\n</form>\n\n// 2. TRONG CONTROLLER: BẮT BUỘC KIỂM TRA TOKEN\n[HttpPost]\n[ValidateAntiForgeryToken]\npublic IActionResult TransferSecure(decimal amount)\n{\n    // Chỉ thực thi khi Token từ Form khớp với Token trong Cookie\n}", size=9.5, color=RGBColor(6, 95, 70), space_before=8, font_name="Courier New")
add_p(tf, "Cơ chế bảo vệ:\nMáy chủ phát hành một cặp token: một lưu trong Cookie và một nhúng vào biểu mẫu. Trang web độc hại bị chặn đọc token bởi chính sách Same-Origin Policy (SOP).", size=11.5, space_before=8)

c18_2 = s18.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.85), Inches(1.35), card_w, Inches(5.4))
c18_2.fill.solid()
c18_2.fill.fore_color.rgb = CARD_BG
c18_2.line.color.rgb = CARD_BORDER
tf = c18_2.text_frame
p = tf.paragraphs[0]
p.text = "Tích Hợp Token Trong Fetch API (Slide Buổi 9)"
p.font.name = "Arial"
p.font.size = Pt(15)
p.font.bold = True
p.font.color.rgb = ACCENT_BLUE
add_p(tf, "// TRÍCH XUẤT VÀ GỬI TOKEN QUA HTTP HEADER:\nconst token = document.querySelector('input[name=\"__RequestVerificationToken\"]').value;\n\nfetch('/Csrf/TransferAjax', {\n    method: 'POST',\n    headers: {\n        'Content-Type': 'application/json',\n        'RequestVerificationToken': token // HEADER BẢO VỆ\n    },\n    body: JSON.stringify({ amount: 5000000 })\n})\n.then(response => response.json());", size=9.5, color=TEXT_HEAD, space_before=8, font_name="Courier New")
add_p(tf, "Ứng dụng trong thực tế:\nĐối với các tác vụ gọi AJAX không tải lại trang, token được đọc từ thẻ ẩn và truyền qua HTTP Header để đảm bảo vừa tiện dụng vừa an toàn.", size=11.5, space_before=8)

# ==============================================================================
# SLIDE 19: LỜI KẾT
# ==============================================================================
s19 = prs.slides.add_slide(blank_layout)
add_header(s19, "Kết Luận Báo Cáo Thực Nghiệm OWASP Top 10")
main_c19 = s19.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.0), Inches(1.5), Inches(11.333), Inches(5.0))
main_c19.fill.solid()
main_c19.fill.fore_color.rgb = CARD_BG
main_c19.line.color.rgb = CARD_BORDER
tf = main_c19.text_frame
p = tf.paragraphs[0]
p.text = "TỔNG KẾT KẾT QUẢ THỰC NGHIỆM CỦA NHÓM WNC.G01"
p.font.name = "Arial"
p.font.size = Pt(16)
p.font.bold = True
p.font.color.rgb = ACCENT_BLUE

add_p(tf, "1. Nguyễn Danh Học (SQL Injection): Phân tích cơ chế tấn công bẻ gãy cú pháp CSDL ở tầng Backend và giải pháp phòng thủ triệt để bằng Parameterized Query / EF Core LINQ.", size=13, space_before=14)
add_p(tf, "2. Nguyễn Thanh Bình (Stored XSS): Phân tích nguy cơ đánh cắp Session Cookie qua dữ liệu lưu trữ CSDL và giải pháp tự động HTML Encoding của Razor View kết hợp cờ HttpOnly.", size=13, space_before=12)
add_p(tf, "3. Nguyễn Minh Cường (CSRF Attack): Phân tích kỹ thuật mượn quyền người dùng qua biểu mẫu ẩn và giải pháp phòng thủ bằng Synchronizer Token Pattern trên cả Form lẫn Fetch API.", size=13, space_before=12)
add_p(tf, "Kính chúc thầy sức khỏe và công tác tốt. Nhóm WNC.G01 xin trân trọng cảm ơn!", size=15, bold=True, color=ACCENT_RED, space_before=24)

out_file = "/media/hocjsoo/New Volume/OWASP_Demo/Slide_OWASP_Top3_Nhom_WNC_G01.pptx"
prs.save(out_file)
print(f"Perfect 19-slide presentation saved to: {out_file}")
